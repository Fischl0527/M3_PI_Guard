"""Create a two-page teacher brief and expand the existing continuous report."""

from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree as ET

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / "harness/workspaces/teacher-brief-20261014"
FULL = ROOT / "docs/M3_PI_Guard_研究介绍_2026-10-14.docx"
BRIEF = ROOT / "docs/M3_PI_Guard_研究介绍_教师交流两页版_2026-10-14.docx"
BACKUP = WORK / "original_detailed_report.docx"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{" + NS["w"] + "}"

# The user's latest brief scope is historical work plus current corrected-v1 only.
BRIEF_CONTENT = [
    ("title", "M3 PI Guard 研究介绍"),
    ("subtitle", "工具型大语言模型的授权与运行时防护"),
    ("body", "本项目研究大语言模型调用工具时，如何在完成用户任务的同时控制提示注入引起的越权操作。前期从权限清单和运行时检查出发，完成公开基准、模块消融、来源对照、困难攻击样本及自适应攻击等研究，逐步形成对授权信息、安全结果与任务完成之间关系的认识。目前在这些历史工作的基础上开展 v1 执行与测量链路修复。"),
    ("heading", "一 研究目标与基本方法"),
    ("body", "用户请求和可信应用声明先被编译为任务权限；模型提出工具名称与参数，独立的 Guard 在执行前检查，给出 ALLOW、ASK 或 DENY。检查不仅覆盖工具，还覆盖收件人、文件路径、动作顺序和数据流。授权不足时补充确认，明确违反规则时拒绝，满足规则时执行并登记结果。"),
    ("body", "例如，用户要求读取 A 并将摘要写入 B，权限应对应“读 A、写 B”。发送邮件时，邮箱格式正确还不够，收件人也必须符合用户授权；收件人获准后，正文是否可以传送仍需结合可信数据分类判断。这样可以把工具使用范围、参数目标和内容保密落实到具体动作上。"),
    ("heading", "二 早期研究怎样展开"),
    ("body", "早期工作的出发点是长任务中权限清单难以持续约束模型的问题。项目逐步将文字约束落实为执行前的程序检查，在 AgentDojo 公开基准及 M3-v2 模拟任务上，以无防护、提示词防御和不同 Guard 配置作对照，分别观察正常任务完成、攻击结果与执行负担。"),
    ("body", "在此基础上，历史增量研究进一步拆分影响因素。仅权限配置主要限制工具和参数范围；完整模块还结合来源、执行顺序、数据流和确认规则。普通提示与安全提示作为另一个因素交叉比较，使“提示词改变了行为”和“执行检查改变了行为”能够分别分析。[1]"),
    ("heading", "三 历史增量实验与主要观察"),
    ("body", "E1 提示词与模块消融：180 个案例在四条件下形成 720 条执行。普通提示下，仅权限和完整模块的正常工具任务完成分别为 23/60、21/60；攻击历史风险标签阳性分别为 1/100、0/100。这一阶段发现，增加模块的效果需要同时结合正常完成与攻击活性评价。[1]"),
    ("body", "E2 来源对照：60 个受控任务保持文本、参数值及位置相同，只改变关键内容的可信来源，完成 240 条端到端执行和 480 次离线 Guard 判定。检查器能够区分部分来源条件，但端到端实际工具动作均为零。这促使项目把机制是否生效与整个任务是否推进分开研究。[1]"),
    ("body", "E3 困难攻击复查：从历史曾出现风险的案例中选取 60 个任务，完成 240 条条件执行，原有风险仅在 1/60 个案例中重现。同时有 17 条解析失败和 38 条提供方失败。结果提示，历史攻击难度与当前运行窗口中的活性可能不同，失败也是解释实验现象的重要线索。[1]"),
    ("body", "E4 第二模型复核：Flash 在 90 个嵌套案例上完成 360 条执行。四条件的正常工具任务完成为 10/30、8/30、10/30、10/30，攻击历史风险标签均为 0/50。该实验补充了同一服务提供方内的模型观察，也显示任务完成与攻击事件的变化需要分别比较。[1]"),
    ("body", "E5 自适应攻击尾部：对预先选定、前 200 次查询尚未出现历史成功的 12 个任务，继续执行第 201 至 500 次查询，共 3,600 次终态查询，未发现新增事件。该实验用于观察持续尝试后的变化，其结果对应这组预选任务的后续区间。[1]"),
    ("heading", "四 历史研究形成的认识"),
    ("body", "前期工作将研究重点逐步收束到授权如何影响安全与任务完成。权限不足会使正常动作停在确认环节，范围过宽又可能允许偏离任务的操作；来源检查具有组件层作用，还需要真实执行链验证。研究因此同时记录提案、检查决定、执行效果、完成状态和失败。历史风险标签按原实验定义保留，后续修复另用清楚的执行证据解释结果。"),
    ("heading", "五 目前研究的 v1 修复与验证"),
    ("body", "修正版 v1 重点修复动作与资源绑定、同工具多条规则的完整匹配、可信数据标签传播和确认状态衔接。Pro 与 Flash 分别完成 48 任务、六配置的开发研究，每模型 288 条条件执行。[2]"),
    ("body", "Pro 的 AgentDojo 普通提示子集中，无防护、仅权限、完整模块的正常完成为 6/8、3/8、3/8，攻击判定阳性为 5/12、2/12、0/12。目前这一步主要为后续研究建立准确的授权、执行与评分链路，并继续解释防护对正常任务的影响。[2]"),
    ("heading", "六 拟与老师讨论的问题"),
    ("body", "希望请教，是否以“授权构造对工具型智能体安全与任务完成的影响”作为论文主线；已有历史实验中，哪些结果最适合作为主要证据；下一步应优先完善权限表达、正常任务的确认交互，还是补充攻击活性更充分的验证场景。"),
    ("source", "资料依据  [1] 项目增量实证强化阶段报告（2026年8月26—28日）；[2] 修正版研究中文稿的 v1 开发研究及对应实验记录。M3-v2 为模拟任务基准名称；v1 为目前介绍的权限与执行链修正版。"),
]
# One citation at the end of each evidence group avoids isolated citation lines.
BRIEF_CONTENT = [
    (kind, text.replace("[1]", "").rstrip()
     if kind == "body" and not text.startswith("E5 ") else text)
    for kind, text in BRIEF_CONTENT
]

# Cover the original E1--E17 programme separately from the later five-study series.
v1_start = next(i for i, (kind, text) in enumerate(BRIEF_CONTENT)
                if kind == "heading" and text.startswith("五 "))
BRIEF_CONTENT = [
    *BRIEF_CONTENT[:6],
    ("heading", "二 早期 E1 至 E17 实验概况"),
    ("body", "防护效果与机制：E1 在相同提示下比较有无完整 Guard；E2 检验单独强化提示词的作用；E3 检验强化提示后 Guard 的增量；E4 通过逐项移除组件定位贡献与代价；E5 比较不同本地模型，考察效果对模型选择的依赖。"),
    ("body", "攻击覆盖：E6 按攻击家族分析残余风险；E7 比较不同语言形式；E8 检验编码与隐藏变体；E9 改变上下文长度和注入位置；E10 考察多轮、延迟与累积攻击；E11 研究恶意内容能否沿读取、检索、邮件和文件等工具链传播。"),
    ("body", "可用性与运行代价：E12 测量正常任务完成及误拦截；E13 在验证集上校准 ASK、DENY 阈值，比较漏放与误阻；E14 测量检查延迟、端到端开销、并发吞吐和显存占用，将模型运行成本与 Guard 自身开销分别观察。"),
    ("body", "稳定性与失败解释：E15 比较随机种子和温度变化；E16 根据前轮响应构造后续攻击，观察自适应绕过；E17 对分层抽样记录开展人工复核，把泄露标签、误拦截、等待确认及运行失败分类。以上构成前期从防护效果、攻击适应性到效用与解释的完整实验布局。[1]"),
    ("heading", "三 代表性历史结果与研究认识"),
    ("body", "历史结果同时呈现保护作用与正常任务代价。例如，E1 的 gemma3:4b 普通提示子集中，攻击历史标签阳性由无 Guard 的 126/189 降至完整 Guard 的 0/189；E12 同模型正常完成由 48/81 降至 42/81。联合观察这两个端点，才便于分析限制是否保留了完成合法任务的能力。[1]"),
    ("body", "E9 的校准上下文实验在历史评价口径下记录无 Guard 36/60、完整 Guard 0/60 的攻击成功数。E14 相同正常路径的配对研究记录 Guard 自身 p95 延迟约 0.45 毫秒、端到端开销中位数约 0.20%。E17 由两位评审完成 360 条分层样本复核，其中 300 条为代表性样本，补充了自动指标背后的失败解释。[1]"),
    ("body", "这些结果把研究重点逐步收束到权限表达和执行过程：权限过窄会妨碍正常任务，范围过宽可能放行偏离任务的操作。历史安全标签按原定义保留；后续将动作提案、Guard 决定、实际执行、任务完成与失败分别记录，以定位差异发生在哪一环节。"),
    ("heading", "四 后续增量研究"),
    ("body", "此后另开展五项增量研究，补充同期提示词与模块因子比较、来源反事实、困难攻击复查、Flash 复核和自适应尾部观察。它们进一步发现，检查器在受控条件下能够区分来源，端到端任务却可能没有推进；历史攻击也不一定稳定重现。这使确认负担、模型失败和攻击活性成为后续分析的重要内容。[2]"),
    *((kind, text.replace("[2]", "[3]"))
      for kind, text in BRIEF_CONTENT[v1_start:-1]),
    ("source", "资料依据  [1] 早期 E1—E17 实验注册表及分析记录；[2] 增量实证强化阶段报告（2026年8月）；[3] 修正版研究中文稿的 v1 开发研究。当前部分仅介绍 v1；M3-v2 为模拟任务基准名称。"),
]

EXPLANATIONS = {
    "N01": ["这条流程把“模型想做什么”和“系统允许做什么”分别交给模型与检查器。比如文档正文要求把资料转发给另一个邮箱，模型可能把它生成成发送提案；但正文属于待处理内容，不能替代用户批准收件人的权限。执行层记录实际发生的状态变化，评价层再判断是否完成原任务，因此一次提案被提出、获得允许和成功执行是三个不同的观察点。"],
    "N02": ["版本变化围绕同一条问题链展开。v1先让动作、资源和检查记录准确对应；v2进一步处理“目标需要限制、内容需要生成”的参数差异；v3尝试直接表达多步骤任务的契约；v4减少模型生成规则的自由度，改为提取意图后由程序生成。前一阶段暴露的问题构成后一阶段的设计动机，版本号本身并不表示实验效果逐次提高。"],
    "N14": ["阅读这组结果需要同时看两类任务。攻击子集用于观察模型是否偏离原任务，正常子集用于观察加入防护后还能完成多少工作。仅比较攻击阳性会遗漏正常完成从6/8降到3/8的代价；联合观察则能提出更具体的问题：哪些限制阻止了攻击，又有哪些限制使正常任务无法继续。该开发研究为后续授权信息对照提供了动机。"],
    "N03": ["图中的每一部分是任务在某种配置下执行的记录数量。同一个任务可以分别在仅权限、完整模块和不同授权条件下运行，所以执行条数大于底层任务数。同期校准复用了主研究中的案例，其作用是补充同一任务上的无防护参照；合成控制则主动设置数据分类和目标，定位某条规则怎样改变决定。三部分回答的问题不同，应分别阅读。"],
    "N15": ["v2要解决的关键歧义是，发送目标应被权限约束，而摘要等内容需要由模型生成。精确目标条件把离线派生的操作目标交给Guard，同时继续让模型产生实际工具参数。AgentDojo完成数改善，说明更多有效的目标信息能够解除一部分阻断；M3-v2变化较小，说明该作用与任务结构有关。它为改进自动编译提供诊断方向，而不是一项可以直接部署的授权来源。"],
    "N04": ["五组柱形保持同一批任务和执行模型，依次改变防护配置与授权信息。最明显的差异发生在AgentDojo：自动授权下两个模块均仅完成8项，明确目标后完成数回升，但仍低于无防护。这个模式把问题从“是否需要更多限制”推进到“如何准确表达本次任务需要的权限”。M3-v2的差异较小，需要结合它的确认任务占比和任务组构成解释。"],
    "N05": ["Flash复核沿用预先选定的任务子集，重点观察方向是否与Pro一致。AgentDojo再次出现自动权限完成较低、精确目标改善部分任务的模式，而M3-v2五条件相同。两个模型的样本量和任务覆盖不同，因此应在各自分母内比较条件，再讨论相同的现象；不能把各条件条数或两种模型的记录简单相加当作独立任务。"],
    "N06": ["配对比较追踪同一个任务在两种授权条件下的变化：原先未完成而随后完成记为改善，反向变化记为损失。图中的点表示净完成差，横线表示已有统计区间。这种比较比单看两根比例柱更贴近授权信息的作用。AgentDojo的两个主比较显示明确改善，因而后续改进应优先检查哪些必要目标在自动编译时缺失、过窄或绑定错误。"],
    "N07": ["这里先检查无防护条件是否真的出现攻击事件，再解释防护条件的变化。若无防护也始终为零，就缺少观察防护降低事件的机会。Pro/AgentDojo提供了7个对照事件，但比较仍受样本量和配对结果限制。安全观察需要与前面的正常任务完成一同呈现，才能判断限制产生的是有用保护，还是主要让系统更少执行动作。"],
    "N08": ["ASK表示当前证据或批准不足，需要额外信息才能继续，不能把它理解为已经成功拦截一次攻击。在没有模拟用户补充批准的实验中，频繁ASK会使本来正常的任务停住。该图把确认负担与完成放在同一任务分母下观察，用于定位授权链条的瓶颈；真正的交互方案还需要考察澄清问题是否准确，以及补充后能否恢复完成。"],
    "N10": ["受控夹具主动提供能够触发某项机制的提案，适合验证检查规则是否按设计工作；自然提案回放使用模型实际提出的动作，观察增量机制在当前任务分布中出现多少作用机会。358对中只有3对改变决定，提示不少提案在基础权限检查阶段就得到相同结论。下一步可以针对已捕获动作分析阻断原因和数据来源，而不能用夹具的触发比例代替自然任务的收益。"],
    "N11": ["编译通过表示草稿能够成为运行时可检查的任务契约，尚未涉及模型是否完成整个用户任务。44次响应中失败较多，说明瓶颈主要出现在自然语言任务到结构化表达的转换。目标、内容和证据引用如果混淆，后续执行器即使正确也无法得到足够明确的授权。分来源和提示阶段查看失败，有助于判断问题来自任务复杂度、表达能力还是编译约束。"],
    "N12": ["v4尝试让模型只提取可定位的任务事实，再由程序物化规则，减少自由生成契约时的格式与语义错误。当前32任务的比较尚未体现这种设计带来的可运行性改善。可继续检查的具体问题包括：用户原话能否准确支持每个参数，对象选择是否超出后端可验证的范围，以及某些需要动态查询的关系是否缺少可表达的有限形式。"],
    "N13": ["覆盖图同时列出已有记录和未执行项，帮助区分实验计划、运行状态与实际效果。因编译失败或需要澄清而终止的记录仍属于已知流程状态，但不会产生完整的模型执行链；Flash后续300项则确实没有运行。这里的数量用于交代开发过程，不能从有多少记录推导有多少次任务成功，也不能把计划覆盖解释为已经完成双模型效果验证。"],
}

HISTORY = [
    ("heading", "2 3 后续五项增量研究与修正动机"),
    ("body", "早期研究关注权限清单在长任务中的约束作用，随后把检查从模型上下文中的文字要求落实到执行前的程序规则。项目先在公开基准与模拟任务中观察不同配置，再通过提示词和模块消融、来源对照及自适应攻击追问现象的来源。由此形成了后续的研究问题：权限从哪里来、怎样绑定本次动作，以及更严格的检查会怎样影响正常任务完成。"),
    ("body", "2026年8月的增量研究包含E1至E5。E1以720条同期2×2条件执行分离提示词与模块因素；E2以240条端到端执行和480次离线Guard判定，在保持文本与参数值相同的条件下改变来源权威性；E3使用240条风险富集挑战执行观察困难样本；E4以360条Flash执行检查另一模型的表现。E2的端到端实际动作均为零，而离线Guard能够区分部分来源条件，说明机制能触发与系统能完成任务需要分别检验。"),
    ("body", "E5对预先选定、前200次查询尚未出现历史风险标签成功的12个任务继续执行第201至500次查询，共3,600个终态查询，未观察到新增事件。这一结果描述的是该组幸存任务的条件尾部。它推动项目同时关注攻击活性和任务选择，避免仅凭大批零事件就推断所有场景都得到保护。历史实验保留原评价标签，后续修正版将实际效果与隔离层拒绝分别记录。"),
    ("source", "上述历史过程依据《M3-PI Guard 增量实证强化阶段报告》（2026年8月26—28日）；历史图表与来源索引见附录及docs/picture图集。"),
]

EARLY_EXPERIMENTS = [
    ("E1 主防护效果", "在相同提示条件下比较不使用Guard与使用完整Guard，观察外部执行检查带来的差异。以gemma3:4b普通提示子集为例，历史攻击标签阳性为126/189与0/189。这里先固定提示因素，再讨论模块作用，避免把提示词变化一并算成执行检查的收益。"),
    ("E2 提示词强化", "关闭Guard，比较普通提示与强化提示，让模型更多地识别不可信内容并遵循原任务。它回答仅靠提示层能产生什么变化，为后续判断执行检查是否仍有增量提供参照；提示词是模型输入的一部分，不能替代实际工具参数的授权检查。"),
    ("E3 强化提示后的增量", "在已经使用强化提示的条件下，再比较有无Guard。该设计与E2衔接，用来判断提示层与执行层的作用是否重叠。分析时要同时保留攻击结果和正常任务表现，才能判断新增限制减少了哪类动作，以及是否影响合法工作。"),
    ("E4 组件消融", "分别移除或单独启用检测、任务动作、数据流等组件，比较完整配置和简化配置。消融帮助定位哪些场景触发了规则，以及组件增加了哪些限制。多个配置都没有事件时，还应检查它们是否收到足以触发差异的动作提案，避免单凭零值给组件排序。"),
    ("E5 跨模型比较", "在不同本地模型上复用任务与防护配置，考察模型能力和输出行为对结果的影响。这一实验用于观察效果是否依赖特定模型，而不是把某一模型上的低攻击率推广到所有模型。比较时，任务完成、解析成功及调用行为同样需要对应各自模型报告。"),
    ("E6 攻击家族分析", "将攻击按家族分层，定位残余风险集中在哪些形式，以及完整模块对哪些模式更容易产生判定差异。家族分析还可以指导后续反例和回归样本设计：如果某家族始终没有可执行提案，需要先区分攻击不活跃与检查真正发挥作用。"),
    ("E7 多语言鲁棒性", "比较不同语言与表达形式下的攻击结果，考察改变表面语言是否影响模型与防护链。其研究价值是把语言变化从攻击目标本身分离出来，检查系统是否过度依赖特定措辞；结果应在各语言的实际样本与失败记录内解释。"),
    ("E8 编码与隐藏变体", "对攻击内容采用编码、字符拆分、混淆或隐藏等变体，考察表面形式变化能否绕过检测和执行约束。它与多语言实验共同覆盖内容形式的变化，同时说明检测到恶意文本和阻止越权工具动作是两个不同层次的观察。"),
    ("E9 长上下文与位置", "改变上下文负载和恶意内容出现的位置，检验任务约束在不同输入条件下的表现。注册表指定校准后的V2 r3为主证据，共120条正式首步记录；历史评价口径下，无Guard与完整Guard攻击成功数为36/60和0/60。读取时需要同时关注解析失败及每个长度、位置单元的样本量。"),
    ("E10 多轮攻击", "考察攻击是否可以通过延迟出现、逐轮累积或分阶段铺垫影响后续动作。该实验把单轮文本攻击扩展到对话过程，关注前文信息怎样进入后续决策，也为执行历史与确认状态需要持续记录提供具体场景。"),
    ("E11 跨工具链传播", "研究读取后发送、检索后调用API、记忆写入文件等链式任务中的攻击传播。一次读取合法，并不表示后续传送目标或写入位置已经得到许可。原实验与后续48条事后修订补充执行分别保留，用于解释不同工具链中的动作与授权关系。"),
    ("E12 正常任务效用", "用正常任务测量完成率和误拦截，使安全观察对应到实际可用性。gemma3:4b普通提示下，无Guard完成48/81，完整Guard完成42/81。该结果与E1同模型的攻击观察并列，能够展示限制产生的收益与代价，但两类任务各有自己的分母。"),
    ("E13 阈值校准", "仅在验证集上扫描ASK与DENY阈值，比较恶意放行、正常误拦截及确认负担，再固定参数。它检验判定规则对安全与效用的取舍；验证集上的分类指标用于选择参数，和后续独立任务的端到端完成属于不同结果。"),
    ("E14 性能与资源", "注册表指定校准后的V2 r3为主证据，包含432条正式记录和144对相同正常调用路径。记录Guard自身p95延迟约0.45毫秒、正常路径端到端开销中位数约0.20%，并观察并发吞吐与显存。显存是设备级观测，混合任务的吞吐也包含模型和调度影响。"),
    ("E15 随机稳定性", "在不同随机种子和温度下重复观察配置表现，检查某个结果是否只出现在一次采样中。分析按预设任务和条件保留重复结构，比较稳定性与失败，而不是把同一个任务的多次运行都作为新的独立任务。"),
    ("E16 自适应红队", "根据前轮响应，从固定变体目录构造后续攻击，检验已知反馈下的连续尝试。42条规范攻击谱系在三个轮次中未观察到历史成功率上升，并保留一个未解决轮次。它补充单次攻击以外的局部适应性观察，和后来扩展到查询预算的增量尾部研究分别报告。"),
    ("E17 人工失败复核", "两位独立盲评者完成360条分层样本，其中300条代表性样本、40条自动分歧和20条稀有复杂记录。复核分类包括泄露标签、误拦截、等待确认、解析失败等，用于解释自动分数背后的行为。总体比例使用代表性层的抽样权重，两个诊断层提供针对性案例分析。"),
]
HISTORY = [
    ("heading", "2 2 早期 E1 至 E17 实验布局与解释"),
    ("body", "早期实验注册表将E1至E17组织为防护效果、攻击覆盖、效用与性能、稳定性及人工解释等方向。下面按每项实验的问题和比较方式概括。这里的E编号属于早期实验系列；随后2026年8月五项增量研究重新使用了E1至E5编号，二者的任务、模型和评价版本分别对应各自记录。"),
    *(("body", name + "。" + text) for name, text in EARLY_EXPERIMENTS),
    ("source", "本节来源为experiments/registry.json、outputs/experiments/E1至E16的分析；E9/E14采用注册表指定的context_correction_v2_r3，E17采用human_stratified_360_v1/review记录。历史结果保持原评价定义。"),
    *HISTORY,
]


def digest(data: bytes) -> str:
    return sha256(data).hexdigest()


def paragraph_text(root):
    return ["".join(p.xpath(".//w:t/text()", namespaces=NS))
            for p in root.findall(".//w:p", NS)]


def drop_forced_pages(root):
    removed = 0
    for element in list(root.iter()):
        if (element.tag in {W + "pageBreakBefore", W + "lastRenderedPageBreak"}
                or (element.tag == W + "br" and element.get(W + "type") == "page")):
            element.getparent().remove(element)
            removed += 1
    # A section never needs a forced new page in the continuous detailed report.
    for element in root.findall(".//w:sectPr/w:type", NS):
        if element.get(W + "val") in {"nextPage", "oddPage", "evenPage"}:
            element.set(W + "val", "continuous")
            removed += 1
    return removed


def blacken(root):
    for element in root.findall(".//w:color", NS):
        element.attrib.clear()
        element.set(W + "val", "000000")


def prepare():
    WORK.mkdir(parents=True, exist_ok=True)
    if not BACKUP.exists():
        shutil.copy2(FULL, BACKUP)
        staged = subprocess.run(
            ["git", "diff", "--cached", "--binary"], cwd=ROOT,
            capture_output=True, check=True,
        ).stdout
        with ZipFile(BACKUP) as z:
            root = ET.fromstring(z.read("word/document.xml"))
            state = {
                "original_sha256": digest(BACKUP.read_bytes()),
                "original_text_paragraphs": paragraph_text(root),
                "media": {n: digest(z.read(n)) for n in z.namelist()
                          if n.startswith("word/media/") and not n.endswith("/")},
                "staged_diff_sha256": digest(staged),
            }
        (WORK / "before.json").write_text(
            json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


def create_xml_paragraph(text, prototype, heading=False):
    p = ET.Element(W + "p")
    ppr = prototype.find(W + "pPr")
    if ppr is not None:
        p.append(deepcopy(ppr))
    drop_forced_pages(p)
    r = ET.SubElement(p, W + "r")
    rpr = ET.SubElement(r, W + "rPr")
    fonts = ET.SubElement(rpr, W + "rFonts")
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(W + name, "黑体" if heading else "宋体")
    color = ET.SubElement(rpr, W + "color")
    color.set(W + "val", "000000")
    size = ET.SubElement(rpr, W + "sz")
    size.set(W + "val", "24" if heading else "23")
    if heading:
        ET.SubElement(rpr, W + "b")
    t = ET.SubElement(r, W + "t")
    t.text = text
    return p


def expand_detailed():
    with ZipFile(BACKUP) as z:
        root = ET.fromstring(z.read("word/document.xml"))
        body = root.find("w:body", NS)
        original_paragraphs = list(body.findall("w:p", NS))
        body_prototype = next(p for p in original_paragraphs
                              if "M3 PI Guard 研究大语言模型" in "".join(p.xpath(".//w:t/text()", namespaces=NS)))
        heading_prototype = next(p for p in original_paragraphs
                                 if "2 1 发送权限与内容保密" in "".join(p.xpath(".//w:t/text()", namespaces=NS)))
        added = []
        for p in original_paragraphs:
            text = "".join(p.xpath(".//w:t/text()", namespaces=NS))
            match = re.match(r"图\s+(N\d+)\s", text)
            if match:
                previous = p
                for explanation in EXPLANATIONS.get(match.group(1), []):
                    element = create_xml_paragraph(explanation, body_prototype)
                    previous.addnext(element)
                    previous = element
                    added.append(explanation)
            if text.startswith("3 v1 至 v4"):
                for kind, explanation in HISTORY:
                    element = create_xml_paragraph(
                        explanation,
                        heading_prototype if kind == "heading" else body_prototype,
                        heading=kind == "heading",
                    )
                    p.addprevious(element)
                    added.append(explanation)
        removed = drop_forced_pages(root)
        blacken(root)
        replacements = {"word/document.xml": ET.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)}
        for name in z.namelist():
            if name.startswith("word/") and name.endswith(".xml") and name != "word/document.xml":
                xml = ET.fromstring(z.read(name))
                before = ET.tostring(xml)
                removed += drop_forced_pages(xml)
                blacken(xml)
                if ET.tostring(xml) != before:
                    replacements[name] = ET.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
        tmp = WORK / "detailed_output.docx"
        with ZipFile(tmp, "w") as out:
            for info in z.infolist():
                out.writestr(info, replacements.get(info.filename, z.read(info.filename)))
    shutil.copy2(tmp, FULL)
    (WORK / "detailed_changes.json").write_text(json.dumps({
        "forced_page_elements_removed": removed,
        "added_paragraphs": added,
        "added_characters": sum(map(len, added)),
        "spacing": "Original paragraph and image spacing preserved; only new paragraphs inserted.",
    }, ensure_ascii=False, indent=2), encoding="utf-8")


def font(run, name="宋体", size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), name)


def create_brief():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin, section.right_margin = Cm(2.15), Cm(2.15)
    section.header_distance, section.footer_distance = Cm(0.75), Cm(0.75)
    for style in doc.styles:
        if style.type == 1 or style.type == 2:
            style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "宋体", Pt(12)
    for kind, text in BRIEF_CONTENT:
        p = doc.add_paragraph(style="Title" if kind == "title" else "Normal")
        pf = p.paragraph_format
        pf.line_spacing = Pt(19)
        pf.space_before, pf.space_after = Pt(0), Pt(4)
        pf.widow_control = True
        pf.keep_together = True
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        font(r)
        if kind == "title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(r, "黑体", 18, True)
            pf.line_spacing, pf.space_after = Pt(23), Pt(4)
            pf.keep_with_next = True
        elif kind == "subtitle":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            font(r, "宋体", 11.5)
            pf.space_after, pf.keep_with_next = Pt(9), True
        elif kind == "heading":
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            font(r, "黑体", 12.5, True)
            pf.space_before, pf.space_after = Pt(7), Pt(4)
            pf.keep_with_next = True
        elif kind == "source":
            font(r, "宋体", 9)
            pf.line_spacing, pf.space_before = Pt(13), Pt(4)
        else:
            pf.first_line_indent = Pt(24)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rpr.extend([color, size])
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)
    doc.core_properties.title = "M3 PI Guard 研究介绍"
    doc.core_properties.subject = "教师交流两页版 研究历史与主要发现"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "M3 PI Guard 授权 提示注入 研究介绍"
    raw = WORK / "brief_raw.docx"
    doc.save(raw)
    with ZipFile(raw) as z, ZipFile(BRIEF, "w") as out:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                xml = ET.fromstring(data)
                blacken(xml)
                drop_forced_pages(xml)
                for element in list(xml.iter()):
                    if element.tag in {W + name for name in ("shd", "highlight", "background", "pBdr", "borders")}:
                        element.getparent().remove(element)
                data = ET.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)
            out.writestr(info, data)
    shutil.copy2(BRIEF, WORK / "teacher_brief.docx")
    shutil.copy2(FULL, WORK / "detailed_report.docx")
    (WORK / "brief_text.txt").write_text("\n\n".join(t for _, t in BRIEF_CONTENT), encoding="utf-8")


if __name__ == "__main__":
    prepare()
    if "--brief-only" not in sys.argv:
        expand_detailed()
    create_brief()
    print(json.dumps({"brief": str(BRIEF), "detailed": str(FULL),
                      "brief_characters": sum(len(t) for _, t in BRIEF_CONTENT),
                      "added_detailed_paragraphs": sum(map(len, EXPLANATIONS.values())) + len(HISTORY)}, ensure_ascii=False))
