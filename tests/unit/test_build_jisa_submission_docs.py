from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from scripts import build_jisa_submission_docs as builder


def _document_text(document: Document) -> str:
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([*paragraphs, *cells])


def _table_widths(table) -> list[int]:
    return [int(grid_col.get(qn("w:w"))) for grid_col in table._tbl.tblGrid]


def test_highlights_are_separate_editable_and_within_jisa_limit() -> None:
    document = Document(builder.OUTPUTS["highlights"])
    items = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "List Bullet"]

    assert len(items) == 5
    assert all(len(item) <= 85 for item in items)
    assert all("[" not in item for item in items)


def test_cover_letter_preserves_bounded_claims_and_author_placeholders() -> None:
    document = Document(builder.OUTPUTS["cover"])
    text = _document_text(document)

    assert "Journal of Information Security and Applications" in text
    assert "did not outperform the policy package" in text
    assert "not a claim of a universally leading defense" in text
    assert "[Submission date]" in text
    assert "[Corresponding author name]" in text


def test_glossary_contains_all_reviewed_terms_without_markdown_tokens() -> None:
    document = Document(builder.OUTPUTS["glossary"])
    text = _document_text(document)

    assert sum(1 for paragraph in document.paragraphs if ": " in paragraph.text) == 15
    assert "exact_gold" in text
    assert "Unsafe effect" in text
    assert "**" not in text and "`" not in text


def test_author_form_keeps_unknown_facts_explicit_and_tables_fixed_width() -> None:
    document = Document(builder.OUTPUTS["author"])
    text = _document_text(document)

    assert len(document.tables) == 3
    assert "[required]" in text
    assert "do not upload with placeholders" in text
    assert "API keys" in text
    assert all(sum(_table_widths(table)) == builder.CONTENT_WIDTH_DXA for table in document.tables)
    assert all(
        int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))) == builder.CONTENT_WIDTH_DXA
        for table in document.tables
    )


def test_all_submission_documents_use_letter_geometry() -> None:
    for output in builder.OUTPUTS.values():
        document = Document(output)
        section = document.sections[0]
        assert round(section.page_width.inches, 3) == 8.5
        assert round(section.page_height.inches, 3) == 11.0
        assert round(section.top_margin.inches, 3) == 1.0
        assert round(section.right_margin.inches, 3) == 1.0
        assert round(section.bottom_margin.inches, 3) == 1.0
        assert round(section.left_margin.inches, 3) == 1.0
