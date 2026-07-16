from __future__ import annotations

import hashlib
import json
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document

from scripts.build_final_analysis import generate_final_analysis
from scripts.build_final_report import (
    FORMAL_PROTOTYPE_NOTE,
    ReportBuildError,
    build_final_report,
    load_report_evidence,
)
from tests.unit.test_build_final_analysis import (
    _install_synthetic_corrected_e17_loader,
    _make_repository,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(autouse=True)
def _synthetic_corrected_e17_loader(monkeypatch: pytest.MonkeyPatch) -> None:
    _install_synthetic_corrected_e17_loader(monkeypatch)


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _write_json(path: Path, payload: dict) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _synthetic_report_fixture(tmp_path: Path) -> dict[str, Path]:
    root, registry = _make_repository(tmp_path)
    analysis = root / "outputs" / "experiments" / "final_analysis"
    analysis.parent.mkdir(parents=True, exist_ok=True)
    generate_final_analysis(registry_path=registry, repository_root=root, output_dir=analysis)
    template = root / "docs" / "report" / "作品报告_模板.docx"
    template.parent.mkdir(parents=True)
    shutil.copy2(PROJECT_ROOT / "docs" / "report" / "作品报告_模板.docx", template)
    return {
        "root": root,
        "registry": registry,
        "analysis": analysis,
        "template": template,
        "output": root / "docs" / "report" / "作品报告_终稿.docx",
        "source": root / "docs" / "report_sections" / "final_report_source.md",
        "quality": root / "outputs" / "quality" / "final_report_format_check.md",
    }


def _build(paths: dict[str, Path], *, output_key: str = "output") -> dict:
    return build_final_report(
        repository_root=paths["root"],
        analysis_dir=paths["analysis"],
        registry_path=paths["registry"],
        template_path=paths["template"],
        output_docx=paths[output_key],
        source_markdown=paths["source"],
        quality_report=paths["quality"],
    )


def test_report_uses_only_verified_final_evidence_and_rejects_drift(tmp_path: Path) -> None:
    paths = _synthetic_report_fixture(tmp_path)
    # The synthetic repository intentionally has no historical Day5 metrics or
    # figures.  A successful build therefore proves they are not report inputs.
    assert not (paths["root"] / "outputs" / "metrics").exists()
    result = _build(paths)
    assert result["status"] == "completed"
    assert result["chart_count"] == 13
    assert result["visual_qa"] == "pending_render_and_every-page_inspection"

    report = Document(str(paths["output"]))
    assert len(report.sections) == 10
    assert len(report.inline_shapes) == 13
    with zipfile.ZipFile(paths["output"], "r") as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        footer_xml = "\n".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/footer") and name.endswith(".xml")
        )
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
    assert document_xml.count("PAGEREF") == 7
    assert footer_xml.count("PAGE") == 2
    assert "updateFields" in settings_xml and 'w:val="true"' in settings_xml
    text = "\n".join(paragraph.text for paragraph in report.paragraphs)
    text += "\n" + "\n".join(cell.text for table in report.tables for row in table.rows for cell in row.cells)
    assert FORMAL_PROTOTYPE_NOTE in text
    assert "Qwen2.5 代表样本" not in text
    assert "baseline 的 ASR 和 DSR 均为 0.5507" not in text
    assert "WRS 从 0.4891 提升至 0.9826" not in text
    assert "最终预注册验收矩阵" in text
    assert "Wilson 95% CI" in text
    assert "结果解读：" in text
    assert "证据与限制：" in text
    assert "表 2 主矩阵解析完整性" in text
    assert "A11 仍仅表示 E1 regular_full" in text
    assert "model-beta" in text
    assert "isolated-tool-sandbox-fixture-v1" in text
    assert "不是 manifest schema version" in text
    assert "E11 范围偏差披露" in text
    assert "pre_registered=false" in text
    assert "不与原始分母合并" in text
    assert "cross_tool_chain_patterns" in text
    assert paths["source"].is_file()
    assert paths["quality"].is_file()
    assert "DOCX visual render QA: PENDING" in paths["quality"].read_text(encoding="utf-8")

    # A modified final-analysis artifact is rejected before any new report is
    # created, even when the synthesis manifest still says completed.
    analysis_md = paths["analysis"] / "analysis.md"
    original_analysis = analysis_md.read_bytes()
    analysis_md.write_bytes(original_analysis + b"\ntamper\n")
    tampered_output = paths["root"] / "docs" / "report" / "tampered.docx"
    paths["tampered"] = tampered_output
    with pytest.raises(ReportBuildError, match="artifact (size|hash) drift"):
        _build(paths, output_key="tampered")
    assert not tampered_output.exists()
    analysis_md.write_bytes(original_analysis)

    # A registry entry that is no longer completed is rejected by the E0-E17
    # bundle validator.  Updating only the admission hash cannot bypass it.
    registry_payload = json.loads(paths["registry"].read_text(encoding="utf-8"))
    registry_payload["experiments"][4]["status"] = "running"
    _write_json(paths["registry"], registry_payload)
    final_manifest_path = paths["analysis"] / "manifest.json"
    original_manifest = final_manifest_path.read_bytes()
    final_manifest = json.loads(original_manifest.decode("utf-8"))
    final_manifest["registry_sha256"] = _hash(paths["registry"])
    _write_json(final_manifest_path, final_manifest)
    incomplete_output = paths["root"] / "docs" / "report" / "incomplete.docx"
    paths["incomplete"] = incomplete_output
    with pytest.raises(ReportBuildError, match="requires E0-E17 completed"):
        _build(paths, output_key="incomplete")
    assert not incomplete_output.exists()


def test_bundle_hash_drift_is_rejected_before_report_creation(tmp_path: Path) -> None:
    paths = _synthetic_report_fixture(tmp_path)
    metrics = paths["root"] / "evidence" / "E5" / "metrics.csv"
    metrics.write_text(metrics.read_text(encoding="utf-8-sig") + "tamper\n", encoding="utf-8")
    drift_output = paths["root"] / "docs" / "report" / "bundle-drift.docx"
    paths["drift"] = drift_output
    with pytest.raises(ReportBuildError, match="artifact (size|hash) drift"):
        _build(paths, output_key="drift")
    assert not drift_output.exists()


def test_report_rejects_live_e11_posthoc_drift(tmp_path: Path) -> None:
    paths = _synthetic_report_fixture(tmp_path)
    metrics = (
        paths["root"]
        / "outputs"
        / "experiments"
        / "E11_posthoc_amendment_v1"
        / "metrics.csv"
    )
    metrics.write_text(metrics.read_text(encoding="utf-8-sig") + "tamper\n", encoding="utf-8")
    drift_output = paths["root"] / "docs" / "report" / "posthoc-drift.docx"
    paths["posthoc_drift"] = drift_output

    with pytest.raises(ReportBuildError, match="post-hoc supplement validation failed"):
        _build(paths, output_key="posthoc_drift")
    assert not drift_output.exists()


def test_evidence_loader_rejects_missing_final_analysis(tmp_path: Path) -> None:
    root = tmp_path / "repo"
    root.mkdir()
    registry = root / "experiments" / "registry.json"
    registry.parent.mkdir()
    registry.write_text("{}\n", encoding="utf-8")
    with pytest.raises(ReportBuildError, match="missing final_analysis"):
        load_report_evidence(
            repository_root=root,
            analysis_dir=root / "outputs" / "experiments" / "final_analysis",
            registry_path=registry,
        )


def test_report_rejects_live_integrity_audit_drift_after_analysis(tmp_path: Path) -> None:
    paths = _synthetic_report_fixture(tmp_path)
    audit = paths["root"] / "agent" / "audits" / "final_result_integrity.json"
    payload = json.loads(audit.read_text(encoding="utf-8"))
    payload["post_analysis_change"] = True
    _write_json(audit, payload)
    output = paths["root"] / "docs" / "report" / "stale-audit.docx"
    paths["stale_audit"] = output

    with pytest.raises(ReportBuildError, match="differs from the final_analysis anchor"):
        _build(paths, output_key="stale_audit")
    assert not output.exists()
