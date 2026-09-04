from __future__ import annotations

import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from scripts import build_lean_paper_docx as builder


def test_prepare_markdown_removes_placeholder_reference_text() -> None:
    source = "# Title\n\nText [@L01].\n\n## References\n\nPlaceholder.\n"

    prepared = builder._prepare_markdown(source)

    assert "Placeholder" not in prepared
    assert "Text [@L01]." in prepared
    assert "## References" in prepared
    assert '<div id="refs"></div>' in prepared


def test_prepare_markdown_requires_reference_boundary() -> None:
    with pytest.raises(builder.DocxBuildError, match="References boundary"):
        builder._prepare_markdown("# Title\n")


def test_column_widths_sum_to_content_width() -> None:
    class Cell:
        def __init__(self, text: str) -> None:
            self.text = text

    class Row:
        def __init__(self, values: tuple[str, ...]) -> None:
            self.cells = [Cell(value) for value in values]

    class Table:
        rows = (
            Row(("Short", "A much longer descriptive value", "n")),
            Row(("A", "Another descriptive value", "40")),
        )

    widths = builder._column_widths(Table())

    assert len(widths) == 3
    assert sum(widths) == builder.CONTENT_WIDTH_DXA
    assert widths[1] > widths[0]
    assert widths[1] > widths[2]
    assert min(widths) >= 720


def test_generated_english_docx_preserves_scientific_content_and_references() -> None:
    document = Document(builder.OUTPUT_PATH)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert len(document.tables) == 7
    assert len(document.inline_shapes) == 5
    assert len(document.sections) == 1
    section = document.sections[0]
    assert round(section.page_width.inches, 3) == 8.5
    assert round(section.page_height.inches, 3) == 11.0
    assert round(section.top_margin.inches, 3) == 1.0
    assert round(section.right_margin.inches, 3) == 1.0
    assert round(section.bottom_margin.inches, 3) == 1.0
    assert round(section.left_margin.inches, 3) == 1.0
    assert "38/600" in text and "13/600" in text
    assert "17/60" in text and "0/12" in text
    assert "30/240" in text
    assert "not estimable" in text
    assert document.core_properties.title == (
        "Evaluating Authorization Regimes and Runtime Guardrail Configurations "
        "for Tool-Using LLM Agents"
    )
    assert "Bagdasaryan" in text
    assert "Bagdasarian" not in text
    assert "[@" not in text
    assert "�" not in text
    assert sum(
        paragraph.style.name == "Bibliography" for paragraph in document.paragraphs
    ) == 21

    with zipfile.ZipFile(builder.OUTPUT_PATH) as archive:
        custom = archive.read("docProps/custom.xml").decode("utf-8")
    assert str(builder.PROJECT_ROOT) not in custom
    assert "references_v1.bib" in custom


def test_generated_english_docx_keeps_headings_left_and_table_rows_together() -> None:
    document = Document(builder.OUTPUT_PATH)

    title = document.paragraphs[0]
    headings = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name.startswith("Heading")
    ]

    assert title.style.name == "Heading 1"
    assert title.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert headings
    assert all(paragraph.alignment == WD_ALIGN_PARAGRAPH.LEFT for paragraph in headings)
    assert all(
        row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None
        for table in document.tables
        for row in table.rows
    )
    assert all(
        paragraph.paragraph_format.keep_with_next
        for table in document.tables
        for cell in table.rows[0].cells
        for paragraph in cell.paragraphs
    )
