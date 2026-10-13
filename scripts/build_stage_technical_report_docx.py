"""构建 M3-PI Guard 2026-08-06 阶段技术报告 Word 文档（DOCX）。

基于阶段性进展与实验里程碑数据，生成格式规范的阶段性技术报告 Word 文档。
"""

from __future__ import annotations

import json
import math
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs/reports/paper_v2/M3_PI_Guard_阶段技术报告_2026-08-06_artifact.json"
OUTPUT = ROOT / "docs/research/paper_v2/reports/M3_PI_Guard_阶段技术状态报告_2026-08-06.docx"
QA_DIR = ROOT / "outputs/report_qa/stage_technical_report_20260806"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B3A"
MUTED = "5D6B78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE_BLUE = "F4F7FB"
RISK = "9B1C1C"
PALE_RED = "FBEAEA"
WHITE = "#FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def create_numbering_id(doc: Document, *, num_format: str = "decimal", level_text: str = "%1.") -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    level.append(fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), level_text)
    level.append(lvl_text)
    jc = OxmlElement("w:lvlJc")
    jc.set(qn("w:val"), "left")
    level.append(jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_numbered_items(doc: Document, items: list[str]) -> None:
    num_id = create_numbering_id(doc)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.1
        p_pr = p._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_pr.append(ilvl)
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.append(num_id_node)
        p_pr.append(num_pr)
        run = p.add_run(item)
        set_run_font(run, size=10.2, color=INK)


def set_table_geometry(table, widths_in: list[float], indent_dxa: int = 120) -> None:
    total_dxa = sum(round(width * 1440) for width in widths_in)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(round(width * 1440)))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width_dxa = round(widths_in[index] * 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width_dxa))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_in[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None,
                 color: str | None = None, bold: bool | None = None,
                 italic: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    tokens = {
        "Normal": ("Microsoft YaHei", 10.5, INK, 0, 6, 1.1),
        "Title": ("Microsoft YaHei", 24, INK, 0, 6, 1.0),
        "Subtitle": ("Microsoft YaHei", 13, MUTED, 0, 14, 1.0),
        "Heading 1": ("Microsoft YaHei", 16, BLUE, 16, 8, 1.0),
        "Heading 2": ("Microsoft YaHei", 13, BLUE, 12, 6, 1.0),
        "Heading 3": ("Microsoft YaHei", 11.5, DARK_BLUE, 8, 4, 1.0),
        "Caption": ("Microsoft YaHei", 8.5, MUTED, 4, 4, 1.0),
        "List Bullet": ("Microsoft YaHei", 10.5, INK, 0, 6, 1.1),
        "List Number": ("Microsoft YaHei", 10.5, INK, 0, 6, 1.1),
    }
    for style_name, (font_name, size, color, before, after, spacing) in tokens.items():
        style = doc.styles[style_name]
        style.font.name = font_name
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font_name)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font_name)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = spacing
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 2"].font.bold = True
    doc.styles["Heading 3"].font.bold = True
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=8.5, color=MUTED)


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    left = p.add_run("M3-PI GUARD  |  阶段技术状态报告")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    tab = p.add_run("\t2026-08-06")
    set_run_font(tab, size=8.5, color=MUTED)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Inches(6.45))

    footer = section.footer
    p = footer.paragraphs[0]
    add_page_number(p)


def add_paragraph(doc: Document, text: str = "", *, bold_prefix: str | None = None,
                  italic: bool = False, color: str = INK, size: float = 10.5,
                  after: float = 6, keep_with_next: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, size=size, color=color, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, size=size, color=color, italic=italic)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, italic=italic)
    return p


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, size={1: 16, 2: 13, 3: 11.5}[level],
                 color=BLUE if level < 3 else DARK_BLUE, bold=True)


def add_callout(doc: Document, title: str, body: str, fill: str = PALE_BLUE,
                accent: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    set_row_cant_split(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11, color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.3, color=INK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_metric_strip(doc: Document, metrics: list[tuple[str, str, str]]) -> None:
    widths = [6.5 / len(metrics)] * len(metrics)
    table = doc.add_table(rows=1, cols=len(metrics))
    set_table_geometry(table, widths)
    for index, (value, label, note) in enumerate(metrics):
        cell = table.cell(0, index)
        set_cell_shading(cell, LIGHT_BLUE if index % 2 == 0 else LIGHT_GRAY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        value_run = p.add_run(value)
        set_run_font(value_run, size=17, color=DARK_BLUE, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(1)
        label_run = p2.add_run(label)
        set_run_font(label_run, size=8.5, color=INK, bold=True)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(0)
        note_run = p3.add_run(note)
        set_run_font(note_run, size=7.4, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Caption")
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    set_run_font(r, size=8.2, color=MUTED, italic=True)


def set_picture_alt(inline_shape, *, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float],
              font_size: float = 8.7) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=font_size, color=DARK_BLUE, bold=True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        if row_index % 2:
            for cell in cells:
                set_cell_shading(cell, "F8FAFC")
        for col_index, value in enumerate(values):
            p = cells[col_index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if col_index in (0, 2):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            set_run_font(r, size=font_size, color=INK)
        for cell in cells:
            set_cell_margins(cell, top=95, bottom=95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def pil_font(size: int, bold: bool = False):
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_harness_chart(rows: list[dict], path: Path) -> None:
    image = Image.new("RGB", (1500, 760), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(42, bold=True)
    label_font = pil_font(30)
    small_font = pil_font(24)
    value_font = pil_font(31, bold=True)
    draw.text((75, 42), "Fresh Harness 聚合阶段对比", font=title_font, fill="#172B3A")
    draw.text((75, 102), "同一13阶段口径；v5因宿主Git所有权门失败，v6全部通过", font=small_font, fill="#5D6B78")
    groups = ["Fresh Harness v5", "Fresh Harness v6"]
    values = {group: {"通过": 0, "失败": 0} for group in groups}
    for row in rows:
        values[row["run"]][row["stage_status"]] = row["stage_count"]
    left, top, right, bottom = 165, 190, 1400, 650
    draw.line((left, bottom, right, bottom), fill="#A7B2BD", width=3)
    for tick in range(0, 14, 2):
        y = bottom - (tick / 13) * (bottom - top)
        draw.line((left, y, right, y), fill="#E3E8ED", width=2)
        draw.text((100, y - 15), str(tick), font=small_font, fill="#5D6B78")
    bar_width = 170
    for index, group in enumerate(groups):
        center = 500 + index * 570
        x1 = center - bar_width / 2
        pass_height = values[group]["通过"] / 13 * (bottom - top)
        fail_height = values[group]["失败"] / 13 * (bottom - top)
        if pass_height:
            draw.rounded_rectangle((x1, bottom - pass_height, x1 + bar_width, bottom), radius=12, fill="#2E74B5")
            draw.text((center - 18, bottom - pass_height - 42), str(values[group]["通过"]), font=value_font, fill="#1F4D78")
        if fail_height:
            draw.rounded_rectangle((x1, bottom - pass_height - fail_height, x1 + bar_width, bottom - pass_height), radius=12, fill="#D98C8C")
            draw.text((center - 18, bottom - pass_height - fail_height - 42), str(values[group]["失败"]), font=value_font, fill="#9B1C1C")
        text_box = draw.textbbox((0, 0), group, font=label_font)
        draw.text((center - (text_box[2] - text_box[0]) / 2, 675), group, font=label_font, fill="#172B3A")
    draw.rounded_rectangle((1060, 55, 1100, 95), radius=6, fill="#2E74B5")
    draw.text((1115, 56), "通过", font=small_font, fill="#172B3A")
    draw.rounded_rectangle((1225, 55, 1265, 95), radius=6, fill="#D98C8C")
    draw.text((1280, 56), "失败", font=small_font, fill="#172B3A")
    image.save(path, quality=95)


def draw_coverage_chart(rows: list[dict], path: Path) -> None:
    image = Image.new("RGB", (1500, 930), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(42, bold=True)
    label_font = pil_font(27)
    small_font = pil_font(22)
    value_font = pil_font(25, bold=True)
    draw.text((70, 38), "正式工作流数量覆盖率", font=title_font, fill="#172B3A")
    draw.text((70, 98), "不同单位仅作排期归一化；不代表科研效果或质量门通过", font=small_font, fill="#5D6B78")
    left, right = 420, 1400
    bar_width = right - left
    top = 180
    row_height = 112
    for index, row in enumerate(rows):
        y = top + index * row_height
        ratio = max(0.0, min(1.0, float(row["coverage_ratio"])))
        label = row["workstream"]
        draw.text((70, y + 13), label, font=label_font, fill="#172B3A")
        draw.rounded_rectangle((left, y + 8, right, y + 54), radius=16, fill="#E6EAF0")
        if ratio > 0:
            fill_width = max(8, math.floor(bar_width * ratio))
            draw.rounded_rectangle((left, y + 8, left + fill_width, y + 54), radius=min(16, fill_width // 2), fill="#2E74B5")
        completed_target = f"{row['completed']}/{row['target']} {row['unit']}"
        percent = f"{ratio * 100:.1f}%"
        draw.text((left, y + 64), completed_target, font=small_font, fill="#5D6B78")
        percent_box = draw.textbbox((0, 0), percent, font=value_font)
        draw.text((right - (percent_box[2] - percent_box[0]), y + 62), percent, font=value_font, fill="#1F4D78")
        draw.text((760, y + 64), row["evidence_class"], font=small_font, fill="#5D6B78")
    image.save(path, quality=95)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("阶段技术状态报告")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph(style="Title")
    p.paragraph_format.keep_with_next = True
    r = p.add_run("M3-PI Guard")
    set_run_font(r, size=27, color=INK, bold=True)

    p = doc.add_paragraph(style="Subtitle")
    r = p.add_run("Fresh Harness v6 收口、现有实验覆盖与结项差距")
    set_run_font(r, size=14, color=MUTED)

    metadata = [
        ("报告日期", "2026年8月6日"),
        ("证据截止", "2026年8月6日 03:40（Asia/Shanghai）"),
        ("项目状态", "in_progress（工程基线绿色，实验未结项）"),
        ("阅读对象", "项目负责人、研究合作者、后续执行者"),
        ("本轮边界", "只读盘点与报告；无新API调用、无新provider admission"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        label_run = p.add_run(f"{label}：")
        set_run_font(label_run, size=10, color=INK, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_callout(
        doc,
        "结论",
        "当前工程与证据 Harness 已通过完整重放，可以安全恢复后续工作；但 Gemini 正式矩阵、三模型正式 Adaptive 和600条独立盲测尚未完成，因此本项目不能标记为实验结项。",
    )
    add_metric_strip(
        doc,
        [
            ("13/13", "Harness 阶段", "Fresh v6 全部通过"),
            ("1,538", "pytest 通过", "另1跳过、3已知警告"),
            ("1,791/1,791", "Flash 公共记录", "物理闭包；结果仍封存"),
            ("0", "活动正式进程", "报告前只读检查"),
        ],
    )
    add_paragraph(
        doc,
        "本报告只陈述证据闭包、执行覆盖与剩余门槛，不判断论文胜负，不把候选、探索性前缀或绿色 Harness 解释为科研结果。",
        color=MUTED,
        italic=True,
        size=9.3,
    )
    doc.add_page_break()


def build_report(data: dict) -> None:
    QA_DIR.mkdir(parents=True, exist_ok=True)
    harness_chart = QA_DIR / "harness_comparison.png"
    coverage_chart = QA_DIR / "workstream_coverage.png"
    datasets = data["snapshot"]["datasets"]
    draw_harness_chart(datasets["harness_comparison"], harness_chart)
    draw_coverage_chart(datasets["workstream_coverage"], coverage_chart)

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_cover(doc)

    add_heading(doc, "1. 技术摘要：工程基线已收口，实验仍未结项", 1)
    add_paragraph(
        doc,
        "Fresh Harness v6 已完整通过，可以把当前工程基线视为绿色、可恢复、可审计。直接证据包括子进程 exit 0、13/13 聚合阶段通过、pytest 1,538 通过/1 跳过/3 条已知警告，以及唯一终态“Verification failed=0”。报告生成前的只读进程盘点未发现遗留正式实验进程。",
        bold_prefix="Fresh Harness v6 已完整通过",
    )
    add_paragraph(
        doc,
        "这不等于实验结项。DeepSeek Flash 的1,791条公共矩阵实现了物理证据闭包，但科研结果仍未独立准入；Gemini Public-v5仍是非授权候选；独立盲测只完成Day-1保管人登记；正式Adaptive仍只有8任务、预算50的探索性前缀。功能状态必须继续保持 in_progress。",
        bold_prefix="这不等于实验结项。",
    )
    add_callout(
        doc,
        "本轮停止点",
        "没有创建新的 provider admission，没有调用 DeepSeek 或 Gemini，没有运行新的 Adaptive 或盲测，也没有读取密钥或案例级盲测标签。",
        fill=LIGHT_GRAY,
        accent=DARK_BLUE,
    )

    add_heading(doc, "2. v6 将验证基线从部分失败修复为完整绿色", 1)
    add_paragraph(
        doc,
        "v5 的 pytest 本身通过，但13个聚合阶段中仅10个通过；后三个证据阶段因冻结 AgentDojo checkout 的 Git ownership 防护而失败。修复只为精确冻结路径加入进程作用域 safe.directory，没有写入 global 或 repository Git config。",
    )
    harness_picture = doc.add_picture(str(harness_chart), width=Inches(6.35))
    set_picture_alt(
        harness_picture,
        title="Fresh Harness 聚合阶段对比",
        description="柱状图：Fresh Harness v5通过10个阶段、失败3个阶段；v6通过13个阶段、失败0个阶段。",
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_source_note(
        doc,
        "图1  Fresh Harness v5/v6 聚合阶段。来源：v6 receipt、v5 failure review 与冻结验证日志；13阶段为相同口径。",
    )
    add_paragraph(
        doc,
        "v6 达到13/13，并保持 verifier/helper 运行前后哈希稳定、冻结 AgentDojo commit 精确匹配。其含义是当前验证基线能够在正确宿主身份下完整重放；它只修复验证环境和证据链，不增加实验结果，也不授权网络调用。",
        bold_prefix="v6 达到13/13",
    )

    add_heading(doc, "3. 剩余结项缺口集中在 Gemini、正式 Adaptive 和盲测", 1)
    add_paragraph(
        doc,
        "DeepSeek Flash 公共矩阵已覆盖1,791/1,791个唯一工作键，16条 diagnostic 全部保留在ITT，终态重试、provider tools和external effects均为零。但一个分片的冻结质量门仍因10条 diagnostic高于上限6而失败，因此该结果只能作为经验性安全—效用证据，不能被包装成干净机制质量。",
    )
    coverage_picture = doc.add_picture(str(coverage_chart), width=Inches(6.35))
    set_picture_alt(
        coverage_picture,
        title="正式工作流数量覆盖率",
        description="横向进度图：Flash公共100%，Gemini公共约0.1%，Pro预算50为20%，Flash预算50、Gemini预算50和独立盲测均为0%。",
    )
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_source_note(
        doc,
        "图2  各正式工作流数量覆盖。不同工作流单位不同，仅用于排期判断；不用于比较模型安全效果。",
    )
    add_paragraph(
        doc,
        "图中最关键的三处空白是：Gemini正式公共矩阵、三模型正式Adaptive面板，以及600条独立盲测。它们也是当前不能宣布实验结项的直接原因。",
        bold_prefix="图中最关键的三处空白是：",
    )

    status_rows = []
    for row in datasets["workstream_status"]:
        progress = f"{row['completed']}/{row['target']} {row['unit']}"
        status_rows.append([
            str(row["sequence"]),
            row["workstream"],
            progress,
            row["status"],
            row["next_gate"],
        ])
    add_table(
        doc,
        ["序", "工作流", "覆盖", "当前状态", "下一门槛"],
        status_rows,
        [0.38, 1.32, 1.02, 1.87, 1.91],
        font_size=7.8,
    )
    add_source_note(doc, "表1  结项工作流状态。覆盖表示数量闭包，不表示科研结果准入。")

    add_heading(doc, "4. 口径与方法：本报告衡量证据闭包，不衡量论文胜负", 1)
    definitions = [
        ("时间窗", "证据截至2026年8月6日03:40（Asia/Shanghai）。"),
        ("物理终态记录", "工作键已进入不可变记录链且尝试状态闭合；它不自动等于科研结果准入。"),
        ("数量覆盖率", "已完成单位除以计划单位。公共矩阵以记录计，Adaptive以到达预算节点的任务计，盲测以完成封存的案例计。"),
        ("ITT", "超时、解析失败、拒答、provider failure和diagnostic均保留在intention-to-treat分母。"),
        ("绿色 Harness", "说明该次工程、冻结输入与验证脚本通过；不授权网络调用、读取结果或跳过单独provider admission。"),
    ]
    for label, value in definitions:
        add_paragraph(doc, f"{label}：{value}", bold_prefix=f"{label}：")

    add_heading(doc, "证据复核采用四层分离", 2)
    add_numbered_items(doc, [
        "工程完整性：进程退出码、阶段标记、日志哈希、源码运行前后哈希和冻结依赖身份。",
        "执行完整性：工作键唯一性、attempt/record/budget/trace闭包、零终态重试，以及失败留在ITT。",
        "安全边界：provider tools、外部连接和真实effect始终关闭；本轮provider/network authorization=false。",
        "科研准入：工程与执行完整性通过后，仍需独立结果admission和统计复算；candidate或merge不能自我授权。",
    ])

    doc.add_page_break()
    add_heading(doc, "5. 可核验证据链已绑定，但权限边界保持关闭", 1)
    evidence_rows = []
    for row in datasets["evidence_register"]:
        evidence_rows.append([
            str(row["sequence"]),
            row["artifact"],
            row["status"],
            row["boundary"],
        ])
    add_table(
        doc,
        ["序", "证据", "已证明", "未授权/边界"],
        evidence_rows,
        [0.38, 1.72, 2.12, 2.28],
        font_size=8.0,
    )
    add_source_note(doc, "表2  核心证据登记。完整SHA-256见附录A。")
    add_paragraph(
        doc,
        "关键原则是：证据文件证明什么，就只声明什么。v6 receipt能证明验证重放通过；Flash merge能证明物理闭包；Gemini candidate只能证明候选被离线构造；Day-1 acceptance只能证明保管人登记链成立。它们均不能相互替代。",
    )

    add_heading(doc, "6. 限制与不确定性决定了当前不能宣布结项", 1)
    limitations = [
        "DeepSeek Flash已有物理闭包，但结果访问仍封存，且分片质量门失败不能被全局诊断率掩盖。",
        "Gemini Public-v5仍只有候选；免费层日配额使canary和1,791条矩阵不能按普通并发矩阵估算工期。",
        "Adaptive的0/8 unsafe effect、Wilson 95% CI 0–32.44%仅适用于探索性8任务预算50前缀，不能证明40任务、多预算或攻击饱和安全。",
        "独立保管人登记已通过，但blind_dataset_authoring_completed=false；当前没有600条候选、独立标签或评分manifest。",
        "v6运行约4.37小时；后续若改变verifier或正式运行时，需重新过门，不能永久复用一次绿色结论。",
        "本报告不包含论文写作、JISA格式、SOTA或first/only声明。",
    ]
    for item in limitations:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=INK)

    add_callout(
        doc,
        "风险提示",
        "如果后续 Gemini 免费层长期不稳定，或盲测保管人不能交付候选/标签分离包，21天计划需要顺延；不得用诊断性结果伪装为正式完成。",
        fill=PALE_RED,
        accent=RISK,
    )

    add_heading(doc, "7. 恢复后应从非授权 gate 与精确 admission 开始", 1)
    next_steps = [
        "重放v6 receipt、v5 failure review、Day-1 acceptance和相关候选，生成不自行授权的fresh closeout gate。",
        "按America/Los_Angeles配额日检查Gemini项目日账本，只为一个未触碰Public-v5 segment创建单次admission；free tier、USD0、retry0、tools/effects0。",
        "canary通过结果盲技术门后，才设计Gemini 1,791条矩阵的日段扩展；失败即封存，不重试终态工作键。",
        "依冻结顺序完成Pro 50→200→500→hardest-20×2000，再运行Flash 50/200和Gemini 50/200。",
        "保管人完成600条候选/标签分离和custody seal后，开发者只导入无标签执行集，运行3,600条六条件记录并交由保管人评分。",
        "最后完成独立结果准入、Wilson/McNemar/10,000次cluster bootstrap/Holm复算、最终Harness与SHA-256结项包。",
    ]
    add_numbered_items(doc, next_steps)
    add_callout(
        doc,
        "本轮执行决定",
        "按用户指示，本轮在报告完成后停止；不启动上述任何新API实验。",
        fill=LIGHT_GRAY,
        accent=DARK_BLUE,
    )

    add_heading(doc, "8. 恢复前需回答的三个问题", 1)
    questions = [
        "Gemini免费层能否在连续配额日稳定提供可审计的官方response ID与精确模型身份？",
        "保管人何时能提交600条候选的无标签执行集、独立标签文件与签名custody manifest？",
        "DeepSeek Flash分片质量门失败后，最终研究是否接受转为经验性安全—效用与授权不确定性研究，而不是干净机制质量或SOTA主张？",
    ]
    for item in questions:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, size=10.2, color=INK)

    doc.add_page_break()
    add_heading(doc, "附录A：核心证据SHA-256", 1)
    add_paragraph(
        doc,
        "以下摘要绑定报告中的核心状态。SHA-256只标识对应证据对象，不扩大该对象的权限或结论范围。",
        color=MUTED,
        size=9.4,
    )
    for row in datasets["evidence_register"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        label = p.add_run(f"{row['sequence']}. {row['artifact']}")
        set_run_font(label, size=9.2, color=DARK_BLUE, bold=True)
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.18)
        p2.paragraph_format.space_after = Pt(3)
        digest = p2.add_run(row["sha256"])
        set_run_font(digest, name="Consolas", size=7.8, color=INK)
        p3 = doc.add_paragraph()
        p3.paragraph_format.left_indent = Inches(0.18)
        p3.paragraph_format.space_after = Pt(2)
        status = p3.add_run(f"状态：{row['status']}；边界：{row['boundary']}。")
        set_run_font(status, size=8.6, color=MUTED)

    doc.add_page_break()
    add_heading(doc, "附录B：报告来源与生成边界", 1)
    add_paragraph(doc, f"规范化来源：{SOURCE.relative_to(ROOT).as_posix()}", bold_prefix="规范化来源：", keep_with_next=True)
    add_paragraph(doc, "来源文件已在生成前通过既有Data Analytics artifact validator与render validator；其文件SHA-256为 c3b992d145c92d4c25f393fe28a8ecb42272cfe939eaefe774b83b4826ee2197。", keep_with_next=True)
    add_paragraph(doc, "生成边界：只读取已存在的公开状态与证据摘要；未读取config/project.local.env，未访问任何密钥，未创建provider admission，未执行网络请求。")

    doc.core_properties.title = "M3-PI Guard 阶段技术状态报告"
    doc.core_properties.subject = "Fresh Harness v6 收口、实验覆盖与结项差距"
    doc.core_properties.author = "M3-PI Guard Project"
    doc.core_properties.keywords = "M3-PI Guard, Harness, experiment closeout, evidence"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def main() -> int:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    with SOURCE.open("r", encoding="utf-8") as handle:
        data = json.load(handle)
    build_report(data)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
