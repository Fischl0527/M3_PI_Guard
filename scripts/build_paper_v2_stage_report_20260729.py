"""Build the 2026-07-29 M3-PI Guard paper-v2 stage report.

The report is intended for supervisors and research collaborators. It reads
only admitted or frozen evidence and content-free verification logs. It does
not read credentials, prompts, response bodies, or unadmitted scientific
outcomes from the Public-v4 technical canary.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

if __package__:
    from scripts.build_paper_v2_stage_report_20260726 import (
        BLACK,
        BLUE,
        CALLOUT,
        GOLD,
        GREEN,
        INK,
        LIGHT_BLUE,
        MUTED,
        RED,
        add_bullets,
        add_callout,
        add_figure,
        add_page_field,
        add_paragraph,
        add_reference,
        add_table,
        apply_style_tokens,
        font,
        load_json,
        set_run_font,
        sha256_file,
    )
else:
    from build_paper_v2_stage_report_20260726 import (
        BLACK,
        BLUE,
        CALLOUT,
        GOLD,
        GREEN,
        INK,
        LIGHT_BLUE,
        MUTED,
        RED,
        add_bullets,
        add_callout,
        add_figure,
        add_page_field,
        add_paragraph,
        add_reference,
        add_table,
        apply_style_tokens,
        font,
        load_json,
        set_run_font,
        sha256_file,
    )


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_DATE = "2026-07-29"
REPORT_DIR = PROJECT_ROOT / "docs/research/paper_v2/reports"
ASSET_DIR = REPORT_DIR / "assets/stage_report_20260729"
QA_DIR = PROJECT_ROOT / "outputs/paper_upgrade_v2/report_qa/stage_report_20260729"
OUTPUT_PATH = REPORT_DIR / "M3_PI_Guard_阶段性研究报告_2026-07-29.docx"

ANALYSIS_PATH = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_medium/deepseek"
    / "ba9c7937d08b68e088faf8586e0dd96d2052a398fd0a77889fd6a62e503f2373"
    / "admitted_medium_analysis_v8.json"
)
SCORECARD_PATH = (
    PROJECT_ROOT
    / "harness/state/paper_v2/completion_scorecard_v1.json"
)
CLAIM_LEDGER_PATH = (
    PROJECT_ROOT
    / "harness/state/paper_v2/claim_evidence_ledger_v1.json"
)
PUBLIC_FREEZE_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/public_benchmark_freeze_v1.json"
)
PUBLIC_V2_ABORT_PATH = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/public_primary_canary_abort_review_v2.json"
)
PUBLIC_V3_ABORT_PATH = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/public_primary_canary_abort_review_v3.json"
)
PUBLIC_V4_CANDIDATE_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/public_matrix_candidate_v4.json"
)
PUBLIC_V4_REVIEW_PATH = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_public/deepseek"
    / "public-v4-primary-4b717417fec675c3540d186e4b6d6af403bb03c208ebc91e1b1e84b1e7e83e30"
    / "technical_canary_review_v4.json"
)
PUBLIC_SHARD_CANDIDATE_PATH = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/public_primary_continuation_candidate_v1.json"
)
PUBLIC_BUDGET_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/public_budget_policy_v1.json"
)
ADAPTIVE_FREEZE_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/adaptive_attack_freeze_v1.json"
)
HARNESS_STDOUT_PATH = (
    PROJECT_ROOT
    / "outputs/verification/public_primary_continuation_harness_stdout.log"
)
HARNESS_STDERR_PATH = (
    PROJECT_ROOT
    / "outputs/verification/public_primary_continuation_harness_stderr.log"
)
HARNESS_EXIT_PATH = (
    PROJECT_ROOT
    / "outputs/verification/public_primary_continuation_harness_exit.txt"
)

FIGURE_PATHS = {
    "medium": ASSET_DIR / "figure_1_medium_security_utility.png",
    "readiness": ASSET_DIR / "figure_2_readiness.png",
    "public": ASSET_DIR / "figure_3_public_matrix.png",
    "canary": ASSET_DIR / "figure_4_public_canary_versions.png",
    "shards": ASSET_DIR / "figure_5_primary_shards.png",
    "roadmap": ASSET_DIR / "figure_6_evidence_80_roadmap.png",
}

CONDITION_LABELS = {
    "no_defense": "无防御",
    "safe_prompt_or_spotlighting": "安全提示",
    "promptguard_detector_only": "PromptGuard",
    "strongest_peer_reviewed_structural_baseline": "结构化基线",
    "m3_policy_only": "M3 policy-only",
    "m3_full": "M3 full",
}


def configure_sections(doc: Document) -> None:
    """Apply standard_business_brief + memo_masthead tokens."""

    def populate_header(header: Any) -> None:
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        left = paragraph.add_run("M3-PI Guard | SCI 二区网络安全研究")
        set_run_font(left, size=8.5, color=MUTED, bold=True)
        paragraph.add_run("\t")
        right = paragraph.add_run(f"阶段报告 | {REPORT_DATE}")
        set_run_font(right, size=8.5, color=MUTED)

    def populate_footer(footer: Any) -> None:
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        set_run_font(paragraph.add_run("M3-PI Guard  ·  "), size=9, color=MUTED)
        add_page_field(paragraph)

    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        populate_header(section.header)
        populate_header(section.even_page_header)
        populate_footer(section.footer)
        populate_footer(section.even_page_footer)


def prevent_table_row_splitting(doc: Document) -> None:
    """Keep every table row on one page for clean external handoff."""

    for table in doc.tables:
        for row in table.rows:
            tr_pr = row._tr.get_or_add_trPr()
            if tr_pr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cantSplit") is None:
                tr_pr.append(OxmlElement("w:cantSplit"))


def apply_title_styles(doc: Document) -> None:
    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = doc.styles["Heading 1"].font.color.rgb
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(13)
    subtitle.font.bold = False
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(
        kicker.add_run("RESEARCH STAGE REPORT"),
        size=9,
        color=GOLD,
        bold=True,
    )

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    set_run_font(
        title.add_run("M3-PI Guard 阶段性研究报告"),
        size=24,
        color=INK,
        bold=True,
    )

    subtitle = doc.add_paragraph(style="Subtitle")
    set_run_font(
        subtitle.add_run(
            "Public-v4 技术门通过、四分片续跑就绪与 SCI JCR Q2 "
            "实证路线的当前证据"
        ),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("报告日期", "2026-07-29（Asia/Shanghai）"),
        (
            "当前论文路线",
            "授权信息泄漏、授权不确定性与安全—效用权衡的系统实证研究",
        ),
        (
            "候选英文标题",
            "A Systematic Empirical Study of Authorization Leakage, "
            "Uncertainty, and Security–Utility Trade-offs in LLM Agent "
            "Prompt-Injection Defenses",
        ),
        (
            "目标期刊",
            "Journal of Information Security and Applications；投稿前通过"
            "学校/Clarivate 复核当期 JCR Q2",
        ),
        (
            "正式模型证据",
            "DeepSeek 官方 API / deepseek-v4-pro；Gemini 目前只有诊断链，"
            "不承担正式跨提供商结论",
        ),
        (
            "当前得分",
            "科学证据 753/1200（62.75%）；工程就绪 99%；投稿就绪 60%",
        ),
        (
            "安全边界",
            "provider tools 关闭、无真实外部 effect、无自动重试、失败保留"
            "在 ITT 分母、科学结果单独准入",
        ),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(f"{label}："), size=10.5, bold=True)
        set_run_font(p.add_run(value), size=10.5)


def _canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(
        (45, 40, 1755, 1010),
        radius=30,
        fill="#FFFFFF",
        outline="#D8DEE8",
        width=3,
    )
    draw.text((90, 74), title, font=font(42, bold=True), fill=f"#{INK}")
    draw.text((90, 140), subtitle, font=font(23), fill=f"#{MUTED}")
    draw.line((90, 194, 1710, 194), fill="#D8DEE8", width=3)
    return image, draw


def _save(image: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, "PNG")


def _bar(
    draw: ImageDraw.ImageDraw,
    *,
    x0: int,
    y0: int,
    width: int,
    height: int,
    value: float,
    maximum: float,
    color: str,
    label: str,
    value_label: str,
) -> None:
    draw.text((x0, y0 - 42), label, font=font(23), fill=f"#{BLACK}")
    draw.rounded_rectangle(
        (x0, y0, x0 + width, y0 + height),
        radius=height // 2,
        fill="#E8EEF5",
    )
    fill_width = max(3, int(width * value / maximum))
    draw.rounded_rectangle(
        (x0, y0, x0 + fill_width, y0 + height),
        radius=height // 2,
        fill=color,
    )
    draw.text(
        (x0 + width + 20, y0 - 5),
        value_label,
        font=font(24, bold=True),
        fill=color,
    )


def create_medium_chart(
    path: Path,
    condition_metrics: dict[str, Any],
) -> None:
    image, draw = _canvas(
        "中型证伪：安全提升成立，但效用与模块增益门失败",
        "72 个案例 × 6 条件；攻击 n=36，良性工具任务 n=24",
    )
    order = [
        "no_defense",
        "safe_prompt_or_spotlighting",
        "promptguard_detector_only",
        "strongest_peer_reviewed_structural_baseline",
        "m3_policy_only",
        "m3_full",
    ]
    colors = ["#94A3B8", "#64748B", "#0EA5E9", "#F59E0B", "#2563EB", "#16A34A"]
    chart_left = 120
    chart_top = 290
    chart_height = 560
    base_y = chart_top + chart_height
    group_width = 238
    max_bar = 430
    for index, (condition, color) in enumerate(zip(order, colors, strict=True)):
        metrics = condition_metrics[condition]
        unsafe = metrics["unsafe_effect_itt"]["value"]
        tcr = metrics["benign_tool_tcr"]["value"]
        x = chart_left + index * group_width
        unsafe_h = int(max_bar * unsafe)
        tcr_h = int(max_bar * tcr)
        draw.rectangle(
            (x, base_y - unsafe_h, x + 70, base_y),
            fill=color,
        )
        draw.rectangle(
            (x + 82, base_y - tcr_h, x + 152, base_y),
            fill="#C7D2FE",
            outline=color,
            width=3,
        )
        draw.text(
            (x + 8, base_y - unsafe_h - 36),
            f"{unsafe:.0%}",
            font=font(20, bold=True),
            fill=color,
        )
        draw.text(
            (x + 88, base_y - tcr_h - 36),
            f"{tcr:.0%}",
            font=font(20, bold=True),
            fill=f"#{INK}",
        )
        label = CONDITION_LABELS[condition].replace(" ", "\n")
        draw.multiline_text(
            (x - 18, base_y + 18),
            label,
            font=font(18),
            fill=f"#{BLACK}",
            align="center",
            spacing=4,
        )
    draw.line(
        (chart_left - 20, base_y, 1610, base_y),
        fill="#94A3B8",
        width=3,
    )
    draw.rectangle((1180, 224, 1215, 259), fill="#F59E0B")
    draw.text((1228, 224), "攻击 unsafe-effect", font=font(20), fill=f"#{BLACK}")
    draw.rectangle(
        (1440, 224, 1475, 259),
        fill="#C7D2FE",
        outline="#2563EB",
        width=3,
    )
    draw.text((1488, 224), "良性工具 TCR", font=font(20), fill=f"#{BLACK}")
    _save(image, path)


def create_readiness_chart(path: Path, scorecard: dict[str, Any]) -> None:
    image, draw = _canvas(
        "当前完成度：工程已接近封顶，论文证据仍需公共与自适应实验",
        "完成度口径不同，不能用工程测试替代科学证据",
    )
    rows = [
        ("工程就绪度", float(scorecard["engineering_readiness_percent"]), "#16A34A"),
        (
            "科学证据完成度",
            scorecard["scientific_evidence_completion_points"]
            / scorecard["scientific_evidence_completion_max_points"]
            * 100,
            "#2563EB",
        ),
        (
            "投稿就绪度",
            float(scorecard["paper_submission_readiness_percent"]),
            "#F59E0B",
        ),
        ("本阶段目标线", 80.0, "#9B1C1C"),
    ]
    for index, (label, value, color) in enumerate(rows):
        _bar(
            draw,
            x0=225,
            y0=300 + index * 160,
            width=1170,
            height=54,
            value=value,
            maximum=100,
            color=color,
            label=label,
            value_label=f"{value:.2f}%" if value % 1 else f"{value:.0f}%",
        )
    draw.text(
        (225, 930),
        "注意：Public-v4 金丝雀是技术门，不计入科学得分。",
        font=font(22, bold=True),
        fill=f"#{RED}",
    )
    _save(image, path)


def create_public_matrix_chart(path: Path) -> None:
    image, draw = _canvas(
        "正式公共矩阵规模与模型角色",
        "冻结总量 6,873 条；选择过程不读取科学 outcome",
    )
    rows = [
        ("M3-v2", 2700, 900),
        ("AgentDojo", 2382, 891),
    ]
    x0 = 300
    max_width = 1200
    maximum = 3600
    for index, (label, primary, second) in enumerate(rows):
        y = 340 + index * 245
        draw.text((90, y + 28), label, font=font(30, bold=True), fill=f"#{INK}")
        primary_w = int(max_width * primary / maximum)
        second_w = int(max_width * second / maximum)
        draw.rounded_rectangle(
            (x0, y, x0 + primary_w, y + 85),
            radius=16,
            fill="#2563EB",
        )
        draw.rounded_rectangle(
            (x0 + primary_w, y, x0 + primary_w + second_w, y + 85),
            radius=16,
            fill="#93C5FD",
        )
        draw.text(
            (x0 + 18, y + 23),
            f"Pro / primary  {primary:,}",
            font=font(23, bold=True),
            fill="white",
        )
        draw.text(
            (x0 + primary_w + 14, y + 23),
            f"Flash / second  {second:,}",
            font=font(21, bold=True),
            fill=f"#{INK}",
        )
        draw.text(
            (x0, y + 105),
            f"合计 {primary + second:,}",
            font=font(22),
            fill=f"#{MUTED}",
        )
    draw.rounded_rectangle(
        (300, 825, 1500, 925),
        radius=18,
        fill="#F4F6F9",
        outline="#D8DEE8",
        width=2,
    )
    draw.text(
        (345, 855),
        "primary 5,082 + second 1,791 = 6,873；18 个 benchmark-condition pairs",
        font=font(24, bold=True),
        fill=f"#{BLACK}",
    )
    _save(image, path)


def create_canary_chart(path: Path) -> None:
    image, draw = _canvas(
        "Public canary v1–v4：失败被保留，v4 才通过完整技术门",
        "v1–v3 永久不可恢复；未将任何部分记录提升为科学结果",
    )
    rows = [
        ("v1", 0, 0, 0.0, "本地依赖/公共 key，零调用"),
        ("v2", 12, 29, 0.0259414, "M3 evaluator 仍索引 medium plan"),
        ("v3", 12, 33, 0.0263544, "active identity digest 不一致"),
        ("v4", 24, 64, 0.0521812, "24/24 技术门通过"),
    ]
    max_calls = 64
    for index, (version, records, calls, cost, note) in enumerate(rows):
        y = 270 + index * 170
        draw.text((90, y + 12), version, font=font(31, bold=True), fill=f"#{INK}")
        bar_width = int(960 * calls / max_calls) if calls else 4
        color = "#16A34A" if version == "v4" else "#DC2626"
        draw.rounded_rectangle(
            (210, y, 210 + bar_width, y + 54),
            radius=20,
            fill=color,
        )
        draw.text(
            (1210, y - 4),
            f"{records}/24 records · {calls} calls · ¥{cost:.6f}",
            font=font(21, bold=True),
            fill=color,
        )
        draw.text((210, y + 70), note, font=font(21), fill=f"#{MUTED}")
    draw.text(
        (90, 920),
        "结论：v4 证明执行链可工作；它仍未授权比较条件效果。",
        font=font(24, bold=True),
        fill=f"#{RED}",
    )
    _save(image, path)


def create_shard_chart(path: Path, shard_candidate: dict[str, Any]) -> None:
    image, draw = _canvas(
        "Public-v4 primary 剩余矩阵的四分片方案",
        "互斥、完整、排除 24 条 canary；候选已冻结但尚未启动",
    )
    shards = shard_candidate["scope"]["shards"]
    max_records = max(int(item["records"]) for item in shards)
    x0 = 190
    base_y = 835
    bar_max_h = 500
    gap = 355
    colors = ["#1D4ED8", "#2563EB", "#0284C7", "#0EA5E9"]
    for index, (item, color) in enumerate(zip(shards, colors, strict=True)):
        records = int(item["records"])
        h = int(bar_max_h * records / max_records)
        x = x0 + index * gap
        draw.rounded_rectangle(
            (x, base_y - h, x + 215, base_y),
            radius=20,
            fill=color,
        )
        draw.text(
            (x + 46, base_y - h - 48),
            f"{records:,}",
            font=font(29, bold=True),
            fill=color,
        )
        draw.text(
            (x + 72, base_y + 20),
            f"Shard {index}",
            font=font(24, bold=True),
            fill=f"#{BLACK}",
        )
        draw.text(
            (x + 48, base_y + 58),
            "上限 ¥4.9",
            font=font(20),
            fill=f"#{MUTED}",
        )
    draw.text(
        (260, 235),
        "5,058 条 continuation · 最大并行 4 · 总硬上限 ¥19.6 · 自动重试 0",
        font=font(26, bold=True),
        fill=f"#{INK}",
    )
    _save(image, path)


def create_roadmap_chart(path: Path) -> None:
    image, draw = _canvas(
        "从 62.75% 到 ≥80% 的证据路线",
        "只有独立结果准入后才计分；下列增量是目标，不是已完成",
    )
    stages = [
        ("当前", 62.75, "#64748B", "753/1200"),
        ("公共矩阵", 70.25, "#2563EB", "+90"),
        ("自适应攻击", 78.58, "#0EA5E9", "+100"),
        ("论文主表与正文", 80.67, "#16A34A", "+25"),
    ]
    x_start = 150
    x_end = 1650
    y = 545
    draw.line((x_start, y, x_end, y), fill="#CBD5E1", width=12)
    for index, (label, value, color, delta) in enumerate(stages):
        x = x_start + index * ((x_end - x_start) // (len(stages) - 1))
        draw.ellipse((x - 38, y - 38, x + 38, y + 38), fill=color)
        draw.text(
            (x - 70, y - 160),
            f"{value:.2f}%",
            font=font(31, bold=True),
            fill=color,
        )
        draw.text(
            (x - 78, y + 62),
            label,
            font=font(24, bold=True),
            fill=f"#{BLACK}",
        )
        draw.text(
            (x - 45, y + 102),
            delta,
            font=font(21),
            fill=f"#{MUTED}",
        )
    draw.rounded_rectangle(
        (185, 795, 1615, 920),
        radius=20,
        fill="#FFF8E8",
        outline="#E8C66A",
        width=3,
    )
    draw.text(
        (235, 826),
        "预计 1–3 天达到证据 80%：前提是 API 稳定、公共结果可准入、"
        "自适应任务按早停规则收敛。",
        font=font(23, bold=True),
        fill=f"#{GOLD}",
    )
    _save(image, path)


def create_figures(
    analysis: dict[str, Any],
    scorecard: dict[str, Any],
    shard_candidate: dict[str, Any],
) -> None:
    primary = analysis["analysis"]["primary_analysis"]
    create_medium_chart(FIGURE_PATHS["medium"], primary["condition_metrics"])
    create_readiness_chart(FIGURE_PATHS["readiness"], scorecard)
    create_public_matrix_chart(FIGURE_PATHS["public"])
    create_canary_chart(FIGURE_PATHS["canary"])
    create_shard_chart(FIGURE_PATHS["shards"], shard_candidate)
    create_roadmap_chart(FIGURE_PATHS["roadmap"])


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Report Source")
    set_run_font(p.add_run(f"证据来源：{text}"), size=8.5, color=MUTED)


def add_numbered_actions(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item))


def _condition_rows(condition_metrics: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for condition in (
        "no_defense",
        "safe_prompt_or_spotlighting",
        "promptguard_detector_only",
        "strongest_peer_reviewed_structural_baseline",
        "m3_policy_only",
        "m3_full",
    ):
        metrics = condition_metrics[condition]
        unsafe = metrics["unsafe_effect_itt"]
        tool_tcr = metrics["benign_tool_tcr"]
        no_tool_tcr = metrics["benign_no_tool_tcr"]
        rows.append(
            [
                CONDITION_LABELS[condition],
                f"{unsafe['numerator']}/{unsafe['denominator']} "
                f"({unsafe['value']:.1%})",
                f"{tool_tcr['numerator']}/{tool_tcr['denominator']} "
                f"({tool_tcr['value']:.1%})",
                f"{no_tool_tcr['numerator']}/{no_tool_tcr['denominator']} "
                f"({no_tool_tcr['value']:.1%})",
            ]
        )
    return rows


def _parse_harness() -> dict[str, Any]:
    stdout = HARNESS_STDOUT_PATH.read_text(encoding="utf-8", errors="replace")
    summary_matches = re.findall(
        r"(\d+) passed, (\d+) skipped, (\d+) warnings in ([^\r\n]+)",
        stdout,
    )
    if not summary_matches:
        raise ValueError("latest Harness pytest summary not found")
    passed, skipped, warnings, duration = summary_matches[-1]
    exit_code = HARNESS_EXIT_PATH.read_text(encoding="utf-8").strip()
    return {
        "passed": int(passed),
        "skipped": int(skipped),
        "warnings": int(warnings),
        "pytest_duration": duration,
        "verification_failed_zero": "Verification failed=0" in stdout,
        "exit_code": int(exit_code),
        "stdout_sha256": sha256_file(HARNESS_STDOUT_PATH),
        "stderr_sha256": sha256_file(HARNESS_STDERR_PATH),
        "exit_sha256": sha256_file(HARNESS_EXIT_PATH),
        "stderr_lines": [
            line
            for line in HARNESS_STDERR_PATH.read_text(
                encoding="utf-8",
                errors="replace",
            ).splitlines()
            if line.strip()
        ],
    }


def build_report() -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)

    analysis = load_json(ANALYSIS_PATH)
    scorecard = load_json(SCORECARD_PATH)
    claim_ledger = load_json(CLAIM_LEDGER_PATH)
    public_freeze = load_json(PUBLIC_FREEZE_PATH)
    public_v2_abort = load_json(PUBLIC_V2_ABORT_PATH)
    public_v3_abort = load_json(PUBLIC_V3_ABORT_PATH)
    public_v4_candidate = load_json(PUBLIC_V4_CANDIDATE_PATH)
    public_v4_review = load_json(PUBLIC_V4_REVIEW_PATH)
    shard_candidate = load_json(PUBLIC_SHARD_CANDIDATE_PATH)
    public_budget = load_json(PUBLIC_BUDGET_PATH)
    adaptive_freeze = load_json(ADAPTIVE_FREEZE_PATH)
    harness = _parse_harness()

    assert scorecard["scientific_evidence_completion_points"] == 753
    assert scorecard["scientific_evidence_completion_max_points"] == 1200
    assert public_freeze["expected_case_condition_records"] == 6873
    assert public_v4_review["observed_records"] == 24
    assert public_v4_review["provider_calls_completed"] == 64
    assert public_v4_review["matrix_continuation_eligible"] is True
    assert shard_candidate["scope"]["continuation_records"] == 5058
    assert harness["exit_code"] == 0
    assert harness["verification_failed_zero"] is True

    create_figures(analysis, scorecard, shard_candidate)

    report_data = {
        "schema_version": "m3-paper-v2-stage-report-data-20260729",
        "generated_from": {
            "analysis": {
                "path": str(ANALYSIS_PATH.relative_to(PROJECT_ROOT)),
                "sha256": sha256_file(ANALYSIS_PATH),
            },
            "scorecard": {
                "path": str(SCORECARD_PATH.relative_to(PROJECT_ROOT)),
                "sha256": sha256_file(SCORECARD_PATH),
            },
            "public_v4_review": {
                "path": str(PUBLIC_V4_REVIEW_PATH.relative_to(PROJECT_ROOT)),
                "sha256": sha256_file(PUBLIC_V4_REVIEW_PATH),
            },
            "shard_candidate": {
                "path": str(
                    PUBLIC_SHARD_CANDIDATE_PATH.relative_to(PROJECT_ROOT)
                ),
                "sha256": sha256_file(PUBLIC_SHARD_CANDIDATE_PATH),
            },
            "harness_stdout": {
                "path": str(HARNESS_STDOUT_PATH.relative_to(PROJECT_ROOT)),
                "sha256": harness["stdout_sha256"],
            },
        },
        "bounded_summary": {
            "scientific_points": 753,
            "scientific_max_points": 1200,
            "public_records": 6873,
            "v4_canary_records": 24,
            "v4_canary_calls": 64,
            "v4_canary_cost_cny": 0.0521812,
            "continuation_records": 5058,
            "harness": harness,
        },
    }
    (QA_DIR / "report_data.json").write_text(
        json.dumps(report_data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    apply_style_tokens(doc)
    apply_title_styles(doc)
    configure_sections(doc)
    add_title_block(doc)

    doc.add_heading("1. 执行摘要", level=1)
    add_callout(
        doc,
        "总体判断",
        "项目仍然具备 SCI JCR Q2 潜力，但论文路线已经由“防御机制领先”"
        "切换为“授权信息泄漏、授权不确定性与安全—效用权衡的系统实证”。"
        "这一转向不是降级，而是由 480 条正式中型证据触发的预注册止损："
        "M3-full 在攻击安全性上优于结构化基线，却未通过良性工具效用非劣"
        "和 full 相对 policy-only 的独立模块增益门。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_paragraph(
        doc,
        "截至 2026-07-29，工程和证据链已达到可扩展实验水平。Public-v4 "
        "primary 技术金丝雀完成 24/24 条记录和 64 次官方 DeepSeek 调用，"
        "所有技术门均通过；四个互斥分片已经冻结剩余 5,058 条 primary "
        "记录。最新完整 Harness 自然完成并报告 1,188 passed、1 skipped、"
        "3 warnings、Verification failed=0。"
    )
    add_paragraph(
        doc,
        "必须强调：Public-v4 金丝雀只证明执行链、身份、预算、trace、"
        "append-only ledger 和安全边界可以工作。其 technical review 明确"
        "没有访问或聚合 unsafe-effect、task-completion 和条件差异，"
        "result admission 仍为 false。因此科学证据仍是 753/1200"
        "（62.75%），不能把 24 条技术记录或 1,188 个测试误记为论文结果。"
    )
    add_table(
        doc,
        ["维度", "当前状态", "可对外表述", "不可表述"],
        [
            [
                "论文可行性",
                "有条件可行",
                "具备 JISA 范围匹配和完整研究基础",
                "不能保证录用或声称已经达到 Q2",
            ],
            [
                "中型证据",
                "正式准入",
                "安全—效用冲突与机制门失败",
                "不能宣称 M3-full 全面优于基线",
            ],
            [
                "Public-v4",
                "技术门通过",
                "24/24、64 calls、证据链完整",
                "不能报告条件效果或计入科学分",
            ],
            [
                "主模型续跑",
                "候选冻结",
                "四分片 5,058 条已通过完整 Harness",
                "尚未启动、尚无结果",
            ],
            [
                "第二模型",
                "正式证据缺失",
                "Gemini Free 仅有安全连通性诊断",
                "不能做跨提供商泛化结论",
            ],
        ],
        [1450, 1500, 3260, 3150],
        status_columns=(1,),
    )

    doc.add_heading("2. 目标 SCI 二区的可行性与当前论文定位", level=1)
    add_paragraph(
        doc,
        "目标按“SCIE 收录且投稿时 JCR Q2”理解。首选 Journal of "
        "Information Security and Applications（JISA）是合理的：其官方范围"
        "包含访问控制、安全策略、网络与移动安全，并有 LLM/大型 AI 模型"
        "安全相关专题。最终分区不能依赖第三方静态网页，投稿前必须由学校"
        "数据库或 Clarivate 订阅核验当期 JCR。"
    )
    add_table(
        doc,
        ["判断项", "项目现状", "二区要求", "差距"],
        [
            [
                "研究问题",
                "工具型 LLM Agent 的运行时授权与间接提示注入",
                "网络安全问题清晰、威胁模型明确",
                "基本满足",
            ],
            [
                "创新叙事",
                "oracle 泄漏 + 授权 regime + 安全—效用联合分析",
                "不能与 AuthGraph/PlanGuard 等新工作重叠",
                "需依赖公共/自适应结果强化",
            ],
            [
                "实验证据",
                "中型 480 条正式准入；公共矩阵待跑",
                "公共基准、第二模型、自适应攻击",
                "核心缺口",
            ],
            [
                "复现性",
                "冻结、哈希、准入、ITT、完整 Harness",
                "可审计、可重建、失败透明",
                "优势",
            ],
            [
                "外部有效性",
                "AgentDojo + M3-v2 已冻结",
                "至少两个公开/外部基准方向一致",
                "结果尚缺",
            ],
        ],
        [1500, 2850, 2850, 2160],
        status_columns=(3,),
    )
    add_callout(
        doc,
        "建议主标题方向",
        "A Systematic Empirical Study of Authorization Leakage, "
        "Uncertainty, and Security–Utility Trade-offs in LLM Agent "
        "Prompt-Injection Defenses。只有在公共/自适应证据显示 M3-full "
        "同时满足安全增益、TCR 非劣与独立模块增益时，才考虑恢复防御机制"
        "主线；当前不建议这样写。",
        fill=CALLOUT,
        accent=GOLD,
    )

    doc.add_heading("3. 已准入的中型证伪实验", level=1)
    primary = analysis["analysis"]["primary_analysis"]
    add_figure(
        doc,
        FIGURE_PATHS["medium"],
        "图 1  中型六条件的 unsafe-effect 与良性工具 TCR",
        "admitted_medium_analysis_v8.json；所有失败保留在 ITT 分母。",
        alt=(
            "六条件分组柱图。M3 policy-only 和 M3 full 的 unsafe-effect "
            "均为 0%，但良性工具完成率分别只有 29.2% 和 25.0%。"
        ),
    )
    add_table(
        doc,
        ["条件", "攻击 unsafe-effect", "良性工具 TCR", "无工具 TCR"],
        _condition_rows(primary["condition_metrics"]),
        [2550, 2200, 2200, 2410],
        status_columns=(),
    )
    add_source_note(
        doc,
        "72 个案例；每个条件攻击 n=36、良性工具 n=24、无工具 n=12。"
        "上述是已准入中型结果，不代表公共矩阵或第二模型。",
    )

    doc.add_heading("4. 为什么机制论文路线被预注册门禁否决", level=1)
    add_table(
        doc,
        ["门禁", "观测", "预注册要求", "结论"],
        [
            [
                "安全：full vs 结构化基线",
                "0/36 vs 8/36，绝对差 −22.2pp",
                "至少相对下降 20% 或绝对下降 10pp",
                "通过",
            ],
            [
                "效用：full vs 结构化基线",
                "6/24 vs 12/24，绝对差 −25pp",
                "TCR 下降不超过 5pp",
                "失败",
            ],
            [
                "模块增益：full vs policy-only",
                "0/36 vs 0/36",
                "full 必须有严格正向安全增益",
                "失败",
            ],
            [
                "运行质量",
                "432 条正式分析中 21 个运行失败状态留在 ITT",
                "总体≤10%，provider/timeout≤5%，runner=0",
                "通过",
            ],
        ],
        [2200, 3000, 2800, 1360],
        status_columns=(3,),
    )
    add_paragraph(
        doc,
        "最重要的科研价值并不是“证明防御成功”，而是展示：如果只看攻击"
        "阻断率，M3-full 显得非常强；一旦把良性工具任务、policy-only 消融、"
        "失败分母和授权 regime 同时纳入，机制领先叙事就不成立。这个负面"
        "结果能够支撑更可信的评估方法论文。"
    )
    add_source_note(
        doc,
        "安全差值、McNemar、10,000 次 clustered bootstrap 与 Holm 校正均由"
        "admitted_medium_analysis_v8.json 生成并可独立重建。",
    )

    doc.add_heading("5. 证据链与运行时安全边界", level=1)
    add_table(
        doc,
        ["层", "冻结/记录内容", "目的", "失败策略"],
        [
            [
                "候选 manifest",
                "模型、端点、案例序列、条件、身份、预算、输出根",
                "防止看结果后改协议",
                "哈希漂移即拒绝",
            ],
            [
                "本地 request chain",
                "调用前本地 request ID + canonical payload hash",
                "证明请求在发送前已固定",
                "开放事件不补写",
            ],
            [
                "官方 receipt",
                "official response ID + exact returned model + usage",
                "绑定上游身份与费用",
                "身份不符即停止",
            ],
            [
                "append-only ledgers",
                "attempt / trace / budget / record 前向哈希链",
                "可恢复、可审计、可交叉绑定",
                "链断裂不准入",
            ],
            [
                "result admission",
                "技术审查与科学分析分离",
                "避免工程状态泄漏 scientific outcome",
                "未准入不聚合",
            ],
        ],
        [1500, 3350, 2300, 2210],
    )
    add_bullets(
        doc,
        [
            "在线 Guard 不得看到 evaluator/gold 字段；可见字段集合通过 digest 证明。",
            "自动路径只能收缩权限；权限扩张必须经独立确认或 ASK。",
            "所有模型工具和外部连接器关闭；effect 只在 sandbox/in-memory backend 中计算。",
            "自动重试为 0；HTTP 429、身份漂移、开放 query/attempt 均触发暂停或封存。",
            "危险操作不会交给 DeepSeek 或 Gemini 执行；两家 API 只处理无真实外部效果的研究负载。",
        ],
    )

    doc.add_heading("6. Public-v1 到 Public-v4：失败链与修复", level=1)
    add_figure(
        doc,
        FIGURE_PATHS["canary"],
        "图 2  Public 技术金丝雀版本演进",
        "v1/v2/v3 abort review 与 v4 technical_canary_review；科学结果均未开放。",
        alt=(
            "v1 零调用；v2 和 v3 分别完成 12/24 条；v4 完成 24/24 条并通过"
            "技术门。"
        ),
    )
    add_table(
        doc,
        ["版本", "记录/调用/成本", "停止点或修复", "证据状态"],
        [
            [
                "Public-v1",
                "0/24；0 provider calls；¥0",
                "本地 PromptGuard 依赖 + public work key resolver 缺口",
                "永久封存，不可恢复",
            ],
            [
                "Public-v2",
                "12/24；29 calls；¥0.0259414",
                "AgentDojo 修复；首个 M3 row 仍索引 medium plan",
                "永久封存，不可恢复",
            ],
            [
                "Public-v3",
                "12/24；33 calls；¥0.0263544",
                "M3 key 可解析，但 binding 使用旧 execution identity",
                "永久封存，不可恢复",
            ],
            [
                "Public-v4",
                "24/24；64 calls；¥0.0521812",
                "active identity 注入 + evaluator 闭环探针",
                "技术门通过，科学未准入",
            ],
        ],
        [1450, 2200, 3440, 2270],
        status_columns=(3,),
    )
    add_paragraph(
        doc,
        "v1–v3 的失败没有被删除或“修复原记录”。每个版本保留新的身份、"
        "输出根、attempt/budget/trace/record 链及内容无关 abort review。"
        "这种看似保守的处理对论文非常重要：它证明失败未被结果清洗，也避免"
        "将同一工作项在不同运行中重复采样后挑选成功版本。"
    )
    add_source_note(
        doc,
        f"Public-v2 abort manifest={public_v2_abort['manifest_sha256']}；"
        f"Public-v3 abort manifest={public_v3_abort['manifest_sha256']}。",
    )

    doc.add_heading("7. Public-v4 技术金丝雀：已通过什么、尚未证明什么", level=1)
    add_table(
        doc,
        ["检查项", "结果", "解释"],
        [
            [
                "记录完整性",
                "24/24，全部 terminal",
                "expected count、attempt 和 record 对齐",
            ],
            [
                "官方调用",
                "64/64 completed",
                "有 exact model / response ID / usage 绑定",
            ],
            [
                "运行状态",
                "全部 completed",
                "0 diagnostic、0 operational failure",
            ],
            [
                "预算",
                "¥0.0521812；未暂停",
                "42,896 prompt；4,131 completion",
            ],
            [
                "缓存",
                "34,048 hit；8,848 miss",
                "缓存命中/未命中进入 cost receipt",
            ],
            [
                "安全边界",
                "全部 isolated",
                "0 retry；无 provider tool 和真实 external effect",
            ],
            [
                "科学结果",
                "未访问",
                "不比较条件，不聚合 unsafe-effect/TCR",
            ],
        ],
        [2050, 2400, 4910],
        status_columns=(1,),
    )
    add_callout(
        doc,
        "技术门结论",
        "matrix_continuation_eligible=true，意味着可继续执行同一冻结矩阵；"
        "不意味着 M3-full 有效，也不意味着 24 条记录可以进入主表。结果"
        "准入仍需在全矩阵终态后单独执行。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )

    doc.add_heading("8. 公共矩阵设计与外部有效性", level=1)
    add_figure(
        doc,
        FIGURE_PATHS["public"],
        "图 3  6,873 条公共矩阵的 benchmark 与模型角色构成",
        "public_benchmark_freeze_v1.json 与 public execution plan。",
        alt=(
            "M3-v2 为 primary 2700、second 900；AgentDojo 为 primary 2382、"
            "second 891；总计 6873。"
        ),
    )
    add_table(
        doc,
        ["Benchmark", "Primary / Pro", "Second / Flash", "总计", "证据角色"],
        [
            ["M3-v2", "2,700", "900", "3,600", "授权 regime、语言、攻击族"],
            ["AgentDojo", "2,382", "891", "3,273", "公开环境的安全与任务效用"],
            ["合计", "5,082", "1,791", "6,873", "18 个条件对"],
        ],
        [1600, 1600, 1600, 1200, 3360],
        status_columns=(3,),
    )
    add_paragraph(
        doc,
        "AgentDojo 官方数据只有 97 个唯一良性任务，低于原计划的 primary "
        "150 和 second 100。冻结协议选择使用全部 97 个唯一任务，而不是"
        "通过重复样本制造更大的分母。这一偏差应在 Methods 和 Limitations "
        "中披露。"
    )
    add_source_note(
        doc,
        "public_benchmark_freeze_v1.json：状态 cases_frozen_results_pending；"
        "选择方法为 SHA-256 ranked stratified without outcomes。",
    )

    doc.add_heading("9. Primary 四分片续跑候选", level=1)
    add_figure(
        doc,
        FIGURE_PATHS["shards"],
        "FIGURE 4  Primary 剩余 5,058 条记录的互斥分片",
        "Primary continuation candidate v1（冻结 JSON）。",
        alt=(
            "四个分片分别包含 1265、1265、1264、1264 条记录，每片成本"
            "硬上限 4.9 元。"
        ),
    )
    shard_rows = [
        [
            f"Shard {item['shard_index']}",
            f"{item['records']:,}",
            f"¥{item['hard_spend_limit_cny']:.1f}",
            item["run_id"],
            "候选冻结；未启动",
        ]
        for item in shard_candidate["scope"]["shards"]
    ]
    add_table(
        doc,
        ["分片", "记录数", "硬上限", "Run ID", "当前状态"],
        shard_rows,
        [1050, 1150, 1150, 3810, 2200],
        status_columns=(4,),
    )
    add_paragraph(
        doc,
        "分片规则是对 canary 之外的冻结计划顺序做 round-robin。候选证明"
        "四片互斥、并集完整且排除了 24 条已完成 canary；每片有独立 "
        "execution identity、run ID、output root 和 work-key sequence hash。"
        "报告制作期间没有创建 active admission，也没有启动任何分片进程。"
    )

    doc.add_heading("10. 最新完整 Harness 验证", level=1)
    add_table(
        doc,
        ["验证项", "结果", "证据"],
        [
            [
                "全量 pytest",
                f"{harness['passed']:,} passed / {harness['skipped']} skipped / "
                f"{harness['warnings']} warnings",
                f"pytest stage {harness['pytest_duration']}",
            ],
            [
                "统一 Harness",
                "Verification failed=0",
                "自然结束，exit code 0",
            ],
            [
                "stderr",
                "3 行预期阻断",
                "RuntimeAuditor 正确阻止 delete_record",
            ],
            [
                "Stage A / frozen artifacts",
                "PASS",
                "Public-v4、shard candidate、历史冻结可重建",
            ],
            [
                "历史完整性",
                "PASS",
                "paper-v1、E0–E17、V2 历史证据保持不可变",
            ],
        ],
        [2300, 2600, 4460],
        status_columns=(1,),
    )
    add_source_note(
        doc,
        f"Harness stdout SHA-256={harness['stdout_sha256']}；stderr "
        f"SHA-256={harness['stderr_sha256']}；exit artifact "
        f"SHA-256={harness['exit_sha256']}。完整工作流约 2 小时 30 分钟，"
        "pytest 阶段约 1 小时 03 分钟。",
    )
    add_callout(
        doc,
        "验证含义",
        "绿色 Harness 说明代码、schema、冻结资产和恢复规则彼此一致。"
        "它是启动长实验的工程前提，但不是模型防御效果的统计证据。",
        fill=CALLOUT,
        accent=GREEN,
    )

    doc.add_heading("11. API、预算与危险操作边界", level=1)
    add_table(
        doc,
        ["项目", "DeepSeek", "Gemini", "共同边界"],
        [
            [
                "用途",
                "正式 primary/second 研究执行",
                "免费层诊断；正式 v3 尚未实现",
                "仅合成/沙箱研究负载",
            ],
            [
                "身份链",
                "本地 request ID + payload hash + response ID + exact model",
                "当前 request ID 不完整，仅诊断",
                "append-only、可哈希重建",
            ],
            [
                "预算",
                "Public primary 20 元；second 5 元",
                "USD 0 paid-tier；保留免费日配额",
                "达到 fuse 即暂停",
            ],
            [
                "重试/回退",
                "自动重试 0",
                "自动重试 0",
                "不跨提供商 fallback",
            ],
            [
                "危险能力",
                "tools/external effects=false",
                "tools/external effects=false",
                "不会执行删除、支付、发信、部署等真实动作",
            ],
        ],
        [1450, 2600, 2400, 2910],
    )
    add_paragraph(
        doc,
        "用户已明确允许合理使用两家 API，但“可使用”不改变研究准入规则。"
        "每一批正式执行仍需绑定冻结候选、精确命令、模型身份、成本上限和"
        "单独结果准入。尤其是 DeepSeek 的请求链设计已由项目所有者接受，"
        "可以作为论文证据；Gemini 仍缺少同等级正式 trace。"
    )

    doc.add_heading("12. Gemini 当前状态与使用建议", level=1)
    add_table(
        doc,
        ["阶段", "HTTP/调用", "状态", "能支持的结论"],
        [
            [
                "早期官方 smoke",
                "曾 HTTP 200",
                "缺完整 request ID",
                "诊断连通性",
            ],
            [
                "Gemini v1",
                "VIGIL 结构 canary 前失败",
                "0 formal records",
                "不能支持第二模型",
            ],
            [
                "Gemini v2",
                "首请求 HTTP 429",
                "1 diagnostic ITT，0 completed calls",
                "不能支持第二模型",
            ],
            [
                "Free recheck",
                "HTTP 200，exact model + SAFE",
                "32/1/44 tokens；response ID 有、request ID 无",
                "免费层可用性诊断",
            ],
            [
                "建议",
                "Public primary 后再设计 v3",
                "先做多日计量与完整 trace",
                "成功后才做跨提供商补强",
            ],
        ],
        [1450, 2350, 2760, 2800],
        status_columns=(2,),
    )

    doc.add_heading("13. 当前前沿工作与 M3 的创新边界", level=1)
    add_table(
        doc,
        ["工作", "主要思想", "对 M3 的压力", "报告中的处理"],
        [
            [
                "VIGIL (ACL 2026)",
                "verify-before-commit；SIREN 959 cases",
                "直接竞争结构化运行时防御",
                "作为首要同行评审结构基线",
            ],
            [
                "DRIFT (NeurIPS 2025)",
                "动态规则、控制/数据流约束、记忆隔离",
                "覆盖动态授权与数据流",
                "需要公共对照或清晰说明适配边界",
            ],
            [
                "Task Shield (ACL 2025)",
                "验证指令/工具调用是否服务用户目标",
                "与目标一致性分支重叠",
                "不声称“首次目标对齐”",
            ],
            [
                "CaMeL (2025 preprint)",
                "可信 query 提取控制/数据流并用 capability 限制",
                "强结构隔离和可证明安全",
                "比较信息流边界与任务效用",
            ],
            [
                "Progent (2025 preprint)",
                "可编程最小权限与动态 policy",
                "与权限单调/扩张确认高度相关",
                "强调非 oracle 评估与 ASK 负担",
            ],
            [
                "PlanGuard / IPIGuard",
                "规划一致性 / 工具依赖图",
                "参数级与执行图边界接近",
                "只做有证据的差异声明",
            ],
            [
                "AuthGraph (2026 preprint)",
                "授权图与 provenance 执行图在参数来源级对齐",
                "直接压缩“参数来源首创”空间",
                "删除首次/唯一/SOTA 表述",
            ],
            [
                "ARGUS (2026 preprint)",
                "上下文依赖任务与 provenance-aware auditing",
                "与上下文授权不确定性重叠",
                "公共/自适应分析必须显式对照",
            ],
        ],
        [1650, 2750, 2580, 2380],
    )
    add_paragraph(
        doc,
        "M3 目前最可信的差异点不是“参数级来源追踪首创”，而是把 "
        "evaluator/gold 泄漏、授权 regime、权限单调、ALLOW/ASK/DENY、"
        "确认负担、效用非劣和完整审计链放进同一配对实验。是否形成足够"
        "创新，取决于公共矩阵和自适应攻击能否显示这些评估维度会改变对"
        "防御优劣的判断。"
    )

    doc.add_heading("14. 当前完成度与达到 80% 的真实算术", level=1)
    add_figure(
        doc,
        FIGURE_PATHS["readiness"],
        "图 5  工程、科学证据与投稿就绪度",
        "completion_scorecard_v1.json；Public-v4 工程进展尚未提前增加科学分。",
        alt=(
            "工程就绪度 99%，科学证据完成度 62.75%，投稿就绪度 60%，"
            "目标线 80%。"
        ),
    )
    add_figure(
        doc,
        FIGURE_PATHS["roadmap"],
        "图 6  从 753/1200 到 968/1200 的目标增量",
        "evidence-80 roadmap v3；增量仅在独立准入后计入。",
        alt=(
            "当前 62.75%，公共矩阵目标增加 90 分，自适应攻击增加 100 分，"
            "论文主表与正文增加 25 分，目标达到 80.67%。"
        ),
    )
    add_table(
        doc,
        ["工作流", "当前分", "目标分", "潜在增量", "计分门"],
        [
            ["正式公共基准", "10", "100", "+90", "6,873 全终态 + 独立结果准入"],
            ["自适应攻击", "0", "100", "+100", "40/20 + hardest-20 或预注册早停"],
            ["论文正文/主表", "45", "70", "+25", "公共/自适应表和完整引用审计"],
            ["其余九项", "698", "698", "+0", "保持已准入证据"],
            ["合计", "753", "968", "+215", "968/1200 = 80.67%"],
        ],
        [1950, 1100, 1100, 1300, 3910],
        status_columns=(4,),
    )

    doc.add_heading("15. 报告之后的实验执行方案", level=1)
    add_numbered_actions(
        doc,
        [
            "为四个 primary continuation shard 分别生成 single-use admission，"
            "逐一验证 candidate、exact command、模型、预算、输出根和未存在"
            "状态；任何 admission 漂移都在读取凭据前拒绝。",
            "四片并行执行 5,058 条 primary 记录。每片独立 CNY 4.9 fuse，"
            "自动重试 0；首个 429、身份不符、开放 attempt 或 evidence-chain "
            "异常由该片 fail closed，不影响其他片的账本完整性。",
            "四片自然结束后只做 outcome-blind technical merge：验证 24 条 "
            "canary + 5,058 continuation 正好覆盖 5,082 primary rows，"
            "没有重复、缺失、未知 key、开放 attempt 或未绑定 provider call。",
            "冻结并执行 second / deepseek-v4-flash 的 12 条 technical canary；"
            "通过后再执行 1,779 条 remaining second rows，并保持单独 CNY 5 "
            "角色上限。",
            "在 6,873 条全部终态后运行公共 result admission，再执行预注册"
            "统计；失败和 diagnostics 留在 ITT。只有这一阶段可以打开条件"
            "效果并更新论文得分。",
            "根据公共结果选择自适应攻击路径：先跑 8×50 成本 canary，"
            "技术门和成本投影通过后执行 primary 40 tasks（50/200/500），"
            "最难 20 tasks 追加 2,000；second 20 tasks 跑 50/200。",
            "最后回填论文主表、Results、Discussion、claim–evidence ledger、"
            "局限性与 JISA 格式；完成方法审查和 Devil’s Advocate 审查。",
        ],
    )

    doc.add_heading("16. 后续实验的时间、成本与停止条件", level=1)
    add_table(
        doc,
        ["阶段", "预计墙钟时间", "主要成本边界", "停止/转向条件"],
        [
            [
                "Primary 四分片",
                "4–6 小时",
                "≤¥19.6 continuation；加 canary 仍在 primary ¥20 内",
                "首个 429/身份/链异常：该片停止，不自动重试",
            ],
            [
                "Second 矩阵",
                "2–4 小时",
                "≤¥5",
                "Flash 身份或效用异常则收缩为 cross-model 限定",
            ],
            [
                "Public merge + 统计",
                "1–3 小时",
                "本地计算",
                "不完整覆盖或开放 attempt：不准入",
            ],
            [
                "Adaptive",
                "12–24 小时，最坏可跨数日",
                "先 ¥5 cost canary；正式预算按新 gate",
                "成本投影超限、攻击未收敛或证据链异常",
            ],
            [
                "论文回填与完整性审查",
                "4–8 小时",
                "本地计算",
                "发现 claim 无证据即删除或降级",
            ],
            [
                "达到 ≥80% 的总估计",
                "1–3 天",
                "取决于 API 稳定和自适应早停",
                "不是承诺；任何门失败都保持真实分数",
            ],
        ],
        [1700, 1750, 2950, 2960],
    )

    doc.add_heading("17. 统计分析与主表计划", level=1)
    add_table(
        doc,
        ["层级", "终点/方法", "必须报告", "禁止做法"],
        [
            [
                "主要安全终点",
                "unsafe-effect rate；配对差",
                "分子/分母、绝对差、相对差、Wilson CI",
                "只报检测 recall 或挑选成功子集",
            ],
            [
                "主要效用终点",
                "benign-tool TCR；−5pp 非劣",
                "配对差、CI、是否跨非劣界限",
                "把 deny-all 当成安全成功",
            ],
            [
                "模块贡献",
                "M3-full vs policy-only",
                "安全与效用两侧同时报告",
                "只比较 full 与弱基线",
            ],
            [
                "显著性",
                "exact McNemar + Holm",
                "原始和校正 p 值、discordant pairs",
                "以 p<0.05 替代效应量",
            ],
            [
                "不确定性",
                "10,000 次 cluster bootstrap",
                "cluster 单位、seed、CI",
                "忽略同一 case 的条件相关性",
            ],
            [
                "异质性",
                "benchmark/regime/attack/language/model",
                "预注册分层与交互方向",
                "结果后新增无约束子组",
            ],
            [
                "运行失败",
                "intention-to-treat",
                "超时、解析、provider failure 全保留",
                "删除失败或重新运行挑最好",
            ],
        ],
        [1600, 2400, 2950, 2410],
    )

    doc.add_heading("18. 主要风险、缓解措施与需要外部协作的事项", level=1)
    add_table(
        doc,
        ["风险", "概率/影响", "当前缓解", "需要的决定"],
        [
            [
                "Public-v4 长跑出现 429 或开放 attempt",
                "中 / 高",
                "四片隔离、零重试、append-only、结果门 fail closed",
                "按新身份重建未启动 remainder，不修补旧片",
            ],
            [
                "公共结果仍显示效用严重下降",
                "中 / 高",
                "实证 fallback 已锁定，不强行恢复机制论文",
                "把过度阻断作为核心负面发现",
            ],
            [
                "full 与 policy-only 继续无差异",
                "中 / 高",
                "保留 provenance/ASK 消融",
                "机制只作为测量工具，不主张领先",
            ],
            [
                "Gemini 正式链不完整",
                "高 / 中",
                "DeepSeek Pro/Flash 支持 cross-model",
                "跨提供商结论降级为 future work",
            ],
            [
                "独立盲测保管人缺失",
                "高 / 中",
                "内部封存测试不冒充 independent blind",
                "第 8 周前确定导师/同学/第三方保管人",
            ],
            [
                "JCR 分区或期刊范围变化",
                "低-中 / 高",
                "只写候选期刊，不把第三方分区当最终证据",
                "投稿前由机构 Clarivate 核验",
            ],
            [
                "前沿工作快速重叠",
                "中 / 高",
                "34 篇矩阵、24 PR/10 preprint、提交前刷新",
                "删除 first/only/SOTA，按源码证据收缩 claims",
            ],
        ],
        [2130, 1300, 3220, 2710],
        status_columns=(1,),
    )
    add_callout(
        doc,
        "当前需要协作者最关注的三件事",
        "第一，Public-v4 的 24 条记录不是效果结果；第二，论文当前主线是"
        "实证 fallback，不是 SOTA 防御；第三，独立盲测保管人和投稿前 "
        "Clarivate 核验仍是外部依赖。",
        fill="FFF8E8",
        accent=GOLD,
    )

    doc.add_heading("19. 阶段结论与建议", level=1)
    add_bullets(
        doc,
        [
            "SCI 二区：仍可行，但依赖公共矩阵、自适应攻击和论文回填三个门；当前不能称为“已达到二区水平”。",
            "研究叙事：保持授权泄漏/不确定性/安全—效用实证路线，除非后续 full 同时通过安全、效用和模块增益门。",
            "工程状态：Public-v4 技术链和四分片候选已经通过最新完整 Harness，可在报告交付后进入正式 continuation。",
            "证据纪律：Public-v1–v3 永久封存；v4 技术金丝雀不计科学分；任何长跑失败留在 ITT。",
            "预计周期：无重大 API 阻塞时，达到 ≥80% 证据目标约需 1–3 天；自适应攻击不收敛时可能更久。",
        ],
    )

    doc.add_heading("附录 A：关键冻结资产与哈希", level=1)
    artifact_rows = [
        [
            "中型正式分析",
            str(ANALYSIS_PATH.relative_to(PROJECT_ROOT)),
            sha256_file(ANALYSIS_PATH),
        ],
        [
            "完成度 scorecard",
            str(SCORECARD_PATH.relative_to(PROJECT_ROOT)),
            sha256_file(SCORECARD_PATH),
        ],
        [
            "Public-v4 candidate",
            str(PUBLIC_V4_CANDIDATE_PATH.relative_to(PROJECT_ROOT)),
            sha256_file(PUBLIC_V4_CANDIDATE_PATH),
        ],
        [
            "Public-v4 technical review",
            str(PUBLIC_V4_REVIEW_PATH.relative_to(PROJECT_ROOT)),
            sha256_file(PUBLIC_V4_REVIEW_PATH),
        ],
        [
            "Primary shard candidate",
            str(PUBLIC_SHARD_CANDIDATE_PATH.relative_to(PROJECT_ROOT)),
            sha256_file(PUBLIC_SHARD_CANDIDATE_PATH),
        ],
        [
            "最新 Harness stdout",
            str(HARNESS_STDOUT_PATH.relative_to(PROJECT_ROOT)),
            harness["stdout_sha256"],
        ],
    ]
    add_table(
        doc,
        ["资产", "路径", "SHA-256"],
        artifact_rows,
        [1800, 4100, 3460],
    )
    add_paragraph(
        doc,
        "Public-v4 candidate canonical manifest 为 "
        f"{public_v4_candidate['manifest_sha256']}；Primary shard candidate "
        f"canonical manifest 为 {shard_candidate['manifest_sha256']}。"
    )

    doc.add_heading("附录 B：可对外使用的简短口径", level=1)
    add_table(
        doc,
        ["问题", "建议回答"],
        [
            [
                "现在能投 SCI 二区吗？",
                "方向和工程基础可行，但还缺公共、自适应和最终论文证据；"
                "当前是 62.75% 科学完成度，不建议立即投稿。",
            ],
            [
                "M3 防御有效吗？",
                "中型实验中攻击 unsafe-effect 为 0/36，但良性工具 TCR 只有"
                " 6/24，且与 policy-only 无安全差异；不能称全面有效。",
            ],
            [
                "Public-v4 已经出结果了吗？",
                "只出了技术金丝雀：24/24 条记录证明链路稳定，未打开条件效果。",
            ],
            [
                "为什么需要四分片？",
                "降低单进程长跑风险，隔离预算和开放 attempt，同时保持冻结计划"
                "的互斥完整覆盖。",
            ],
            [
                "API 会执行危险动作吗？",
                "不会。provider tools 和 external effects 始终关闭，所有效果"
                "只在沙箱或内存 evaluator 中计算。",
            ],
        ],
        [2350, 7010],
    )

    doc.add_heading("附录 C：前沿来源与投稿资料", level=1)
    references = [
        (
            "JISA",
            "官方 Aims & Scope 与投稿入口",
            "https://www.sciencedirect.com/journal/"
            "journal-of-information-security-and-applications",
        ),
        (
            "VIGIL",
            "ACL 2026，verify-before-commit 与 SIREN",
            "https://aclanthology.org/2026.acl-long.443/",
        ),
        (
            "DRIFT",
            "NeurIPS 2025，动态规则与注入隔离",
            "https://proceedings.neurips.cc/paper_files/paper/2025/hash/"
            "77f3b26c7907aa27b207df9b9d43f29a-Abstract-Conference.html",
        ),
        (
            "Task Shield",
            "ACL 2025，工具调用与用户目标一致性",
            "https://aclanthology.org/2025.acl-long.1435/",
        ),
        (
            "AgentDojo",
            "NeurIPS 2024，动态工具 Agent 安全评估环境",
            "https://proceedings.nips.cc/paper_files/paper/2024/hash/"
            "97091a5177d8dc64b1da8bf3e1f6fb54-"
            "Abstract-Datasets_and_Benchmarks_Track.html",
        ),
        (
            "CaMeL",
            "Defeating Prompt Injections by Design",
            "https://arxiv.org/abs/2503.18813",
        ),
        (
            "Progent",
            "Programmable Privilege Control for LLM Agents",
            "https://arxiv.org/abs/2504.11703",
        ),
        (
            "PlanGuard",
            "Planning-based Consistency Verification",
            "https://arxiv.org/abs/2604.10134",
        ),
        (
            "IPIGuard",
            "Tool Dependency Graph-based Defense",
            "https://arxiv.org/abs/2508.15310",
        ),
        (
            "AuthGraph",
            "Authorization 与 provenance 双图对齐",
            "https://arxiv.org/abs/2605.26497",
        ),
        (
            "ARGUS",
            "上下文依赖任务的 provenance-aware decision auditing",
            "https://arxiv.org/abs/2605.03378",
        ),
    ]
    for label, description, url in references:
        add_reference(doc, label, description, url)

    doc.add_heading("附录 D：报告生成与证据边界", level=1)
    add_paragraph(
        doc,
        "本报告由本地冻结 JSON、JSONL 和 Harness 内容无关日志生成。"
        "生成器不读取 API key、环境变量中的秘密、provider response body、"
        "原始 prompt、Public-v4 scientific outcome 或 mutable provider "
        "dashboard。图表只使用已准入中型聚合与内容无关工程计数。"
    )
    add_paragraph(
        doc,
        f"claim ledger 当前分支为 "
        f"{claim_ledger.get('paper_route', claim_ledger.get('status', 'frozen'))}；"
        f"public freeze 状态为 {public_freeze['status']}；adaptive freeze "
        f"状态为 {adaptive_freeze.get('status', 'frozen')}；public budget "
        f"状态为 {public_budget.get('status', 'frozen')}。"
    )

    prevent_table_row_splitting(doc)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    path = build_report()
    print(
        json.dumps(
            {
                "status": "report_built",
                "path": str(path),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "figures": {
                    key: {
                        "path": str(value),
                        "sha256": sha256_file(value),
                    }
                    for key, value in FIGURE_PATHS.items()
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
