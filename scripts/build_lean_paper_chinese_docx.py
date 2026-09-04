"""Build the Simplified Chinese comparison draft of the lean JISA manuscript."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    from scripts import build_lean_paper_docx as english_builder
except ImportError:  # pragma: no cover - supports direct script execution
    import build_lean_paper_docx as english_builder


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = PROJECT_ROOT / "docs/research/paper_v2/manuscript"
MARKDOWN_PATH = MANUSCRIPT_DIR / "manuscript_zh_v1.md"
BIBLIOGRAPHY_PATH = MANUSCRIPT_DIR / "references_v1.bib"
REFERENCE_DOCX_PATH = MANUSCRIPT_DIR / "M3_PI_Guard_JISA_working_draft.docx"
OUTPUT_PATH = MANUSCRIPT_DIR / "M3_PI_Guard_JISA_chinese_comparison.docx"

LATIN_FONT = "Times New Roman"
BODY_CJK_FONT = "SimSun"
HEADING_CJK_FONT = "Microsoft YaHei"


class ChineseDocxBuildError(RuntimeError):
    """Raised when the Chinese comparison draft cannot be built safely."""


def _prepare_markdown(text: str) -> str:
    marker = "\n## 参考文献\n"
    if marker not in text:
        raise ChineseDocxBuildError("Chinese manuscript lacks the reference boundary")
    body = text.split(marker, 1)[0].rstrip()
    return body + "\n\n## 参考文献\n\n<div id=\"refs\"></div>\n"


def _set_run_fonts(run, *, cjk_font: str) -> None:
    run.font.name = LATIN_FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), cjk_font)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _set_style_fonts(style, *, cjk_font: str) -> None:
    style.font.name = LATIN_FONT
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), LATIN_FONT)
    r_fonts.set(qn("w:hAnsi"), LATIN_FONT)
    r_fonts.set(qn("w:eastAsia"), cjk_font)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")


def _apply_chinese_fonts(path: Path) -> None:
    document = Document(path)
    for style in document.styles:
        if style.type != 1:
            continue
        cjk_font = (
            HEADING_CJK_FONT
            if style.name == "Title" or style.name.startswith("Heading")
            else BODY_CJK_FONT
        )
        _set_style_fonts(style, cjk_font=cjk_font)

    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "工具型 LLM 智能体的运行时授权"
        for run in header.runs:
            _set_run_fonts(run, cjk_font=HEADING_CJK_FONT)

    for paragraph in document.paragraphs:
        cjk_font = (
            HEADING_CJK_FONT
            if paragraph.style.name == "Title"
            or paragraph.style.name.startswith("Heading")
            else BODY_CJK_FONT
        )
        for run in paragraph.runs:
            _set_run_fonts(run, cjk_font=cjk_font)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        _set_run_fonts(run, cjk_font=BODY_CJK_FONT)

    document.core_properties.title = (
        "工具型大语言模型智能体的授权制度与运行时 Guard 配置评估"
    )
    document.core_properties.subject = "JISA 中文内部对照稿"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.save(path)


def build_docx() -> Path:
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise ChineseDocxBuildError("pandoc is not available")
    for required in (MARKDOWN_PATH, BIBLIOGRAPHY_PATH, REFERENCE_DOCX_PATH):
        if not required.is_file():
            raise ChineseDocxBuildError(f"Required input is missing: {required}")

    markdown = _prepare_markdown(MARKDOWN_PATH.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="m3-paper-zh-docx-") as temp_name:
        input_path = Path(temp_name) / "manuscript_zh.md"
        input_path.write_text(markdown, encoding="utf-8", newline="\n")
        command = [
            pandoc,
            str(input_path),
            "--from=markdown+citations",
            "--to=docx",
            "--citeproc",
            "--bibliography=references_v1.bib",
            f"--reference-doc={REFERENCE_DOCX_PATH}",
            "--metadata=reference-section-title:参考文献",
            f"--resource-path={MANUSCRIPT_DIR}",
            f"--output={OUTPUT_PATH}",
        ]
        completed = subprocess.run(
            command,
            cwd=MANUSCRIPT_DIR,
            capture_output=True,
            check=False,
            text=True,
        )
        if completed.returncode != 0:
            raise ChineseDocxBuildError(
                "pandoc failed: " + (completed.stderr.strip() or "unknown error")
            )

    english_builder._apply_academic_style(OUTPUT_PATH)
    _apply_chinese_fonts(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    output = build_docx()
    document = Document(output)
    print(
        f"wrote={output.relative_to(PROJECT_ROOT)} "
        f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} "
        f"figures={len(document.inline_shapes)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
