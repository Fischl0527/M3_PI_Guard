"""Build a single-column Word review draft of the lean JISA manuscript."""

from __future__ import annotations

import math
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = PROJECT_ROOT / "docs/research/paper_v2/manuscript"
MARKDOWN_PATH = MANUSCRIPT_DIR / "manuscript_deepseek_v1.md"
BIBLIOGRAPHY_PATH = MANUSCRIPT_DIR / "references_v1.bib"
OUTPUT_PATH = MANUSCRIPT_DIR / "M3_PI_Guard_JISA_working_draft.docx"

PAGE_WIDTH_DXA = 12_240
PAGE_HEIGHT_DXA = 15_840
MARGIN_DXA = 1_440
CONTENT_WIDTH_DXA = 9_360
TABLE_INDENT_DXA = 120
CELL_MARGIN_VERTICAL_DXA = 80
CELL_MARGIN_HORIZONTAL_DXA = 120

FONT_NAME = "Times New Roman"
BODY_SIZE_PT = 11
BODY_AFTER_PT = 6
BODY_LINE_SPACING = 1.15
HEADING_SPECS = {
    "Title": (16, 0, 12),
    "Heading 1": (15, 14, 6),
    "Heading 2": (12, 10, 4),
    "Heading 3": (11, 8, 3),
    "Heading 4": (10, 6, 2),
}


class DocxBuildError(RuntimeError):
    """Raised when an editable review draft cannot be built safely."""


def _prepare_markdown(text: str) -> str:
    marker = "\n## References\n"
    if marker not in text:
        raise DocxBuildError("Reader manuscript lacks the References boundary")
    body = text.split(marker, 1)[0].rstrip()
    return body + "\n\n## References\n\n<div id=\"refs\"></div>\n"


def _font_style(style, *, size: int, bold: bool = False) -> None:
    style.font.name = FONT_NAME
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", CELL_MARGIN_VERTICAL_DXA),
        ("bottom", CELL_MARGIN_VERTICAL_DXA),
        ("start", CELL_MARGIN_HORIZONTAL_DXA),
        ("end", CELL_MARGIN_HORIZONTAL_DXA),
    ):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_width(element, tag: str, width: int) -> None:
    child = element.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        element.append(child)
    child.set(qn("w:w"), str(width))
    child.set(qn("w:type"), "dxa")


def _column_widths(table) -> list[int]:
    if not table.rows:
        raise DocxBuildError("Encountered a table without rows")
    column_count = len(table.rows[0].cells)
    if column_count == 0:
        raise DocxBuildError("Encountered a table without columns")
    minimum = 600 if column_count >= 7 else 720
    if minimum * column_count >= CONTENT_WIDTH_DXA:
        minimum = CONTENT_WIDTH_DXA // column_count

    weights: list[float] = []
    for column_index in range(column_count):
        longest = max(
            len(cell.text.strip()) for cell in (row.cells[column_index] for row in table.rows)
        )
        weights.append(math.sqrt(max(6, min(longest, 80))))

    remaining = CONTENT_WIDTH_DXA - minimum * column_count
    weight_total = sum(weights)
    widths = [minimum + int(remaining * weight / weight_total) for weight in weights]
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B7C3D0")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _prevent_table_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _format_table(table) -> None:
    widths = _column_widths(table)
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    _set_table_borders(table)
    _set_width(tbl_pr, "w:tblW", CONTENT_WIDTH_DXA)
    table_indent = tbl_pr.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        tbl_pr.append(table_indent)
    table_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    table_indent.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    table_font_size = 7.5 if len(widths) >= 7 else 8.5
    for row_index, row in enumerate(table.rows):
        _prevent_table_row_split(row)
        if row_index == 0:
            _repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            _set_width(tc_pr, "w:tcW", widths[column_index])
            if row_index == 0:
                _shade_cell(cell, "E8EEF5")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                if column_index > 1 and len(cell.text) < 28:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = FONT_NAME
                    run.font.size = Pt(table_font_size)
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
                    if row_index == 0:
                        run.bold = True


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, value, end):
        run._r.append(element)
    run.font.name = FONT_NAME
    run.font.size = Pt(9)


def _apply_academic_style(path: Path) -> None:
    document = Document(path)
    for section in document.sections:
        section.start_type = WD_SECTION.CONTINUOUS
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header.paragraphs[0]
        header.text = "Runtime Authorization for Tool-Using LLM Agents"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(90, 90, 90)
        footer = section.footer.paragraphs[0]
        footer.clear()
        _add_page_number(footer)

    normal = document.styles["Normal"]
    _font_style(normal, size=BODY_SIZE_PT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_AFTER_PT)
    normal.paragraph_format.line_spacing = BODY_LINE_SPACING
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name in ("Body Text", "First Paragraph"):
        if style_name in document.styles:
            style = document.styles[style_name]
            _font_style(style, size=BODY_SIZE_PT)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(BODY_AFTER_PT)
            style.paragraph_format.line_spacing = BODY_LINE_SPACING
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, (size, before, after) in HEADING_SPECS.items():
        if style_name not in document.styles:
            continue
        style = document.styles[style_name]
        _font_style(style, size=size, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if style_name == "Title":
            style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if paragraph.style.name == "Title":
            paragraph.paragraph_format.page_break_before = False
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for table in document.tables:
        _format_table(table)

    maximum_width = Inches(6.3)
    for shape in document.inline_shapes:
        if shape.width > maximum_width:
            ratio = maximum_width / shape.width
            shape.width = maximum_width
            shape.height = int(shape.height * ratio)

    document.core_properties.title = (
        "Evaluating Runtime Authorization Under Context and Provenance Uncertainty "
        "for Tool-Using LLM Agents"
    )
    document.core_properties.subject = "JISA working manuscript"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.save(path)


def build_docx() -> Path:
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise DocxBuildError("pandoc is not available")
    if not MARKDOWN_PATH.is_file() or not BIBLIOGRAPHY_PATH.is_file():
        raise DocxBuildError("Manuscript or bibliography is missing")

    markdown = _prepare_markdown(MARKDOWN_PATH.read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="m3-paper-docx-") as temp_name:
        input_path = Path(temp_name) / "manuscript.md"
        input_path.write_text(markdown, encoding="utf-8", newline="\n")
        command = [
            pandoc,
            str(input_path),
            "--from=markdown+citations",
            "--to=docx",
            "--citeproc",
            f"--bibliography={BIBLIOGRAPHY_PATH}",
            "--metadata=reference-section-title:References",
            f"--resource-path={MANUSCRIPT_DIR}",
            f"--output={OUTPUT_PATH}",
        ]
        completed = subprocess.run(
            command,
            cwd=MANUSCRIPT_DIR,
            capture_output=True,
            check=False,
            text=True,
        )
        if completed.returncode != 0:
            raise DocxBuildError(
                "pandoc failed: " + (completed.stderr.strip() or "unknown error")
            )
    _apply_academic_style(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    output = build_docx()
    document = Document(output)
    print(
        f"wrote={output.relative_to(PROJECT_ROOT)} "
        f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} "
        f"figures={len(document.inline_shapes)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
