"""Build the separate editable supplementary-table package for the lean paper."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document

from scripts.build_lean_paper_docx import _apply_academic_style
from scripts.build_lean_paper_manuscript import (
    MANUSCRIPT_DIR,
    PROJECT_ROOT,
    TITLE,
    _drop_wrapper,
    _read,
    _split_supplementary_tables,
)


MARKDOWN_PATH = MANUSCRIPT_DIR / "supplementary_tables_v1.md"
DOCX_PATH = MANUSCRIPT_DIR / "M3_PI_Guard_JISA_supplementary_tables.docx"


class SupplementBuildError(RuntimeError):
    """Raised when the supplementary package cannot be built safely."""


def build_supplementary_markdown() -> str:
    results = _drop_wrapper(
        _read("results_shell.md"),
        start_heading="## Medium falsification study",
    )
    _, tables = _split_supplementary_tables(results)
    parts = [
        "# Supplementary material",
        f"**Article:** {TITLE}",
        (
            "These six tables provide detailed results referenced by the main "
            "manuscript. They introduce no additional hypothesis family or "
            "claim beyond the analyses reported in the article."
        ),
    ]
    for number, title, table in tables:
        parts.append(f"## Supplementary Table S{number}. {title}\n\n{table}")
    return "\n\n".join(parts).strip() + "\n"


def build_supplementary_docx(markdown: str) -> Path:
    pandoc = shutil.which("pandoc")
    if pandoc is None:
        raise SupplementBuildError("pandoc is not available")

    with tempfile.TemporaryDirectory(prefix="m3-paper-supplement-") as temp_name:
        input_path = Path(temp_name) / "supplement.md"
        input_path.write_text(markdown, encoding="utf-8", newline="\n")
        completed = subprocess.run(
            [
                pandoc,
                str(input_path),
                "--from=markdown",
                "--to=docx",
                f"--output={DOCX_PATH}",
            ],
            cwd=MANUSCRIPT_DIR,
            capture_output=True,
            check=False,
            text=True,
        )
    if completed.returncode != 0:
        raise SupplementBuildError(
            "pandoc failed: " + (completed.stderr.strip() or "unknown error")
        )

    _apply_academic_style(DOCX_PATH)
    document = Document(DOCX_PATH)
    document.core_properties.title = f"{TITLE}: Supplementary material"
    document.core_properties.subject = "JISA supplementary tables"
    document.save(DOCX_PATH)
    return DOCX_PATH


def main() -> int:
    markdown = build_supplementary_markdown()
    MARKDOWN_PATH.write_text(markdown, encoding="utf-8", newline="\n")
    output = build_supplementary_docx(markdown)
    document = Document(output)
    print(
        f"wrote={MARKDOWN_PATH.relative_to(PROJECT_ROOT)} "
        f"and={output.relative_to(PROJECT_ROOT)} "
        f"paragraphs={len(document.paragraphs)} tables={len(document.tables)}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
