"""构建 2026-07-26 阶段性 paper-v2 进展报告（Word 文档）。

汇总 paper-v2 阶段截至 2026-07-26 的实验设计、指标数据与阶段性产物。
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any, Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = PROJECT_ROOT / "docs/research/paper_v2/reports"
QA_DIR = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/report_qa/stage_report_20260726"
)
OUTPUT_PATH = REPORT_DIR / "M3_PI_Guard_阶段性研究报告_2026-07-26.docx"

V7_ABORT_REVIEW = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_medium/deepseek"
    / "a34fe5b87607cb48144b0d36da282accc25d33cd18706351eb39dd25ebe04c75"
    / "formal_matrix_abort_review_v3.json"
)
V8_CANDIDATE = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/deepseek_medium_matrix_candidate_v8.json"
)
V8_CANARY_ADMISSION = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/admissions"
    / "deepseek_medium_matrix_canary_admission_v8.json"
)
V8_CONTINUATION_ADMISSION = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/admissions"
    / "deepseek_medium_matrix_continuation_admission_v8.json"
)
RESULT_GATE_QUALIFICATION = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/formal_result_gate_qualification_v5.json"
)
RESULT_GATE_CANDIDATE = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/deepseek_medium_result_gate_candidate_v5.json"
)
V8_CANARY_SUMMARY = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/verification/v8_deepseek_canary"
    / "canary_stdout.log"
)
FULL_HARNESS_LOG = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/verification/pre_v8_canary_full"
    / "harness_stdout.log"
)

FONT_REGULAR = Path("C:/Windows/Fonts/msyh.ttc")
FONT_BOLD = Path("C:/Windows/Fonts/msyhbd.ttc")

INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "596579"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
GOLD = "7A5A00"
RED = "9B1C1C"
GREEN = "1F6B4F"
WHITE = "FFFFFF"
BLACK = "111827"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 100, "bottom": 100, "start": 120, "end": 120}


def load_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError(f"{path}: expected object")
    return value


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_last_json_object(path: Path) -> dict[str, Any]:
    text = path.read_text(encoding="utf-8")
    start = text.find("{")
    if start < 0:
        raise ValueError(f"{path}: missing JSON object")
    value = json.loads(text[start:])
    if not isinstance(value, dict):
        raise TypeError(f"{path}: expected object")
    return value


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(
    run: Any,
    *,
    latin: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_border(
    cell: Any,
    *,
    color: str = "D8DEE8",
    size: int = 6,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_repeat_table_header(row: Any) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_table_geometry(
    table: Any,
    widths_dxa: Sequence[int],
    *,
    indent_dxa: int = TABLE_INDENT_DXA,
) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(cell_mar)
    for name, amount in CELL_MARGINS_DXA.items():
        node = cell_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            cell_mar.append(node)
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)


def set_image_alt(paragraph: Any, *, title: str, description: str) -> None:
    drawings = paragraph._p.xpath(".//w:drawing")
    if not drawings:
        return
    doc_pr_nodes = drawings[-1].xpath(".//wp:docPr")
    if not doc_pr_nodes:
        return
    doc_pr = doc_pr_nodes[0]
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def apply_style_tokens(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.0

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9.5)
    caption.font.color.rgb = rgb(DARK_BLUE)
    caption.font.bold = True
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    if "Report Source" not in [style.name for style in styles]:
        source = styles.add_style("Report Source", WD_STYLE_TYPE.PARAGRAPH)
    else:
        source = styles["Report Source"]
    source.font.name = "Calibri"
    source.font.size = Pt(8.5)
    source.font.color.rgb = rgb(MUTED)
    source._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(4)
    source.paragraph_format.line_spacing = 1.0

    if "Table Text" not in [style.name for style in styles]:
        table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    else:
        table_text = styles["Table Text"]
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(9)
    table_text.font.color.rgb = rgb(BLACK)
    table_text._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    table_text.paragraph_format.space_before = Pt(0)
    table_text.paragraph_format.space_after = Pt(0)
    table_text.paragraph_format.line_spacing = 1.05


def configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        left = paragraph.add_run("M3-PI Guard | SCI 二区网络安全研究")
        set_run_font(left, size=8.5, color=MUTED, bold=True)
        paragraph.add_run("\t")
        right = paragraph.add_run("阶段报告 | 2026-07-26")
        set_run_font(right, size=8.5, color=MUTED)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        prefix = fp.add_run("M3-PI Guard  ·  ")
        set_run_font(prefix, size=9, color=MUTED)
        add_page_field(fp)


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    run = kicker.add_run("RESEARCH STAGE REPORT")
    set_run_font(run, size=9, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("M3-PI Guard 阶段性研究报告")
    set_run_font(run, size=24, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run(
        "从 v7 fail-closed 审计到 v8 六条件金丝雀通过："
        "面向 SCI 二区网络安全论文的证据、风险与后续实验"
    )
    set_run_font(run, size=13, color=MUTED)

    metadata = [
        ("报告日期", "2026-07-26（Asia/Shanghai）"),
        ("研究定位", "评估方法 + provenance-aware runtime authorization 防御机制"),
        ("目标期刊", "Journal of Information Security and Applications（投稿前复核 JCR Q2）"),
        ("当前模型", "DeepSeek 官方 API / deepseek-v4-pro；Gemini 保留为独立泛化矩阵"),
        ("证据边界", "本报告只读取技术/运行元数据；未开启 6 条金丝雀的科学 outcome 分析"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        label_run = p.add_run(f"{label}：")
        set_run_font(label_run, size=10.5, color=BLACK, bold=True)
        value_run = p.add_run(value)
        set_run_font(value_run, size=10.5, color=BLACK)


def add_callout(
    doc: Document,
    label: str,
    text: str,
    *,
    fill: str = CALLOUT,
    accent: str = BLUE,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.14)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    label_run = p.add_run(f"{label}  ")
    set_run_font(label_run, size=10.5, color=accent, bold=True)
    body_run = p.add_run(text)
    set_run_font(body_run, size=10.5, color=BLACK)


def add_paragraph(
    doc: Document,
    text: str,
    *,
    bold_lead: str | None = None,
    style: str | None = None,
) -> Any:
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def fill_table_cell(
    cell: Any,
    text: str,
    *,
    bold: bool = False,
    color: str = BLACK,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.style = "Table Text"
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=9, color=color, bold=bold)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_dxa: Sequence[int],
    *,
    status_columns: Sequence[int] = (),
) -> Any:
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        fill_table_cell(
            cell,
            label,
            bold=True,
            color=INK,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if index in status_columns
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            color = BLACK
            if index in status_columns:
                normalized = value.lower()
                if "pass" in normalized or "通过" in value or "就绪" in value:
                    color = GREEN
                elif "失败" in value or "阻塞" in value:
                    color = RED
                elif "待" in value or "未" in value:
                    color = GOLD
            fill_table_cell(
                row.cells[index],
                value,
                color=color,
                bold=index in status_columns,
                align=align,
            )
    set_table_geometry(table, widths_dxa)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    return table


def add_figure(
    doc: Document,
    image_path: Path,
    caption: str,
    source: str,
    *,
    alt: str,
) -> None:
    caption_p = doc.add_paragraph(style="Caption")
    run = caption_p.add_run(caption)
    set_run_font(run, size=9.5, color=DARK_BLUE, bold=True)
    image_p = doc.add_paragraph()
    image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_p.paragraph_format.space_before = Pt(0)
    image_p.paragraph_format.space_after = Pt(2)
    image_p.paragraph_format.keep_together = True
    image_p.add_run().add_picture(str(image_path), width=Inches(6.15))
    set_image_alt(image_p, title=caption, description=alt)
    source_p = doc.add_paragraph(style="Report Source")
    source_run = source_p.add_run(f"来源：{source}")
    set_run_font(source_run, size=8.5, color=MUTED)


def font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def draw_centered(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int],
    text: str,
    *,
    font_value: ImageFont.FreeTypeFont,
    fill: str,
) -> None:
    x, y = xy
    bbox = draw.textbbox((0, 0), text, font=font_value)
    width = bbox[2] - bbox[0]
    draw.text((x - width / 2, y), text, font=font_value, fill=fill)


def create_progress_chart(path: Path) -> None:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 55), "论文证据工作流：当前完成度", font=font(48, bold=True), fill=f"#{INK}")
    draw.text(
        (90, 120),
        "百分比表示已形成可进入该工作流的正式证据，不代表论文录用概率",
        font=font(25),
        fill=f"#{MUTED}",
    )
    rows = [
        ("工程与协议冻结", 100.0, "v8 + result gate v5"),
        ("DeepSeek 中型矩阵", 1.25, "6 / 480"),
        ("Gemini 跨模型验证", 0.0, "独立矩阵待冻结"),
        ("公共基准外部有效性", 0.0, "AgentDojo / SIREN 待运行"),
        ("自适应攻击", 0.0, "预算 50/200/500/2000 待运行"),
        ("独立盲测", 0.0, "保管人仍待落实"),
    ]
    y = 210
    for label, value, note in rows:
        draw.text((90, y), label, font=font(28, bold=True), fill=f"#{BLACK}")
        x0, x1 = 610, 1390
        draw.rounded_rectangle((x0, y + 2, x1, y + 36), 17, fill=f"#{LIGHT_GRAY}")
        if value > 0:
            fill_x = x0 + max(18, int((x1 - x0) * value / 100))
            draw.rounded_rectangle((x0, y + 2, fill_x, y + 36), 17, fill=f"#{BLUE}")
        draw.text((610, y + 48), note, font=font(22), fill=f"#{MUTED}")
        draw.text((1410, y), f"{value:.2f}%", font=font(25, bold=True), fill=f"#{DARK_BLUE}")
        y += 105
    image.save(path, dpi=(160, 160))


def create_cost_call_chart(
    path: Path,
    *,
    v7_calls: int,
    v7_cost: float,
    v8_calls: int,
    v8_cost: float,
) -> None:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 55), "官方 API 技术运行：调用量与增量成本", font=font(48, bold=True), fill=f"#{INK}")
    draw.text(
        (90, 120),
        "仅统计 v7 失败金丝雀与 v8 成功金丝雀；均无自动重试、无 Gemini 回退",
        font=font(25),
        fill=f"#{MUTED}",
    )
    panels = [
        ("官方响应次数", [("v7", v7_calls), ("v8", v8_calls)], max(v8_calls, 1), "次"),
        ("增量成本", [("v7", v7_cost), ("v8", v8_cost)], max(v8_cost, 1e-6), "元"),
    ]
    panel_x = [90, 850]
    colors = [f"#{RED}", f"#{BLUE}"]
    for panel_index, (title, values, maximum, unit) in enumerate(panels):
        x = panel_x[panel_index]
        draw.rounded_rectangle((x, 205, x + 660, 790), 24, outline="#D8DEE8", width=3, fill="#FBFCFE")
        draw.text((x + 35, 235), title, font=font(32, bold=True), fill=f"#{DARK_BLUE}")
        baseline = 690
        for item_index, (label, value) in enumerate(values):
            bar_x = x + 115 + item_index * 260
            height = int(360 * float(value) / float(maximum))
            draw.rounded_rectangle(
                (bar_x, baseline - height, bar_x + 145, baseline),
                18,
                fill=colors[item_index],
            )
            formatted = (
                f"{int(value)} {unit}"
                if unit == "次"
                else f"¥{float(value):.7f}"
            )
            draw_centered(
                draw,
                (bar_x + 72, baseline - height - 47),
                formatted,
                font_value=font(24, bold=True),
                fill=colors[item_index],
            )
            draw_centered(
                draw,
                (bar_x + 72, baseline + 25),
                label,
                font_value=font(29, bold=True),
                fill=f"#{BLACK}",
            )
    image.save(path, dpi=(160, 160))


def create_token_chart(
    path: Path,
    *,
    v7_hit: int,
    v7_miss: int,
    v7_completion: int,
    v8_hit: int,
    v8_miss: int,
    v8_completion: int,
) -> None:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 55), "金丝雀 token 组成（技术账本）", font=font(48, bold=True), fill=f"#{INK}")
    draw.text(
        (90, 120),
        "DeepSeek cache hit 显著占优；v8 六条件金丝雀总 token 增量为 7,841",
        font=font(25),
        fill=f"#{MUTED}",
    )
    series = [
        ("v7 失败段", v7_hit, v7_miss, v7_completion),
        ("v8 6/6", v8_hit, v8_miss, v8_completion),
    ]
    maximum = max(sum(values[1:]) for values in series)
    segments = [
        ("缓存命中输入", f"#{BLUE}"),
        ("缓存未命中输入", "#D5A021"),
        ("输出", "#6B7280"),
    ]
    y_positions = [315, 555]
    for (label, hit, miss, completion), y in zip(series, y_positions):
        draw.text((90, y - 70), label, font=font(31, bold=True), fill=f"#{BLACK}")
        x = 320
        total = hit + miss + completion
        values = [hit, miss, completion]
        for value, (_, color) in zip(values, segments):
            width = max(2, int(1050 * value / maximum))
            draw.rounded_rectangle((x, y, x + width, y + 90), 12, fill=color)
            if width > 105:
                draw_centered(
                    draw,
                    (x + width // 2, y + 27),
                    f"{value:,}",
                    font_value=font(23, bold=True),
                    fill="white",
                )
            x += width
        draw.text((1390, y + 25), f"{total:,}", font=font(25, bold=True), fill=f"#{DARK_BLUE}")
    legend_y = 760
    x = 320
    for label, color in segments:
        draw.rounded_rectangle((x, legend_y, x + 30, legend_y + 30), 7, fill=color)
        draw.text((x + 43, legend_y - 2), label, font=font(22), fill=f"#{BLACK}")
        x += 310
    image.save(path, dpi=(160, 160))


def create_lifecycle_chart(path: Path) -> None:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.text((90, 55), "v7 → v8 证据生命周期", font=font(48, bold=True), fill=f"#{INK}")
    draw.text(
        (90, 120),
        "开放 attempt 不允许原地重试；修复必须使用新 run / matrix identity",
        font=font(25),
        fill=f"#{MUTED}",
    )
    steps = [
        ("v7 准入", "全量门禁通过", BLUE),
        ("v7 fail-closed", "lazy binding\n0 result / 1 open", RED),
        ("Abort review v3", "2 response IDs\noutcome closed", GOLD),
        ("v8 新身份", "full _execute binding\ndelayed test pass", DARK_BLUE),
        ("v8 金丝雀", "6/6 completed\n13 calls / 0 failure", GREEN),
        ("下一门", "474 continuation\nresult gate closed", BLUE),
    ]
    start_x = 65
    box_w = 225
    gap = 28
    y = 290
    for index, (title, note, color) in enumerate(steps):
        x = start_x + index * (box_w + gap)
        draw.rounded_rectangle(
            (x, y, x + box_w, y + 290),
            22,
            fill="#FBFCFE",
            outline=f"#{color}",
            width=4,
        )
        draw.rounded_rectangle((x, y, x + box_w, y + 64), 20, fill=f"#{color}")
        draw_centered(
            draw,
            (x + box_w // 2, y + 14),
            title,
            font_value=font(25, bold=True),
            fill="white",
        )
        lines = note.split("\n")
        ty = y + 105
        for line in lines:
            draw_centered(
                draw,
                (x + box_w // 2, ty),
                line,
                font_value=font(18, bold=False),
                fill=f"#{BLACK}",
            )
            ty += 48
        if index < len(steps) - 1:
            ax = x + box_w + 5
            ay = y + 145
            draw.line((ax, ay, ax + gap - 10, ay), fill="#94A3B8", width=5)
            draw.polygon(
                [(ax + gap - 10, ay), (ax + gap - 23, ay - 10), (ax + gap - 23, ay + 10)],
                fill="#94A3B8",
            )
    draw.rounded_rectangle((90, 680, 1510, 805), 20, fill=f"#{CALLOUT}")
    draw.text(
        (130, 710),
        "核心审计结论：v7 失败没有被隐藏或重试；v8 只继承冻结计划，不继承开放 attempt。",
        font=font(26, bold=True),
        fill=f"#{INK}",
    )
    draw.text(
        (130, 757),
        "这条证据链可直接进入论文 Methods / Reproducibility，而不能替代安全效果统计。",
        font=font(23),
        fill=f"#{MUTED}",
    )
    image.save(path, dpi=(160, 160))


def add_hyperlink(paragraph: Any, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_reference(
    doc: Document,
    label: str,
    description: str,
    url: str,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"{label}：{description} ")
    set_run_font(run, size=9.5)
    add_hyperlink(p, "查看来源", url)


def build_report() -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)

    abort = load_json(V7_ABORT_REVIEW)
    candidate = load_json(V8_CANDIDATE)
    canary_admission = load_json(V8_CANARY_ADMISSION)
    continuation_admission = load_json(V8_CONTINUATION_ADMISSION)
    gate_candidate = load_json(RESULT_GATE_CANDIDATE)
    canary_summary = read_last_json_object(V8_CANARY_SUMMARY)
    canary_review_path = (
        PROJECT_ROOT
        / candidate["scope"]["outputs"]["root"]
        / candidate["scope"]["canary"]["review_filename"]
    )
    canary_review = load_json(canary_review_path)

    v7_usage = abort["operational_summary"]["project_cumulative_usage"]
    v7_prior = {
        "prompt_cache_hit_tokens": 472832,
        "prompt_cache_miss_tokens": 162536,
        "completion_tokens": 67873,
        "prompt_tokens": 635368,
    }
    v7_hit = v7_usage["prompt_cache_hit_tokens"] - v7_prior["prompt_cache_hit_tokens"]
    v7_miss = (
        v7_usage["prompt_cache_miss_tokens"]
        - v7_prior["prompt_cache_miss_tokens"]
    )
    v7_completion = v7_usage["completion_tokens"] - v7_prior["completion_tokens"]
    v7_cost = abort["operational_summary"]["incremental_cost_cny"]

    v8_prior = candidate["scope"]["budget"]["prior_project_usage"]
    v8_usage = canary_summary["project_cumulative_usage"]
    v8_hit = v8_usage["prompt_cache_hit_tokens"] - v8_prior["prompt_cache_hit_tokens"]
    v8_miss = (
        v8_usage["prompt_cache_miss_tokens"]
        - v8_prior["prompt_cache_miss_tokens"]
    )
    v8_completion = v8_usage["completion_tokens"] - v8_prior["completion_tokens"]
    v8_cost = (
        canary_summary["project_cumulative_cost_cny"]
        - candidate["scope"]["budget"]["prior_project_cost_cny"]
    )

    lifecycle_chart = QA_DIR / "fig01_v7_v8_lifecycle.png"
    cost_chart = QA_DIR / "fig02_api_cost_calls.png"
    token_chart = QA_DIR / "fig03_token_composition.png"
    progress_chart = QA_DIR / "fig04_research_progress.png"
    create_lifecycle_chart(lifecycle_chart)
    create_cost_call_chart(
        cost_chart,
        v7_calls=abort["operational_summary"]["provider_calls_completed"],
        v7_cost=v7_cost,
        v8_calls=canary_summary["provider_calls_completed"],
        v8_cost=v8_cost,
    )
    create_token_chart(
        token_chart,
        v7_hit=v7_hit,
        v7_miss=v7_miss,
        v7_completion=v7_completion,
        v8_hit=v8_hit,
        v8_miss=v8_miss,
        v8_completion=v8_completion,
    )
    create_progress_chart(progress_chart)

    doc = Document()
    apply_style_tokens(doc)
    configure_sections(doc)
    props = doc.core_properties
    props.title = "M3-PI Guard 阶段性研究报告（2026-07-26）"
    props.subject = "SCI 二区网络安全论文工程与实验阶段报告"
    props.author = "M3-PI Guard Research Team"
    props.keywords = "LLM agent security, prompt injection, authorization, provenance"
    props.comments = "Generated from content-free technical evidence."

    add_title_block(doc)
    add_callout(
        doc,
        "阶段结论",
        "SCI 二区方向仍然有条件可行。工程主线已从 v7 的 fail-closed "
        "故障收敛到 v8 六条件金丝雀 6/6 技术通过；但正式科学证据仍只有 "
        "6/480 条，尚不能声称安全效果、非劣性或 SOTA。",
        fill=LIGHT_BLUE,
        accent=DARK_BLUE,
    )
    add_callout(
        doc,
        "证据边界",
        "本报告只读取 operational status、result tier、调用/成本/token、哈希与"
        "授权状态。unsafe_effect、task_completion 的六条件对比尚未进入分析；"
        "result admission 和 analysis authorization 均为 false。",
        fill="FFF8E8",
        accent=GOLD,
    )
    doc.add_page_break()

    doc.add_heading("1. 执行摘要", level=1)
    add_paragraph(
        doc,
        "本阶段完成了三件对投稿价值高于“继续堆代码”的工作。第一，v7 金丝雀"
        "在真实官方 API 响应后因 adapter 生命周期接线错误 fail-closed，失败被"
        "完整保留而没有原地重试。第二，v8 以新 run 和 matrix identity 修复"
        " lazy backend 解析时序，并通过离线回归、冻结候选、六条件金丝雀。"
        "第三，result gate v5 在科学 outcome 打开前冻结，确保后续分析必须经过"
        "独立准入。",
    )
    add_bullets(
        doc,
        [
            "v7：0 条结果记录、1 个开放 attempt、2 次官方响应、增量成本 "
            f"¥{v7_cost:.7f}；永久不可恢复。",
            "v8：6/6 条技术记录 completed + formal_candidate，13 次官方响应，"
            f"0 diagnostic、0 operational failure、增量成本 ¥{v8_cost:.7f}。",
            "当前累计项目 API 成本账本为 "
            f"¥{canary_summary['project_cumulative_cost_cny']:.7f}，"
            "距离 CNY 27 abort fuse 仍有约 "
            f"¥{27 - canary_summary['project_cumulative_cost_cny']:.4f}。",
            "v8 continuation admission 已生成，可执行剩余 474 条；"
            "result admission 仍关闭。",
        ],
    )

    doc.add_heading("2. 当前研究定位与 SCI 二区可行性", level=1)
    add_paragraph(
        doc,
        "论文定位保持为“评估方法 + 防御机制”双主线：一方面量化逐样本答案、"
        "gold authorization 或过度明确策略进入在线 Guard 后造成的评估高估；"
        "另一方面验证非 oracle、权限单调、参数来源感知与 ALLOW/ASK/DENY "
        "机制能否在上下文依赖任务上取得真实安全—效用优势。",
    )
    add_table(
        doc,
        ["判断维度", "当前证据", "对 SCI Q2 的含义"],
        [
            [
                "工程与审计",
                "新旧 run 隔离、append-only 请求链、失败留在分母、准入分层均已实现",
                "具备可复现 Methods / Artifact 基础",
            ],
            [
                "机制增益",
                "尚未打开 480 条完整配对结果",
                "不能声称优于 policy-only 或结构化基线",
            ],
            [
                "外部有效性",
                "AgentDojo/SIREN/InjecAgent 正式结果为空",
                "当前不能支撑广泛泛化声明",
            ],
            [
                "统计与效用",
                "−5pp TCR 非劣与配对安全差尚未计算",
                "这是防御机制论文能否成立的决定性门槛",
            ],
            [
                "期刊匹配",
                "网络安全、访问控制、安全策略与 LLM agent 应用安全",
                "JISA 仍为首选；投稿前正式复核 SCIE/JCR Q2",
            ],
        ],
        [1700, 3600, 4060],
    )
    add_callout(
        doc,
        "可行性判断",
        "“做成 SCI 二区”仍然可行，但条件从工程完成转为实验达标："
        "M3-full 必须在至少两个公共/外部基准上对最强结构化基线产生实质安全增益，"
        "且 TCR 不劣于 −5pp、full 独立优于 policy-only、自适应攻击不推翻主结论。",
    )

    doc.add_heading("3. 本阶段证据边界与不变量", level=1)
    add_table(
        doc,
        ["边界", "冻结要求", "当前状态"],
        [
            ["在线可见性", "evaluator/gold 字段不得进入在线状态", "已由投影与测试约束"],
            ["权限扩张", "自动路径只能收缩；扩张需独立确认", "协议与 runtime 接口保留"],
            ["真实效果", "模型输出仅是不可信 proposal；无外部 connector/tool", "v8 金丝雀满足"],
            ["失败处理", "超时/解析/runner failure 留在 ITT；开放 attempt 不重试", "v7 证明 fail-closed"],
            ["请求身份", "本地 request ID + payload hash + response ID + 精确模型 + hash-chain", "v7/v8 均保留"],
            ["结果访问", "技术审查与科学分析分离", "result gate v5 已预先冻结"],
        ],
        [1700, 4080, 3580],
        status_columns=[2],
    )
    add_figure(
        doc,
        lifecycle_chart,
        "图 1  v7 到 v8 的证据生命周期与授权分界",
        "formal_matrix_abort_review_v3、v8 candidate/canary review、result gate v5",
        alt="从 v7 准入、失败封存、v8 新身份到六条件金丝雀和下一门禁的流程图。",
    )

    doc.add_heading("4. v7 失败审计：为什么这是有价值的负面工程证据", level=1)
    add_paragraph(
        doc,
        "v7 在预金丝雀全量 Harness 通过后只执行一次 admitted canary。"
        "首个 work item 完成两次官方 provider 响应，随后 v3 offline evaluator "
        "拒绝 material，错误为 boundary adapter version missing/mismatched。"
        "由于结果记录尚未 append，账本形成 1 个开放 attempt。",
    )
    add_table(
        doc,
        ["项目", "v7 观测值", "审计解释"],
        [
            ["结果记录", "0", "没有不完整结果进入效果表"],
            ["开放 attempt", "1", "禁止自动恢复，必须新版本"],
            ["官方响应", "2", "response identity 与预算链完整"],
            ["trace 文件", "1", "与 open work key 一一绑定"],
            ["自动重试", "0", "未掩盖失败率"],
            ["外部效果", "0", "无工具、connector 或真实动作"],
            ["增量成本", f"¥{v7_cost:.7f}", "成本很小，但证据必须保留"],
        ],
        [1900, 2200, 5260],
        status_columns=[1],
    )
    add_paragraph(
        doc,
        "根因不是模型或 API 不稳定，而是作用域错误：v7 只在 dispatcher 构造时"
        "临时替换 v3 backend 类；真正 backend factory 在运行阶段延迟解析模块"
        "全局变量，此时绑定已经恢复为旧类。v3 evaluator 正确拒绝了旧 material。"
    )
    add_callout(
        doc,
        "审计价值",
        "如果原地重试，论文会丢失一个真实 operational failure，并破坏"
        " intention-to-treat 与请求链可审计性。保留 v7、创建 v8 是正确的"
        "科学工程选择。",
        fill="FDECEC",
        accent=RED,
    )

    doc.add_heading("5. v8 修复与新增回归", level=1)
    add_paragraph(
        doc,
        "v8 不修改冻结 v7，而是以新 run ID、matrix identity、candidate 和"
        " canary admission 重建证据链。唯一机制修复是让 v3 adapter binding "
        "覆盖完整 _execute 生命周期，使 backend factory 的延迟解析仍指向 v3。"
    )
    add_table(
        doc,
        ["修复项", "v7", "v8", "验证方式"],
        [
            ["adapter binding 作用域", "dispatcher 构造期", "完整 _execute", "延迟 backend 解析断言"],
            ["旧开放 attempt", "1，禁止恢复", "0，新输出根", "候选输出 freshness + state replay"],
            ["资格场景", "backend→evaluator 直接", "6 个 delayed resolution 场景", "3/3 时序聚焦测试"],
            ["版本声明", "v3 主字段但旧元数据残留", "qualification/failure adapter 均 v3", "validator 强制一致"],
            ["结果门禁", "v4 绑定 v7", "v5 绑定 v8 生命周期修复", "预金丝雀 qualification"],
        ],
        [2060, 1900, 2080, 3320],
    )
    add_paragraph(
        doc,
        "v8 candidate manifest 为 "
        f"{candidate['manifest_sha256']}；matrix identity 为 "
        f"{candidate['scope']['matrix_identity']}。candidate 本身不可执行，"
        "canary 与 continuation 通过独立 admission 分层。",
    )

    doc.add_heading("6. 验证证据与未关闭事项", level=1)
    add_table(
        doc,
        ["验证层", "结果", "说明"],
        [
            ["预 v7 全量 Harness", "PASS", "Verification failed=0，provider calls=0"],
            ["v8 生命周期聚焦", "PASS", "3/3；延迟 backend 解析覆盖"],
            ["v8 冻结 artifact", "PASS", "candidate/admission 复现 1/1"],
            ["result gate v5", "PASS", "4/4；746.04 秒"],
            ["预 v8 全量 Harness", "1072 PASS / 1 SKIP / 1 stale fail", "唯一失败为旧 v7 freshness 断言"],
            ["stale test 修复", "PASS", "v7 永久中止聚焦 1/1，143.55 秒"],
            ["完整全量复跑", "待执行", "长矩阵 continuation 前必须补跑"],
        ],
        [2200, 2800, 4360],
        status_columns=[1],
    )
    add_callout(
        doc,
        "透明披露",
        "当前不能把预 v8 全量 Harness 写成完全绿色：它有 1 个陈旧生命周期"
        "测试失败。对应测试已被新永久不变量替换并聚焦通过，但 474 条 continuation "
        "启动前仍应再跑一次全量 Harness，取得 Verification failed=0。",
        fill="FFF8E8",
        accent=GOLD,
    )

    doc.add_heading("7. 官方 API 调用、成本与缓存", level=1)
    add_figure(
        doc,
        cost_chart,
        "图 2  v7 与 v8 金丝雀的官方响应次数及增量成本",
        "append-only matrix budget ledgers；成本按冻结 DeepSeek 费率计算",
        alt="两组柱状图比较 v7 和 v8 的官方响应次数与人民币增量成本。",
    )
    add_figure(
        doc,
        token_chart,
        "图 3  v7 与 v8 金丝雀的 token 组成",
        "项目预算账本增量；输入 token = cache hit + cache miss",
        alt="堆叠条形图显示 v7 与 v8 的缓存命中输入、未命中输入和输出 token。",
    )
    add_table(
        doc,
        ["指标", "v7 失败段", "v8 6 条金丝雀"],
        [
            ["官方响应", str(abort["operational_summary"]["provider_calls_completed"]), str(canary_summary["provider_calls_completed"])],
            ["缓存命中输入", f"{v7_hit:,}", f"{v8_hit:,}"],
            ["缓存未命中输入", f"{v7_miss:,}", f"{v8_miss:,}"],
            ["输出 token", f"{v7_completion:,}", f"{v8_completion:,}"],
            ["总增量 token", f"{v7_hit + v7_miss + v7_completion:,}", f"{v8_hit + v8_miss + v8_completion:,}"],
            ["增量成本", f"¥{v7_cost:.7f}", f"¥{v8_cost:.7f}"],
            ["自动重试 / fallback", "0 / 无", "0 / 无"],
        ],
        [3000, 3180, 3180],
        status_columns=[1, 2],
    )
    add_paragraph(
        doc,
        "v8 cache-hit 输入占其输入 token 的 "
        f"{v8_hit / (v8_hit + v8_miss) * 100:.1f}%。这验证了用户关于 DeepSeek "
        "缓存命中率较高的判断，也意味着 474 条 continuation 的实际成本更可能"
        "受输出长度与 cache-miss 波动影响，而不是简单按全 miss 上界线性增长。",
    )

    doc.add_heading("8. v8 六条件金丝雀：能证明什么，不能证明什么", level=1)
    add_table(
        doc,
        ["技术项目", "观测值", "判定"],
        [
            ["预定条件", "6：no defense / safe prompt / PromptGuard / structural / policy-only / full", "覆盖通过"],
            ["记录数", "6 / 6", "通过"],
            ["operational status", "全部 completed", "通过"],
            ["result tier", "全部 formal_candidate", "通过"],
            ["provider calls", str(canary_review["provider_calls_completed"]), "账本完整"],
            ["diagnostic records", "0", "通过"],
            ["operational failures", "0", "通过"],
            ["rejected effects executed", "0", "通过"],
            ["scientific outcomes reviewed", "false", "保持关闭"],
        ],
        [2400, 4480, 2480],
        status_columns=[2],
    )
    add_callout(
        doc,
        "不能外推",
        "6 条金丝雀只验证运行链、依赖、身份、成本、隔离与记录层。"
        "它不能证明 M3-full 安全率更低、TCR 非劣、policy-only 增益、"
        "跨基准泛化或自适应攻击鲁棒性。",
        fill="FDECEC",
        accent=RED,
    )

    doc.add_heading("9. 当前授权状态", level=1)
    add_table(
        doc,
        ["Artifact", "Manifest / SHA-256", "授权状态"],
        [
            ["v8 candidate", candidate["manifest_sha256"], "candidate-only"],
            ["v8 canary admission", canary_admission["manifest_sha256"], "已消费：6 条"],
            ["v8 canary review", sha256_file(canary_review_path), "technical gate pass"],
            ["result gate v5 candidate", gate_candidate["manifest_sha256"], "科学 outcome 前冻结"],
            ["v8 continuation admission", continuation_admission["manifest_sha256"], "已授权：剩余 474 条"],
            ["result admission", "不存在", "未授权"],
            ["scientific analysis", "不存在", "未授权"],
        ],
        [2200, 4700, 2460],
        status_columns=[2],
    )
    add_paragraph(
        doc,
        "result gate v5 candidate 明确绑定 480 条预期记录、canary review hash、"
        "technical policy 与 planned outputs。当前证据状态允许执行，不允许读"
        "科学 outcome，更不允许生成论文主结果表。",
    )

    doc.add_heading("10. 完整实验的下一步执行方案", level=1)
    add_figure(
        doc,
        progress_chart,
        "图 4  面向 SCI 二区论文的证据工作流完成度",
        "按已形成正式证据的工作流计数；不是论文录用概率",
        alt="横向进度条显示工程协议已冻结、DeepSeek 中型矩阵 6/480，其余外部验证尚待完成。",
    )

    doc.add_heading("10.1 Gate A：长矩阵前离线复核", level=2)
    add_bullets(
        doc,
        [
            "重跑完整 Harness，要求 exit 0 且 Verification failed=0；"
            "不得只引用聚焦测试。",
            "复核 v8 status：6 records、474 remaining、0 open attempts、"
            "budget not paused、result admission false。",
            "复核 official config 只显示 redacted:set，不输出密钥；"
            "Gemini 不并发、不作为 fallback。",
            "确认 result-gate candidate、canary review、continuation admission "
            "哈希均与本报告一致。",
        ],
    )

    doc.add_heading("10.2 Gate B：执行剩余 474 条 DeepSeek 矩阵", level=2)
    add_paragraph(
        doc,
        "建议使用 dedicated PromptGuard Python 3.13.5 环境、concurrency=1、"
        "automatic_retries=0。精确命令为：",
    )
    code = doc.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.25)
    code.paragraph_format.right_indent = Inches(0.25)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    run = code.add_run(
        r"E:\M3_PI_Guard_runtime\promptguard\Scripts\python.exe "
        r"-m experiments.paper_v2.formal_matrix_controller_v8 "
        r"--execute-continuation"
    )
    set_run_font(run, latin="Consolas", east_asia="Microsoft YaHei", size=8.8, color=INK)
    add_bullets(
        doc,
        [
            "每 30–60 秒只监控 record_count、remaining、open attempts、"
            "cumulative cost、process health；不读取科学 outcome。",
            "若出现 open attempt、unexpected adapter mismatch、身份漂移、"
            "预算 pause 或非零外部效果，立即停止并创建新 abort review；"
            "不得原地 retry。",
            "DeepSeek CNY 27 abort fuse 保持不变；当前账本距 fuse 约 "
            f"¥{27 - canary_summary['project_cumulative_cost_cny']:.4f}。",
        ],
    )

    doc.add_heading("10.3 Gate C：480 条完成后的技术审查与结果准入", level=2)
    add_table(
        doc,
        ["顺序", "动作", "允许读取", "通过条件"],
        [
            ["1", "生成 formal result technical review v5", "状态、tier、trace、成本、隔离", "480 records、0 remaining/open、身份链完整"],
            ["2", "若技术门通过，生成 result analysis admission", "仍不读效果字段", "技术 review 与 candidate 哈希绑定"],
            ["3", "执行 admitted analysis", "unsafe_effect、TCR、ASK 等", "显式 result admission 已激活"],
            ["4", "生成主表、CI、McNemar、bootstrap", "完整科学 outcome", "统计脚本从 manifest 可重建"],
        ],
        [800, 3300, 2500, 2760],
        status_columns=[0],
    )

    doc.add_heading("10.4 Gate D：机制论文成败判据", level=2)
    add_table(
        doc,
        ["判据", "预注册门槛", "未满足时动作"],
        [
            ["安全效应", "相对最强结构化基线降低 ≥20% 或绝对 ≥10pp", "转测量/实证论文"],
            ["合法效用", "TCR 差异的非劣界限 ≥−5pp", "停止机制优越性主张"],
            ["模块增益", "M3-full − policy-only unsafe-effect 严格为负", "删除 provenance 模块独立贡献"],
            ["拒绝退化", "benign 初始 ASK+DENY <0.90", "视为近 deny-all，门禁失败"],
            ["运行失败", "自动通过要求 0；非零需根因审查", "失败留 ITT，必要时新版本"],
        ],
        [1750, 4000, 3610],
    )

    doc.add_heading("11. DeepSeek 后续与 Gemini 独立矩阵", level=1)
    add_paragraph(
        doc,
        "DeepSeek 480 条完成并通过技术门后，再冻结 Gemini 独立矩阵。"
        "Gemini 不能作为 DeepSeek 单项失败的即时 fallback，否则会破坏配对、"
        "供应商身份和预算归属。Gemini 应拥有独立 run ID、candidate、admission、"
        "quota snapshot、request chain 与 result gate。",
    )
    add_table(
        doc,
        ["模型", "计划规模", "条件", "主要目的"],
        [
            ["DeepSeek v8", "480 条（已完成 6）", "全部 6 核心条件 + regime probes", "主中型证伪与功效估计"],
            ["Gemini", "至少 300 条起步", "最强结构化 / policy-only / full", "跨模型方向一致性"],
            ["公共基准主模型", "≥300 attack + 150 benign / 基准", "核心条件", "外部有效性"],
            ["第二模型公共基准", "≥200 attack + 100 benign", "3 条核心条件", "泛化复核"],
        ],
        [1700, 2200, 2820, 2640],
    )

    doc.add_heading("12. 公共基准、自适应攻击与盲测", level=1)
    add_bullets(
        doc,
        [
            "AgentDojo：主公共基准，必须统一 unsafe-effect / task-completion / "
            "operational failure 分母语义。",
            "SIREN：优先用于上下文依赖与工具流注入外部有效性；若适配失败，"
            "不得用私有任务替代而不披露。",
            "InjecAgent：作为补充泛化，不替代主公共基准。",
            "自适应攻击：主模型 40 任务，预算 50/200/500；最难 20 任务追加 2000；"
            "攻击者知道防御结构并优化 unsafe effect。",
            "独立盲测：第 8 周前确定标签保管人；否则只能称内部 sealed holdout，"
            "不能称独立盲测。",
        ],
    )
    add_table(
        doc,
        ["实验", "主要终点", "统计", "最低交付"],
        [
            ["480 条中型矩阵", "unsafe-effect；TCR；full-policy gain", "Wilson、exact McNemar、10k cluster bootstrap、Holm", "原始 n/N + 绝对/相对差 + CI"],
            ["跨模型", "方向一致性与异质性", "分层/cluster bootstrap", "至少一个可审计第二模型"],
            ["自适应攻击", "预算—ASR 饱和曲线", "按任务聚类", "防御感知而非静态攻击"],
            ["盲测", "sealed unsafe-effect/TCR", "独立 custody chain", "SHA-256 + 标签保管链"],
        ],
        [1800, 2700, 2800, 2060],
    )

    doc.add_heading("13. 风险、止损与工程债务", level=1)
    add_table(
        doc,
        ["风险", "当前级别", "触发信号", "处置"],
        [
            ["full 不优于 policy-only", "高", "模块增益≥0", "切换授权泄漏/不确定性实证论文"],
            ["安全提升来自过度阻断", "高", "TCR 非劣失败或 ASK+DENY≈deny-all", "停止 SOTA 防御主张"],
            ["最强结构化基线更好", "中高", "配对 unsafe-effect 不占优", "保留 M3 作为测量工具"],
            ["外部基准方向不一致", "高", "AgentDojo/SIREN 异质性显著", "收缩外部有效性声明"],
            ["result gate v5 构建过慢", "中", "candidate 构建约 33 分钟", "当前冻结不改；下一版本做 memoization 设计审计"],
            ["API/身份漂移", "中", "model ID、response chain 或 gateway 声明不符", "记录降级 diagnostic，停止正式矩阵"],
        ],
        [2050, 1200, 2870, 3240],
        status_columns=[1],
    )

    doc.add_heading("14. 建议排期（从当前 checkpoint 起）", level=1)
    add_table(
        doc,
        ["时间", "工作包", "出口条件"],
        [
            ["第 1 周", "全量 Harness 复跑；执行 DeepSeek 剩余 474 条；技术 review", "480/480、0 open、result gate 技术通过"],
            ["第 2 周", "admitted analysis；功效分析；主结果表", "机制/实证分支自动判定"],
            ["第 3–4 周", "Gemini 独立矩阵；AgentDojo 主公共基准", "第二模型与主公共基准方向证据"],
            ["第 5–6 周", "SIREN/InjecAgent；异质性分析", "至少两个外部基准方向可解释"],
            ["第 7–8 周", "自适应攻击；盲测保管链", "攻击预算曲线与 sealed 数据"],
            ["第 9–10 周", "Methods/Results、claim–evidence ledger、统计审查", "所有数字可从 manifest 重建"],
            ["第 11–12 周", "Introduction/Related Work、JISA 格式、引用与图表核验", "投稿包或自动切换实证论文"],
        ],
        [1300, 4300, 3760],
    )

    doc.add_heading("15. 论文写作建议", level=1)
    add_bullets(
        doc,
        [
            "先写 Methods / Results，再写 Introduction；把 v7 失败与 v8 新身份"
            "作为 artifact integrity / failure handling 的方法证据。",
            "主标题暂保留 provenance-aware runtime authorization，"
            "但是否使用 “Beyond Answer-Key Authorization” 取决于源码审计是否"
            "证明 oracle leakage 具有系统性。",
            "不使用“首次、唯一、SOTA”；与 VIGIL、DRIFT、Task Shield、CaMeL、"
            "ARGUS、AuthGraph 明确列差异和重叠。",
            "表格先报告原始分子/分母，再报告效应量与 CI；"
            "p < 0.05 不能替代实质门槛。",
            "若前三个机制门槛失败，自动切换题目为 authorization leakage、"
            "uncertainty 与 security–utility trade-off 的系统实证研究。",
        ],
    )

    doc.add_heading("16. 可复现性检查清单", level=1)
    checklist = [
        ("v7 abort review", abort["manifest_sha256"], "已冻结"),
        ("v8 candidate", candidate["manifest_sha256"], "已冻结"),
        ("v8 canary admission", canary_admission["manifest_sha256"], "已消费"),
        ("v8 canary review", sha256_file(canary_review_path), "技术门通过"),
        ("result gate v5 qualification", sha256_file(RESULT_GATE_QUALIFICATION), "预金丝雀冻结"),
        ("result gate v5 candidate", gate_candidate["manifest_sha256"], "科学 outcome 前冻结"),
        ("v8 continuation admission", continuation_admission["manifest_sha256"], "已授权、未执行"),
    ]
    add_table(
        doc,
        ["证据", "SHA-256 / manifest", "状态"],
        [[name, digest, status] for name, digest, status in checklist],
        [2300, 4800, 2260],
        status_columns=[2],
    )
    add_callout(
        doc,
        "下一位执行者的唯一安全起点",
        "先读取 AGENTS.md、current-task、architecture、progress；"
        "然后重跑全量 Harness。只有 Verification failed=0，才可使用现有 "
        "v8 continuation admission 执行剩余 474 条。",
        fill=LIGHT_BLUE,
        accent=DARK_BLUE,
    )

    doc.add_heading("17. 关键前沿工作与投稿资料", level=1)
    add_reference(
        doc,
        "VIGIL",
        "ACL 2026，verify-before-commit 与 SIREN 工具流注入基准",
        "https://aclanthology.org/2026.acl-long.443/",
    )
    add_reference(
        doc,
        "DRIFT",
        "NeurIPS 2025，动态规则、控制/数据流约束与记忆隔离",
        "https://proceedings.neurips.cc/paper_files/paper/2025/hash/77f3b26c7907aa27b207df9b9d43f29a-Abstract-Conference.html",
    )
    add_reference(
        doc,
        "Task Shield",
        "ACL 2025，工具调用与用户目标一致性验证",
        "https://aclanthology.org/2025.acl-long.1435/",
    )
    add_reference(
        doc,
        "ARGUS",
        "2026 预印本，上下文依赖任务与 provenance-aware decision auditing",
        "https://arxiv.org/abs/2605.03378",
    )
    add_reference(
        doc,
        "AuthGraph",
        "2026 预印本，授权图与参数来源级执行图",
        "https://arxiv.org/abs/2605.26497",
    )
    add_reference(
        doc,
        "JISA",
        "Journal of Information Security and Applications 官方范围与投稿入口",
        "https://www.sciencedirect.com/journal/journal-of-information-security-and-applications",
    )
    add_paragraph(
        doc,
        "注：期刊 JCR 分区具有时效性，正式投稿前必须通过学校或 Clarivate "
        "订阅再次确认 SCIE 与 JCR Q2；本报告不把第三方分区页当作最终投稿证据。",
    )

    doc.add_heading("附录 A：本阶段关键数值", level=1)
    add_table(
        doc,
        ["字段", "值"],
        [
            ["v7 abort review manifest", abort["manifest_sha256"]],
            ["v8 candidate manifest", candidate["manifest_sha256"]],
            ["v8 matrix identity", candidate["scope"]["matrix_identity"]],
            ["v8 canary review SHA-256", sha256_file(canary_review_path)],
            ["result gate v5 candidate manifest", gate_candidate["manifest_sha256"]],
            ["v8 continuation admission manifest", continuation_admission["manifest_sha256"]],
            ["v8 canary cumulative cost CNY", f"{canary_summary['project_cumulative_cost_cny']:.7f}"],
            ["v8 canary incremental cost CNY", f"{v8_cost:.7f}"],
            ["v8 canary provider calls", str(canary_summary["provider_calls_completed"])],
            ["v8 canary records / remaining", f"{canary_summary['record_count']} / {canary_summary['remaining_work_items']}"],
        ],
        [3400, 5960],
    )

    doc.add_heading("附录 B：报告生成边界", level=1)
    add_paragraph(
        doc,
        "报告生成脚本只读取冻结 JSON、content-free canary summary、验证日志和"
        "哈希；未读取密钥、prompt、response content、unsafe_effect 或 "
        "task_completion 的条件对比。图 1–4 均由上述技术元数据生成。"
    )
    add_paragraph(
        doc,
        "本报告是阶段性工程与实验计划文件，不是论文 Results 草稿。"
        "任何安全/效用结论必须等待 result admission 激活后的正式统计分析。"
    )

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    output = build_report()
    print(
        json.dumps(
            {
                "output": str(output),
                "bytes": output.stat().st_size,
                "sha256": sha256_file(output),
                "qa_dir": str(QA_DIR),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
