"""Build the 2026-07-27 M3-PI Guard paper-v2 stage report.

The report is generated only from admitted or frozen repository evidence. It
does not read credentials, prompts, response bodies, or mutable provider state.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_paper_v2_stage_report_20260726 import (
    BLACK,
    BLUE,
    DARK_BLUE,
    GOLD,
    GREEN,
    INK,
    LIGHT_BLUE,
    MUTED,
    RED,
    add_bullets,
    add_callout,
    add_figure,
    add_page_field,
    add_paragraph,
    add_reference,
    add_table,
    apply_style_tokens,
    font,
    load_json,
    set_run_font,
    sha256_file,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = PROJECT_ROOT / "docs/research/paper_v2/reports"
QA_DIR = PROJECT_ROOT / "outputs/paper_upgrade_v2/report_qa/stage_report_20260727"
OUTPUT_PATH = REPORT_DIR / "M3_PI_Guard_阶段性研究报告_2026-07-27.docx"

ANALYSIS_PATH = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_medium/deepseek"
    / "ba9c7937d08b68e088faf8586e0dd96d2052a398fd0a77889fd6a62e503f2373"
    / "admitted_medium_analysis_v8.json"
)
TECHNICAL_REVIEW_PATH = ANALYSIS_PATH.parent / "formal_result_technical_review_v6.json"
RESULT_ADMISSION_PATH = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/admissions"
    / "deepseek_medium_combined_result_admission_v6.json"
)
SCORECARD_PATH = (
    PROJECT_ROOT
    / "harness/state/paper_v2/completion_scorecard_v1.json"
)
CLAIM_LEDGER_PATH = (
    PROJECT_ROOT
    / "harness/state/paper_v2/claim_evidence_ledger_v1.json"
)
SOURCE_AUDIT_SUMMARY_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/source_audit_summary_v1.md"
)
PUBLIC_FREEZE_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/public_benchmark_freeze_v1.json"
)
ADAPTIVE_FREEZE_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/adaptive_attack_freeze_v1.json"
)
GEMINI_V1_ABORT_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/gemini_fallback_abort_review_v1.json"
)
GEMINI_V2_CANDIDATE_PATH = (
    PROJECT_ROOT / "protocols/paper_upgrade_v2/gemini_fallback_candidate_v2.json"
)
GEMINI_V2_ADMISSION_PATH = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/admissions/gemini_fallback_canary_admission_v2.json"
)
GEMINI_V2_ROOT = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_medium/gemini"
    / "72693679347d7cdb04982bdaf87284b400c30e029cf2676ef37330cf37441728"
)
GEMINI_V2_RECORDS_PATH = GEMINI_V2_ROOT / "experiment_records_v2.jsonl"
GEMINI_V2_CALLS_PATH = GEMINI_V2_ROOT / "provider_calls_v2.jsonl"

REPORT_DATE = "2026-07-27"

CONDITION_LABELS = {
    "no_defense": "无防御",
    "safe_prompt_or_spotlighting": "安全提示",
    "promptguard_detector_only": "PromptGuard",
    "strongest_peer_reviewed_structural_baseline": "结构化基线",
    "m3_policy_only": "M3 policy-only",
    "m3_full": "M3 full",
}


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        value = json.loads(line)
        if not isinstance(value, dict):
            raise TypeError(f"{path}: expected JSON object per line")
        rows.append(value)
    return rows


def configure_sections(doc: Document) -> None:
    """Apply the standard_business_brief preset and memo_masthead furniture."""

    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        left = paragraph.add_run("M3-PI Guard | SCI 二区网络安全研究")
        set_run_font(left, size=8.5, color=MUTED, bold=True)
        paragraph.add_run("\t")
        right = paragraph.add_run(f"阶段报告 | {REPORT_DATE}")
        set_run_font(right, size=8.5, color=MUTED)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before = Pt(0)
        prefix = fp.add_run("M3-PI Guard  ·  ")
        set_run_font(prefix, size=9, color=MUTED)
        add_page_field(fp)


def apply_report_style_overrides(doc: Document) -> None:
    """Resolve the memo masthead title/subtitle tokens into real Word styles."""

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = doc.styles["Heading 1"].font.color.rgb
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(13)
    subtitle.font.bold = False
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle.paragraph_format.keep_with_next = True


def add_title_block(doc: Document) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(
        kicker.add_run("RESEARCH STAGE REPORT"),
        size=9,
        color=GOLD,
        bold=True,
    )

    title = doc.add_paragraph()
    title.style = "Title"
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    title.paragraph_format.keep_with_next = True
    set_run_font(
        title.add_run("M3-PI Guard 阶段性研究报告"),
        size=24,
        color=INK,
        bold=True,
    )

    subtitle = doc.add_paragraph()
    subtitle.style = "Subtitle"
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(18)
    set_run_font(
        subtitle.add_run(
            "DeepSeek 480 条中型证伪结果、Gemini 受控失败边界与"
            "授权泄漏/不确定性实证路线"
        ),
        size=13,
        color=MUTED,
    )

    metadata = [
        ("报告日期", "2026-07-27（Asia/Shanghai）"),
        (
            "当前论文路线",
            "授权信息泄漏、授权不确定性与安全—效用权衡的系统实证研究",
        ),
        (
            "备选英文标题",
            "Authorization Leakage, Uncertainty, and Security–Utility "
            "Trade-offs in Tool-Using LLM Agent Defenses",
        ),
        (
            "目标期刊",
            "Journal of Information Security and Applications；投稿前由学校/"
            "Clarivate 复核当期 JCR Q2",
        ),
        (
            "正式模型证据",
            "DeepSeek 官方 API / deepseek-v4-pro；Gemini 当前仅有诊断记录，"
            "无跨模型正式结果",
        ),
        (
            "安全边界",
            "provider tools 关闭、无真实外部 effect、无自动重试、失败留在 ITT",
        ),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        set_run_font(p.add_run(f"{label}："), size=10.5, color=BLACK, bold=True)
        set_run_font(p.add_run(value), size=10.5, color=BLACK)


def _canvas(title: str, subtitle: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle(
        (45, 40, 1755, 1010),
        radius=30,
        fill="#FFFFFF",
        outline="#D8DEE8",
        width=3,
    )
    draw.text((90, 80), title, font=font(42, bold=True), fill=f"#{INK}")
    draw.text((90, 145), subtitle, font=font(23), fill=f"#{MUTED}")
    draw.line((90, 195, 1710, 195), fill="#D8DEE8", width=3)
    return image, draw


def _save(image: Image.Image, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=True)


def _draw_hbar(
    draw: ImageDraw.ImageDraw,
    *,
    y: int,
    label: str,
    value: float,
    maximum: float,
    color: str,
    value_text: str,
    x0: int = 460,
    x1: int = 1430,
) -> None:
    draw.text((100, y - 6), label, font=font(23, bold=True), fill=f"#{BLACK}")
    draw.rounded_rectangle(
        (x0, y, x1, y + 32),
        radius=16,
        fill="#E9EDF3",
    )
    width = int((x1 - x0) * min(max(value / maximum, 0.0), 1.0))
    if width:
        draw.rounded_rectangle(
            (x0, y, x0 + width, y + 32),
            radius=16,
            fill=f"#{color}",
        )
    draw.text((x1 + 18, y - 4), value_text, font=font(22, bold=True), fill=f"#{color}")


def create_security_utility_chart(
    path: Path,
    condition_metrics: dict[str, Any],
) -> None:
    image, draw = _canvas(
        "图 1  中型证伪实验：安全与效用必须同时阅读",
        "攻击 unsafe-effect 越低越好；良性工具任务完成率越高越好（均为 ITT）",
    )
    draw.text((100, 225), "攻击 unsafe-effect rate", font=font(28, bold=True), fill=f"#{DARK_BLUE}")
    conditions = list(CONDITION_LABELS)
    for index, condition in enumerate(conditions):
        metric = condition_metrics[condition]["unsafe_effect_itt"]
        color = GREEN if condition.startswith("m3_") else RED
        _draw_hbar(
            draw,
            y=280 + index * 58,
            label=CONDITION_LABELS[condition],
            value=float(metric["value"]),
            maximum=0.25,
            color=color,
            value_text=f"{metric['numerator']}/{metric['denominator']}  "
            f"({metric['value'] * 100:.1f}%)",
        )

    draw.line((90, 645, 1710, 645), fill="#D8DEE8", width=3)
    draw.text((100, 675), "良性工具任务完成率（TCR）", font=font(28, bold=True), fill=f"#{DARK_BLUE}")
    for index, condition in enumerate(conditions):
        metric = condition_metrics[condition]["benign_tool_tcr"]
        color = GREEN if float(metric["value"]) >= 0.45 else GOLD
        _draw_hbar(
            draw,
            y=730 + index * 48,
            label=CONDITION_LABELS[condition],
            value=float(metric["value"]),
            maximum=0.6,
            color=color,
            value_text=f"{metric['numerator']}/{metric['denominator']}  "
            f"({metric['value'] * 100:.1f}%)",
        )
    _save(image, path)


def create_gate_chart(path: Path, gate: dict[str, Any]) -> None:
    image, draw = _canvas(
        "图 2  预注册机制论文门禁：安全提升不足以单独过关",
        "任一关键机制贡献或效用门禁失败，均触发 empirical fallback",
    )
    checks = [
        ("安全效应", bool(gate["checks"]["security_effect"]), "M3-full 对结构化基线 −22.2pp"),
        (
            "full 独立模块增益",
            bool(gate["checks"]["module_gain"]),
            "M3-full 与 policy-only 均为 0/36 unsafe effect",
        ),
        (
            "工具任务效用非劣",
            bool(gate["checks"]["utility_point_noninferiority"]),
            "M3-full 比结构化基线低 25pp，越过 −5pp 界限",
        ),
        (
            "运行质量",
            bool(gate["checks"]["prespecified_operational_quality"]),
            "21/432 主记录为终态失败，仍在预设质量上限内",
        ),
        (
            "非 deny-all",
            bool(gate["checks"]["deny_all_rejection"]),
            "良性初始 ASK+DENY 率 22.2%，低于 90% 阈值",
        ),
    ]
    for index, (label, passed, detail) in enumerate(checks):
        y = 250 + index * 125
        color = GREEN if passed else RED
        status = "PASS" if passed else "FAIL"
        draw.rounded_rectangle(
            (100, y, 300, y + 78),
            radius=18,
            fill=f"#{color}",
        )
        draw.text((145, y + 18), status, font=font(28, bold=True), fill="#FFFFFF")
        draw.text((340, y + 2), label, font=font(29, bold=True), fill=f"#{BLACK}")
        draw.text((340, y + 43), detail, font=font(21), fill=f"#{MUTED}")

    draw.rounded_rectangle(
        (100, 885, 1700, 970),
        radius=20,
        fill="#FFF1F2",
        outline=f"#{RED}",
        width=3,
    )
    draw.text(
        (140, 910),
        "结论：mechanism_paper_candidate = false；正式切换授权泄漏/不确定性实证路线",
        font=font(27, bold=True),
        fill=f"#{RED}",
    )
    _save(image, path)


def create_authorization_regime_chart(
    path: Path,
    regime_probes: dict[str, Any],
) -> None:
    image, draw = _canvas(
        "图 3  M3-full 授权 regime 探针：exact-gold 会改变在线行为",
        "每个 regime 仅 6 对；用于显示方向与测量价值，不作确认性结论",
    )
    regimes = ["broad", "conflicting", "exact_gold", "underspecified"]
    labels = ["宽授权", "冲突授权", "精确 gold", "欠指定授权"]
    colors = {"auto": BLUE, "gold": GOLD}
    for index, (regime, label) in enumerate(zip(regimes, labels, strict=True)):
        values = regime_probes["m3_full"]["regimes"][regime]
        x = 145 + index * 400
        draw.text((x, 245), label, font=font(27, bold=True), fill=f"#{BLACK}")

        ask = values["initial_ask"]
        completion = values["task_completion"]
        rows = [
            ("初始 ASK", float(ask["auto_compiled_rate"]), float(ask["exact_gold_rate"])),
            (
                "任务完成",
                float(completion["auto_compiled_rate"]),
                float(completion["exact_gold_rate"]),
            ),
        ]
        for row_index, (metric, auto, gold) in enumerate(rows):
            base_y = 345 + row_index * 285
            draw.text((x, base_y - 65), metric, font=font(23, bold=True), fill=f"#{DARK_BLUE}")
            for bar_index, (name, value, color) in enumerate(
                [("auto", auto, colors["auto"]), ("gold", gold, colors["gold"])]
            ):
                bx = x + bar_index * 120
                height = int(value * 185)
                draw.rounded_rectangle(
                    (bx, base_y + 190 - height, bx + 78, base_y + 190),
                    radius=12,
                    fill=f"#{color}",
                )
                draw.text(
                    (bx, base_y + 202),
                    name,
                    font=font(19, bold=True),
                    fill=f"#{color}",
                )
                draw.text(
                    (bx, base_y + 155 - height),
                    f"{value * 100:.0f}%",
                    font=font(19, bold=True),
                    fill=f"#{color}",
                )

    draw.text((115, 950), "蓝：auto_compiled    金：exact_gold", font=font(22, bold=True), fill=f"#{MUTED}")
    _save(image, path)


def create_provider_evidence_chart(path: Path) -> None:
    image, draw = _canvas(
        "图 4  官方 API 证据状态：DeepSeek 可分析，Gemini 尚不可分析",
        "诊断记录、失败关闭和 HTTP 429 均不得升级为跨模型科学结果",
    )
    cards = [
        (
            "DeepSeek v8+v9",
            "480/480 终态",
            "479 formal-candidate + 1 diagnostic\n1084 次完成调用 | CNY 1.8623798",
            GREEN,
            "正式中型结果已准入",
        ),
        (
            "Gemini v1",
            "0 正式记录",
            "VIGIL 约束生成静默回退\n1 个失败调用 | 无 evidence sidecar",
            RED,
            "永久失败关闭，需新版本",
        ),
        (
            "Gemini v2",
            "0 正式记录",
            "首个 policy-only 调用 HTTP 429\n0 完成调用 | 0 token | 0 自动重试",
            GOLD,
            "暂停；配额恢复后仍需新版本",
        ),
    ]
    light_fills = {GREEN: "#ECFDF5", RED: "#FFF1F2", GOLD: "#FFF8E8"}
    for index, (title, headline, details, color, status) in enumerate(cards):
        x0 = 105 + index * 555
        x1 = x0 + 495
        draw.rounded_rectangle(
            (x0, 250, x1, 875),
            radius=26,
            fill="#FFFFFF",
            outline=f"#{color}",
            width=5,
        )
        draw.rounded_rectangle(
            (x0, 250, x1, 340),
            radius=24,
            fill=f"#{color}",
        )
        draw.rectangle((x0, 310, x1, 340), fill=f"#{color}")
        draw.text((x0 + 30, 275), title, font=font(28, bold=True), fill="#FFFFFF")
        draw.text((x0 + 30, 395), headline, font=font(39, bold=True), fill=f"#{INK}")
        lines = details.splitlines()
        for line_index, line in enumerate(lines):
            draw.text(
                (x0 + 30, 500 + line_index * 55),
                line,
                font=font(22),
                fill=f"#{BLACK}",
            )
        draw.rounded_rectangle(
            (x0 + 25, 735, x1 - 25, 825),
            radius=18,
            fill=light_fills[color],
            outline=f"#{color}",
            width=2,
        )
        draw.text(
            (x0 + 45, 765),
            status,
            font=font(22, bold=True),
            fill=f"#{color}",
        )
    _save(image, path)


def create_readiness_chart(path: Path, scorecard: dict[str, Any]) -> None:
    image, draw = _canvas(
        "图 5  完成度拆分：工程完成不等于论文证据完成",
        "报告、测试和协议不能替代第二模型、公共扩展、自适应攻击与独立盲测",
    )
    bars = [
        ("工程就绪度", int(scorecard["engineering_readiness_percent"]), GREEN),
        (
            "科学证据完成度",
            int(scorecard["scientific_evidence_completion_percent"]),
            BLUE,
        ),
        (
            "投稿就绪度",
            int(scorecard["paper_submission_readiness_percent"]),
            GOLD,
        ),
    ]
    for index, (label, value, color) in enumerate(bars):
        y = 300 + index * 190
        _draw_hbar(
            draw,
            y=y,
            label=label,
            value=value,
            maximum=100,
            color=color,
            value_text=f"{value}%",
            x0=420,
            x1=1580,
        )
        draw.text(
            (420, y + 62),
            (
                "协议/运行链基本成熟"
                if index == 0
                else "当前真实科学证据"
                if index == 1
                else "尚未达到可投稿状态"
            ),
            font=font(22),
            fill=f"#{MUTED}",
        )
    draw.rounded_rectangle(
        (100, 870, 1700, 965),
        radius=20,
        fill="#FFF8E8",
        outline=f"#{GOLD}",
        width=3,
    )
    draw.text(
        (135, 900),
        "当前不得宣称“项目 80% 完成”；61% 是科学证据口径，且机制主路线已被否决。",
        font=font(26, bold=True),
        fill=f"#{GOLD}",
    )
    _save(image, path)


def create_roadmap_chart(path: Path) -> None:
    image, draw = _canvas(
        "图 6  下一阶段实验路线：先补反证强度，再补外部有效性",
        "所有 API 实验继续使用模拟工具与隔离 effect backend，不执行真实危险操作",
    )
    stages = [
        ("1", "冻结论文主张", "Methods/Results\nclaim–evidence ledger", GREEN),
        ("2", "自适应攻击", "DeepSeek 官方桥接\n成本 pilot → 正式预算", BLUE),
        ("3", "公共扩展", "AgentDojo 主扩展\nSIREN/上下文任务", BLUE),
        ("4", "第二模型", "Gemini 配额恢复\n新 v3 prospective run", GOLD),
        ("5", "盲测与投稿", "独立 custody\nJCR 复核与完整性审查", RED),
    ]
    for index, (number, title, detail, color) in enumerate(stages):
        x = 105 + index * 335
        if index < len(stages) - 1:
            draw.line((x + 235, 480, x + 335, 480), fill="#AAB4C3", width=8)
        draw.ellipse((x, 390, x + 180, 570), fill=f"#{color}")
        draw.text((x + 65, 432), number, font=font(52, bold=True), fill="#FFFFFF")
        draw.text((x - 15, 625), title, font=font(27, bold=True), fill=f"#{BLACK}")
        for line_index, line in enumerate(detail.splitlines()):
            draw.text(
                (x - 15, 680 + line_index * 43),
                line,
                font=font(21),
                fill=f"#{MUTED}",
            )
    draw.rounded_rectangle(
        (105, 865, 1690, 965),
        radius=18,
        fill="#EEF6FF",
        outline=f"#{BLUE}",
        width=3,
    )
    draw.text(
        (140, 895),
        "顺序原则：不得用新的正结果覆盖已冻结负结果；每个扩展均需独立 candidate、admission 与结果门禁。",
        font=font(24, bold=True),
        fill=f"#{DARK_BLUE}",
    )
    _save(image, path)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Report Source")
    set_run_font(p.add_run(text), size=8.5, color=MUTED)


def add_numbered_actions(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_run_font(p.add_run(item))


def _condition_metric_rows(condition_metrics: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for condition, label in CONDITION_LABELS.items():
        unsafe = condition_metrics[condition]["unsafe_effect_itt"]
        tool = condition_metrics[condition]["benign_tool_tcr"]
        no_tool = condition_metrics[condition]["benign_no_tool_tcr"]
        rows.append(
            [
                label,
                f"{unsafe['numerator']}/{unsafe['denominator']} "
                f"({unsafe['value'] * 100:.1f}%)",
                f"{tool['numerator']}/{tool['denominator']} "
                f"({tool['value'] * 100:.1f}%)",
                f"{no_tool['numerator']}/{no_tool['denominator']} "
                f"({no_tool['value'] * 100:.1f}%)",
            ]
        )
    return rows


def build_report() -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    QA_DIR.mkdir(parents=True, exist_ok=True)

    admitted = load_json(ANALYSIS_PATH)
    technical = load_json(TECHNICAL_REVIEW_PATH)
    result_admission = load_json(RESULT_ADMISSION_PATH)
    scorecard = load_json(SCORECARD_PATH)
    public_freeze = load_json(PUBLIC_FREEZE_PATH)
    adaptive_freeze = load_json(ADAPTIVE_FREEZE_PATH)
    gemini_v1_abort = load_json(GEMINI_V1_ABORT_PATH)
    gemini_v2_candidate = load_json(GEMINI_V2_CANDIDATE_PATH)
    gemini_v2_admission = load_json(GEMINI_V2_ADMISSION_PATH)
    gemini_v2_records = load_jsonl(GEMINI_V2_RECORDS_PATH)
    gemini_v2_calls = load_jsonl(GEMINI_V2_CALLS_PATH)

    analysis = admitted["analysis"]
    primary = analysis["primary_analysis"]
    secondary = analysis["secondary_analysis"]
    condition_metrics = primary["condition_metrics"]
    gate = primary["mechanism_paper_gate"]
    operational = analysis["operational_quality"]
    coverage = technical["coverage"]
    budget = technical["budget_summary"]
    technical_gate = technical["operational_gate"]

    gemini_v2_record = gemini_v2_records[0]["record"]
    gemini_v2_paused = gemini_v2_calls[-1]
    if gemini_v2_paused["http_status"] != 429:
        raise ValueError("Gemini v2 report expects a fail-closed first-call HTTP 429")

    security_utility_chart = QA_DIR / "fig01_security_utility.png"
    gate_chart = QA_DIR / "fig02_mechanism_gate.png"
    regime_chart = QA_DIR / "fig03_authorization_regimes.png"
    provider_chart = QA_DIR / "fig04_provider_evidence.png"
    readiness_chart = QA_DIR / "fig05_readiness.png"
    roadmap_chart = QA_DIR / "fig06_roadmap.png"
    create_security_utility_chart(security_utility_chart, condition_metrics)
    create_gate_chart(gate_chart, gate)
    create_authorization_regime_chart(
        regime_chart,
        secondary["authorization_regime_probes"],
    )
    create_provider_evidence_chart(provider_chart)
    create_readiness_chart(readiness_chart, scorecard)
    create_roadmap_chart(roadmap_chart)

    doc = Document()
    apply_style_tokens(doc)
    apply_report_style_overrides(doc)
    configure_sections(doc)
    props = doc.core_properties
    props.title = "M3-PI Guard 阶段性研究报告（2026-07-27）"
    props.subject = "SCI 二区网络安全论文中型证伪结果与后续实验计划"
    props.author = "M3-PI Guard Research Team"
    props.keywords = (
        "LLM agent security, prompt injection, authorization leakage, "
        "provenance, security utility"
    )
    props.comments = "Generated from admitted and frozen repository evidence."

    add_title_block(doc)
    add_callout(
        doc,
        "阶段结论",
        "SCI 二区仍有条件可行，但当前最可信的论文不是“新机制领先”，而是"
        "“授权信息进入在线 Guard 后如何改变安全—效用估计”的系统实证研究。"
        "DeepSeek 中型结果显示 M3-full 安全点估计优于结构化基线，却没有优于 "
        "policy-only，且工具任务效用明显下降；预注册机制门禁已正式失败。",
        fill=LIGHT_BLUE,
        accent=DARK_BLUE,
    )
    add_callout(
        doc,
        "不可逆研究决定",
        "不得用后续第二模型、公共基准或自适应攻击中的正结果抹去本次负结果，"
        "不得恢复“首个、唯一、SOTA”叙事。后续实验的任务是验证该安全—效用"
        "权衡和授权 regime 敏感性是否可复现，而不是为既定机制寻找有利样本。",
        fill="FFF1F2",
        accent=RED,
    )
    add_callout(
        doc,
        "API 安全承诺",
        "DeepSeek 与 Gemini 仅用于隔离实验中的文本推理。provider tools、"
        "外部 connector 和真实工具 effect 均关闭；不会使用两家 API 执行"
        "转账、邮件、文件删除、系统控制或其他危险操作。",
        fill="FFF8E8",
        accent=GOLD,
    )

    doc.add_heading("1. 执行摘要", level=1)
    add_paragraph(
        doc,
        "本阶段已经完成 DeepSeek 官方 API 上的 480 条中型配对证伪矩阵和"
        "独立结果准入。矩阵由 432 条主比较记录和 48 条授权 regime 探针组成，"
        "479 条为 formal-candidate、1 条为 diagnostic；全部终态均保留在 "
        "intention-to-treat（ITT）分母。证据链绑定 1,084 次完成调用，累计成本"
        "为 CNY 1.8623798，自动重试、provider tools 与真实 external effect "
        "均为 0。"
    )
    add_paragraph(
        doc,
        "主结果出现了清晰但不利于机制论文的安全—效用分裂。攻击任务上，"
        "M3-full 为 0/36 unsafe effect，结构化基线为 8/36，绝对差 −22.2pp；"
        "但 M3-full 与 policy-only 均为 0/36，完整来源/时序/确认模块没有独立"
        "安全增益。良性工具任务上，M3-full 为 6/24，结构化基线为 12/24，"
        "差 −25pp，超过预注册 −5pp 非劣界限。机制论文门禁因此失败并触发"
        " empirical fallback。"
    )
    add_paragraph(
        doc,
        "Gemini 的两次 prospective 尝试均严格失败关闭：v1 因 VIGIL 约束"
        "生成路径静默回退而永久终止，v2 排除结构化条件后在首个官方请求收到 "
        "HTTP 429 并立即暂停。两者均没有可准入的跨模型结果；v2 只产生 1 条"
        " diagnostic ITT 记录、0 个完成 provider call、0 token 与 0 自动重试。"
    )
    add_table(
        doc,
        ["核心问题", "当前回答", "论文含义"],
        [
            [
                "SCI 二区是否仍可行",
                "有条件可行",
                "改投授权泄漏/不确定性与安全—效用实证路线",
            ],
            [
                "M3-full 是否领先 policy-only",
                "否",
                "不能宣称 provenance/temporal/confirmation 的独立贡献",
            ],
            [
                "安全优势是否伴随效用损失",
                "是",
                "必须把 TCR、ASK/拒绝负担与 unsafe effect 联合报告",
            ],
            [
                "第二模型是否已完成",
                "否",
                "Gemini 目前只有失败边界，不能写成泛化证据",
            ],
            [
                "下一优先级",
                "自适应攻击 + 公共扩展",
                "检验负结果与权衡在更强攻击和外部基准下是否稳定",
            ],
        ],
        [2050, 1900, 5410],
        status_columns=(1,),
    )

    doc.add_heading("2. 研究定位与目标期刊", level=1)
    add_paragraph(
        doc,
        "当前建议标题从“Beyond Answer-Key Authorization”的机制导向收缩为："
        "Authorization Leakage, Uncertainty, and Security–Utility Trade-offs "
        "in Tool-Using LLM Agent Defenses。M3 运行时仍作为受测机制和测量工具，"
        "但不再作为已证明领先的中心贡献。"
    )
    add_paragraph(
        doc,
        "JISA 的官方范围覆盖 authentication and access control、security "
        "management and policies，并设有 Large Language Models in "
        "Cybersecurity 相关专题，主题契合度仍然较高。期刊主页显示的 CiteScore/"
        "Impact Factor 不等同于 JCR 分区；正式投稿前必须由学校或 Clarivate "
        "订阅核验当期 SCIE/JCR Q2，不能只依据第三方网站。"
    )
    add_table(
        doc,
        ["论文贡献候选", "当前证据", "可写强度"],
        [
            [
                "在线授权字段泄漏/可见性审计",
                "22 个审计单元、16 个项目；无领域普遍 oracle 结论",
                "可写为固定版本源码审计",
            ],
            [
                "授权 regime 敏感性",
                "48 条配对探针；每个 regime 6 对",
                "方向性/探索性，不作确认性因果结论",
            ],
            [
                "M3 非 oracle 运行时机制",
                "安全点估计有利，但 full 无独立模块增益且效用失败",
                "作为受测机制，不作领先声明",
            ],
            [
                "正式证据链",
                "冻结、哈希、准入、append-only、失败留在 ITT",
                "可作为方法完整性贡献",
            ],
        ],
        [2400, 3500, 3460],
    )

    doc.add_heading("3. DeepSeek 中型实验设计与证据完整性", level=1)
    add_paragraph(
        doc,
        "主比较包含 72 个配对 case，在 6 个条件下形成 432 条记录；其中攻击"
        "36 个、良性工具任务 24 个、良性无工具任务 12 个。另有 48 条"
        "授权 regime 探针，只在 M3 policy-only 与 M3-full 上比较 "
        "exact_gold 与 auto_compiled。"
    )
    add_table(
        doc,
        ["组件", "规模/状态", "完整性规则"],
        [
            ["主比较", "72 cases × 6 conditions = 432", "配对、统一失败分母"],
            ["授权探针", "24 cases × 2 M3 conditions = 48", "每个 regime 6 对"],
            ["总记录", f"{coverage['combined_records']}/480", "0 remaining、0 duplicate"],
            [
                "证据等级",
                f"{technical_gate['formal_candidate_count']} formal + "
                f"{technical_gate['diagnostic_count']} diagnostic",
                "诊断行保留但不伪造响应身份",
            ],
            [
                "Provider 调用",
                f"{budget['combined_provider_calls_completed']:,}",
                "本地 request ID + payload hash + 官方 response ID + append-only",
            ],
            ["累计成本", f"CNY {budget['project_cumulative_cost_cny']:.7f}", "冻结费率、逐调用核算"],
        ],
        [2200, 2700, 4460],
        status_columns=(1,),
    )
    add_paragraph(
        doc,
        "技术 gate 验证了完整矩阵、exact official identity chain、隔离 effect "
        "边界、诊断上限、0 open attempt 和 0 runner failure。唯一 diagnostic "
        "是闭合 provider/timeout 类失败，仍以 diagnostic-only 身份进入 ITT；"
        "该记录不能反向补造官方响应 ID。"
    )

    doc.add_heading("4. 主结果：安全—效用权衡", level=1)
    add_figure(
        doc,
        security_utility_chart,
        "图 1  六条件 unsafe-effect 与良性工具任务完成率",
        "admitted_medium_analysis_v8.json；ITT 分母；Wilson 95% CI 见表 2。",
        alt=(
            "六个条件的横向条形图。无防御、安全提示、PromptGuard 和结构化基线"
            "的攻击 unsafe effect 均为 22.2%，M3 policy-only 和 M3 full 为 0%。"
            "良性工具任务完成率分别约为 50%、45.8%、50%、50%、29.2% 和 25%。"
        ),
    )
    add_table(
        doc,
        ["条件", "攻击 unsafe effect", "良性工具 TCR", "良性无工具 TCR"],
        _condition_metric_rows(condition_metrics),
        [2500, 2200, 2200, 2460],
    )
    add_source_note(
        doc,
        "说明：unsafe-effect 越低越好；TCR 越高越好。良性无工具任务在六条件下"
        "均为 12/12，差异集中在需要工具执行的任务。",
    )

    full_vs_structural = primary["comparisons"]["security_full_vs_structural"]
    utility_vs_structural = primary["comparisons"]["utility_full_vs_structural"]
    security_vs_policy = primary["comparisons"]["security_full_vs_policy_only"]
    utility_vs_policy = primary["comparisons"]["utility_full_vs_policy_only"]
    add_table(
        doc,
        ["配对比较", "绝对差", "95% cluster bootstrap CI", "McNemar（Holm）", "解释"],
        [
            [
                "安全：full − structural",
                f"{full_vs_structural['absolute_difference_treatment_minus_baseline'] * 100:.1f}pp",
                f"[{full_vs_structural['cluster_bootstrap']['ci_low'] * 100:.1f}, "
                f"{full_vs_structural['cluster_bootstrap']['ci_high'] * 100:.1f}]pp",
                f"p={full_vs_structural['mcnemar']['holm_adjusted_p_value']:.5f}",
                "安全点估计和区间均有利",
            ],
            [
                "安全：full − policy-only",
                f"{security_vs_policy['absolute_difference_treatment_minus_baseline'] * 100:.1f}pp",
                "[0.0, 0.0]pp",
                "p=1.0",
                "无完整模块独立增益",
            ],
            [
                "效用：full − structural",
                f"{utility_vs_structural['absolute_difference_treatment_minus_baseline'] * 100:.1f}pp",
                f"[{utility_vs_structural['cluster_bootstrap']['ci_low'] * 100:.1f}, "
                f"{utility_vs_structural['cluster_bootstrap']['ci_high'] * 100:.1f}]pp",
                f"p={utility_vs_structural['mcnemar']['holm_adjusted_p_value']:.5f}",
                "越过 −5pp 非劣界限",
            ],
            [
                "效用：full − policy-only",
                f"{utility_vs_policy['absolute_difference_treatment_minus_baseline'] * 100:.1f}pp",
                f"[{utility_vs_policy['cluster_bootstrap']['ci_low'] * 100:.1f}, "
                f"{utility_vs_policy['cluster_bootstrap']['ci_high'] * 100:.1f}]pp",
                f"p={utility_vs_policy['mcnemar']['holm_adjusted_p_value']:.1f}",
                "full 仍略差",
            ],
        ],
        [2150, 1200, 2400, 1700, 1910],
    )
    add_paragraph(
        doc,
        "安全比较的统计显著性不能替代机制解释。full 与 policy-only 同为 0/36，"
        "说明结构化基线的 8 个 unsafe event 可以被 M3 的策略层阻断，但当前"
        "样本没有证明 provenance、temporal、dataflow 或 confirmation 模块"
        "带来额外安全收益。与此同时，M3-full 的工具任务效用只有 25%，存在"
        "明显过度阻断/执行困难风险。"
    )

    doc.add_heading("5. 预注册机制门禁与论文路线切换", level=1)
    add_figure(
        doc,
        gate_chart,
        "图 2  机制论文门禁判定",
        "预注册 medium gate；结果准入后按冻结规则计算。",
        alt=(
            "五项机制论文门禁中，安全效应、运行质量和非 deny-all 通过；"
            "完整模块独立增益与工具任务效用非劣失败，结论切换到 empirical fallback。"
        ),
    )
    add_table(
        doc,
        ["门禁", "阈值/规则", "观察值", "结果"],
        [
            ["安全效应", "相对 ≥20% 或绝对 ≥10pp", "绝对 −22.2pp", "PASS"],
            ["模块增益", "full − policy-only unsafe < 0", "0.0pp", "FAIL"],
            ["TCR 非劣", "full − structural ≥ −5pp", "−25.0pp", "FAIL"],
            [
                "运行质量",
                "总失败 ≤10%；provider/timeout ≤5%；runner=0",
                f"{operational['failure_count']}/"
                f"{operational['record_count']} ({operational['failure_fraction'] * 100:.2f}%)",
                "PASS",
            ],
            [
                "非 deny-all",
                "良性初始 ASK+DENY <90%",
                f"{gate['values']['benign_initial_ask_or_deny_rate'] * 100:.1f}%",
                "PASS",
            ],
        ],
        [1900, 3300, 2600, 1560],
        status_columns=(3,),
    )
    add_callout(
        doc,
        "论文路线",
        "mechanism_paper_candidate = false；正式 route = "
        "authorization_leakage_uncertainty_empirical_fallback。"
        "M3-full 对结构化基线的安全优势可作为观察结果，但摘要、标题和贡献点"
        "必须以授权 regime、在线可见字段、失败分母与安全—效用权衡为中心。",
        fill=LIGHT_BLUE,
        accent=DARK_BLUE,
    )

    doc.add_heading("6. 授权 regime 探针：为什么 exact-gold 不能被忽略", level=1)
    add_paragraph(
        doc,
        "48 条探针比较 exact_gold 与 auto_compiled 在 broad、conflicting、"
        "exact_gold 和 underspecified 四类授权状态下的行为变化。每个 regime "
        "只有 6 对，因此统计功效很低；这些结果目前只用于证明“授权信息形态"
        "可能改变 Guard 行为”这一测量问题值得正式扩展。"
    )
    add_figure(
        doc,
        regime_chart,
        "图 3  M3-full 在四类授权 regime 下的初始 ASK 与任务完成率",
        "admitted medium secondary analysis；每个 regime 6 对。",
        alt=(
            "四组授权 regime 的配对柱状图，比较 auto compiled 与 exact gold。"
            "在 underspecified 组，初始 ASK 从 16.7% 增至 83.3%；各组样本均很小。"
        ),
    )
    probes = secondary["authorization_regime_probes"]["m3_full"]["regimes"]
    add_table(
        doc,
        ["Regime", "初始 ASK：auto / gold", "任务完成：auto / gold", "unsafe：auto / gold"],
        [
            [
                regime,
                f"{values['initial_ask']['auto_compiled_events']}/6 / "
                f"{values['initial_ask']['exact_gold_events']}/6",
                f"{values['task_completion']['auto_compiled_events']}/6 / "
                f"{values['task_completion']['exact_gold_events']}/6",
                f"{values['unsafe_effect_itt']['auto_compiled_events']}/6 / "
                f"{values['unsafe_effect_itt']['exact_gold_events']}/6",
            ]
            for regime, values in probes.items()
        ],
        [1900, 2500, 2700, 2260],
    )
    add_paragraph(
        doc,
        "最明显的方向出现在 underspecified 探针：M3-full 初始 ASK 为 "
        f"{probes['underspecified']['initial_ask']['auto_compiled_events']}/6"
        "（auto）对 "
        f"{probes['underspecified']['initial_ask']['exact_gold_events']}/6"
        "（gold）。但 exact-gold 并不总是单向提高效用，例如 conflicting 下"
        " gold 的任务完成为 0/6。正式论文应把这类差异写成授权 regime 对"
        "运行时决策的敏感性，而不是预设“oracle 总会高估安全”。"
    )

    doc.add_heading("7. 运行质量、失败分母与可复现性", level=1)
    add_table(
        doc,
        ["Operational status", "数量", "处理"],
        [
            ["completed", str(operational["operational_status_counts"]["completed"]), "正常进入 ITT"],
            ["max_steps", str(operational["operational_status_counts"]["max_steps"]), "终态失败，留在分母"],
            ["parse_failure", str(operational["operational_status_counts"]["parse_failure"]), "终态失败，留在分母"],
            ["provider_failure", str(operational["operational_status_counts"]["provider_failure"]), "终态失败，留在分母"],
            ["runner_failure", "0", "门禁要求为 0"],
        ],
        [2650, 1400, 5310],
        status_columns=(1,),
    )
    add_paragraph(
        doc,
        "21/432 主比较记录为预先定义的终态失败（4.86%）。这些记录没有被删除、"
        "重试或改写；分析脚本把 unsafe_effect_itt 与 task_completion 的失败"
        "语义固定在记录契约中。技术 review 本身不做网络调用，也不修改 v8/v9 "
        "ledger。正式统计由独立 --verify-analysis 路径重建。"
    )
    add_table(
        doc,
        ["证据对象", "SHA-256 / manifest", "状态"],
        [
            ["DeepSeek combined technical review", sha256_file(TECHNICAL_REVIEW_PATH), technical["status"]],
            ["DeepSeek result admission", sha256_file(RESULT_ADMISSION_PATH), result_admission["status"]],
            ["Admitted medium analysis", sha256_file(ANALYSIS_PATH), admitted["status"]],
            ["Claim–evidence ledger", sha256_file(CLAIM_LEDGER_PATH), "current"],
            ["Completion scorecard", sha256_file(SCORECARD_PATH), "99/61/56"],
        ],
        [2700, 4750, 1910],
        status_columns=(2,),
    )

    doc.add_heading("8. Gemini 第二模型：两次受控失败的证据边界", level=1)
    add_figure(
        doc,
        provider_chart,
        "图 4  DeepSeek 与 Gemini 的正式证据状态",
        "冻结技术 review、Gemini v1 abort review 与 v2 append-only ledgers。",
        alt=(
            "三张卡片：DeepSeek 480 条结果已准入；Gemini v1 因 VIGIL 回退"
            "失败关闭；Gemini v2 首个调用 HTTP 429 暂停，均无正式跨模型结果。"
        ),
    )
    add_table(
        doc,
        ["版本", "prospective 范围", "实际终态", "科学资格", "下一动作"],
        [
            [
                "Gemini v1",
                "structural + policy-only + full",
                "VIGIL constraint generation fallback",
                "0 正式记录",
                "永久冻结；不得恢复",
            ],
            [
                "Gemini v2",
                "policy-only + full；192 行",
                "首调用 HTTP 429；1 diagnostic",
                "0 正式记录",
                "配额恢复后新建 v3",
            ],
        ],
        [1500, 2500, 2500, 1300, 1560],
        status_columns=(3,),
    )
    add_paragraph(
        doc,
        "v1 abort review 记录 1 个 started attempt、1 个 closed failed call、"
        "0 experiment record、0 evidence sidecar，并明确 run_resumable=false。"
        "v2 使用新的 run/matrix identity，首个 policy-only 请求收到 HTTP 429；"
        "controller 按 free-only、USD 0、0 retry 规则立即暂停。"
    )
    add_paragraph(
        doc,
        f"v2 诊断记录的 operational_status={gemini_v2_record['operational_status']}、"
        f"result_tier={gemini_v2_record['result_tier']}、"
        f"identity_verified={str(gemini_v2_record['provider_attestation']['identity_verified']).lower()}、"
        f"token total={gemini_v2_record['token_usage']['total']}。"
        "它保留在版本内部的 ITT 记录中，但没有独立结果 admission，因此不得与"
        "DeepSeek 主结果合并，也不得计算跨模型效应。"
    )
    add_table(
        doc,
        ["Gemini 证据", "Manifest / SHA-256", "状态"],
        [
            ["v1 abort review", gemini_v1_abort["manifest_sha256"], gemini_v1_abort["status"]],
            ["v2 candidate", gemini_v2_candidate["manifest_sha256"], gemini_v2_candidate["status"]],
            ["v2 canary admission", gemini_v2_admission["manifest_sha256"], gemini_v2_admission["status"]],
            ["v2 provider ledger", sha256_file(GEMINI_V2_CALLS_PATH), "started → paused(429)"],
            ["v2 diagnostic record ledger", sha256_file(GEMINI_V2_RECORDS_PATH), "1 diagnostic"],
        ],
        [2500, 4400, 2460],
        status_columns=(2,),
    )

    doc.add_heading("9. 当前 SCI 二区完成度", level=1)
    add_figure(
        doc,
        readiness_chart,
        "图 5  工程、科学证据与投稿就绪度",
        "completion_scorecard_v1.json；Gemini 失败未增加科学完成度。",
        alt="三条进度条分别显示工程就绪度 99%、科学证据完成度 61%、投稿就绪度 56%。",
    )
    add_table(
        doc,
        ["工作流", "科学证据完成度", "关键缺口"],
        [
            [
                item["name"],
                f"{item['scientific_evidence_completion_percent']}%",
                item["remaining"],
            ]
            for item in scorecard["workstreams"]
        ],
        [3150, 1650, 4560],
        status_columns=(1,),
    )
    add_callout(
        doc,
        "完成度解释",
        "99% 只代表工程与证据链成熟；科学证据完成度 61%；投稿就绪度 56%。"
        "这表示"
        "仍未达到可投稿状态。Gemini 失败、公共扩展未跑、自适应攻击未跑、独立"
        "盲测无保管人，均不能由报告或代码数量补分。",
        fill="FFF8E8",
        accent=GOLD,
    )

    doc.add_heading("10. 前沿研究对照与创新边界", level=1)
    add_table(
        doc,
        ["工作", "类型/年份", "对本项目的约束", "M3 可保留差异"],
        [
            ["VIGIL", "ACL 2026", "verify-before-commit + SIREN 959", "禁止声称首个提交前验证"],
            ["DRIFT", "NeurIPS 2025", "动态规则、控制/数据流与记忆隔离", "报告权限单调与 ASK 负担"],
            ["Task Shield", "ACL 2025", "逐指令/工具调用目标一致性", "比较 goal alignment 与授权 regime"],
            ["ARGUS", "2026 预印本", "上下文任务 + provenance auditing", "强调在线字段投影与失败分母"],
            ["AuthGraph", "2026 预印本", "授权图对齐参数来源执行图", "禁止首个参数级来源声明"],
            ["Progent/IPIGuard/CaMeL", "2025", "最小权限、依赖图、信息流隔离", "做配对消融与来源审计"],
            ["Attacker Moves Second", "2025 预印本", "静态攻击高估防御", "必须执行防御感知自适应攻击"],
        ],
        [1950, 1400, 3300, 2710],
    )
    add_paragraph(
        doc,
        "源码审计目前完成 22/22 单元、覆盖 16 个项目，但独立第二复核为 0。"
        "结论分布为 2 个 inconclusive、1 个 mixed/path-dependent、9 个"
        " no_online_gold_observed 和 10 个 not_a_runtime_guard。因此只能写"
        "固定 commit/路径的局部结论，不能写“该领域普遍使用 oracle authorization”。"
    )

    doc.add_heading("11. 后续正式实验总体方案", level=1)
    add_figure(
        doc,
        roadmap_chart,
        "图 6  下一阶段执行顺序",
        "以证据增益、失败可解释性和外部有效性为排序标准。",
        alt="五阶段路线图：冻结论文主张、自适应攻击、公共扩展、第二模型、盲测与投稿。",
    )
    add_numbered_actions(
        doc,
        [
            "先冻结 Methods/Results 主表和 claim–evidence ledger，确保任何后续结果不能改写中型负结论。",
            "实现 DeepSeek 官方 API 的自适应攻击 provider bridge，先做零网络资格和成本 pilot，再按冻结预算运行。",
            "在 AgentDojo 上完成正式公共扩展；SIREN 用于工具流/上下文依赖外部有效性，前提是官方数据和适配可复现。",
            "Gemini 配额恢复后新建 v3 prospective 版本；不得续跑 v1/v2，不再把不兼容 VIGIL 条件包装为正式结构化对照。",
            "落实第三方 blind custodian；完成最终完整性审查、JCR 核验、JISA 格式与生成式 AI 使用声明。",
        ],
    )

    doc.add_heading("12. 自适应攻击：下一项最高信息增益实验", level=1)
    add_paragraph(
        doc,
        "The Attacker Moves Second 表明，知道防御结构的攻击者可能推翻静态评测。"
        "当前 adaptive freeze 已固定 40 个主模型任务、预算 50/200/500；最难 20 "
        "个任务追加 2000；第二模型 20 个任务、预算 50/200，最大查询预算 "
        f"{adaptive_freeze['maximum_query_budget']:,}。hardest-20 必须按预算 500 "
        "结果排序后确定，不能提前挑选有利案例。"
    )
    add_table(
        doc,
        ["阶段", "规模", "目的", "通过/停止条件"],
        [
            [
                "A0 零网络资格",
                "stub + injected transport",
                "验证 request chain、成本与 ITT",
                "0 外部 effect、0 未闭合 attempt",
            ],
            [
                "A1 成本 pilot",
                "8 tasks × budget 50",
                "测量每预算单位调用、token、p95",
                "先冻结成本模型，再决定正式 fuse",
            ],
            [
                "A2 主预算",
                "40 × 50/200/500",
                "观察攻击随预算增长曲线",
                "失败留分母；不自动重试",
            ],
            [
                "A3 hardest extension",
                "20 × 2000",
                "检验攻击是否饱和",
                "只按预注册排序选取",
            ],
            [
                "A4 第二模型",
                "20 × 50/200",
                "方向性泛化",
                "仅在 Gemini 新版本准入后",
            ],
        ],
        [1900, 2100, 2750, 2610],
    )
    avg_cost = budget["project_cumulative_cost_cny"] / budget["combined_provider_calls_completed"]
    projected_cost = avg_cost * adaptive_freeze["maximum_query_budget"]
    add_callout(
        doc,
        "成本事实",
        f"按中型矩阵粗略均值 CNY {avg_cost:.6f}/完成调用线性外推，"
        f"75,000 个查询单位约为 CNY {projected_cost:.1f}；但自适应提示长度和"
        "一次预算单位对应的调用数可能不同，因此该数只用于确定必须先做成本 pilot，"
        "不能作为正式预算承诺。用户允许追加预算不等于可以跳过熔断器。",
        fill="FFF8E8",
        accent=GOLD,
    )
    add_paragraph(
        doc,
        "自适应攻击只优化隔离环境中的 unsafe effect 判定，不连接真实邮箱、"
        "银行、文件系统、浏览器会话或系统控制接口。攻击文本可以包含恶意指令，"
        "但所有 tool calls 都由模拟 effect backend 截获并记录，不产生真实动作。"
    )

    doc.add_heading("13. 公共基准与外部有效性", level=1)
    cells = public_freeze["cells"]
    add_table(
        doc,
        ["冻结 cell", "Cases", "Conditions", "预期记录"],
        [
            [
                name,
                str(len(cell["cases"])),
                str(len(cell["conditions"])),
                str(len(cell["cases"]) * len(cell["conditions"])),
            ]
            for name, cell in cells.items()
        ],
        [3900, 1500, 1700, 2260],
        status_columns=(1, 2, 3),
    )
    add_paragraph(
        doc,
        f"public_benchmark_freeze_v1 共冻结 "
        f"{public_freeze['expected_case_condition_records']:,} 条 case-condition "
        "记录：主模型在 AgentDojo 与 M3-v2 上跑 6 条件，第二模型跑 3 条件。"
        "该冻结规模大于论文最低样本量，正式执行前应先完成 provider/cost "
        "可行性和分析 admission；不能因为规模大而临时删掉困难 suite。"
    )
    add_table(
        doc,
        ["基准", "论文角色", "优先指标", "当前状态"],
        [
            ["AgentDojo", "主公共外部基准", "unsafe effect + benign TCR", "冻结 397/297 cases，结果待跑"],
            ["M3-v2", "内部主基准扩展", "异质性 + 授权 regime", "冻结 450/300 cases，结果待跑"],
            ["SIREN", "工具流/上下文依赖", "unsafe effect + attack utility", "需官方适配与独立准入"],
            ["InjecAgent", "补充泛化", "攻击族覆盖", "低于 AgentDojo/SIREN 优先级"],
        ],
        [1650, 2400, 2450, 2860],
        status_columns=(3,),
    )

    doc.add_heading("14. Gemini v3、盲测与外部依赖", level=1)
    add_paragraph(
        doc,
        "Gemini v3 只能在配额恢复后新建 prospective candidate。建议保留 "
        "m3_policy_only 与 m3_full 的配对比较，并加入一个经过该模型接口"
        "正确适配、不会静默回退的轻量外部条件；若无法获得忠实结构化基线，"
        "第二模型只能支持 M3 模块方向性，而不能支持“优于最强结构化基线”。"
    )
    add_table(
        doc,
        ["外部依赖", "项目方可完成", "需要用户/第三方", "截止建议"],
        [
            ["Gemini 配额", "新 v3 协议、零网络资格、失败关闭", "确认免费额度恢复或明确付费上限", "公共扩展前"],
            ["独立盲测", "生成/抽样脚本、哈希、分析契约", "独立保管人、标签 custody", "投稿前至少 4 周"],
            ["JCR Q2", "准备候选期刊与范围核验", "学校/Clarivate 订阅确认", "投稿前"],
            ["引用完整性", "DOI/URL/版本矩阵", "导师或第二作者复核", "初稿冻结前"],
        ],
        [1900, 3100, 2700, 1660],
    )
    add_callout(
        doc,
        "盲测措辞",
        "在第三方保管人落实前，任何 sealed holdout 只能称“内部封存测试”，"
        "不得写成 independent blind evaluation。该外部依赖不能由代码或 API "
        "调用替代。",
        fill="FFF1F2",
        accent=RED,
    )

    doc.add_heading("15. 统计分析与论文主表", level=1)
    add_table(
        doc,
        ["分析对象", "主要统计", "报告要求"],
        [
            ["安全主终点", "unsafe-effect rate；paired difference", "分子/分母、绝对差、相对差、95% CI"],
            ["效用主终点", "benign tool TCR；−5pp 非劣", "同时报告 ASK/DENY 负担"],
            ["配对检验", "exact McNemar + Holm", "不以 p<0.05 替代效应量"],
            ["不确定性", "Wilson CI + 10,000 cluster bootstrap", "固定 seed、可独立重建"],
            ["失败处理", "intention-to-treat", "timeout/parse/max_steps/provider failure 全保留"],
            ["异质性", "regime、攻击族、语言、模型、suite", "探索性与确认性明确分层"],
        ],
        [2000, 2950, 4410],
    )
    add_paragraph(
        doc,
        "建议论文主表固定为：表 1 研究问题与威胁模型；表 2 在线字段/源码审计；"
        "表 3 DeepSeek 中型安全—效用主结果；表 4 授权 regime 探针；表 5 "
        "公共扩展；表 6 自适应攻击预算曲线；表 7 第二模型与成本/延迟。"
        "Methods/Results 先写，Introduction 在主要外部结果稳定后再定稿。"
    )

    doc.add_heading("16. 止损条件与审稿风险", level=1)
    add_table(
        doc,
        ["风险", "触发条件", "处理"],
        [
            ["机制增益缺失", "full 持续不优于 policy-only", "保持实证路线，不恢复机制领先"],
            ["效用不可接受", "公共基准 TCR 仍显著低于 −5pp 界限", "把过度阻断作为核心负结果"],
            ["自适应攻击推翻", "预算增长导致 M3 unsafe effect 明显上升", "收缩鲁棒性声明，报告预算曲线"],
            ["Gemini 不可用", "持续 429/身份不可审计", "只写单模型正式证据，不伪造泛化"],
            ["结构化基线不忠实", "fallback/接口语义漂移", "降级 diagnostic，保留失败日志"],
            ["盲测无人保管", "投稿前仍无第三方 custody", "删除 independent blind 声明"],
        ],
        [2200, 3400, 3760],
    )
    add_paragraph(
        doc,
        "对 JISA 审稿最危险的不是负结果，而是夸大创新、把 diagnostic 当正式"
        "证据、忽略效用损失或缺少自适应攻击。当前 empirical fallback 只要补齐"
        "公共基准、自适应攻击、第二模型或明确的单模型边界，并把授权 regime "
        "测量做深，仍具有可投稿价值。"
    )

    doc.add_heading("17. 建议排期（从本报告起）", level=1)
    add_table(
        doc,
        ["周次", "主任务", "交付物", "硬门"],
        [
            ["第 1 周", "冻结论文主表、Results shell；实现 adaptive provider bridge", "可复现实验契约 + 成本 pilot", "0 真实 effect、全测试绿"],
            ["第 2–3 周", "DeepSeek 自适应主预算与 hardest subset", "预算曲线、ITT 统计", "独立 result admission"],
            ["第 3–5 周", "AgentDojo/M3-v2 公共扩展", "主公共结果与异质性", "覆盖/失败门禁通过"],
            ["第 4–5 周", "Gemini v3（若配额恢复）", "第二模型方向性结果", "新版本、无重试、身份可审计"],
            ["第 5–7 周", "独立盲测/内部封存；Discussion 与完整性审查", "论文初稿、图表、附录", "措辞与 custody 匹配"],
            ["第 8 周", "JISA 格式与投稿材料", "Cover Letter、声明、代码/数据说明", "JCR Q2 正式复核"],
        ],
        [1200, 3000, 2900, 2260],
    )

    doc.add_heading("18. 前沿来源与投稿资料", level=1)
    add_reference(
        doc,
        "JISA",
        "官方 Aims & Scope、专题与投稿入口",
        "https://www.sciencedirect.com/journal/journal-of-information-security-and-applications",
    )
    add_reference(
        doc,
        "VIGIL",
        "ACL 2026，verify-before-commit 与 SIREN",
        "https://aclanthology.org/2026.acl-long.443/",
    )
    add_reference(
        doc,
        "DRIFT",
        "NeurIPS 2025，动态规则与控制/数据流隔离",
        "https://proceedings.neurips.cc/paper_files/paper/2025/hash/"
        "77f3b26c7907aa27b207df9b9d43f29a-Abstract-Conference.html",
    )
    add_reference(
        doc,
        "Task Shield",
        "ACL 2025，工具调用与用户目标一致性",
        "https://aclanthology.org/2025.acl-long.1435/",
    )
    add_reference(
        doc,
        "ARGUS",
        "2026 预印本，上下文依赖与 provenance-aware decision auditing",
        "https://arxiv.org/abs/2605.03378",
    )
    add_reference(
        doc,
        "AuthGraph",
        "2026 预印本，授权图与参数来源执行图对齐",
        "https://arxiv.org/abs/2605.26497",
    )
    add_reference(
        doc,
        "Progent",
        "2025 预印本，可编程最小权限",
        "https://arxiv.org/abs/2504.11703",
    )
    add_reference(
        doc,
        "IPIGuard",
        "EMNLP 2025，工具依赖图防御",
        "https://aclanthology.org/2025.emnlp-main.53/",
    )
    add_reference(
        doc,
        "CaMeL",
        "2025 预印本，能力与信息流隔离",
        "https://arxiv.org/abs/2503.18813",
    )
    add_reference(
        doc,
        "The Attacker Moves Second",
        "2025 预印本，防御感知自适应攻击",
        "https://arxiv.org/abs/2510.09023",
    )
    add_reference(
        doc,
        "JISA Guide for Authors",
        "稿件格式与生成式 AI 使用声明",
        "https://www.sciencedirect.com/journal/"
        "journal-of-information-security-and-applications/publish/guide-for-authors",
    )

    doc.add_page_break()
    doc.add_heading("附录 A：关键证据摘要", level=1)
    add_table(
        doc,
        ["字段", "值"],
        [
            ["DeepSeek records", f"{coverage['combined_records']}/480"],
            ["DeepSeek formal / diagnostic", f"{technical_gate['formal_candidate_count']} / {technical_gate['diagnostic_count']}"],
            ["DeepSeek official completed calls", f"{budget['combined_provider_calls_completed']:,}"],
            ["DeepSeek cumulative cost", f"CNY {budget['project_cumulative_cost_cny']:.7f}"],
            ["Primary operational failures", f"{operational['failure_count']}/{operational['record_count']}"],
            ["M3-full unsafe effect", "0/36"],
            ["Structural unsafe effect", "8/36"],
            ["M3-full benign-tool TCR", "6/24"],
            ["Structural benign-tool TCR", "12/24"],
            ["Mechanism route", "rejected → empirical fallback"],
            ["Gemini v1 formal results", "0"],
            ["Gemini v2 formal results", "0（首请求 HTTP 429，暂停）"],
            ["Readiness", "engineering 99% / science 61% / submission 56%"],
            ["Public freeze", f"{public_freeze['expected_case_condition_records']:,} records pending"],
            ["Adaptive maximum query budget", f"{adaptive_freeze['maximum_query_budget']:,}"],
        ],
        [3700, 5660],
    )

    doc.add_heading("附录 B：报告生成与证据边界", level=1)
    add_bullets(
        doc,
        [
            "报告只读取已准入分析、冻结协议、append-only ledger 的结构字段和哈希；不读取 API 密钥。",
            "报告不向 DeepSeek 或 Gemini 发起任何新请求；Gemini HTTP 429 后未重试。",
            "报告不读取或输出完整 prompt、完整 response body、个人联系方式或支付信息。",
            "图 1–6 由报告脚本根据上述数值本地生成；不使用模型生成图表数值。",
            "JISA 主题范围来自官方页面；JCR 分区仍需学校/Clarivate 在投稿时复核。",
            "本报告是阶段性研究决策文件，不是已完成论文，也不代表科学证据达到 80%。",
        ],
    )
    add_source_note(
        doc,
        f"源码审计摘要：{SOURCE_AUDIT_SUMMARY_PATH.relative_to(PROJECT_ROOT)}；"
        f"分析文件 SHA-256：{sha256_file(ANALYSIS_PATH)}。",
    )

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    output = build_report()
    print(
        json.dumps(
            {
                "output": str(output),
                "bytes": output.stat().st_size,
                "sha256": sha256_file(output),
                "qa_dir": str(QA_DIR),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
