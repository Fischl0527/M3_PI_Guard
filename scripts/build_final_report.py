"""Build the competition report from completed, hash-verified E0-E17 evidence.

The report builder deliberately has no Day4/Day5/Day7 metric or figure inputs.
Formal claims come from ``outputs/experiments/final_analysis``, the exact
E0-E17 bundles, and the independently denominated E11 post-hoc supplement
admitted by that synthesis.  E0, Stub, smoke and the historical
``FPR=0`` claim are mentioned only as excluded prototype/diagnostic evidence.

The retained competition template remains the layout authority.  Content is
inserted into its existing section slots, preserving the cover, instructions,
TOC styles, section breaks, page geometry, headers and footers.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import shutil
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from experiments.registry import EXPERIMENT_IDS
from m3_eval.figure_manifest import load_figure_manifest
from scripts.build_final_analysis import (
    Bundle,
    FinalAnalysisError,
    build_e11_scope_boundary,
    build_e16_provenance,
    build_main_parse_quality,
    validate_e11_supplement,
    validate_inputs,
    validate_result_integrity_audit,
)


DEFAULT_TEMPLATE = ROOT / "docs" / "report" / "作品报告_模板.docx"
DEFAULT_OUT_DOCX = ROOT / "docs" / "report" / "作品报告_终稿.docx"
DEFAULT_SOURCE_MD = ROOT / "docs" / "report_sections" / "final_report_source.md"
DEFAULT_QUALITY_MD = ROOT / "outputs" / "quality" / "final_report_format_check.md"
DEFAULT_ANALYSIS_DIR = ROOT / "outputs" / "experiments" / "final_analysis"
DEFAULT_REGISTRY = ROOT / "experiments" / "registry.json"

REQUIRED_ANALYSIS_FILES = {
    "acceptance_matrix.csv",
    "analysis.md",
    "chart_map.csv",
    "executive_summary.json",
    "failure_summary.csv",
    "figure_manifest.json",
    "input_artifact_ledger.csv",
    "result_integrity_audit.json",
    "main_parse_quality.json",
    "e16_provenance.json",
}
REQUIRED_ACCEPTANCE_IDS = tuple(f"A{index:02d}" for index in range(12))

FORMAL_PROTOTYPE_NOTE = (
    "早期原型中的 Stub 结果、smoke 结果与旧 FPR=0 表述仅用于回归和测量口径诊断，"
    "不属于正式效果结论，也未进入本报告摘要、验收表或综合图的效果分母。"
)

FORBIDDEN_TERMS = (
    "AI" + "agent",
    "AI" + "Agent",
    "Co" + "dex",
    "Co" + "pilot",
    "Chat" + "GPT 生成",
    "AI " + "生成",
    "你是" + "队员",
    "请只" + "修改",
    "完成后请" + "输出",
    "给 AI" + "agent 的指令",
)

LEGACY_FORMAL_CLAIMS = (
    "Qwen2.5 代表样本",
    "baseline 的 ASR 和 DSR 均为 0.5507",
    "WRS 从 0.4891 提升至 0.9826",
)


class ReportBuildError(RuntimeError):
    """Raised before report publication when evidence is incomplete or drifts."""


@dataclass(frozen=True)
class ChartEvidence:
    chart_id: str
    title: str
    source_experiments: str
    caption: str
    png_path: Path
    rows: tuple[dict[str, str], ...]
    takeaway: str
    limitations: str


@dataclass(frozen=True)
class ReportEvidence:
    repository_root: Path
    analysis_dir: Path
    registry_path: Path
    analysis_manifest: dict[str, Any]
    executive: dict[str, Any]
    acceptance_rows: tuple[dict[str, str], ...]
    failure_rows: tuple[dict[str, str], ...]
    charts: tuple[ChartEvidence, ...]
    bundles: Mapping[str, Bundle]
    main_parse_quality: dict[str, Any]
    e16_provenance: dict[str, Any]
    e11_scope_boundary: dict[str, Any]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_json(path: Path, *, label: str) -> dict[str, Any]:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ReportBuildError(f"cannot read {label}: {path}: {exc}") from exc
    if not isinstance(payload, dict):
        raise ReportBuildError(f"{label} must be a JSON object: {path}")
    return payload


def _read_csv(path: Path, *, label: str) -> list[dict[str, str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    except (OSError, UnicodeDecodeError, csv.Error) as exc:
        raise ReportBuildError(f"cannot read {label}: {path}: {exc}") from exc


def _strict_child(path: Path, parent: Path, *, label: str) -> Path:
    resolved = path.resolve()
    try:
        resolved.relative_to(parent.resolve())
    except ValueError as exc:
        raise ReportBuildError(f"{label} escapes its allowed directory: {path}") from exc
    if not resolved.is_file():
        raise ReportBuildError(f"missing {label}: {resolved}")
    return resolved


def _verify_analysis_manifest(analysis_dir: Path) -> dict[str, Any]:
    manifest_path = _strict_child(analysis_dir / "manifest.json", analysis_dir, label="final analysis manifest")
    manifest = _read_json(manifest_path, label="final analysis manifest")
    if manifest.get("artifact_type") != "final_analysis_synthesis":
        raise ReportBuildError("final_analysis manifest has the wrong artifact_type")
    if manifest.get("status") != "completed":
        raise ReportBuildError("final_analysis is not completed")
    if tuple(manifest.get("registry_experiments", ())) != EXPERIMENT_IDS:
        raise ReportBuildError("final_analysis does not declare the exact E0-E17 sequence")
    if manifest.get("synthesis_id") != "E18":
        raise ReportBuildError("final_analysis synthesis_id must be E18")

    artifacts = manifest.get("artifacts")
    if not isinstance(artifacts, list) or not artifacts:
        raise ReportBuildError("final_analysis manifest has no hashed artifacts")
    declared: set[str] = set()
    for index, item in enumerate(artifacts):
        if not isinstance(item, Mapping):
            raise ReportBuildError(f"final_analysis artifact[{index}] is not an object")
        relative = str(item.get("path", ""))
        if not relative or relative in declared:
            raise ReportBuildError(f"invalid or duplicate final_analysis artifact path: {relative!r}")
        declared.add(relative)
        path = _strict_child(analysis_dir / relative, analysis_dir, label=f"final_analysis artifact[{index}]")
        expected_bytes = item.get("bytes")
        expected_hash = str(item.get("sha256", "")).casefold()
        if not isinstance(expected_bytes, int) or expected_bytes <= 0 or len(expected_hash) != 64:
            raise ReportBuildError(f"invalid final_analysis artifact ledger entry: {relative}")
        if path.stat().st_size != expected_bytes:
            raise ReportBuildError(f"final_analysis artifact size drift: {relative}")
        if _sha256(path) != expected_hash:
            raise ReportBuildError(f"final_analysis artifact hash drift: {relative}")
    missing = REQUIRED_ANALYSIS_FILES - declared
    if missing:
        raise ReportBuildError(f"final_analysis manifest omits required artifacts: {sorted(missing)}")
    anchor = manifest.get("result_integrity_audit")
    if not isinstance(anchor, Mapping):
        raise ReportBuildError("final_analysis manifest lacks the result-integrity audit anchor")
    required_anchor_fields = {
        "source_path",
        "source_bytes",
        "source_sha256",
        "snapshot_path",
        "snapshot_bytes",
        "snapshot_sha256",
        "program_manifest_path",
        "program_manifest_sha256",
        "blind_freeze_path",
        "blind_freeze_sha256",
        "thresholds_path",
        "thresholds_sha256",
        "stages",
    }
    if required_anchor_fields - set(anchor):
        raise ReportBuildError("final_analysis result-integrity audit anchor is incomplete")
    if anchor.get("snapshot_path") != "result_integrity_audit.json":
        raise ReportBuildError("final_analysis audit snapshot path is unexpected")
    if anchor.get("source_sha256") != anchor.get("snapshot_sha256"):
        raise ReportBuildError("final_analysis source/snapshot audit hashes disagree")
    if anchor.get("source_bytes") != anchor.get("snapshot_bytes"):
        raise ReportBuildError("final_analysis source/snapshot audit byte sizes disagree")
    if tuple(anchor.get("stages", ())) != ("main", "derived", "stability"):
        raise ReportBuildError("final_analysis audit anchor does not cover all final stages")
    diagnostics = manifest.get("derived_diagnostics")
    if not isinstance(diagnostics, Mapping) or diagnostics.get("main_parse_quality") != "main_parse_quality.json" or diagnostics.get("e16_provenance") != "e16_provenance.json":
        raise ReportBuildError("final_analysis derived-diagnostic anchors are incomplete")
    return manifest


def _normalize_ledger(rows: Iterable[Mapping[str, Any]]) -> list[tuple[str, str, str, int, str]]:
    normalized: list[tuple[str, str, str, int, str]] = []
    for row in rows:
        try:
            item = (
                str(row["experiment_id"]),
                str(row["source"]),
                str(row["path"]).replace("\\", "/"),
                int(row["bytes"]),
                str(row["sha256"]).casefold(),
            )
        except (KeyError, TypeError, ValueError) as exc:
            raise ReportBuildError(f"invalid input artifact ledger row: {row}") from exc
        normalized.append(item)
    return sorted(normalized)


def _parse_analysis_narrative(markdown: str) -> dict[str, tuple[str, str]]:
    result: dict[str, tuple[str, str]] = {}
    matches = list(re.finditer(r"^###\s+\d+\.\s+(.+?)\s*$", markdown, flags=re.MULTILINE))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        section = markdown[match.end() : end]
        takeaway_match = re.search(r"\*\*结果解读：\*\*\s*(.+?)(?:\n\n|$)", section, flags=re.DOTALL)
        limitation_match = re.search(r"\*\*限制：\*\*\s*(.+?)(?:\n\n|$)", section, flags=re.DOTALL)
        if takeaway_match:
            result[match.group(1).strip()] = (
                " ".join(takeaway_match.group(1).split()),
                " ".join(limitation_match.group(1).split()) if limitation_match else "",
            )
    return result


def _load_charts(
    analysis_dir: Path,
    manifest: Mapping[str, Any],
    analysis_markdown: str,
) -> tuple[ChartEvidence, ...]:
    try:
        figure_manifest = load_figure_manifest(
            analysis_dir / "figure_manifest.json", expected_experiment_id="E18"
        )
    except Exception as exc:
        raise ReportBuildError(f"final figure manifest failed verification: {exc}") from exc
    if figure_manifest.get("evidence_tier") != "formal_synthesis":
        raise ReportBuildError("final figure manifest is not formal_synthesis evidence")

    chart_map = _read_csv(analysis_dir / "chart_map.csv", label="final chart map")
    figures = figure_manifest.get("figures", [])
    if len(chart_map) != manifest.get("chart_count") or len(figures) != len(chart_map):
        raise ReportBuildError("final chart count is inconsistent across manifests")
    figure_by_id = {str(item.get("id")): item for item in figures if isinstance(item, Mapping)}
    narrative = _parse_analysis_narrative(analysis_markdown)
    charts: list[ChartEvidence] = []
    for row in sorted(chart_map, key=lambda item: int(item.get("order") or 0)):
        chart_id = str(row.get("chart_id", "")).strip()
        title = str(row.get("title", "")).strip()
        figure = figure_by_id.get(chart_id)
        if not chart_id or not title or not isinstance(figure, Mapping):
            raise ReportBuildError(f"chart map identity mismatch: {chart_id!r}")
        files = figure.get("files")
        if not isinstance(files, Mapping):
            raise ReportBuildError(f"{chart_id} figure manifest lacks files")
        expected_files = {
            "data": str(row.get("data", "")),
            "png": str(row.get("png", "")),
            "svg": str(row.get("svg", "")),
        }
        declared_file_paths = {
            key: str(value.get("path", "")) if isinstance(value, Mapping) else str(value)
            for key, value in files.items()
        }
        if any(declared_file_paths.get(key, "") != value for key, value in expected_files.items()):
            raise ReportBuildError(f"{chart_id} chart map and figure manifest disagree")
        data_path = _strict_child(analysis_dir / expected_files["data"], analysis_dir, label=f"{chart_id} data")
        png_path = _strict_child(analysis_dir / expected_files["png"], analysis_dir, label=f"{chart_id} PNG")
        chart_rows = _read_csv(data_path, label=f"{chart_id} data")
        if not chart_rows:
            raise ReportBuildError(f"{chart_id} has no chart data")
        takeaway, parsed_limitations = narrative.get(
            title,
            ("该图按最终分析包中的原始分母和不确定性字段呈现正式证据。", ""),
        )
        warnings = [str(item) for item in figure.get("warnings", []) if str(item).strip()]
        row_limits = sorted({str(item.get("limitation", "")).strip() for item in chart_rows if str(item.get("limitation", "")).strip()})
        limitations = "；".join(dict.fromkeys([parsed_limitations, *row_limits, *warnings]))
        charts.append(
            ChartEvidence(
                chart_id=chart_id,
                title=title,
                source_experiments=str(row.get("source_experiments", "")),
                caption=str(row.get("caption", "")),
                png_path=png_path,
                rows=tuple(chart_rows),
                takeaway=takeaway,
                limitations=limitations,
            )
        )
    if tuple(chart.chart_id for chart in charts) != tuple(figure_by_id):
        raise ReportBuildError("figure manifest order/identity differs from the final chart map")
    return tuple(charts)


def load_report_evidence(
    *,
    repository_root: str | Path = ROOT,
    analysis_dir: str | Path = DEFAULT_ANALYSIS_DIR,
    registry_path: str | Path = DEFAULT_REGISTRY,
) -> ReportEvidence:
    """Validate the complete evidence chain before returning report inputs."""

    root = Path(repository_root).resolve()
    analysis = Path(analysis_dir).resolve()
    registry = Path(registry_path).resolve()
    try:
        analysis.relative_to(root)
        registry.relative_to(root)
    except ValueError as exc:
        raise ReportBuildError("analysis and registry must remain inside the repository root") from exc
    if not analysis.is_dir():
        raise ReportBuildError(f"missing final_analysis directory: {analysis}")

    manifest = _verify_analysis_manifest(analysis)
    if manifest.get("registry_sha256") != _sha256(registry):
        raise ReportBuildError("registry hash differs from the final_analysis admission hash")
    anchor = manifest["result_integrity_audit"]
    try:
        integrity = validate_result_integrity_audit(
            root / str(anchor["source_path"]),
            program_manifest_path=root / str(anchor["program_manifest_path"]),
            repository_root=root,
        )
    except FinalAnalysisError as exc:
        raise ReportBuildError(f"final result-integrity audit validation failed: {exc}") from exc
    if integrity.audit_sha256 != anchor.get("source_sha256") or integrity.audit_bytes != anchor.get("source_bytes"):
        raise ReportBuildError("live result-integrity audit differs from the final_analysis anchor")
    snapshot = analysis / str(anchor["snapshot_path"])
    if _sha256(snapshot) != integrity.audit_sha256 or snapshot.stat().st_size != integrity.audit_bytes:
        raise ReportBuildError("final_analysis audit snapshot differs from its live admitted source")
    if (
        _sha256(integrity.program_manifest_path) != anchor.get("program_manifest_sha256")
        or _sha256(integrity.blind_freeze_path) != anchor.get("blind_freeze_sha256")
        or _sha256(integrity.thresholds_path) != anchor.get("thresholds_sha256")
    ):
        raise ReportBuildError("final_analysis audit-scope hashes differ from live bound inputs")
    try:
        e17_policy = str((manifest.get("formal_evidence_policy") or {}).get("E17", ""))
        if e17_policy.startswith("two_human_stratified_sample_"):
            from scripts.build_final_analysis_sampled_e17 import validate_inputs_sampled_e17

            _registry, bundles, live_ledger = validate_inputs_sampled_e17(
                registry, repository_root=root
            )
        else:
            _registry, bundles, live_ledger = validate_inputs(registry, repository_root=root)
    except FinalAnalysisError as exc:
        raise ReportBuildError(f"E0-E17 bundle validation failed: {exc}") from exc
    recorded_boundary = manifest.get("e11_scope_boundary")
    if not isinstance(recorded_boundary, Mapping):
        raise ReportBuildError("final_analysis lacks the structured E11 scope boundary")
    supplement_record = recorded_boundary.get("posthoc_supplement")
    if not isinstance(supplement_record, Mapping):
        raise ReportBuildError("final_analysis lacks the E11 post-hoc supplement boundary")
    supplement_manifest = _strict_child(
        root / str(supplement_record.get("manifest_path", "")),
        root,
        label="E11 post-hoc supplement manifest",
    )
    if supplement_manifest.name != "manifest.json":
        raise ReportBuildError("E11 post-hoc boundary must reference manifest.json")
    try:
        supplement, supplement_ledger = validate_e11_supplement(
            supplement_manifest.parent, repository_root=root
        )
        expected_boundary = build_e11_scope_boundary(
            bundles["E11"], supplement, repository_root=root
        )
    except FinalAnalysisError as exc:
        raise ReportBuildError(f"E11 post-hoc supplement validation failed: {exc}") from exc
    if dict(recorded_boundary) != expected_boundary:
        raise ReportBuildError("E11 scope boundary differs from the live hash-verified evidence")

    live_ledger = [
        *live_ledger,
        *supplement_ledger,
        {
            "experiment_id": "PROGRAM",
            "source": "result_integrity_audit",
            "path": integrity.audit_path.relative_to(root).as_posix(),
            "bytes": integrity.audit_bytes,
            "sha256": integrity.audit_sha256,
        },
    ]
    stored_ledger = _read_csv(analysis / "input_artifact_ledger.csv", label="input artifact ledger")
    if _normalize_ledger(stored_ledger) != _normalize_ledger(live_ledger):
        raise ReportBuildError("E0-E17 input ledger drifted after final_analysis generation")
    if manifest.get("input_artifact_count") != len(live_ledger):
        raise ReportBuildError("final_analysis input artifact count is inconsistent")

    executive = _read_json(analysis / "executive_summary.json", label="executive summary")
    acceptance = _read_csv(analysis / "acceptance_matrix.csv", label="acceptance matrix")
    failures = _read_csv(analysis / "failure_summary.csv", label="failure summary")
    if tuple(row.get("target_id") for row in acceptance) != REQUIRED_ACCEPTANCE_IDS:
        raise ReportBuildError("acceptance matrix does not contain the exact A00-A11 sequence")
    status = executive.get("overall_acceptance_status")
    if status not in {"pass", "fail", "not_evaluable"} or status != manifest.get("acceptance_status"):
        raise ReportBuildError("executive and final_analysis acceptance statuses disagree")
    target_statuses = {str(item.get("target_id")): str(item.get("status")) for item in executive.get("targets", []) if isinstance(item, Mapping)}
    if target_statuses != {str(row["target_id"]): str(row["status"]) for row in acceptance}:
        raise ReportBuildError("executive targets disagree with acceptance_matrix.csv")
    a11 = next(row for row in acceptance if row.get("target_id") == "A11")
    main_parse_quality = _read_json(
        analysis / "main_parse_quality.json", label="main parse quality"
    )
    if main_parse_quality.get("source_audit_sha256") != integrity.audit_sha256:
        raise ReportBuildError("main parse quality is bound to a different integrity audit")
    try:
        expected_parse_quality = build_main_parse_quality(
            integrity, reference_threshold=float(a11["threshold"])
        )
    except (FinalAnalysisError, TypeError, ValueError) as exc:
        raise ReportBuildError(f"cannot rederive main parse quality: {exc}") from exc
    parse_without_time = {
        key: value for key, value in main_parse_quality.items() if key != "generated_at"
    }
    expected_parse_without_time = {
        key: value for key, value in expected_parse_quality.items() if key != "generated_at"
    }
    if parse_without_time != expected_parse_without_time:
        raise ReportBuildError("main parse quality differs from the audited counters")

    e16_provenance = _read_json(
        analysis / "e16_provenance.json", label="E16 provenance"
    )
    try:
        expected_e16_provenance = build_e16_provenance(
            bundles["E16"], repository_root=root
        )
    except FinalAnalysisError as exc:
        raise ReportBuildError(f"cannot rederive E16 provenance: {exc}") from exc
    e16_without_time = {
        key: value for key, value in e16_provenance.items() if key != "generated_at"
    }
    expected_e16_without_time = {
        key: value for key, value in expected_e16_provenance.items() if key != "generated_at"
    }
    if e16_without_time != expected_e16_without_time:
        raise ReportBuildError("E16 provenance differs from the hash-bound raw log")
    markdown = (analysis / "analysis.md").read_text(encoding="utf-8")
    charts = _load_charts(analysis, manifest, markdown)
    return ReportEvidence(
        repository_root=root,
        analysis_dir=analysis,
        registry_path=registry,
        analysis_manifest=manifest,
        executive=executive,
        acceptance_rows=tuple(acceptance),
        failure_rows=tuple(failures),
        charts=charts,
        bundles=bundles,
        main_parse_quality=main_parse_quality,
        e16_provenance=e16_provenance,
        e11_scope_boundary=expected_boundary,
    )


def _optional_float(value: Any) -> float | None:
    if value is None or str(value).strip() in {"", "NA", "None", "null"}:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ReportBuildError(f"expected a finite number, got {value!r}") from exc
    if not math.isfinite(number):
        raise ReportBuildError(f"expected a finite number, got {value!r}")
    return number


def _optional_int(value: Any) -> int | None:
    number = _optional_float(value)
    if number is None:
        return None
    if not number.is_integer():
        raise ReportBuildError(f"expected an integer, got {value!r}")
    return int(number)


def _percent(value: Any) -> str:
    number = _optional_float(value)
    return "NA" if number is None else f"{number * 100:.1f}%"


def _count_text(row: Mapping[str, Any]) -> str:
    numerator = _optional_int(row.get("numerator"))
    denominator = _optional_int(row.get("denominator"))
    return "NA" if numerator is None or denominator is None else f"{numerator}/{denominator}"


def _ci_text(row: Mapping[str, Any], *, percent: bool = True) -> str:
    low = _optional_float(row.get("ci_low"))
    high = _optional_float(row.get("ci_high"))
    if low is None or high is None:
        return "NA"
    if percent:
        return f"[{low * 100:.1f}%, {high * 100:.1f}%]"
    return f"[{low:.3f}, {high:.3f}]"


def _acceptance(evidence: ReportEvidence, target_id: str) -> dict[str, str]:
    for row in evidence.acceptance_rows:
        if row.get("target_id") == target_id:
            return row
    raise ReportBuildError(f"acceptance target is missing: {target_id}")


def _wilson(numerator: int, denominator: int) -> tuple[float, float]:
    if denominator <= 0:
        raise ReportBuildError("Wilson interval requires a positive denominator")
    z = 1.959963984540054
    proportion = numerator / denominator
    denominator_adjusted = 1.0 + z * z / denominator
    center = (proportion + z * z / (2.0 * denominator)) / denominator_adjusted
    half = z * math.sqrt((proportion * (1.0 - proportion) + z * z / (4.0 * denominator)) / denominator) / denominator_adjusted
    return max(0.0, center - half), min(1.0, center + half)


def _aggregate_chart_metric(
    chart: ChartEvidence,
    *,
    metric: str,
    condition_id: str,
) -> dict[str, Any]:
    rows = [
        row
        for row in chart.rows
        if row.get("metric") == metric
        and row.get("condition_id") == condition_id
        and row.get("value_status", "observed") == "observed"
    ]
    numerator = sum(_optional_int(row.get("numerator")) or 0 for row in rows)
    denominator = sum(_optional_int(row.get("denominator")) or 0 for row in rows)
    if not rows or denominator <= 0:
        raise ReportBuildError(f"{chart.chart_id} lacks {metric}/{condition_id} raw counts")
    low, high = _wilson(numerator, denominator)
    return {
        "value": numerator / denominator,
        "numerator": numerator,
        "denominator": denominator,
        "ci_low": low,
        "ci_high": high,
    }


def _format_estimate(row: Mapping[str, Any]) -> str:
    value = _optional_float(row.get("estimate"))
    if value is None:
        return "NA"
    unit = str(row.get("unit", ""))
    if unit in {"proportion", "proportion_difference"}:
        return f"{value * 100:.1f}%"
    if unit == "ms":
        return f"{value:.2f} ms"
    if unit == "boolean":
        return "满足" if value == 1 else "不满足"
    return f"{value:.3f}"


def _format_threshold(row: Mapping[str, Any]) -> str:
    threshold = str(row.get("threshold", "")).strip()
    try:
        value = float(threshold)
    except ValueError:
        value = None
    if value is not None and math.isfinite(value) and str(row.get("unit")) in {"proportion", "proportion_difference"}:
        threshold = f"{value * 100:.1f}%"
    return f"{row.get('comparator', '')} {threshold}".strip()


def _submission_month(manifest: Mapping[str, Any]) -> str:
    raw = str(manifest.get("completed_at") or manifest.get("created_at") or "")
    try:
        parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    except ValueError:
        return ""
    return f"{parsed.year}年{parsed.month}月"


def _formal_models(evidence: ReportEvidence) -> list[str]:
    models: set[str] = set()
    for experiment_id, bundle in evidence.bundles.items():
        if experiment_id == "E0":
            continue
        for model in bundle.manifest.get("models", []):
            text = str(model).strip()
            if text and not any(token in text.casefold() for token in ("stub", "smoke", "synthetic")):
                models.add(text)
    return sorted(models)


def _acceptance_table(evidence: ReportEvidence) -> list[list[str]]:
    rows = [["ID/状态", "验收项", "估计（k/n）", "95% CI", "阈值/来源"]]
    for row in evidence.acceptance_rows:
        unit = str(row.get("unit", ""))
        rows.append(
            [
                f"{row.get('target_id')} / {row.get('status')}",
                str(row.get("criterion", "")),
                f"{_format_estimate(row)}（{_count_text(row)}）",
                _ci_text(row, percent=unit in {"proportion", "proportion_difference"}),
                f"{_format_threshold(row)}；{row.get('source_experiments')}",
            ]
        )
    return rows


def _failure_table(evidence: ReportEvidence) -> list[list[str]]:
    rows = [["来源", "分类", "计数/分母", "95% CI", "证据范围"]]
    for row in evidence.failure_rows:
        count = _optional_int(row.get("count"))
        denominator = _optional_int(row.get("denominator"))
        rows.append(
            [
                str(row.get("source_experiment", "")),
                str(row.get("failure_class", "")),
                "NA" if count is None or denominator is None else f"{count}/{denominator}",
                _ci_text(row),
                f"{row.get('source_scope', '')}；权威={row.get('authoritative', '')}",
            ]
        )
    return rows


def _parse_integrity_table(evidence: ReportEvidence) -> list[list[str]]:
    quality = evidence.main_parse_quality
    reference = float(quality["reference"]["threshold"])

    def metric_text(metric: Mapping[str, Any]) -> str:
        return (
            f"{int(metric['successes'])}/{int(metric['denominator'])}="
            f"{float(metric['estimate']) * 100:.1f}%\n"
            f"95% CI {float(metric['ci_low']) * 100:.1f}%–{float(metric['ci_high']) * 100:.1f}%\n"
            f"{metric['status_against_reference']}"
        )

    rows = [["范围", "Step-row 解析完整性", "Run-level 解析完整性", "诊断参考线"]]
    rows.append(
        [
            "总体（主矩阵全条件）",
            metric_text(quality["overall"]["step_row"]),
            metric_text(quality["overall"]["run"]),
            f">= {reference * 100:.1f}%",
        ]
    )
    for item in quality["by_model"]:
        rows.append(
            [
                str(item["model"]),
                metric_text(item["step_row"]),
                metric_text(item["run"]),
                f">= {reference * 100:.1f}%",
            ]
        )
    return rows


def build_content(evidence: ReportEvidence) -> dict[str, Any]:
    main_chart = next((chart for chart in evidence.charts if chart.chart_id == "main_defense_effect"), None)
    if main_chart is None:
        raise ReportBuildError("final_analysis has no main_defense_effect chart")
    baseline_asr = _aggregate_chart_metric(main_chart, metric="ASR", condition_id="regular_none")
    guarded_asr = _aggregate_chart_metric(main_chart, metric="ASR", condition_id="regular_full")
    reduction = (baseline_asr["value"] - guarded_asr["value"]) / baseline_asr["value"] if baseline_asr["value"] else None
    dsr = _acceptance(evidence, "A02")
    precision = _acceptance(evidence, "A03")
    recall = _acceptance(evidence, "A04")
    fpr = _acceptance(evidence, "A05")
    legal_tcr = _acceptance(evidence, "A07")
    parse = _acceptance(evidence, "A11")
    status = str(evidence.executive["overall_acceptance_status"])
    counts = evidence.executive.get("acceptance_status_counts", {})
    models = _formal_models(evidence)
    model_text = "、".join(models) if models else "正式 bundle 声明的本地模型"
    parse_quality = evidence.main_parse_quality
    step_parse = parse_quality["overall"]["step_row"]
    run_parse = parse_quality["overall"]["run"]
    parse_reference = float(parse_quality["reference"]["threshold"])
    failing_step = list(parse_quality["models_below_reference"]["step_row"])
    failing_run = list(parse_quality["models_below_reference"]["run"])
    runtime = str(
        evidence.e16_provenance["runtime_observation"]["sandbox_runtime_version"]
    )
    audit_hash = str(evidence.analysis_manifest["result_integrity_audit"]["source_sha256"])

    def parse_summary(metric: Mapping[str, Any]) -> str:
        return (
            f"{int(metric['successes'])}/{int(metric['denominator'])}="
            f"{float(metric['estimate']) * 100:.1f}%（Wilson 95% CI "
            f"{float(metric['ci_low']) * 100:.1f}%–{float(metric['ci_high']) * 100:.1f}%，"
            f"相对 {parse_reference * 100:.1f}% 参考线为 {metric['status_against_reference']}）"
        )

    parse_disclosure = (
        f"A11 仍仅表示 E1 regular_full 调用解析率；主矩阵全条件的补充诊断为："
        f"step-row {parse_summary(step_parse)}，run-level {parse_summary(run_parse)}。"
        f"step-row 低于参考线的模型为 {'、'.join(failing_step) or '无'}；"
        f"run-level 低于参考线的模型为 {'、'.join(failing_run) or '无'}。"
    )
    e16_disclosure = (
        f"E16 的哈希绑定原始日志逐行确认沙箱运行时为 {runtime}；该字符串不是 manifest schema version。"
        "E16 使用独立 validation 自适应分母，排除在主矩阵、E14 性能与延迟汇总之外。"
    )
    e11_supplement = evidence.e11_scope_boundary["posthoc_supplement"]
    e11_scope_disclosure = (
        "E11 范围偏差披露：原始冻结 v3 E11 的 4 个实现型链路中，只有 "
        "read_file_to_send_email 与计划的 read→send 直接对应；search→API、memory→file 和 "
        "tool-output→next-action 未进入原始 E11 分母。"
        f"补充协议 {e11_supplement['protocol_id']} 以 24 cases / 48 logical runs 独立覆盖这 4 个请求模式，"
        f"并且只准入 visualization revision {e11_supplement['visualization_revision']} 的全新证据；"
        "它仍是发现偏差后追加的 post-hoc amendment（pre_registered=false）。"
        "补充结果不替换原始 E11 manifest，不与原始分母合并，也不改变最终验收矩阵。"
    )

    abstract = (
        "随着大模型智能体具备文件、邮件与接口等工具能力，提示注入风险已从文本偏移扩展为外部行为风险。"
        "本作品构建 M3-PI Guard，在工具执行前融合规范化、结构化注入、多语言混淆、敏感数据流、策略和意图-动作一致性信号，"
        "输出 allow、ask、deny 及可解释原因。终稿仅使用完成态 E0-E17 bundle、独立登记且明确未预注册的 E11 post-hoc supplement，以及最终综合分析中经字节数和 SHA-256 复验的证据。"
        f"同提示盲测中，未启用 Guard 的 ASR 为 {baseline_asr['numerator']}/{baseline_asr['denominator']}="
        f"{baseline_asr['value'] * 100:.1f}%（Wilson 95% CI {baseline_asr['ci_low'] * 100:.1f}%–{baseline_asr['ci_high'] * 100:.1f}%），"
        f"full Guard 为 {guarded_asr['numerator']}/{guarded_asr['denominator']}={guarded_asr['value'] * 100:.1f}%"
        f"（Wilson 95% CI {guarded_asr['ci_low'] * 100:.1f}%–{guarded_asr['ci_high'] * 100:.1f}%），"
        f"相对下降 {'NA' if reduction is None else f'{reduction * 100:.1f}%'}。"
        f"正式验收状态为 {status}；DSR={_format_estimate(dsr)}（{_count_text(dsr)}），"
        f"恶意调用 Precision={_format_estimate(precision)}、Recall={_format_estimate(recall)}，"
        f"正常调用 FPR={_format_estimate(fpr)}（{_count_text(fpr)}），A11 regular_full ToolCall 解析成功率={_format_estimate(parse)}。"
        f"{parse_disclosure}"
        f"{e11_scope_disclosure}"
        "点估计、分母、置信区间、专项限制与未达标项均在正文保留，不以 Stub 或历史原型结果替代。"
    )

    result_paragraphs = [
        str(evidence.executive.get("headline", "")),
        (
            f"同提示主比较的 baseline ASR 为 {baseline_asr['numerator']}/{baseline_asr['denominator']}={baseline_asr['value'] * 100:.1f}%，"
            f"full Guard 为 {guarded_asr['numerator']}/{guarded_asr['denominator']}={guarded_asr['value'] * 100:.1f}%，"
            f"相对下降 {'NA' if reduction is None else f'{reduction * 100:.1f}%'}。"
            f"full Guard DSR 为 {_format_estimate(dsr)}（{_count_text(dsr)}，95% CI {_ci_text(dsr)}）。"
        ),
        (
            f"调用级 Precision={_format_estimate(precision)}（{_count_text(precision)}），Recall={_format_estimate(recall)}（{_count_text(recall)}），"
            f"正常调用 FPR={_format_estimate(fpr)}（{_count_text(fpr)}，95% CI {_ci_text(fpr)}）；"
            f"合法工具任务 TCR={_format_estimate(legal_tcr)}（{_count_text(legal_tcr)}），解析成功率={_format_estimate(parse)}（{_count_text(parse)}）。"
        ),
        parse_disclosure,
        e11_scope_disclosure,
        (
            f"12 项验收检查中 pass={counts.get('pass', 0)}、fail={counts.get('fail', 0)}、"
            f"not_evaluable={counts.get('not_evaluable', 0)}。状态仅表示预注册规则下的判定，零观测事件仍保留非零置信区间上界。"
        ),
    ]
    limitations = [str(item) for item in evidence.executive.get("limitations", [])]
    failed = [row for row in evidence.acceptance_rows if row.get("status") == "fail"]
    unavailable = [row for row in evidence.acceptance_rows if row.get("status") == "not_evaluable"]
    if failed:
        next_work = "未达标项需按来源实验完整复盘并复跑受影响条件：" + "；".join(
            f"{row['target_id']} {row['criterion']}" for row in failed
        ) + "。"
    elif unavailable:
        next_work = "不可评估项必须补齐原始分母或统计量后再生成报告：" + "；".join(
            f"{row['target_id']} {row['criterion']}" for row in unavailable
        ) + "。"
    else:
        next_work = "当前预注册目标均达到，后续仍需扩大独立模板、量化版本与硬件环境复测，避免把本地结果无条件外推。"

    return {
        "title": "M3-PI Guard：面向工具调用智能体的提示注入行为监督系统",
        "submit_date": _submission_month(evidence.analysis_manifest),
        "abstract": abstract,
        "chapters": [
            {
                "title": "第一章 作品概述",
                "sections": [
                    ("1.1 背景分析", [
                        "大模型正在从问答系统演进为可调用工具的智能体。智能体能够读写文件、发送邮件、调用接口或生成自动化动作，这提升了任务能力，也使提示注入可能直接转化为越权读取、外发和跨工具攻击链。",
                        "因此，本作品把安全目标从文本拒答扩展到工具执行前的行为监督，并要求每项效果结论可追溯到冻结配置、原始分母、统计检验与哈希证据。",
                    ]),
                    ("1.2 作品目标", [
                        "作品目标包括：构建可拦截危险工具调用的 Guard；形成覆盖九类攻击与正常任务的可审计实验集；报告 ASR、UTR、DSR、BSR、TCR、Precision、Recall、FPR、ASK、解析与性能指标；提供 Dashboard 回放与最终综合分析。",
                    ]),
                    ("1.3 特色与应用前景", [
                        "系统将安全裁决绑定到工具执行边界，并把阻断转化为明确原因码。实验工程同时提供断点续跑、失败保留、模型/数据/阈值身份记录与图表重建入口，可用于智能体安全评测、教学演示和内部工具代理审计。",
                    ]),
                ],
            },
            {
                "title": "第二章 作品设计与实现",
                "sections": [
                    ("2.1 系统方案", [
                        "系统由攻击样本、智能体运行器、隔离工具沙箱、行为监督 Guard、实验注册表、指标统计与 Dashboard 组成。运行器读取 AttackCase，真实本地模型输出 ToolCall，Guard 在工具执行前给出 allow、ask 或 deny。",
                        "工具层仅模拟本地文件、邮件和 mock API；运行级临时 sandbox 相互隔离，失败、拒答、解析错误、超时和最大步数均保留在原始证据中。",
                    ]),
                    ("2.2 实现原理", [
                        "Guard 采用分层检测：normalizer 处理 Unicode 与零宽字符；detector 识别指令型文本、工具诱导、结构化注入和多语言混淆；dataflow 标记敏感数据流；policy_engine 执行工具策略；intent-action 模块检查用户意图与模型动作的一致性。",
                        "提示词配置与 Guard 模块正交，支持普通/强化提示、keyword、policy、detector、完整 Guard 及逐组件消融，避免把提示词变化误计为 Guard 效果。",
                    ]),
                    ("2.3 指标与统计", [
                        "案例级指标包括 ASR、UTR、DSR、BSR 与正常任务 TCR；调用级指标包括 Precision、Recall、F1、FPR 与合法工具误拦；工程指标包括解析率、ASK 率、超时、p50/p95 延迟、吞吐与显存。",
                        "比例指标报告 k/n 与 Wilson 95% 置信区间；配对二元结果使用 McNemar，差值使用按案例配对 bootstrap，延迟使用中位数、p95、Wilcoxon 与效应量，多重比较采用 Holm 校正。",
                    ]),
                ],
            },
            {
                "title": "第三章 作品测试与分析",
                "sections": [
                    ("3.1 测试环境与证据准入", [
                        f"正式证据覆盖 E0-E17，涉及本地模型 {model_text}。报告生成前重新验证 registry 完成态、每个 bundle 的运行计数、artifact 字节数和 SHA-256、figure manifest，以及 final_analysis 对输入 ledger 的绑定。最终 main/derived/stability 完整性审计 SHA-256 为 {audit_hash}。",
                        "E1-E12 与 E15 属于正式真实模型效果层；E13 是验证集阈值校准；E14、E16、E17 分别使用工程、自适应和人工裁决专项分母，不能与主矩阵相加。",
                        f"E11 post-hoc 补充 manifest 为 {e11_supplement['manifest_path']}，其哈希及冻结协议、正式 execution manifest、48 条运行、统计检验、visualization revision {e11_supplement['visualization_revision']} 图表与布局 QA 均在报告准入时重新验证。",
                    ]),
                    ("3.2 测试方案与原型隔离", [
                        "主矩阵按同案例、同模型、同提示配置、同种子配对比较六个核心条件；消融、多语言、编码规避、长上下文、多轮、性能、稳定性与自适应攻击使用各自冻结子矩阵。跨工具链的原始 v3 与 post-hoc 请求模式补充严格分面展示。",
                        FORMAL_PROTOTYPE_NOTE,
                        e11_scope_disclosure,
                    ]),
                    ("3.3 正式验收结果", result_paragraphs),
                    ("3.4 多维结果分析", [
                        "下列综合图全部来自 final_analysis 的同一批已验 CSV/PNG/SVG；图注给出来源实验，图后保留结果解读与限制。精确 k/n 与置信区间同时保留在验收表和图源 CSV 中。",
                    ]),
                    ("3.5 失败分类与限制", [
                        "E17 的主分类只在其双人复核和独立裁决 cohort 内权威；E14/E16 的工程或自适应诊断信号可能重叠，不能相加为总体失败率。",
                        parse_disclosure,
                        e16_disclosure,
                        e11_scope_disclosure,
                        *limitations,
                    ]),
                ],
                "tables": [
                    ("表 1 最终预注册验收矩阵", _acceptance_table(evidence), [900, 2350, 1650, 1350, 2450]),
                    ("表 2 主矩阵解析完整性", _parse_integrity_table(evidence), [1700, 2750, 2750, 1500]),
                    ("表 3 失败分类与专项诊断分母", _failure_table(evidence), [1100, 1450, 1300, 1300, 3550]),
                ],
                "figures": list(evidence.charts),
            },
            {
                "title": "第四章 创新性说明",
                "sections": [
                    ("4.1 行为边界监督", ["作品将提示注入防护点前移到工具执行边界，直接约束越权读取、外发与代码类危险动作。"]),
                    ("4.2 多信号可解释裁决", ["系统把结构化注入、多语言混淆、敏感数据流、工具策略和意图-动作不一致合并为可解释原因，同时保留 ASK 作为不确定情形的人工确认通道。"]),
                    ("4.3 可重建实验闭环", ["注册表、原子 manifest、冻结阈值、失败恢复、统计检验、图表 manifest、Dashboard、报告与演示材料共用同一证据链；未完成或哈希漂移时报告生成器拒绝产出。"]),
                ],
            },
            {
                "title": "第五章 总结",
                "sections": [
                    ("5.1 工作总结", [
                        f"本作品完成了面向工具调用智能体的行为监督 Guard 与 E0-E17 实验闭环。最终验收状态为 {status}：pass={counts.get('pass', 0)}、fail={counts.get('fail', 0)}、not_evaluable={counts.get('not_evaluable', 0)}。",
                        f"同提示主比较中 ASR 从 {baseline_asr['numerator']}/{baseline_asr['denominator']}={baseline_asr['value'] * 100:.1f}% 变为 {guarded_asr['numerator']}/{guarded_asr['denominator']}={guarded_asr['value'] * 100:.1f}%；安全、可用性、解析与工程指标的精确结论以表 1 为准。",
                        "跨工具链请求模式已通过单独的 post-hoc 补充图透明呈现，但该补充不追溯改变原始 E11 或验收状态。",
                    ]),
                    ("5.2 局限与后续工作", [next_work, parse_disclosure, e16_disclosure, e11_scope_disclosure, *limitations, FORMAL_PROTOTYPE_NOTE]),
                ],
            },
        ],
        "references": [
            "[1] Perez F, Ribeiro I. Ignore Previous Prompt: Attack Techniques for Language Models[C]//NeurIPS ML Safety Workshop. 2022.",
            "[2] Greshake K, Abdelnabi S, Mishra S, et al. Not What You Have Signed Up For: Compromising Real-World LLM-Integrated Applications with Indirect Prompt Injection[C]//AISec. 2023.",
            "[3] OWASP Foundation. OWASP Top 10 for Large Language Model Applications[EB/OL]. 2025.",
            "[4] Yao S, Zhao J, Yu D, et al. ReAct: Synergizing Reasoning and Acting in Language Models[C]//ICLR. 2023.",
            "[5] Wei A, Haghtalab N, Steinhardt J. Jailbroken: How Does LLM Safety Training Fail?[C]//NeurIPS. 2023.",
            "[6] 中国国家标准化管理委员会. 信息与文献 参考文献著录规则: GB/T 7714-2015[S]. 北京: 中国标准出版社, 2015.",
        ],
    }


def _format_run(run: Any, font: str = "宋体", size: float = 12, *, bold: bool = False) -> None:
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold


def _replace_paragraph_text(paragraph: Any, text: str, *, font: str = "宋体", size: float = 12, bold: bool = False) -> None:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(text)
    _format_run(run, font, size, bold=bold)


def _has_section_break(paragraph_element: Any) -> bool:
    p_pr = paragraph_element.find(qn("w:pPr"))
    return p_pr is not None and p_pr.find(qn("w:sectPr")) is not None


def _find_section_slot(doc: Document, heading_text: str) -> tuple[Any, Any]:
    matches = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() == heading_text and paragraph.style.name.startswith("Heading")
    ]
    if len(matches) != 1:
        raise ReportBuildError(f"template must contain exactly one heading slot {heading_text!r}, found {len(matches)}")
    heading = matches[0]._p
    current = heading.getnext()
    while current is not None:
        if current.tag == qn("w:p") and _has_section_break(current):
            return heading, current
        if current.tag == qn("w:sectPr"):
            return heading, current
        current = current.getnext()
    raise ReportBuildError(f"template section slot has no terminating section marker: {heading_text}")


class _SectionCursor:
    def __init__(self, doc: Document, marker: Any) -> None:
        self.doc = doc
        self.marker = marker

    def paragraph(self, text: str = "", *, style: str | None = None) -> Any:
        paragraph = self.doc.add_paragraph(style=style)
        self.marker.addprevious(paragraph._p)
        if text:
            run = paragraph.add_run(text)
            _format_run(run)
        return paragraph

    def table(self, rows: int, cols: int) -> Any:
        table = self.doc.add_table(rows=rows, cols=cols)
        self.marker.addprevious(table._tbl)
        return table


def _clear_section_slot(doc: Document, heading_text: str) -> _SectionCursor:
    heading, marker = _find_section_slot(doc, heading_text)
    current = heading.getnext()
    while current is not None and current is not marker:
        following = current.getnext()
        current.getparent().remove(current)
        current = following
    return _SectionCursor(doc, marker)


def _add_paragraph(cursor: _SectionCursor, text: str, *, first_line: bool = True, bold_label: str | None = None) -> Any:
    paragraph = cursor.paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.first_line_indent = Pt(24) if first_line else Pt(0)
    if bold_label and text.startswith(bold_label):
        first = paragraph.add_run(bold_label)
        _format_run(first, bold=True)
        rest = paragraph.add_run(text[len(bold_label) :])
        _format_run(rest)
    else:
        run = paragraph.add_run(text)
        _format_run(run)
    return paragraph


def _add_heading(cursor: _SectionCursor, text: str, level: int = 2) -> Any:
    style_name = f"Heading {level}"
    try:
        cursor.doc.styles[style_name]
    except KeyError:
        style_name = None
    paragraph = cursor.paragraph(style=style_name)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    run = paragraph.add_run(text)
    _format_run(run, "黑体", 14 if level == 2 else 12, bold=True)
    return paragraph


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, *, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _apply_table_geometry(table: Any, widths: Sequence[int]) -> None:
    if len(widths) != len(table.columns) or any(width <= 0 for width in widths):
        raise ReportBuildError("table geometry width count/value is invalid")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[index]))
            _set_cell_margins(cell)


def _add_caption(cursor: _SectionCursor, text: str) -> Any:
    paragraph = cursor.paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    _format_run(run, size=10.5, bold=True)
    return paragraph


def _add_table(cursor: _SectionCursor, rows: Sequence[Sequence[str]], widths: Sequence[int]) -> None:
    if not rows or not rows[0] or any(len(row) != len(rows[0]) for row in rows):
        raise ReportBuildError("report table must be a non-empty rectangle")
    table = cursor.table(len(rows), len(rows[0]))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index in {0, 2, 3} else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(1.5)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(str(value))
            _format_run(run, size=9.0 if len(rows[0]) >= 5 else 10.5, bold=row_index == 0)
            _set_cell_shading(cell, "D9EAF7" if row_index == 0 else "FFFFFF")
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    header_properties.append(repeat)
    _apply_table_geometry(table, widths)
    cursor.paragraph().paragraph_format.space_after = Pt(3)


def _add_figure(cursor: _SectionCursor, chart: ChartEvidence, number: int) -> None:
    paragraph = cursor.paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    run.add_picture(str(chart.png_path), width=Inches(5.75))
    _add_caption(cursor, f"图 {number} {chart.title}（来源：{chart.source_experiments}）")
    _add_paragraph(cursor, f"结果解读：{chart.takeaway}", first_line=False, bold_label="结果解读：")
    _add_paragraph(
        cursor,
        f"证据与限制：图源为 final_analysis/{chart.chart_id}.csv，保留 k/n、置信区间与 NA；{chart.limitations or chart.caption}",
        first_line=False,
        bold_label="证据与限制：",
    )


def _update_template_front_matter(doc: Document, content: Mapping[str, Any]) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("□命题赛道") or text.startswith("☑命题赛道"):
            _replace_paragraph_text(paragraph, "☑命题赛道             □自由赛道", size=12)
        elif text.startswith("作品名称："):
            _replace_paragraph_text(paragraph, f"作品名称：{content['title']}", size=12)
        elif text.startswith("电子邮箱："):
            _replace_paragraph_text(paragraph, "电子邮箱：", size=12)
        elif text.startswith("提交日期："):
            _replace_paragraph_text(paragraph, f"提交日期：{content['submit_date']}", size=12)


def _set_update_fields(doc: Document) -> None:
    """Request a Word/WPS field refresh without materializing or deleting fields."""

    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _populate_docx(doc: Document, content: Mapping[str, Any]) -> None:
    _update_template_front_matter(doc, content)
    abstract_cursor = _clear_section_slot(doc, "摘要")
    _add_paragraph(abstract_cursor, str(content["abstract"]))
    _add_paragraph(abstract_cursor, "关键词：提示注入；智能体安全；工具调用；行为监督；风险评分", first_line=False)

    figure_number = 1
    for chapter in content["chapters"]:
        cursor = _clear_section_slot(doc, chapter["title"])
        for heading, paragraphs in chapter["sections"]:
            _add_heading(cursor, heading, 2)
            for text in paragraphs:
                _add_paragraph(cursor, str(text))
        for caption, rows, widths in chapter.get("tables", []):
            _add_caption(cursor, caption)
            _add_table(cursor, rows, widths)
        for chart in chapter.get("figures", []):
            _add_figure(cursor, chart, figure_number)
            figure_number += 1

    references_cursor = _clear_section_slot(doc, "参考文献")
    for item in content["references"]:
        _add_paragraph(references_cursor, str(item), first_line=False)


def _scrub_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "M3-PI Guard作品报告"
    props.subject = "智能体工具调用安全作品报告"
    props.comments = ""
    props.keywords = "提示注入;智能体安全;工具调用;行为监督"


def _scrub_docx_package(path: Path) -> None:
    temporary = path.with_suffix(".scrubbed.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if (item.filename.endswith(".xml") or item.filename.endswith(".rels")) and (
                item.filename.startswith("docProps/") or item.filename.startswith("word/")
            ):
                text = data.decode("utf-8", errors="ignore")
                text = re.sub(r'w:rsid\w+="[^"]+"', "", text)
                data = text.encode("utf-8")
            target.writestr(item, data)
    os.replace(temporary, path)


def _document_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        parts.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def _field_code_counts(path: Path) -> dict[str, int]:
    text_parts: list[str] = []
    with zipfile.ZipFile(path, "r") as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                text_parts.append(archive.read(name).decode("utf-8", errors="ignore"))
    text = "\n".join(text_parts)
    return {
        "PAGEREF": len(re.findall(r"(?<![A-Z])PAGEREF\b", text)),
        "PAGE": len(re.findall(r"(?<![A-Z])PAGE\b", text)),
    }


def _structural_audit(path: Path, *, template: Path, expected_chart_count: int) -> list[str]:
    findings: list[str] = []
    final = Document(str(path))
    source = Document(str(template))
    if len(final.sections) != len(source.sections):
        findings.append(f"section count drift: template={len(source.sections)}, report={len(final.sections)}")
    else:
        geometry_fields = (
            "page_width",
            "page_height",
            "top_margin",
            "bottom_margin",
            "left_margin",
            "right_margin",
            "header_distance",
            "footer_distance",
            "start_type",
            "different_first_page_header_footer",
        )
        for index, (source_section, final_section) in enumerate(zip(source.sections, final.sections)):
            for field in geometry_fields:
                if getattr(source_section, field) != getattr(final_section, field):
                    findings.append(
                        f"section {index + 1} {field} drift: "
                        f"template={getattr(source_section, field)}, report={getattr(final_section, field)}"
                    )
    if len(final.inline_shapes) != expected_chart_count:
        findings.append(f"inline figure count mismatch: expected={expected_chart_count}, report={len(final.inline_shapes)}")
    if _field_code_counts(path) != _field_code_counts(template):
        findings.append(
            f"field-code drift: template={_field_code_counts(template)}, report={_field_code_counts(path)}"
        )
    text = _document_text(path)
    for term in (*FORBIDDEN_TERMS, *LEGACY_FORMAL_CLAIMS):
        if term in text:
            findings.append(f"forbidden or legacy trace: {term}")
    for heading in ("摘要", "第一章 作品概述", "第二章 作品设计与实现", "第三章 作品测试与分析", "第四章 创新性说明", "第五章 总结", "参考文献"):
        if heading not in text:
            findings.append(f"missing required section: {heading}")
    if FORMAL_PROTOTYPE_NOTE not in text:
        findings.append("missing prototype/diagnostic evidence separation note")
    return findings


def render_source_markdown(content: Mapping[str, Any]) -> str:
    lines = [f"# {content['title']}", "", "## 摘要", "", str(content["abstract"]), ""]
    for chapter in content["chapters"]:
        lines.extend([f"## {chapter['title']}", ""])
        for heading, paragraphs in chapter["sections"]:
            lines.extend([f"### {heading}", "", *[str(item) for item in paragraphs], ""])
        for caption, rows, _widths in chapter.get("tables", []):
            lines.extend([f"### {caption}", ""])
            lines.extend(" | ".join(str(value) for value in row) for row in rows)
            lines.append("")
        for chart in chapter.get("figures", []):
            lines.extend(
                [
                    f"### {chart.title}",
                    "",
                    f"来源：{chart.source_experiments}；图源：final_analysis/{chart.chart_id}.png",
                    "",
                    f"结果解读：{chart.takeaway}",
                    "",
                    f"限制：{chart.limitations}",
                    "",
                ]
            )
    lines.extend(["## 参考文献", "", *content["references"], ""])
    return "\n".join(lines)


def _quality_markdown(
    *,
    evidence: ReportEvidence,
    report_path: Path,
    template_path: Path,
    structural_findings: Sequence[str],
) -> str:
    checks = [
        f"E0-E17 registry and bundle validation: {'PASS' if not structural_findings else 'see structural findings'}.",
        f"final_analysis manifest SHA-256 chain: PASS ({_sha256(evidence.analysis_dir / 'manifest.json')}).",
        f"Final main/derived/stability result-integrity audit: PASS ({evidence.analysis_manifest['result_integrity_audit']['source_sha256']}).",
        f"Main parse diagnostic rederived from the bound audit: PASS ({evidence.main_parse_quality['source_audit_sha256']}).",
        f"E16 sandbox runtime provenance rederived from the bound raw log: PASS ({evidence.e16_provenance['runtime_observation']['sandbox_runtime_version']}).",
        f"Template section preservation: template={len(Document(str(template_path)).sections)}, report={len(Document(str(report_path)).sections)}.",
        f"Formal comprehensive figures inserted inline: {len(evidence.charts)}.",
        "Template PAGE/PAGEREF fields are preserved and w:updateFields=true requests refresh when opened in Word/WPS.",
        "Legacy Day5 CSV/figures are not report inputs; Stub and old FPR=0 appear only in the explicit prototype-diagnostic exclusion note.",
        "DOCX visual render QA: PENDING. The builder performs structural checks only; final delivery still requires render_docx.py and inspection of every page, or an explicit LibreOffice/soffice-unavailable disclosure.",
    ]
    lines = ["# Final Report Evidence And Format Check", "", "## Automated Checks", *[f"- {item}" for item in checks], "", "## Structural Findings"]
    if structural_findings:
        lines.extend(f"- FAIL: {item}" for item in structural_findings)
    else:
        lines.append("- PASS: no structural, legacy-result, identity or prohibited-trace finding.")
    lines.extend(
        [
            "",
            "## Evidence Scope",
            f"- Overall acceptance status: {evidence.executive['overall_acceptance_status']}",
            f"- Registry SHA-256: {_sha256(evidence.registry_path)}",
            f"- Final analysis artifact count: {len(evidence.analysis_manifest['artifacts'])}",
            f"- Final figure count: {len(evidence.charts)}",
            f"- E16 denominator scope: {evidence.e16_provenance['denominator_scope']} (excluded_from_main_matrix={evidence.e16_provenance['excluded_from_main_matrix']})",
            "",
            "## Required Final Visual Gate",
            "- Render the generated DOCX with the bundled documents-skill renderer when LibreOffice/soffice is available.",
            "- Inspect every page at 100% for clipping, overlap, table wrapping, caption separation, image clarity, headers, footers and pagination.",
            "- Open once in Word/WPS and update all fields so the retained PAGEREF/PAGE values match final pagination before the render review.",
            "- If soffice remains unavailable, record that visual QA could not be completed and inspect in Word/WPS before submission; do not claim the render gate passed.",
            "",
        ]
    )
    return "\n".join(lines)


def build_final_report(
    *,
    repository_root: str | Path = ROOT,
    analysis_dir: str | Path = DEFAULT_ANALYSIS_DIR,
    registry_path: str | Path = DEFAULT_REGISTRY,
    template_path: str | Path = DEFAULT_TEMPLATE,
    output_docx: str | Path = DEFAULT_OUT_DOCX,
    source_markdown: str | Path = DEFAULT_SOURCE_MD,
    quality_report: str | Path = DEFAULT_QUALITY_MD,
) -> dict[str, Any]:
    """Validate evidence, then transactionally build the DOCX and audit files."""

    evidence = load_report_evidence(
        repository_root=repository_root,
        analysis_dir=analysis_dir,
        registry_path=registry_path,
    )
    root = Path(repository_root).resolve()
    template = Path(template_path).resolve()
    output = Path(output_docx).resolve()
    source = Path(source_markdown).resolve()
    quality = Path(quality_report).resolve()
    for path, label in ((template, "template"), (output, "output DOCX"), (source, "source Markdown"), (quality, "quality report")):
        try:
            path.relative_to(root)
        except ValueError as exc:
            raise ReportBuildError(f"{label} must remain inside the repository root: {path}") from exc
    if not template.is_file():
        raise ReportBuildError(f"missing retained report template: {template}")

    content = build_content(evidence)
    output.parent.mkdir(parents=True, exist_ok=True)
    source.parent.mkdir(parents=True, exist_ok=True)
    quality.parent.mkdir(parents=True, exist_ok=True)
    stage_dir = Path(tempfile.mkdtemp(prefix=".final-report-stage-", dir=output.parent))
    stage_docx = stage_dir / output.name
    stage_source = stage_dir / source.name
    stage_quality = stage_dir / quality.name
    try:
        doc = Document(str(template))
        _populate_docx(doc, content)
        _set_update_fields(doc)
        _scrub_core_properties(doc)
        doc.save(str(stage_docx))
        _scrub_docx_package(stage_docx)
        findings = _structural_audit(stage_docx, template=template, expected_chart_count=len(evidence.charts))
        if findings:
            raise ReportBuildError("report structural audit failed: " + "; ".join(findings))
        stage_source.write_text(render_source_markdown(content), encoding="utf-8")
        stage_quality.write_text(
            _quality_markdown(
                evidence=evidence,
                report_path=stage_docx,
                template_path=template,
                structural_findings=findings,
            ),
            encoding="utf-8",
        )
        os.replace(stage_docx, output)
        os.replace(stage_source, source)
        os.replace(stage_quality, quality)
    finally:
        shutil.rmtree(stage_dir, ignore_errors=True)
    return {
        "status": "completed",
        "output_docx": str(output),
        "source_markdown": str(source),
        "quality_report": str(quality),
        "acceptance_status": evidence.executive["overall_acceptance_status"],
        "chart_count": len(evidence.charts),
        "report_sha256": _sha256(output),
        "visual_qa": "pending_render_and_every-page_inspection",
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Build the final DOCX from completed, hash-verified E0-E17 evidence"
    )
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--analysis-dir", type=Path, default=DEFAULT_ANALYSIS_DIR)
    parser.add_argument("--registry", type=Path, default=DEFAULT_REGISTRY)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--out-docx", type=Path, default=DEFAULT_OUT_DOCX)
    parser.add_argument("--source-md", type=Path, default=DEFAULT_SOURCE_MD)
    parser.add_argument("--quality-md", type=Path, default=DEFAULT_QUALITY_MD)
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    result = build_final_report(
        repository_root=args.root,
        analysis_dir=args.analysis_dir,
        registry_path=args.registry,
        template_path=args.template,
        output_docx=args.out_docx,
        source_markdown=args.source_md,
        quality_report=args.quality_md,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
