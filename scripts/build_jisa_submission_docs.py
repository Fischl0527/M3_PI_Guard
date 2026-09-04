"""Build the small editable JISA submission attachments.

The scientific manuscript is intentionally not rewritten here.  This builder
only converts already-reviewed Markdown attachments into Word files and keeps
all author-owned facts as visible placeholders.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT_DIR = PROJECT_ROOT / "docs/research/paper_v2/manuscript"

CONTENT_WIDTH_DXA = 9_360
TABLE_INDENT_DXA = 120
CELL_MARGIN_VERTICAL_DXA = 80
CELL_MARGIN_HORIZONTAL_DXA = 120
FONT_NAME = "Calibri"
HEADING_BLUE = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x66, 0x66)

OUTPUTS = {
    "highlights": MANUSCRIPT_DIR / "M3_PI_Guard_JISA_highlights.docx",
    "cover": MANUSCRIPT_DIR / "M3_PI_Guard_JISA_cover_letter.docx",
    "glossary": MANUSCRIPT_DIR / "M3_PI_Guard_JISA_glossary.docx",
    "author": MANUSCRIPT_DIR / "M3_PI_Guard_JISA_author_input_form.docx",
}


def _set_run_font(run, *, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _configure_document(document: Document, *, running_label: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(11)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, before, after, color in (
        ("Title", 20, 0, 12, RGBColor(0, 0, 0)),
        ("Heading 1", 16, 16, 8, HEADING_BLUE),
        ("Heading 2", 13, 12, 6, HEADING_BLUE),
        ("Heading 3", 12, 8, 4, HEADING_DARK_BLUE),
    ):
        style = document.styles[name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = document.styles["List Bullet"]
    list_style.font.name = FONT_NAME
    list_style.font.size = Pt(11)
    list_style.paragraph_format.left_indent = Inches(0.5)
    list_style.paragraph_format.first_line_indent = Inches(-0.25)
    list_style.paragraph_format.space_after = Pt(8)
    list_style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.text = running_label
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        _set_run_font(run, size=8.5)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("Page ")
    _set_run_font(run, size=8.5)
    run.font.color.rgb = MUTED
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _clean_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "").replace("*", "")
    return re.sub(r"\s+", " ", text).strip()


def _logical_blocks(lines: list[str]) -> list[str]:
    blocks: list[str] = []
    current: list[str] = []
    for raw in lines:
        line = raw.strip()
        if not line:
            if current:
                blocks.append(" ".join(current))
                current = []
            continue
        current.append(line)
    if current:
        blocks.append(" ".join(current))
    return blocks


def _add_title(document: Document, title: str, subtitle: str | None = None) -> None:
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.add_run(title)
    if subtitle:
        sub = document.add_paragraph()
        sub.paragraph_format.space_after = Pt(14)
        run = sub.add_run(subtitle)
        _set_run_font(run, size=10.5)
        run.font.color.rgb = MUTED


def _set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (
        ("top", CELL_MARGIN_VERTICAL_DXA),
        ("bottom", CELL_MARGIN_VERTICAL_DXA),
        ("start", CELL_MARGIN_HORIZONTAL_DXA),
        ("end", CELL_MARGIN_HORIZONTAL_DXA),
    ):
        element = tc_mar.find(qn(f"w:{name}"))
        if element is None:
            element = OxmlElement(f"w:{name}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_width(parent, tag: str, width: int) -> None:
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    child.set(qn("w:w"), str(width))
    child.set(qn("w:type"), "dxa")


def _format_table(table, widths: list[int]) -> None:
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError("Table widths must sum to the 9360-DXA content width")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    _set_width(tbl_pr, "w:tblW", CONTENT_WIDTH_DXA)
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    indent.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row_index, row in enumerate(table.rows):
        if row_index == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            row._tr.get_or_add_trPr().append(header)
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            _set_width(cell._tc.get_or_add_tcPr(), "w:tcW", widths[column_index])
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "F2F4F7")
                cell._tc.get_or_add_tcPr().append(shading)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    _set_run_font(run, size=8 if len(widths) >= 7 else 9)
                    if row_index == 0:
                        run.bold = True


def build_highlights() -> Path:
    source = (MANUSCRIPT_DIR / "highlights_v1.md").read_text(encoding="utf-8")
    highlights = [_clean_markdown(line[2:]) for line in source.splitlines() if line.startswith("- ")]
    if len(highlights) != 5 or any(len(item) > 85 for item in highlights):
        raise ValueError("JISA highlights must contain five items of at most 85 characters")

    document = Document()
    _configure_document(document, running_label="JISA submission | Highlights")
    document.core_properties.title = "Highlights: Runtime Authorization for Tool-Using LLM Agents"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    _add_title(document, "Highlights", "Separate editable upload file")
    for item in highlights:
        document.add_paragraph(item, style="List Bullet")
    document.save(OUTPUTS["highlights"])
    return OUTPUTS["highlights"]


def build_cover_letter() -> Path:
    source = (MANUSCRIPT_DIR / "cover_letter_draft.md").read_text(encoding="utf-8")
    body = source.split("Dear Editors", 1)[1]
    blocks = _logical_blocks(("Dear Editors" + body).splitlines())

    document = Document()
    _configure_document(document, running_label="JISA submission | Cover letter")
    cover_body = document.styles["Normal"]
    cover_body.font.size = Pt(10.5)
    cover_body.paragraph_format.space_after = Pt(4)
    cover_body.paragraph_format.line_spacing = 1.0
    document.core_properties.title = (
        "Cover Letter: Evaluating Authorization Regimes and Runtime Guardrail "
        "Configurations for Tool-Using LLM Agents"
    )
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    _add_title(document, "Cover Letter", "Journal of Information Security and Applications")
    date = document.add_paragraph("[Submission date]")
    date.paragraph_format.space_after = Pt(14)
    for block in blocks:
        paragraph = document.add_paragraph(_clean_markdown(block))
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    signature = document.add_paragraph()
    signature.paragraph_format.space_before = Pt(12)
    signature.add_run("[Corresponding author name]\n")
    signature.add_run("[Affiliation]\n")
    signature.add_run("[Email and postal address]")
    document.save(OUTPUTS["cover"])
    return OUTPUTS["cover"]


def build_glossary() -> Path:
    source = (MANUSCRIPT_DIR / "glossary_v1.md").read_text(encoding="utf-8")
    entries = re.findall(r"^- \*\*(.+?):\*\*\s+(.+?)(?=\n- \*\*|\Z)", source, flags=re.M | re.S)
    if len(entries) != 15:
        raise ValueError(f"Expected 15 glossary entries, found {len(entries)}")

    document = Document()
    _configure_document(document, running_label="JISA submission | Glossary")
    _add_title(document, "Glossary", "Field-specific terms used in the manuscript")
    for term, definition in entries:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        term_run = paragraph.add_run(f"{_clean_markdown(term)}: ")
        _set_run_font(term_run, bold=True)
        definition_run = paragraph.add_run(_clean_markdown(definition))
        _set_run_font(definition_run)
    document.save(OUTPUTS["glossary"])
    return OUTPUTS["glossary"]


def _add_label_value(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    label_run = paragraph.add_run(f"{label}: ")
    _set_run_font(label_run, bold=True)
    value_run = paragraph.add_run(value)
    _set_run_font(value_run)


def build_author_form() -> Path:
    document = Document()
    _configure_document(document, running_label="M3-PI Guard | Author completion form")
    _add_title(document, "JISA Author and Submission Inputs", "Internal completion form - do not upload with placeholders")

    warning = document.add_paragraph()
    warning.paragraph_format.space_after = Pt(12)
    warning_run = warning.add_run(
        "Complete every bracketed field before submission. Do not include API keys, .env files, "
        "unredacted traces, or this unfinished form in an Editorial Manager upload."
    )
    _set_run_font(warning_run, bold=True)
    warning_run.font.color.rgb = RGBColor(0x9B, 0x1C, 0x1C)

    document.add_heading("Article identity", level=1)
    _add_label_value(document, "Title", "Evaluating Authorization Regimes and Runtime Guardrail Configurations for Tool-Using LLM Agents")
    _add_label_value(document, "Article type", "[confirm in Editorial Manager]")
    _add_label_value(document, "Submission date", "[YYYY-MM-DD]")

    document.add_heading("Authors and order", level=1)
    document.add_paragraph("Add one row per author. Names must match Editorial Manager.")
    table = document.add_table(rows=2, cols=7)
    headers = ["Order", "Given name(s)", "Family name", "Affiliation", "Email", "ORCID", "Approval"]
    values = ["1", "[required]", "[required]", "[a]", "[required]", "[optional]", "[yes/no]"]
    for idx, value in enumerate(headers):
        table.cell(0, idx).text = value
    for idx, value in enumerate(values):
        table.cell(1, idx).text = value
    _format_table(table, [700, 1_400, 1_400, 1_100, 2_100, 1_500, 1_160])

    document.add_heading("Affiliations and corresponding author", level=1)
    _add_label_value(document, "Affiliation a", "[Department, Institution, Street/City/Postcode, Country]")
    _add_label_value(document, "Corresponding author", "[full name]")
    _add_label_value(document, "Email", "[required]")
    _add_label_value(document, "Postal address", "[required]")
    _add_label_value(document, "Telephone", "[required by the submission checklist]")
    _add_label_value(document, "Present/permanent address", "[if applicable; otherwise Not applicable]")

    document.add_heading("CRediT contributions", level=1)
    document.add_paragraph(
        "Use only roles actually performed: Conceptualization; Data curation; Formal analysis; "
        "Funding acquisition; Investigation; Methodology; Project administration; Resources; "
        "Software; Supervision; Validation; Visualization; Writing - original draft; Writing - review and editing."
    )
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Author"
    table.cell(0, 1).text = "Approved CRediT roles"
    table.cell(1, 0).text = "[name]"
    table.cell(1, 1).text = "[roles]"
    _format_table(table, [2_000, 7_360])

    sections = {
        "Competing interests and funding": [
            ("Competing interests", "[confirm no conflict or replace the draft statement]"),
            ("Funders and grant numbers", "[required, or confirm no specific funding]"),
            ("Sponsor role", "[required, or confirm no involvement]"),
        ],
        "Acknowledgements and institutional review": [
            ("Non-author assistance", "[names/contributions, or none]"),
            ("Ethics/review/exemption wording", "[exact institutional wording, or confirmed not applicable]"),
        ],
        "Data and code release": [
            ("Public repository release/tag", "[required before publication]"),
            ("Archival DOI or immutable identifier", "[preferred]"),
            ("Provider-derived record restrictions", "[confirm]"),
        ],
        "Biographies and photographs": [
            ("Author", "[name]"),
            ("Biography", "[100 English words maximum]"),
            ("Photograph filename", "[required]"),
            ("Permission to publish photograph", "[yes/no]"),
        ],
    }
    for heading, fields in sections.items():
        document.add_heading(heading, level=1)
        for label, value in fields:
            _add_label_value(document, label, value)

    document.add_heading("Final author confirmations", level=1)
    confirmations = [
        "Author order matches Editorial Manager.",
        "Every author approves the submitted manuscript and CRediT statement.",
        "Competing-interests and funding statements are accurate.",
        "Corresponding-author contact details are current.",
        "The AI-use declaration covers every coauthor's actual tool use.",
        "The data/code statement points to the intended public release.",
        "The submission-date institutional SCIE/JCR Q2 check is complete.",
        "The final Word/WPS page-by-page inspection and similarity check are complete.",
    ]
    checklist = document.add_table(rows=len(confirmations), cols=2)
    for row, item in zip(checklist.rows, confirmations, strict=True):
        row.cells[0].text = "[ ]"
        row.cells[1].text = item
    _format_table(checklist, [650, 8_710])

    source = document.add_paragraph()
    source.paragraph_format.space_before = Pt(12)
    source_run = source.add_run(
        "Source basis: JISA Guide for Authors checked 2026-08-20. Confirm live Editorial Manager fields at submission."
    )
    _set_run_font(source_run, size=9)
    source_run.font.color.rgb = MUTED

    document.save(OUTPUTS["author"])
    return OUTPUTS["author"]


def main() -> None:
    outputs = [build_highlights(), build_cover_letter(), build_glossary(), build_author_form()]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
