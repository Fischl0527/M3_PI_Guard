"""Structural QA for the 2026-07-30 M3-PI Guard stage report."""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from lxml import etree

from scripts.build_paper_v2_stage_report_20260726 import sha256_file
from scripts.validate_paper_v2_stage_report_20260727 import (
    NS,
    _check_clean_content,
    _check_images,
    _check_page_geometry,
    _check_table_geometry,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = (
    PROJECT_ROOT
    / "docs/research/paper_v2/reports"
    / "M3_PI_Guard_阶段性研究报告_2026-07-30.docx"
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/report_qa/stage_report_20260730"
    / "structural_qa.json"
)
DELIVERY_META_DENYLIST = (
    "current script matched",
    "latest PPT",
    "submission folder",
    "workspace",
    "build script",
    "for editing and rehearsal",
    "speaker notes were removed",
)
REQUIRED_MARKERS = (
    "5,082 条 Public-primary",
    "empirical fallback",
    "0/8",
    "独立分析复核",
    "80% 目标审计",
    "Wilson 95% CI",
    "exact McNemar",
    "provider tools 关闭",
    "无真实外部 effect",
    "独立盲测保管人",
    "JCR Q2",
)


def validate(report_path: Path) -> dict[str, Any]:
    if not report_path.is_file():
        raise FileNotFoundError(report_path)
    with zipfile.ZipFile(report_path) as archive:
        corrupt_member = archive.testzip()
        document_xml = archive.read("word/document.xml")
        styles_xml = archive.read("word/styles.xml")
        numbering_xml = archive.read("word/numbering.xml")
        table_geometry = _check_table_geometry(document_xml)
        page_geometry = _check_page_geometry(document_xml)
        images = _check_images(archive, document_xml)
        images["passed"] = (
            images["inline_image_count"] == 6
            and images["media_png_count"] == 6
            and not images["missing_alt"]
            and all(
                item["width"] >= 1200 and item["height"] >= 800
                for item in images["dimensions"]
            )
        )
        root = etree.fromstring(document_xml)
        hyperlinks = len(root.xpath(".//w:hyperlink", namespaces=NS))
        style_root = etree.fromstring(styles_xml)
        numbering_root = etree.fromstring(numbering_xml)
        list_style_num_pr = len(
            style_root.xpath(
                ".//w:style[@w:styleId='ListBullet']/w:pPr/w:numPr | "
                ".//w:style[@w:styleId='ListNumber']/w:pPr/w:numPr",
                namespaces=NS,
            )
        )
        numbering_instances = len(
            numbering_root.xpath(".//w:num", namespaces=NS)
        )

    doc = Document(report_path)
    visible_text = "\n".join(
        paragraph.text for paragraph in doc.paragraphs
    )
    visible_text += "\n" + "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    clean_content = _check_clean_content(document_xml, visible_text)
    delivery_meta_hits = [
        term
        for term in DELIVERY_META_DENYLIST
        if term.casefold() in visible_text.casefold()
    ]
    missing_markers = [
        marker for marker in REQUIRED_MARKERS if marker not in visible_text
    ]
    heading_count = sum(
        paragraph.style.name.startswith("Heading ")
        for paragraph in doc.paragraphs
    )
    title_count = sum(
        paragraph.style.name == "Title" for paragraph in doc.paragraphs
    )
    subtitle_count = sum(
        paragraph.style.name == "Subtitle"
        for paragraph in doc.paragraphs
    )
    checks = {
        "zip_integrity": corrupt_member is None,
        "page_geometry": page_geometry["passed"],
        "table_geometry": table_geometry["passed"],
        "images_and_alt_text": images["passed"],
        "clean_content": clean_content["passed"],
        "delivery_clean": not delivery_meta_hits,
        "real_list_numbering": (
            list_style_num_pr == 2 and numbering_instances > 0
        ),
        "hyperlinks": hyperlinks >= 7,
        "headings": heading_count >= 18,
        "title_styles": title_count == 1 and subtitle_count == 1,
        "tables": len(doc.tables) >= 16,
        "figures": len(doc.inline_shapes) == 6,
        "required_content": not missing_markers,
    }
    return {
        "schema_version": (
            "m3-paper-v2-stage-report-structural-qa-20260730"
        ),
        "report": str(report_path),
        "report_sha256": sha256_file(report_path),
        "report_bytes": report_path.stat().st_size,
        "checks": checks,
        "passed": all(checks.values()),
        "counts": {
            "paragraphs": len(doc.paragraphs),
            "headings": heading_count,
            "tables": len(doc.tables),
            "inline_shapes": len(doc.inline_shapes),
            "hyperlinks": hyperlinks,
            "title_style": title_count,
            "subtitle_style": subtitle_count,
            "list_style_num_pr": list_style_num_pr,
            "numbering_instances": numbering_instances,
        },
        "page_geometry": page_geometry,
        "table_geometry": table_geometry,
        "images": images,
        "clean_content": clean_content,
        "delivery_meta_hits": delivery_meta_hits,
        "missing_required_markers": missing_markers,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "report",
        nargs="?",
        type=Path,
        default=DEFAULT_REPORT,
    )
    parser.add_argument("--out", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    result = validate(args.report.resolve())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
