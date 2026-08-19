"""Structural and content QA for the external-custodian Word SOP."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZipFile

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REPORT = (
    PROJECT_ROOT
    / "docs/research/paper_v2/reports/"
    "M3_PI_Guard_第三方独立盲测保管人操作手册_2026-08-02.docx"
)
EXPECTED_SLOTS_SHA256 = (
    "4ec78a15bf87473d56d9c4a1bb59f6fadd0d8fed76adf66a567c5e794b79a0a8"
)
REQUIRED_TEXT = (
    "第 1 天只要求完成“真实保管人登记”",
    "600 条独立盲测 × 6 条件 = 3,600 条正式记录",
    r"D:\M3_Blind_Custodian_v2",
    r"D:\M3_Blind_Custodian_PrivateKey",
    "M3_PI_Guard_Blind_Custodian_Toolkit_v2_final",
    str(PROJECT_ROOT),
    "blind_execution_v2.jsonl",
    "blind_labels_v2.jsonl",
    "pre_release_review_v2.json",
    "blind_custody_v2.json",
    "custody_signature_verifier_receipt_v2.json",
    "execution_output_seal_v2.json",
    "custodian_scoring_result_seal_v2.json",
    "custodian_scoring_attestation_v2.json",
    "custodian_ed25519_private.pem",
    "developer_handoff_v2",
    "scoring_handoff_v2",
    "00_VERIFY_PRISTINE_TOOLKIT.ps1",
    "18_STAGE_SCORING_HANDOFF.ps1",
    EXPECTED_SLOTS_SHA256,
    "execution_authorized=true",
    "FAIL-CLOSED",
    "experiments.paper_v2.independent_blind_closeout_v2",
)
FORBIDDEN_TEXT = (
    "NotImplementedError",
    "TBD",
    "TODO",
    "<real independent custodian identity>",
    "<required detached signature",
    "AQ.",
    "DEEPSEEK_API_KEY=",
    "GEMINI_API_KEY=",
    r"D:\M3_Blind_Custodian_v2\toolkit",
    "使用首次评分记录的同一 scored_at",
    "使用登记文件中的同一 custodian_id",
    "blind_scoring_result_v2.json",
    "authored_pairs_v2.jsonl",
)


class ManualValidationError(ValueError):
    """Raised when the Word SOP is incomplete or structurally unsafe."""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _all_text(doc: Document) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def _table_geometry_issues(path: Path) -> list[str]:
    issues: list[str] = []
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

    def attr(name: str) -> str:
        return f"{{{ns['w']}}}{name}"

    with ZipFile(path) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    for table_index, table in enumerate(root.findall(".//w:tbl", ns), start=1):
        table_width = table.find("w:tblPr/w:tblW", ns)
        table_indent = table.find("w:tblPr/w:tblInd", ns)
        layout = table.find("w:tblPr/w:tblLayout", ns)
        grid = [
            int(node.get(attr("w"), "0"))
            for node in table.findall("w:tblGrid/w:gridCol", ns)
        ]
        if (
            table_width is None
            or table_width.get(attr("type")) != "dxa"
            or int(table_width.get(attr("w"), "0")) != 9360
            or table_indent is None
            or table_indent.get(attr("type")) != "dxa"
            or int(table_indent.get(attr("w"), "-1")) != 120
            or layout is None
            or layout.get(attr("type")) != "fixed"
            or sum(grid) != 9360
        ):
            issues.append(f"table {table_index}: table geometry drift")
        rows = table.findall("w:tr", ns)
        for row_index, row in enumerate(rows, start=1):
            if row.find("w:trPr/w:trHeight", ns) is not None:
                issues.append(f"table {table_index} row {row_index}: fixed height")
            cells = row.findall("w:tc", ns)
            if len(cells) != len(grid):
                issues.append(f"table {table_index} row {row_index}: merged cell")
                continue
            for cell_index, (cell, expected) in enumerate(zip(cells, grid), start=1):
                width = cell.find("w:tcPr/w:tcW", ns)
                margins = cell.find("w:tcPr/w:tcMar", ns)
                observed_width = int(width.get(attr("w"), "0")) if width is not None else 0
                if width is None or width.get(attr("type")) != "dxa" or observed_width != expected:
                    issues.append(
                        f"table {table_index} row {row_index} cell {cell_index}: width"
                    )
                expected_margins = {"top": 80, "bottom": 80, "start": 120, "end": 120}
                for side, value in expected_margins.items():
                    node = margins.find(f"w:{side}", ns) if margins is not None else None
                    if node is None or int(node.get(attr("w"), "-1")) != value:
                        issues.append(
                            f"table {table_index} row {row_index} cell {cell_index}: {side} margin"
                        )
    return issues


def validate_report(path: Path) -> dict[str, object]:
    if not path.is_file() or path.stat().st_size < 80_000:
        raise ManualValidationError("manual DOCX missing or unexpectedly small")
    doc = Document(path)
    text = _all_text(doc)
    missing = [item for item in REQUIRED_TEXT if item not in text]
    forbidden = [item for item in FORBIDDEN_TEXT if item in text]
    if re.search(r"(?<![A-Za-z0-9])sk-[A-Za-z0-9_-]{16,}", text):
        forbidden.append("secret-like sk key")
    if missing:
        raise ManualValidationError(f"required content missing: {missing}")
    if forbidden:
        raise ManualValidationError(f"forbidden content present: {forbidden}")

    headings = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name.startswith("Heading ")
    ]
    heading_levels = [int(paragraph.style.name.rsplit(" ", 1)[1]) for paragraph in headings]
    for previous, current in zip(heading_levels, heading_levels[1:]):
        if current > previous + 1:
            raise ManualValidationError("heading hierarchy skips a level")
    if len(headings) < 28 or len(doc.tables) < 20 or len(doc.inline_shapes) != 1:
        raise ManualValidationError(
            "manual is missing expected headings, tables, or process figure"
        )
    if len(doc.sections) != 1:
        raise ManualValidationError("manual must use one consistent section")
    section = doc.sections[0]
    geometry = {
        "page_width_inches": section.page_width.inches,
        "page_height_inches": section.page_height.inches,
        "top_margin_inches": section.top_margin.inches,
        "right_margin_inches": section.right_margin.inches,
        "bottom_margin_inches": section.bottom_margin.inches,
        "left_margin_inches": section.left_margin.inches,
        "header_distance_inches": section.header_distance.inches,
        "footer_distance_inches": section.footer_distance.inches,
    }
    expected_geometry = {
        "page_width_inches": 8.5,
        "page_height_inches": 11.0,
        "top_margin_inches": 1.0,
        "right_margin_inches": 1.0,
        "bottom_margin_inches": 1.0,
        "left_margin_inches": 1.0,
        "header_distance_inches": 0.492,
        "footer_distance_inches": 0.492,
    }
    if any(abs(geometry[key] - expected) > 0.01 for key, expected in expected_geometry.items()):
        raise ManualValidationError(f"page geometry drift: {geometry}")

    table_issues = _table_geometry_issues(path)
    if table_issues:
        raise ManualValidationError(f"table geometry issues: {table_issues[:8]}")

    with ZipFile(path) as archive:
        names = set(archive.namelist())
        document_xml = archive.read("word/document.xml").decode("utf-8")
        core_xml = archive.read("docProps/core.xml").decode("utf-8")
    if "word/comments.xml" in names or "<w:ins" in document_xml or "<w:del" in document_xml:
        raise ManualValidationError("comments or tracked changes remain")
    if "descr=\"独立盲测" not in document_xml:
        raise ManualValidationError("process figure alt text missing")
    if re.search(r"<dc:creator>[^<]+", core_xml) or re.search(
        r"<cp:lastModifiedBy>[^<]+", core_xml
    ):
        raise ManualValidationError("personal author metadata remains")
    if text.count(".\\scripts\\") < 19:
        raise ManualValidationError("00-18 wrapper command reference is incomplete")

    return {
        "status": "passed",
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": _sha256_file(path),
        "paragraphs": len(doc.paragraphs),
        "headings": len(headings),
        "tables": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
        "page_geometry": geometry,
        "required_text_checks": len(REQUIRED_TEXT),
        "forbidden_text_checks": len(FORBIDDEN_TEXT),
        "table_geometry_issues": 0,
        "credentials_read": False,
        "blind_labels_read": 0,
        "network_calls_made": 0,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", nargs="?", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args()
    print(json.dumps(validate_report(args.path), ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
