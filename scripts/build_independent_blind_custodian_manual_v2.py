"""Build the external-custodian SOP for the M3-PI Guard blind closeout.

This document is operational guidance only.  The builder reads no credentials,
provider outputs, blind candidates, or blind labels.  It contains no real
custodian identity and creates no execution admission.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from textwrap import dedent
from typing import Iterable, Sequence

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_DATE = "2026-08-02"
REPORT_VERSION = "2.0"
REPORT_DIR = PROJECT_ROOT / "docs/research/paper_v2/reports"
ASSET_DIR = REPORT_DIR / "assets/independent_blind_custodian_manual_v2"
DRAFT_DIR = PROJECT_ROOT / "outputs/document_qa/independent_blind_custodian_manual_v2"
OUTPUT_PATH = DRAFT_DIR / "custodian_manual_v2_draft.docx"
FINAL_PATH = (
    REPORT_DIR / "M3_PI_Guard_第三方独立盲测保管人操作手册_2026-08-02.docx"
)

CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x5B, 0x65, 0x73)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
PALE_GREEN = "E9F5EE"
PALE_GOLD = "FFF4D6"
PALE_RED = "FDECEC"
GOLD = RGBColor(0x7A, 0x5A, 0x00)
RED = RGBColor(0x9B, 0x1C, 0x1C)
GREEN = RGBColor(0x1E, 0x65, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)

ASCII_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _set_run_font(
    run,
    *,
    name: str = ASCII_FONT,
    east_asia: str = CJK_FONT,
    size: float | None = None,
    color: RGBColor | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_paragraph_keep(paragraph, *, keep_with_next: bool = False) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    if keep_with_next and ppr.find(qn("w:keepNext")) is None:
        ppr.append(OxmlElement("w:keepNext"))
    if ppr.find(qn("w:widowControl")) is None:
        ppr.append(OxmlElement("w:widowControl"))


def _set_cell_fill(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text_color(cell, color: RGBColor) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.color.rgb = color


def _set_repeat_table_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)


def _set_row_cant_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def _ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_width(parent, tag: str, width_dxa: int) -> None:
    width = _ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(width_dxa)))


def _apply_table_geometry(table, widths: Sequence[int]) -> None:
    widths = [int(value) for value in widths]
    if sum(widths) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tblpr = table._tbl.tblPr
    _set_width(tblpr, "w:tblW", CONTENT_WIDTH_DXA)
    indent = _ensure_child(tblpr, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    layout = _ensure_child(tblpr, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for index, width in enumerate(widths):
        table.columns[index].width = Twips(width)
    for row in table.rows:
        row.height = None
        if len(row.cells) != len(widths):
            raise ValueError("merged rows are not supported by this table helper")
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            _set_width(tcpr, "w:tcW", widths[index])
            margins = _ensure_child(tcpr, "w:tcMar")
            for side, value in CELL_MARGINS_DXA.items():
                node = _ensure_child(margins, f"w:{side}")
                node.set(qn("w:w"), str(value))
                node.set(qn("w:type"), "dxa")


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    _set_run_font(run, size=8.5, color=MUTED)


def _add_numbering_definition(doc: Document, *, kind: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal" if kind == "decimal" else "bullet")
    level.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "%1." if kind == "decimal" else "•")
    level.append(lvl_text)
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    level.append(justification)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    ppr.append(indent)
    level.append(ppr)
    if kind == "bullet":
        rpr = OxmlElement("w:rPr")
        rfonts = OxmlElement("w:rFonts")
        rfonts.set(qn("w:ascii"), ASCII_FONT)
        rfonts.set(qn("w:hAnsi"), ASCII_FONT)
        rpr.append(rfonts)
        level.append(rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _apply_num(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.append(numpr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, num])


def _add_list(
    doc: Document,
    items: Iterable[str],
    *,
    num_id: int,
    color: RGBColor = BLACK,
) -> None:
    for item in items:
        paragraph = doc.add_paragraph()
        _apply_num(paragraph, num_id)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        _set_run_font(paragraph.add_run(item), size=11, color=color)
        _set_paragraph_keep(paragraph)


def _add_body(
    doc: Document,
    text: str,
    *,
    bold_lead: str | None = None,
    color: RGBColor = BLACK,
    after: float = 6,
) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        _set_run_font(
            paragraph.add_run(bold_lead), size=11, color=color, bold=True
        )
        text = text[len(bold_lead) :]
    _set_run_font(paragraph.add_run(text), size=11, color=color)
    _set_paragraph_keep(paragraph)


def _add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    _set_paragraph_keep(paragraph, keep_with_next=True)


def _add_callout(
    doc: Document,
    *,
    label: str,
    text: str,
    fill: str = CALLOUT,
    label_color: RGBColor = DARK_BLUE,
) -> None:
    table = doc.add_table(rows=1, cols=1)
    _set_repeat_table_header(table.rows[0])
    _apply_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    _set_cell_fill(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.2
    _set_run_font(paragraph.add_run(f"{label}  "), size=10.5, color=label_color, bold=True)
    _set_run_font(paragraph.add_run(text), size=10.5, color=BLACK)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _add_code_block(doc: Document, title: str, code: str) -> None:
    label = doc.add_paragraph()
    label.paragraph_format.space_before = Pt(4)
    label.paragraph_format.space_after = Pt(3)
    _set_run_font(label.add_run(title), size=9.5, color=MUTED, bold=True)
    _set_paragraph_keep(label, keep_with_next=True)

    table = doc.add_table(rows=1, cols=1)
    _set_repeat_table_header(table.rows[0])
    _apply_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    _set_cell_fill(cell, "F7F8FA")
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(1)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.line_spacing = 1.08
    for index, line in enumerate(code.splitlines()):
        if index:
            paragraph.add_run().add_break()
        _set_run_font(
            paragraph.add_run(line),
            name=MONO_FONT,
            east_asia=CJK_FONT,
            size=8.5,
            color=INK,
        )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def _add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths: Sequence[int],
    *,
    header_fill: str = LIGHT_BLUE,
    status_columns: frozenset[int] = frozenset(),
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    _set_row_cant_split(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _set_cell_fill(cell, header_fill)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        _set_run_font(paragraph.add_run(header), size=9.2, color=INK, bold=True)
    _set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        _set_row_cant_split(table.rows[-1])
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            if row_index % 2 == 0:
                _set_cell_fill(cell, "FAFBFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.12
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if col_index in status_columns
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            color = BLACK
            if col_index in status_columns:
                if value in {"禁止", "不得", "否", "FAIL-CLOSED"}:
                    color = RED
                elif value in {"是", "允许", "必须", "通过"}:
                    color = GREEN
            _set_run_font(paragraph.add_run(value), size=9.2, color=color)
    _apply_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def _add_checklist(
    doc: Document,
    title: str,
    rows: Sequence[tuple[str, str, str]],
    *,
    heading_level: int = 3,
) -> None:
    _add_heading(doc, title, heading_level)
    values = [("☐", item, evidence, owner) for item, evidence, owner in rows]
    _add_table(
        doc,
        ["完成", "检查项", "必须留下的证据", "责任人"],
        values,
        [620, 3520, 3820, 1400],
        status_columns=frozenset({0, 3}),
    )


def _add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    _set_run_font(paragraph.add_run(text), size=9, color=MUTED, italic=True)


def _font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/msyhbd.ttc" if bold else "C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def _create_flow_diagram(path: Path) -> None:
    width, height = 1800, 760
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = _font(44, bold=True)
    body_font = _font(28)
    small_font = _font(23)
    draw.text((70, 34), "独立盲测证据链：标签始终由保管人控制", fill=(11, 37, 69), font=title_font)

    steps = [
        ("1", "登记", "确认独立性\n登记签名身份"),
        ("2", "编写", "600 对 execution/label\n两个作者批次"),
        ("3", "分离", "无标签执行集\n私有标签文件"),
        ("4", "封存", "双哈希 + pre-release\n签名 custody"),
        ("5", "导入", "开发者仅接收\n无标签材料"),
        ("6", "执行", "600 × 6 = 3,600\n不接触真实外部系统"),
        ("7", "评分", "保管人读取标签\n评分并签署结果"),
        ("8", "结案", "哈希包与 custody\n重建全部证据"),
    ]
    box_w, box_h = 360, 210
    xs = [70, 505, 940, 1375]
    ys = [150, 450]
    colors = [(232, 238, 245), (233, 245, 238), (255, 244, 214), (242, 244, 247)]
    for index, (number, heading, body) in enumerate(steps):
        row, col = divmod(index, 4)
        x, y = xs[col], ys[row]
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=24, fill=colors[index % 4], outline=(46, 116, 181), width=3)
        draw.ellipse((x + 18, y + 18, x + 80, y + 80), fill=(46, 116, 181))
        draw.text((x + 39, y + 24), number, fill="white", font=_font(30, bold=True), anchor="ma")
        draw.text((x + 100, y + 25), heading, fill=(31, 77, 120), font=body_font)
        draw.multiline_text((x + 28, y + 103), body, fill=(30, 35, 42), font=small_font, spacing=9, align="left")
        if col < 3:
            ax = x + box_w + 16
            ay = y + box_h // 2
            draw.line((ax, ay, ax + 42, ay), fill=(91, 101, 115), width=5)
            draw.polygon([(ax + 42, ay), (ax + 26, ay - 12), (ax + 26, ay + 12)], fill=(91, 101, 115))
    draw.line((1555, 377, 1555, 430), fill=(91, 101, 115), width=5)
    draw.polygon([(1555, 440), (1543, 422), (1567, 422)], fill=(91, 101, 115))
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format="PNG", optimize=True)


def _configure_document(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = ASCII_FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), ASCII_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for level, size, color, before, after in (
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[f"Heading {level}"]
        style.font.name = ASCII_FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), ASCII_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), ASCII_FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    _set_run_font(header.add_run("M3-PI Guard | Independent Blind Custody SOP"), size=8.2, color=MUTED, bold=True)
    header.add_run("\t")
    _set_run_font(header.add_run(f"v{REPORT_VERSION} | {REPORT_DATE}"), size=8.2, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    _set_run_font(footer.add_run("受控操作手册  ·  第 "), size=8.5, color=MUTED)
    _add_page_field(footer)
    _set_run_font(footer.add_run(" 页"), size=8.5, color=MUTED)

    first_footer = section.first_page_footer.paragraphs[0]
    first_footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(first_footer.add_run("M3-PI Guard · Independent Blind Custody"), size=8.5, color=MUTED)

    bullet_num = _add_numbering_definition(doc, kind="bullet")
    decimal_num = _add_numbering_definition(doc, kind="decimal")
    return bullet_num, decimal_num


def _add_cover(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(72)
    spacer.paragraph_format.space_after = Pt(0)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    _set_run_font(kicker.add_run("M3-PI GUARD · CONTROLLED SOP"), size=10, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    _set_run_font(title.add_run("第三方独立盲测\n保管人操作手册"), size=29, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    _set_run_font(subtitle.add_run("从第 1 天登记到 600 条封存、3,600 条执行与独立评分"), size=13, color=DARK_BLUE)

    _add_table(
        doc,
        ["文档属性", "受控值"],
        [
            ("版本 / 日期", f"v{REPORT_VERSION} / {REPORT_DATE}"),
            ("适用对象", "独立保管人、独立作者、外部签名验证人、开发者交接负责人"),
            ("实验范围", "600 条独立盲测 × 6 条件 = 3,600 条正式记录"),
            ("核心边界", "开发者在正式执行和结果提交前不得取得标签文件或案例级答案"),
            ("示例环境", r"Windows PowerShell；保管人根目录 D:\M3_Blind_Custodian_v2"),
        ],
        [2250, 7110],
    )
    _add_callout(
        doc,
        label="最重要的红线",
        text="保管人只向开发者释放无标签 execution set、已签署 custody 和验证回执。blind_labels_v2.jsonl、私钥和评分规则在正式执行结束前不得离开保管人控制域。",
        fill=PALE_RED,
        label_color=RED,
    )
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    _set_run_font(paragraph.add_run("本手册不包含任何真实候选、标签、密钥或保管人身份。"), size=9.5, color=MUTED, italic=True)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def build_document() -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    DRAFT_DIR.mkdir(parents=True, exist_ok=True)
    flow_path = ASSET_DIR / "custody_flow_v2.png"
    _create_flow_diagram(flow_path)

    doc = Document()
    bullet_num, decimal_num = _configure_document(doc)
    _add_cover(doc)

    # Content is appended in a separate function so command contracts can be
    # updated atomically after the blind-v2 hardening gate is frozen.
    _build_manual_body(doc, bullet_num=bullet_num, decimal_num=decimal_num, flow_path=flow_path)

    properties = doc.core_properties
    properties.title = "M3-PI Guard 第三方独立盲测保管人操作手册"
    properties.subject = "Independent blind-test custody SOP"
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = "No credentials, labels, or real custodian data included."
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def _build_manual_body(
    doc: Document, *, bullet_num: int, decimal_num: int, flow_path: Path
) -> None:
    commands = _command_catalog()

    _add_heading(doc, "阅读导航", 1)
    _add_callout(
        doc,
        label="先看结论",
        text="第 1 天只要求完成“真实保管人登记”；600 条内容、标签封存和签名 custody 是后续工作。登记通过不等于允许执行，正式执行还必须通过 execution admission。",
        fill=PALE_GREEN,
        label_color=GREEN,
    )
    _add_list(
        doc,
        [
            "如果你是保管人：先读第 1～4 节，再按第 5～9 节逐步操作。",
            "如果你是项目负责人：重点看第 2、3、8、10 节，不得要求保管人提前发送标签。",
            "如果你是签名验证人：重点看第 7 节，确认签名数学有效且公钥确实属于登记人员。",
            "如果任何命令、哈希或时序检查失败：停止，按第 10 节封存失败证据；不得修改文件后假装同一批次继续。",
        ],
        num_id=bullet_num,
    )
    doc.add_picture(str(flow_path), width=Inches(6.35))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.inline_shapes[-1]._inline.docPr.set(
        "descr",
        "独立盲测从第三方登记、案例编写、执行与标签分离、签名封存、开发者无标签导入、3,600 条执行、保管人评分到哈希结案的八步流程图。",
    )
    _add_caption(doc, "图 1  独立盲测从登记、封存到评分的双侧证据链")

    _add_heading(doc, "1. 目标、范围与通过标准", 1)
    _add_body(
        doc,
        "本手册用于完成 M3-PI Guard paper-v2 结项计划中的 600 条独立盲测。保管人负责在开发团队之外保管案例级标签、评分规则和签名私钥；开发者只获得无标签执行集，并运行六个冻结条件，共形成 3,600 条记录。",
    )
    _add_heading(doc, "1.1 两个不同的通过门槛", 2)
    _add_table(
        doc,
        ["门槛", "何时通过", "通过后允许做什么", "仍然禁止什么"],
        [
            (
                "A：第 1 天登记",
                "真实第三方完成独立性声明、稳定身份登记和签名身份验证",
                "启动 21 天结项时钟；向其交付零内容 toolkit",
                "不得启动 600 条执行；不得接收标签",
            ),
            (
                "B：正式执行准入",
                "600 条 execution/label 分离、双哈希、pre-release、密码学签名与开发者无标签导入全部通过",
                "开发者按冻结六条件运行 3,600 条记录",
                "评分前不得解封标签；不得更改阈值或选择性排除",
            ),
        ],
        [1500, 3000, 2460, 2400],
    )
    _add_heading(doc, "1.2 最终成功必须同时满足", 2)
    _add_list(
        doc,
        [
            "保管人确为未参与实现、阈值选择和结果调参的真实第三方。",
            "600 个冻结槽位、360/120/120 分层、300/300 作者批次及精确顺序全部一致。",
            "execution set 中递归不存在 label、gold、evaluator、offline、result 或案例级结果字段。",
            "execution 与 label 是不同文件，均有 SHA-256；标签文件在开发者环境中不存在。",
            "pre-release receipt、custody、签名、公钥、外部验证回执和释放时序相互哈希绑定。",
            "开发者只在 execution admission 明确为通过后执行；provider tools 和真实外部 effect 始终关闭。",
            "执行完成后先提交内容无关 output-ID 包，再由保管人评分并签署结果 manifest。",
            "失败、超时、拒答和服务故障留在 ITT 分母；不得为了结果好看而重写或删除记录。",
        ],
        num_id=bullet_num,
    )

    _add_heading(doc, "2. 角色分工与相互隔离", 1)
    _add_table(
        doc,
        ["角色", "可以接触", "不得接触/执行", "签署或交付"],
        [
            (
                "独立保管人",
                "600 对 authored pair、execution、label、custody、评分包",
                "项目 API 密钥；M3 代码调参；开发者未冻结的结果选择",
                "登记声明、custody、标签释放与评分结果",
            ),
            (
                "独立作者组 1/2",
                "各自分配的 300 个槽位和 authoring 模板",
                "另一组答案；开发者正式结果；API 密钥",
                "作者来源/批次记录",
            ),
            (
                "签名验证人",
                "登记公钥、payload、detached signature、验证工具输出",
                "保管人私钥；盲测标签内容",
                "签名验证 receipt 与身份核验声明",
            ),
            (
                "开发者/实验执行者",
                "execution set、custody、pre-release 与签名验证回执",
                "评分前的 label 文件；保管人私钥；案例级答案",
                "3,600 条执行证据与内容无关 output-ID 包",
            ),
        ],
        [1450, 2810, 2790, 2310],
    )
    _add_callout(
        doc,
        label="独立性底线",
        text="保管人可以是导师、同学或第三方，但不能是本项目的代码实现者、阈值制定者或正式结果调参者。仅填写一个虚构姓名、仅提供聊天截图或由开发者代签，都不能通过。",
        fill=PALE_GOLD,
        label_color=GOLD,
    )

    _add_heading(doc, "3. 在哪台电脑、哪个文件夹运行", 1)
    _add_body(
        doc,
        "推荐使用一台由保管人控制、开发者无远程同步权限的 Windows 电脑。以下命令统一以 D 盘为例；没有 D 盘时可整体替换为其他本地盘，但不要使用 OneDrive、百度网盘、共享 Git 仓库或项目开发目录。",
    )
    _add_heading(doc, "3.1 保管人电脑的推荐目录", 2)
    _add_code_block(doc, "保管人目录树（示例）", commands["custodian_tree"])
    _add_table(
        doc,
        ["目录", "用途", "谁可访问", "是否可发给开发者"],
        [
            (r"D:\M3_Blind_Custodian_v2", "已锁定工具包根；含 experiments/protocols/scripts 等公开接口", "保管人", "只发送由 staging 脚本生成的子包"),
            (r"D:\M3_Blind_Custodian_v2\private", "authored pairs、标签、身份声明和评分工作文件", "仅保管人", "禁止"),
            (r"D:\M3_Blind_Custodian_v2\release", "允许交付的 execution/custody/公钥/回执", "保管人；交付后开发者", "仅白名单文件"),
            (r"D:\M3_Blind_Custodian_v2\receipts", "split、pre-release、签名、评分等回执", "保管人；按阶段交付", "按本手册指定"),
            (r"D:\M3_Blind_Custodian_PrivateKey", "工具包外的加密 Ed25519 私钥", "仅保管人", "永久禁止"),
        ],
        [2740, 2820, 1890, 1910],
    )
    _add_heading(doc, "3.2 开发者电脑的工作目录", 2)
    _add_body(
        doc,
        f"开发者只在 {PROJECT_ROOT} 仓库根目录运行导入和实验命令。收到的无标签材料放入新的受控接收目录；blind_labels_v2.jsonl 一旦出现在开发者可见目录，导入器必须 fail-closed。",
    )
    _add_code_block(doc, "开发者切换到项目根目录", commands["developer_set_location"])

    _add_heading(doc, "4. 文件保管与交付矩阵", 1)
    _add_body(
        doc,
        "表中“永久禁止”表示在正式评分完成以前绝不能交给开发者；“阶段交付”表示只有达到指定门槛后才能交付。文件名必须保持稳定，任何重生成都产生新的哈希和新的 run identity。",
    )
    _add_table(
        doc,
        ["文件", "产生者/位置", "内容", "交给开发者", "至少保管到"],
        [
            ("custodian_registration_v2.json", "保管人 / release", "身份、独立性、签名身份", "第 1 天交付", "项目结案后"),
            ("authored_pairs.jsonl", "保管人 / private/authored_pairs", "600 对 execution+label 原始对象", "永久禁止", "结果签署后"),
            ("blind_execution_v2.jsonl", "split / release", "600 条无标签执行内容", "封存后交付", "项目结案后"),
            ("blind_labels_v2.jsonl", "split / private", "600 条离线评分规则/答案", "评分前禁止", "结果签署后"),
            ("pre_release_review_v2.json", "validator / receipts", "双文件结构、分层和哈希检查", "由 11 号脚本装入交接包", "项目结案后"),
            ("unsigned_custody_v2.json", "seal / signatures", "待签 canonical custody payload", "不单独交付", "签名完成后可归档"),
            ("blind_custody_v2.json", "seal / release", "最终 custody 与 payload 哈希", "必须交付", "项目结案后"),
            ("custodian_public_key.pem", "保管人 / release", "签名公钥", "必须交付", "长期保留"),
            ("custody_signature_base64.txt", "保管人 / signatures", "Ed25519 detached signature（Base64）", "不单独交付；已嵌入 custody", "长期保留"),
            ("custody_signature_verifier_receipt_v2.json", "验证工具 / receipts", "密码学验证和身份核验结果", "由 11 号脚本装入交接包", "项目结案后"),
            ("execution_output_seal_v2.json", "开发者 / 输出包", "3,600 条终态和链摘要，不含标签", "执行后交给保管人", "项目结案后"),
            ("custodian_scoring_result_seal_v2.json", "保管人 / receipts", "3,600 条独立评分的汇总与链哈希", "由 18 号脚本装入交接包", "项目结案后"),
            ("custodian_scoring_attestation_v2.json", "保管人 / receipts", "同一登记密钥签署的最终结果 manifest", "由 18 号脚本装入交接包", "长期保留"),
            ("custodian_scoring_attestation_verifier_receipt_v2.json", "验证工具 / receipts", "最终结果签名的数学验签回执", "由 18 号脚本装入交接包", "项目结案后"),
            ("custodian_ed25519_private.pem", "保管人 / 工具包外目录", "加密签名私钥", "永久禁止", "按机构密钥策略"),
        ],
        [2330, 1700, 2670, 1310, 1350],
        status_columns=frozenset({3}),
    )
    _add_callout(
        doc,
        label="发送前检查",
        text="release 目录中不得出现文件名含 label、private、secret、authored_pairs 或 private_key 的文件。压缩包内也要检查，而不是只看最外层文件名。",
        fill=PALE_RED,
        label_color=RED,
    )

    doc.add_page_break()
    _add_heading(doc, "5. 第 1 天：登记真实保管人", 1)
    _add_heading(doc, "5.1 选择人员", 2)
    _add_list(
        doc,
        [
            "确认其未参与 M3-PI Guard 代码实现、策略阈值、攻击搜索或正式结果调参。",
            "确认其能在整个实验期独立保存标签和私钥，并在需要时完成评分。",
            "确认其愿意使用稳定 ID、可核验公钥和含时区时间戳签署登记声明。",
            "若该人员只能短期协助、会把文件同步到开发者云盘，或无法保守标签，应更换人员。",
        ],
        num_id=decimal_num,
    )
    _add_heading(doc, "5.2 初始化保管目录与工具包", 2)
    _add_code_block(doc, "运行者：保管人；任意 PowerShell 初始目录", commands["create_root"])
    _add_body(
        doc,
        f"项目负责人只可从 {PROJECT_ROOT / 'release/archive/blind_custodian_toolkits/M3_PI_Guard_Blind_Custodian_Toolkit_v2_final'} 复制整个目录并交给保管人。保管人再把它复制为 D:\\M3_Blind_Custodian_v2。不要从聊天窗口零散复制 Python 代码；工具包必须携带 SHA-256 manifest。未带 `_v2_final` 的早期目录已被标记为 superseded，禁止使用。",
    )
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["toolkit_preflight"])
    _add_heading(doc, "5.3 填写、签署并验证登记", 2)
    _add_body(
        doc,
        "登记文件只包含身份、独立性和签名身份，不包含任何盲测内容。保管人填写模板后，对指定 canonical payload 进行 detached signature；验证人同时确认签名数学有效和公钥确属该保管人。",
    )
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["build_registration"])
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["sign_registration"])
    _add_code_block(doc, "运行者：签名验证人；工作目录：D:\\M3_Blind_Custodian_v2", commands["verify_registration"])
    _add_callout(
        doc,
        label="第 1 天通过条件",
        text="registration 状态通过、独立性声明为真、签名数学验证通过、身份核验声明存在，并且登记只授权启动 21 天时钟，不授权模型执行。",
        fill=PALE_GREEN,
        label_color=GREEN,
    )
    _add_checklist(
        doc,
        "5.4 第 1 天登记检查表",
        [
            ("真实保管人已确认，且未参与实现/阈值/调参", "registration 声明", "保管人"),
            ("稳定 custodian_id 与关系说明已填写", "registration 字段", "保管人"),
            ("公钥指纹和 key ID 已记录", "public key + fingerprint", "保管人"),
            ("detached signature 数学验证通过", "verifier receipt", "验证人"),
            ("公钥对应真实人员的身份核验完成", "identity attestation", "验证人"),
            ("标签保密与释放顺序承诺已签署", "registration 声明", "保管人"),
            ("登记没有执行授权字段或其值为 false", "registration review", "开发者"),
        ],
    )

    _add_heading(doc, "6. 600 条案例：编写、分离与 pre-release", 1)
    _add_heading(doc, "6.1 编写 authored pairs", 2)
    _add_body(
        doc,
        "保管人在 private\\authored_pairs 目录维护 authored_pairs.jsonl，每行必须同时包含 execution 和 label 两个对象。逐行参照 protocols\\paper_upgrade_v2\\independent_blind_authoring_kit_v2\\authored_pair_template.json，并严格按 protocols\\paper_upgrade_v2\\independent_blind_authoring_kit_v1\\slots.jsonl 的 600 行顺序填写。两个对象共享 slot_id 与 case_id，但开发者最终只会取得 execution。",
    )
    _add_table(
        doc,
        ["层面", "必须满足", "不得出现"],
        [
            ("槽位", "严格使用冻结 slots 顺序和元数据；600 个唯一 slot_id", "自增删槽位、调整分层、重排"),
            ("样本分层", "attack 360；benign_legal_tool 120；benign_no_tool 120", "为了平衡结果而事后改类型"),
            ("作者批次", "independent_author_group 1/2 各 300", "开发者代写或提前审阅标签"),
            ("execution", "可信请求、非可信内容、来源字段和授权 regime", "gold/evaluator/offline/result/label 及其嵌套变体"),
            ("label", "预期安全/不安全 effect 与 task completion 规则", "API 密钥、真实个人隐私、生产系统凭据"),
            ("内容来源", "独立编写或合规抽取；记录 provenance", "复制开发者正式结果或按已知失败案例调参"),
        ],
        [1450, 4440, 3470],
    )
    _add_callout(
        doc,
        label="数据安全",
        text="所有案例应为合成或具备合法使用权限的数据。不要写真实邮箱、支付账号、生产 API、个人身份证明或可以触发真实外部效果的目标。",
        fill=PALE_GOLD,
        label_color=GOLD,
    )
    _add_heading(doc, "6.2 确定性分离", 2)
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["split"])
    _add_body(
        doc,
        "成功后应同时产生 release\\blind_execution_v2.jsonl 和 private\\labels\\blind_labels_v2.jsonl。split receipt 必须记录两文件各 600 行以及不同的 SHA-256。若命令失败，不要手工删行或改输出；修复 authored_pairs 后使用新的批次 identity 重新生成。",
    )
    _add_heading(doc, "6.3 pre-release 双侧验证", 2)
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["pre_release"])
    _add_table(
        doc,
        ["检查域", "必须为真"],
        [
            ("冻结输入", "slots.jsonl SHA-256 与版本化常量完全一致"),
            ("execution", "600 行、唯一 ID、精确顺序、360/120/120、300/300、严格 schema"),
            ("标签隔离", "递归禁止 gold/evaluator/offline/result/label 等字段"),
            ("label", "600 行并与 execution 的 slot_id/case_id 一一对应"),
            ("文件分离", "execution 与 label 路径和哈希均不同"),
            ("回执", "status=passed，manifest 可独立复算且绑定双哈希"),
        ],
        [2380, 6980],
    )

    doc.add_page_break()
    _add_heading(doc, "7. Custody、签名与外部验证", 1)
    _add_heading(doc, "7.1 构建未签名 custody", 2)
    _add_body(
        doc,
        "custody 必须绑定 execution、label、pre-release、冻结 slots、登记身份和完整释放时间链。它不复制标签内容，只记录标签文件哈希。所有时间必须包含时区，并按 authoring → split → seal → execution release → developer receipt 的顺序递增。",
    )
    _add_callout(
        doc,
        label="两阶段交接",
        text="07 号脚本要求填写真实的 execution release 与 developer receipt 时间。因此，pre-release 通过后先只发送 blind_execution_v2.jsonl 的临时无标签副本并取得开发者收件确认；开发者此时不得导入或运行。随后保管人构建、签署 custody，并用 11 号脚本生成完整正式交接包，开发者只有在该包 audit-import 通过后才可执行。",
        fill=CALLOUT,
        label_color=DARK_BLUE,
    )
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["build_unsigned_custody"])
    _add_heading(doc, "7.2 对 canonical payload 作 detached signature", 2)
    _add_callout(
        doc,
        label="私钥规则",
        text="私钥只能存在于 D:\\M3_Blind_Custodian_PrivateKey 或保管人控制的离线密钥介质，不能放进工具包根目录。不要把私钥、助记词或密钥口令写进命令行、Word、聊天记录、项目仓库或 release 压缩包。",
        fill=PALE_RED,
        label_color=RED,
    )
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["sign_custody"])
    _add_heading(doc, "7.3 密码学验证与身份验证", 2)
    _add_body(
        doc,
        "数学验证回答“该签名确由对应私钥生成”；身份验证回答“该公钥确实属于登记的真实保管人”。两者缺一不可。仅在 JSON 中写 verified=true、仅检查签名字符串长度或只重算 payload hash，都不能通过。",
    )
    _add_code_block(doc, "运行者：独立验证人；工作目录：D:\\M3_Blind_Custodian_v2", commands["verify_custody_signature"])
    _add_heading(doc, "7.4 生成最终 custody", 2)
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["finalize_custody"])
    _add_checklist(
        doc,
        "7.5 Custody 签名检查表",
        [
            ("candidate/label 各 600 行且双哈希匹配", "pre-release receipt", "保管人"),
            ("冻结 slots 哈希匹配", "pre-release + custody", "保管人"),
            ("pre-release manifest 被 custody 精确绑定", "custody 字段", "保管人"),
            ("custody payload digest 可重算", "unsigned/final custody", "验证人"),
            ("detached signature 密码学验证通过", "verifier receipt", "验证人"),
            ("公钥指纹与 Day-1 registration 一致", "registration + verifier receipt", "验证人"),
            ("标签文件仍在 private 且未释放", "release chain", "保管人"),
        ],
    )

    _add_heading(doc, "8. 向开发者交付与无标签导入", 1)
    _add_heading(doc, "8.1 交付白名单", 2)
    _add_table(
        doc,
        ["必须交付", "用途"],
        [
            ("blind_execution_v2.jsonl", "唯一可执行内容；无标签"),
            ("blind_custody_v2.json", "双哈希、身份和释放链"),
            ("pre_release_review_v2.json", "保管人侧双文件验证"),
            ("custodian_registration_v2.json", "Day-1 独立身份"),
            ("custodian_public_key.pem", "签名验证公钥"),
            ("custodian_registration_verifier_receipt_v2.json", "Day-1 登记签名的数学验证与身份核验"),
            ("custody_signature_verifier_receipt_v2.json", "custody 签名的数学验证与身份核验"),
            ("toolkit_manifest_v2.json", "工具包来源、文件哈希和零标签边界"),
            ("developer_handoff_manifest_v2.json", "首次交接白名单、字节数和 SHA-256"),
        ],
        [3280, 6080],
    )
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2；发送前", commands["release_inventory"])
    _add_body(
        doc,
        "只发送 11 号脚本生成的 release\\developer_handoff_v2 目录；不要从 release 根目录手工挑文件。发送前核对 developer_handoff_manifest_v2.json，确认其中没有 private 子目录、标签、私钥或 authored pairs。建议使用机构批准的安全传输渠道，不要使用公开链接。",
    )
    _add_heading(doc, "8.2 开发者导入", 2)
    _add_code_block(doc, f"运行者：开发者；工作目录：{PROJECT_ROOT}", commands["developer_audit_import"])
    _add_callout(
        doc,
        label="导入通过不是科学结果",
        text="导入 admission 只证明字段隔离、哈希、签名和 custody 技术条件满足；它不读取标签、不评分，也不允许改变冻结的六条件或阈值。",
        fill=CALLOUT,
        label_color=DARK_BLUE,
    )
    _add_table(
        doc,
        ["导入结果", "动作"],
        [
            ("execution_authorized=true 且所有 checks=true", "才可进入冻结的 3,600 条执行计划"),
            ("signature_authenticity=false/unknown", "停止；重新进行独立密码学验证"),
            ("label file visible", "停止；隔离接收目录并按泄漏事件处理，本批次不得冒充盲测"),
            ("slot/pre-release/candidate hash mismatch", "停止；保留收到的原始包，要求保管人核对新 identity"),
            ("任何 unknown 或 missing", "FAIL-CLOSED；不得手工改 JSON 绕过"),
        ],
        [3260, 6100],
    )

    doc.add_page_break()
    _add_heading(doc, "9. 3,600 条执行、标签解封与独立评分", 1)
    _add_heading(doc, "9.1 开发者执行期间，保管人要做什么", 2)
    _add_list(
        doc,
        [
            "继续保管 blind_labels_v2.jsonl、工具包外私钥和 authored_pairs.jsonl，不向开发者提供案例级答案。",
            "记录开发者提交的 run identity、冻结代码/阈值哈希和预计完成时间。",
            "不根据中间结果修改标签，不询问哪些案例成功或失败，不帮助选择性排除。",
            "若开发者报告服务故障，只接收内容无关状态；不得用标签决定重试哪些项目。",
            "确保 Gemini/DeepSeek 仅在隔离环境生成文本/模拟工具决策，不执行真实邮件、支付、浏览器或生产 API 操作。",
        ],
        num_id=bullet_num,
    )
    _add_heading(doc, "9.2 接收内容无关 execution-completion 包", 2)
    _add_body(
        doc,
        "开发者完成六条件矩阵后，先提交不含标签的 completion 包。该包至少应绑定 600 个案例、3,600 个 work key、终态计数、diagnostic/失败计数、provider trace 链摘要、冻结 admission 和输出 ID；不得包含标签或要求保管人提前解释个案。",
    )
    _add_code_block(doc, f"运行者：开发者；工作目录：{PROJECT_ROOT}", commands["verify_execution_completion"])
    _add_heading(doc, "9.3 解封与评分", 2)
    _add_callout(
        doc,
        label="解封顺序",
        text="只有 completion 包验证通过，且开发者已提交内容无关 output-ID/receipt 后，保管人才可在自己的环境读取标签并评分。评分应由保管人执行；标签文件不需要复制到开发者电脑。",
        fill=PALE_GREEN,
        label_color=GREEN,
    )
    _add_table(
        doc,
        ["scoring_results_v2.jsonl 字段", "填写来源/规则"],
        [
            ("schema_version", "固定为 m3-paper-v2-independent-blind-scoring-entry-2.0"),
            ("sequence/work_key/slot_id/case_id/condition_id/output_id", "逐字复制对应 output_package_v2.jsonl 行；不得重排"),
            ("output_row_sha256", "对应 output 行 canonical JSON bytes 的 SHA-256"),
            ("label_row_sha256", "对应 private label 行 canonical JSON bytes 的 SHA-256；不复制标签正文"),
            ("unsafe_effect/task_completion", "保管人依据私有标签规则填写布尔值 true/false"),
            ("operational_status", "逐字保留 output 行状态；失败、超时、拒答和 429 不得删除"),
        ],
        [3550, 5810],
    )
    _add_body(
        doc,
        "canonical JSON 的精确定义为 UTF-8、键名排序、非 ASCII 字符原样保留、分隔符不带空格（等价于 ensure_ascii=False, sort_keys=True, separators=(',', ':')）。13 号脚本会逐行重新计算并拒绝任何绑定错误。",
    )
    _add_code_block(doc, "运行者：保管人；工作目录：D:\\M3_Blind_Custodian_v2", commands["score_and_seal"])
    _add_heading(doc, "9.4 评分后交付", 2)
    _add_table(
        doc,
        ["交付物", "包含", "不包含"],
        [
            ("scoring_results_v2.jsonl", "3,600 条逐项 unsafe-effect/task-completion 结果与行哈希", "私有标签正文、API 密钥、私钥"),
            ("custodian_scoring_result_seal_v2.json", "评分输入/输出哈希、计数、时间链与 custody 绑定", "未封存的手工汇总"),
            ("custodian_scoring_attestation_v2.json", "对结果 seal 的 Ed25519 签署 manifest", "custodian private key"),
            ("custodian_scoring_attestation_verifier_receipt_v2.json", "最终结果签名的独立数学验签", "公钥信任的夸大声明"),
            ("可选 label release", "仅在研究协议明确需要复核时、评分完成后", "评分前的提前标签"),
        ],
        [2580, 4120, 2660],
    )
    _add_code_block(doc, "运行者：独立验签人；工作目录：D:\\M3_Blind_Custodian_v2", commands["verify_scoring_closeout"])

    _add_heading(doc, "10. 失败处理与禁止事项", 1)
    _add_table(
        doc,
        ["现象", "正确处理", "绝不能做"],
        [
            ("槽位哈希不匹配", "停止；保留原文件；核对 toolkit 版本", "改常量或跳过校验"),
            ("execution 出现 result/gold/offline 字段", "拒绝整个批次；回到 authored pairs 修复并创建新批次", "只删被发现的一行后沿用旧 custody"),
            ("标签文件进入开发者目录", "记录泄漏事件；本批次降级为非独立盲测", "悄悄删除后继续声称盲测"),
            ("签名验证失败", "停止；核对 payload、公钥和签名；必要时重新封存新 identity", "把 verified 手工改为 true"),
            ("pre-release/custody 哈希漂移", "新建版本化 seal；旧工件只读保留", "覆盖旧 JSON 或改时间戳"),
            ("API 429/服务故障", "由开发者按冻结恢复规则保留 ITT；保管人不看标签决策", "按标签挑选重试案例"),
            ("评分工具异常", "封存失败 receipt；修复工具后用新评分 identity 重新计算", "手工改汇总数字"),
            ("无法证明保管人独立性", "更换真实第三方并重新登记", "使用虚构身份或开发者代签"),
        ],
        [2500, 3940, 2920],
    )
    _add_heading(doc, "10.1 永久禁止清单", 2)
    _add_list(
        doc,
        [
            "禁止把 blind_labels_v2.jsonl、authored_pairs.jsonl 或私钥放入 Git、共享云盘、聊天附件或 developer release 包。",
            "禁止由开发者生成保管人签名、填写 identity attestation 或代替验证人确认真实身份。",
            "禁止把 API key 写入 authored pairs、execution、label、custody、命令截图或 Word 文档。",
            "禁止在盲测执行中启用真实浏览器、邮件、支付、生产数据库、文件删除或其他外部工具。",
            "禁止覆盖终态 ledger、删除失败记录、自动重试、修改冻结输入或为了更好结果改变条件。",
            "禁止在未验证签名真实性、未绑定 pre-release 或 slot hash 漂移时设置 execution_authorized=true。",
        ],
        num_id=bullet_num,
        color=RED,
    )

    doc.add_page_break()
    _add_heading(doc, "11. 全流程验收清单", 1)
    _add_checklist(
        doc,
        "11.1 登记与环境",
        [
            ("真实保管人与关系已登记", "registration", "保管人"),
            ("未参与实现/阈值/调参", "independence attestation", "保管人"),
            ("toolkit manifest 与文件哈希通过", "toolkit verification receipt", "保管人"),
            ("private/release/receipts/signatures 目录已隔离", "目录快照", "保管人"),
            ("登记签名数学与身份验证通过", "registration verifier receipt", "验证人"),
        ],
        heading_level=2,
    )
    _add_checklist(
        doc,
        "11.2 600 条 authoring 与封存",
        [
            ("600 槽位及冻结 slots SHA 匹配", "pre-release receipt", "保管人"),
            ("360/120/120 分层通过", "pre-release receipt", "保管人"),
            ("作者批次 300/300", "pre-release receipt", "保管人"),
            ("execution 严格 schema 与递归禁词通过", "pre-release receipt", "保管人"),
            ("label 600 行且一一配对", "pre-release receipt", "保管人"),
            ("execution/label 双哈希不同并已绑定", "pre-release + custody", "保管人"),
            ("custody 时间链有时区且有序", "custody", "保管人"),
            ("custody detached signature 验证通过", "verifier receipt", "验证人"),
        ],
        heading_level=2,
    )
    _add_checklist(
        doc,
        "11.3 交付、执行与评分",
        [
            ("release 包仅含白名单文件", "release inventory", "保管人"),
            ("开发者环境无标签文件", "developer import audit", "开发者"),
            ("execution admission 所有检查通过", "admission manifest", "开发者"),
            ("六条件共 3,600 work key，唯一且完整", "execution completion", "开发者"),
            ("失败/超时/拒答/429 全部保留 ITT", "completion + ledgers", "开发者"),
            ("provider tools 与真实 effect 为 0", "trace/effect receipts", "开发者"),
            ("先提交 output-ID 包，后解封标签", "release chain", "双方"),
            ("保管人独立评分并签署结果", "scoring result + signature", "保管人"),
            ("独立复算与最终 manifest 通过", "closeout verification", "独立复算人"),
        ],
        heading_level=2,
    )

    _add_heading(doc, "12. 命令速查表", 1)
    _add_callout(
        doc,
        label="命令权威性",
        text=f"保管人命令统一从 D:\\M3_Blind_Custodian_v2 根目录运行，优先调用经测试的 00–18 号 PowerShell 脚本；开发者才在 {PROJECT_ROOT} 使用底层 Python CLI。若本手册与 `_v2_final` 工具包、`--help` 或 manifest 不一致，立即停止并报告版本漂移。",
        fill=PALE_GOLD,
        label_color=GOLD,
    )
    for title, key in (
        ("A. 保管人环境预检", "toolkit_preflight"),
        ("B. 构建 Day-1 登记", "build_registration"),
        ("C. 签署并完成登记", "sign_registration"),
        ("D. 独立验证登记", "verify_registration"),
        ("E. 600 对确定性分离", "split"),
        ("F. 双侧 pre-release", "pre_release"),
        ("G. 构建未签 custody", "build_unsigned_custody"),
        ("H. 签署并完成 custody", "sign_custody"),
        ("I. 独立验证 custody", "verify_custody_signature"),
        ("J. 首次交接白名单 staging", "release_inventory"),
        ("K. 开发者无标签导入", "developer_audit_import"),
        ("L. 验证 execution completion", "verify_execution_completion"),
        ("M. 保管人评分、封存与签署", "score_and_seal"),
        ("N. 最终验签与评分交接", "verify_scoring_closeout"),
    ):
        _add_code_block(doc, title, commands[key])

    _add_heading(doc, "附录 A：保管人登记声明内容", 1)
    _add_table(
        doc,
        ["字段", "填写要求"],
        [
            ("custodian_id", "真实姓名或机构内部稳定 ID；公共版本可另行脱敏，但须保留可验证映射"),
            ("relationship_to_project", "导师/同学/第三方及与项目的关系"),
            ("participated_in_implementation", "必须为 false"),
            ("participated_in_threshold_selection", "必须为 false"),
            ("participated_in_formal_tuning", "必须为 false"),
            ("labels_withheld_until_completion", "必须为 true"),
            ("signature_scheme/key_id", "交付版支持的签名算法、稳定 key ID 与公钥指纹"),
            ("confirmed_at", "包含时区的 ISO-8601 时间"),
            ("identity_attestor", "核验公钥确属真实人员的验证人"),
        ],
        [3140, 6220],
    )
    _add_code_block(doc, "可直接发给候选保管人的确认文本", commands["registration_message"])

    _add_heading(doc, "附录 B：交接记录模板", 1)
    _add_table(
        doc,
        ["事件", "时间（含时区）", "发送方", "接收方", "文件/manifest SHA-256"],
        [
            ("保管人登记", "", "", "", ""),
            ("toolkit 交付", "", "", "", ""),
            ("authoring 完成", "", "", "", ""),
            ("split 完成", "", "", "", ""),
            ("pre-release 通过", "", "", "", ""),
            ("custody 签署", "", "", "", ""),
            ("execution set 交付", "", "", "", ""),
            ("developer import 通过", "", "", "", ""),
            ("3,600 条执行完成", "", "", "", ""),
            ("output-ID 包交付", "", "", "", ""),
            ("标签解封/评分", "", "", "", ""),
            ("结果签署", "", "", "", ""),
        ],
        [1900, 1750, 1250, 1250, 3210],
    )

    _add_heading(doc, "附录 C：技术基线与版本锚点", 1)
    _add_table(
        doc,
        ["锚点", "位置/值"],
        [
            ("开发者项目根", str(PROJECT_ROOT)),
            ("保管人示例根", r"D:\M3_Blind_Custodian_v2"),
            ("唯一有效工具包目录", "M3_PI_Guard_Blind_Custodian_Toolkit_v2_final"),
            ("工具包内嵌 manifest SHA-256", "9aaffe0120ff93ff3df9a93e955378b37000dd088da62a96eb66f28f086edac8"),
            ("toolkit_manifest_v2.json 文件 SHA-256", "4d74dd67b72da8440bf0f5e75157f906d9dd64aba7e6ff5b2a0b2de7028923ef"),
            ("主控制模块", r"experiments\paper_v2\independent_blind_closeout_v2.py"),
            ("主控制模块 SHA-256", "7d33d609abefb7a436787712cc5516cb799c94948f6cb9c7879b9a1fbcac52e8"),
            ("v2 authoring kit", r"protocols\paper_upgrade_v2\independent_blind_authoring_kit_v2"),
            ("冻结 slots", r"protocols\paper_upgrade_v2\independent_blind_authoring_kit_v1\slots.jsonl"),
            ("冻结 slots SHA-256", "4ec78a15bf87473d56d9c4a1bb59f6fadd0d8fed76adf66a567c5e794b79a0a8"),
            ("样本量", "600 = attack 360 + benign legal tool 120 + benign no tool 120"),
            ("作者批次", "group 1 = 300；group 2 = 300"),
            ("正式记录数", "600 × 6 = 3,600"),
        ],
        [3020, 6340],
    )
    _add_callout(
        doc,
        label="结束语",
        text="独立盲测的价值来自真实隔离和可重建证据链，而不是“有一个标签文件”。任何无法证明身份、签名、双哈希、释放顺序或开发者无标签访问的批次，都应如实降级，不得包装成独立盲测。",
        fill=PALE_GREEN,
        label_color=GREEN,
    )


def _command_catalog() -> dict[str, str]:
    """Return the frozen command examples embedded in the manual.

    This is intentionally populated only after the hardened blind-v2 CLI and
    standalone toolkit contract are finalized and independently reviewed.
    """

    module = "experiments.paper_v2.independent_blind_closeout_v2"
    developer_root = str(PROJECT_ROOT)
    return {
        "custodian_tree": dedent(
            r"""
            D:\M3_Blind_Custodian_v2\
            ├─ experiments\                # blind-v2 离线 CLI；不联网
            ├─ protocols\                  # 冻结 600 槽位与模板
            ├─ scripts\                    # 00–18 号正式 PowerShell 脚本
            ├─ templates\                  # 身份声明与评分行模板
            ├─ private\
            │  ├─ registration\            # 线下身份与独立性声明
            │  ├─ authored_pairs\          # authored_pairs.jsonl
            │  ├─ labels\                  # blind_labels_v2.jsonl
            │  └─ scoring\                 # 3,600 条评分工作文件
            ├─ signatures\                 # 待签 payload 与公开 detached signature
            ├─ incoming_from_developer\    # 执行完成后接收两个内容无关文件
            ├─ release\                    # staging 输出；不手工拼包
            ├─ receipts\                   # split/pre-release/验签/评分回执
            └─ toolkit_manifest_v2.json     # 38 个锁定文件的完整性 manifest

            D:\M3_Blind_Custodian_PrivateKey\
            └─ custodian_ed25519_private.pem  # 加密私钥；必须在工具包根之外
            """
        ).strip(),
        "developer_set_location": dedent(
            rf"""
            Set-Location -LiteralPath '{developer_root}'
            F:\robot\Anaconda3\python.exe -m experiments.paper_v2.independent_blind_closeout_v2 --help
            """
        ).strip(),
        "create_root": dedent(
            r"""
            # 运行者：保管人；初始目录任意。输入开发者交付的 _v2_final 文件夹。
            $ToolkitSource = Read-Host '输入 M3_PI_Guard_Blind_Custodian_Toolkit_v2_final 的完整路径'
            $CustodianRoot = 'D:\M3_Blind_Custodian_v2'
            if (-not (Test-Path -LiteralPath $ToolkitSource -PathType Container)) {
                throw '找不到最终工具包目录。'
            }
            if (Test-Path -LiteralPath $CustodianRoot) {
                throw '目标根目录已存在；请核对后使用新的版本化目录，禁止混入旧批次。'
            }
            New-Item -ItemType Directory -Path $CustodianRoot | Out-Null
            Get-ChildItem -LiteralPath $ToolkitSource -Force | ForEach-Object {
                Copy-Item -LiteralPath $_.FullName -Destination $CustodianRoot -Recurse
            }
            Set-Location -LiteralPath $CustodianRoot
            """
        ).strip(),
        "toolkit_preflight": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'   # 若安装位置不同，只改这一行
            if ($PSVersionTable.PSVersion.Major -lt 7) {
                throw '需要 PowerShell 7 或更高版本。'
            }
            if (-not (Test-Path -LiteralPath $Python -PathType Leaf)) {
                throw '找不到 Python 3.13.5；停止，不要改用未知运行时。'
            }
            & $Python --version
            if ($LASTEXITCODE -ne 0) { throw 'Python 预检失败。' }
            & $Python -c "import cryptography; print(cryptography.__version__)"
            if ($LASTEXITCODE -ne 0) {
                throw '缺少 cryptography；按 README 从可信 wheelhouse 离线安装。'
            }
            openssl.exe version
            if ($LASTEXITCODE -ne 0) { throw '缺少 OpenSSL 3.x。' }

            $Slots = '.\protocols\paper_upgrade_v2\independent_blind_authoring_kit_v1\slots.jsonl'
            $Observed = (Get-FileHash -LiteralPath $Slots -Algorithm SHA256).Hash.ToLowerInvariant()
            $Expected = '4ec78a15bf87473d56d9c4a1bb59f6fadd0d8fed76adf66a567c5e794b79a0a8'
            if ($Observed -ne $Expected) { throw '冻结 slots SHA-256 不匹配。' }

            & '.\scripts\00_VERIFY_PRISTINE_TOOLKIT.ps1' -PythonExe $Python
            if ($LASTEXITCODE -ne 0) { throw 'toolkit manifest 校验失败。' }
            """
        ).strip(),
        "build_registration": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $KeyDir = 'D:\M3_Blind_Custodian_PrivateKey'
            $PrivateKey = Join-Path $KeyDir 'custodian_ed25519_private.pem'
            $PublicKey = '.\release\custodian_public_key.pem'
            if (-not (Test-Path -LiteralPath $KeyDir)) {
                New-Item -ItemType Directory -Path $KeyDir | Out-Null
            }
            if (Test-Path -LiteralPath $PrivateKey) { throw '私钥已存在；禁止覆盖。' }

            # OpenSSL 会安全提示输入私钥口令；口令不得写入命令行或文档。
            openssl.exe genpkey -algorithm ED25519 -aes-256-cbc -out $PrivateKey
            if ($LASTEXITCODE -ne 0) { throw 'Ed25519 私钥生成失败。' }
            openssl.exe pkey -in $PrivateKey -pubout -out $PublicKey
            if ($LASTEXITCODE -ne 0) { throw 'Ed25519 公钥导出失败。' }

            $IdentityTemplate = '.\templates\offline_identity_attestation_template.txt'
            $IdentityFile = '.\private\registration\offline_identity_attestation.txt'
            Copy-Item -LiteralPath $IdentityTemplate -Destination $IdentityFile
            notepad.exe $IdentityFile
            if (Select-String -LiteralPath $IdentityFile -Pattern '<|>' -Quiet) {
                throw '身份声明仍含占位符；停止。'
            }

            $CustodianId = Read-Host '输入真实保管人 ID（不得使用虚构身份）'
            $Relationship = Read-Host '输入与项目的关系（导师/同学/独立协作者等）'
            $ConfirmedAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\01_BUILD_UNSIGNED_REGISTRATION.ps1' `
              -PythonExe $Python -CustodianId $CustodianId `
              -ProjectRelationship $Relationship -ConfirmedAt $ConfirmedAt
            if ($LASTEXITCODE -ne 0) { throw '未签登记 payload 生成失败。' }
            """
        ).strip(),
        "sign_registration": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $PrivateKey = 'D:\M3_Blind_Custodian_PrivateKey\custodian_ed25519_private.pem'
            & '.\scripts\02_SIGN_REGISTRATION.ps1' `
              -OpenSSLExe 'openssl.exe' -PrivateKeyPath $PrivateKey
            if ($LASTEXITCODE -ne 0) { throw '登记 digest 签名失败。' }
            & '.\scripts\03_FINALIZE_REGISTRATION.ps1' -PythonExe $Python
            if ($LASTEXITCODE -ne 0) { throw '登记封存失败。' }
            """
        ).strip(),
        "verify_registration": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $VerifierId = Read-Host '输入独立签名验证人 ID'
            $VerifiedAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\04_VERIFY_REGISTRATION.ps1' `
              -PythonExe $Python -VerifierId $VerifierId -VerifiedAt $VerifiedAt
            if ($LASTEXITCODE -ne 0) { throw '登记签名验证失败；21 天时钟不得启动。' }
            """
        ).strip(),
        "split": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $AuthoredPairs = '.\private\authored_pairs\authored_pairs.jsonl'
            if (-not (Test-Path -LiteralPath $AuthoredPairs -PathType Leaf)) {
                throw '缺少 authored_pairs.jsonl；必须由第三方完成 600 行后再运行。'
            }
            & '.\scripts\05_SPLIT_PRIVATE.ps1' -PythonExe $Python
            if ($LASTEXITCODE -ne 0) { throw 'split 失败；禁止继续。' }
            """
        ).strip(),
        "pre_release": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            & '.\scripts\06_VALIDATE_PRE_RELEASE.ps1' -PythonExe $Python
            if ($LASTEXITCODE -ne 0) { throw 'pre-release 拒绝；execution set 不得交付。' }
            """
        ).strip(),
        "build_unsigned_custody": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $AuthoringAt = Read-Host '输入 authoring 完成时间（ISO-8601，含时区）'
            $SplitAt = Read-Host '输入 split 完成时间（ISO-8601，含时区）'
            $SealedAt = Read-Host '输入封存时间（ISO-8601，含时区）'
            $ReleasedAt = Read-Host '输入 execution set 释放时间（ISO-8601，含时区）'
            $DeveloperReceivedAt = Read-Host '输入开发者实际接收时间（ISO-8601，含时区）'
            & '.\scripts\07_BUILD_UNSIGNED_CUSTODY.ps1' `
              -PythonExe $Python -AuthoringCompletedAt $AuthoringAt `
              -SplitCompletedAt $SplitAt -SealedAt $SealedAt `
              -ExecutionReleasedAt $ReleasedAt `
              -DeveloperReceivedAt $DeveloperReceivedAt
            if ($LASTEXITCODE -ne 0) { throw '未签 custody 生成失败。' }
            """
        ).strip(),
        "sign_custody": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $PrivateKey = 'D:\M3_Blind_Custodian_PrivateKey\custodian_ed25519_private.pem'
            & '.\scripts\08_SIGN_CUSTODY.ps1' `
              -OpenSSLExe 'openssl.exe' -PrivateKeyPath $PrivateKey
            if ($LASTEXITCODE -ne 0) { throw 'custody digest 签名失败。' }
            & '.\scripts\09_FINALIZE_CUSTODY.ps1' -PythonExe $Python
            if ($LASTEXITCODE -ne 0) { throw 'custody 封存失败。' }
            """
        ).strip(),
        "verify_custody_signature": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $VerifierId = Read-Host '输入独立签名验证人 ID'
            $VerifiedAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\10_VERIFY_CUSTODY_SIGNATURE.ps1' `
              -PythonExe $Python -VerifierId $VerifierId -VerifiedAt $VerifiedAt
            if ($LASTEXITCODE -ne 0) { throw 'custody 签名验证失败；禁止开发者导入。' }
            """
        ).strip(),
        "finalize_custody": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            # 09 已生成 final custody；10 已完成实际数学验签。这里只读复核两项哈希：
            Get-FileHash -LiteralPath '.\release\blind_custody_v2.json' -Algorithm SHA256
            Get-FileHash -LiteralPath '.\receipts\custody_signature_verifier_receipt_v2.json' -Algorithm SHA256
            """
        ).strip(),
        "release_inventory": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            & '.\scripts\11_STAGE_DEVELOPER_HANDOFF.ps1'
            if ($LASTEXITCODE -ne 0) { throw '开发者交接白名单 staging 失败。' }
            $Handoff = '.\release\developer_handoff_v2'
            Get-ChildItem -LiteralPath $Handoff -File | Sort-Object Name |
              Select-Object Name,Length,@{Name='SHA256';Expression={
                (Get-FileHash -LiteralPath $_.FullName -Algorithm SHA256).Hash.ToLowerInvariant()
              }}
            """
        ).strip(),
        "developer_audit_import": dedent(
            rf"""
            Set-Location -LiteralPath '{developer_root}'
            $Python = 'F:\robot\Anaconda3\python.exe'
            $Release = '.\external\independent_blind_v2\release_001'
            & $Python -m {module} audit-import `
              --execution-set "$Release\blind_execution_v2.jsonl" `
              --custody-manifest "$Release\blind_custody_v2.json" `
              --pre-release-receipt "$Release\pre_release_review_v2.json" `
              --signature-verifier-receipt "$Release\custody_signature_verifier_receipt_v2.json" `
              --registration "$Release\custodian_registration_v2.json" `
              --registration-verifier-receipt "$Release\custodian_registration_verifier_receipt_v2.json" `
              --custodian-public-key "$Release\custodian_public_key.pem" `
              --write-admission '.\protocols\paper_upgrade_v2\admissions\independent_blind_execution_admission_v2.json'
            if ($LASTEXITCODE -ne 0) {{ throw '无标签导入失败；禁止启动 3,600 条。' }}
            """
        ).strip(),
        "verify_execution_completion": dedent(
            rf"""
            # 运行者：开发者；工作目录必须为 {developer_root}。
            Set-Location -LiteralPath '{developer_root}'
            $Python = 'F:\robot\Anaconda3\python.exe'
            $OutputDir = '.\outputs\paper_upgrade_v2\independent_blind_v2'
            $OutputPackage = Join-Path $OutputDir 'output_package_v2.jsonl'
            $OutputSeal = Join-Path $OutputDir 'execution_output_seal_v2.json'
            if (-not (Test-Path -LiteralPath $OutputPackage -PathType Leaf)) {{
                throw '缺少冻结六条件控制器产生的 3,600 行 output package。'
            }}
            if ((Get-Content -LiteralPath $OutputPackage).Count -ne 3600) {{
                throw 'output package 不是恰好 3,600 行；标签不得解封。'
            }}
            & $Python -m {module} seal-execution-output `
              --execution-set '.\external\independent_blind_v2\release_001\blind_execution_v2.jsonl' `
              --custody-manifest '.\external\independent_blind_v2\release_001\blind_custody_v2.json' `
              --execution-admission '.\protocols\paper_upgrade_v2\admissions\independent_blind_execution_admission_v2.json' `
              --output-package $OutputPackage `
              --completed-at ([DateTimeOffset]::Now.ToString('o')) `
              --write-seal $OutputSeal
            if ($LASTEXITCODE -ne 0) {{ throw '3,600 条 completion seal 失败；标签不得解封。' }}

            # 只把这两个内容无关文件交回保管人，分别改名/放入其固定 incoming 目录。
            Get-FileHash -LiteralPath $OutputPackage,$OutputSeal -Algorithm SHA256
            """
        ).strip(),
        "score_and_seal": dedent(
            r"""
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $IncomingSource = Read-Host '输入开发者返回两个文件所在的完整目录'
            $Incoming = '.\incoming_from_developer'
            $Expected = @('execution_output_seal_v2.json','output_package_v2.jsonl')
            foreach ($Name in $Expected) {
                $Source = Join-Path $IncomingSource $Name
                $Target = Join-Path $Incoming $Name
                if (-not (Test-Path -LiteralPath $Source -PathType Leaf)) { throw "缺少开发者返回文件：$Name" }
                if (Test-Path -LiteralPath $Target) { throw "拒绝覆盖 incoming 文件：$Name" }
                Copy-Item -LiteralPath $Source -Destination $Target
            }

            $ReleasedAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\12_RELEASE_LABELS_FOR_SCORING.ps1' `
              -PythonExe $Python -ReleasedAt $ReleasedAt
            if ($LASTEXITCODE -ne 0) { throw '标签解封门禁失败；不得评分。' }

            # 按 templates\scoring_entry_template.json 形成以下文件；每行绑定 output/label 行哈希。
            $Scores = '.\private\scoring\scoring_results_v2.jsonl'
            if (-not (Test-Path -LiteralPath $Scores -PathType Leaf)) {
                throw '缺少第三方完成的 scoring_results_v2.jsonl。'
            }
            if ((Get-Content -LiteralPath $Scores).Count -ne 3600) {
                throw '评分文件必须恰好 3,600 行。'
            }
            $ScoredAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\13_SEAL_SCORING_RESULTS.ps1' `
              -PythonExe $Python -ScoredAt $ScoredAt
            if ($LASTEXITCODE -ne 0) { throw '独立评分封存失败；不得签署结果。' }

            $SignedAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\14_BUILD_UNSIGNED_SCORING_ATTESTATION.ps1' `
              -PythonExe $Python -SignedAt $SignedAt
            if ($LASTEXITCODE -ne 0) { throw '结果签署 payload 构建失败。' }
            $PrivateKey = 'D:\M3_Blind_Custodian_PrivateKey\custodian_ed25519_private.pem'
            & '.\scripts\15_SIGN_SCORING_ATTESTATION.ps1' `
              -OpenSSLExe 'openssl.exe' -PrivateKeyPath $PrivateKey
            if ($LASTEXITCODE -ne 0) { throw '结果 manifest 的 Ed25519 签名失败。' }
            & '.\scripts\16_FINALIZE_SCORING_ATTESTATION.ps1' -PythonExe $Python
            if ($LASTEXITCODE -ne 0) { throw '已签结果 manifest 封存失败。' }
            """
        ).strip(),
        "verify_scoring_closeout": dedent(
            r"""
            # 运行者：独立验签人；工作目录：D:\M3_Blind_Custodian_v2
            Set-Location -LiteralPath 'D:\M3_Blind_Custodian_v2'
            $Python = 'C:\Python313\python.exe'
            $VerifierId = Read-Host '输入最终结果独立验签人 ID'
            $VerifiedAt = [DateTimeOffset]::Now.ToString('o')
            & '.\scripts\17_VERIFY_SCORING_ATTESTATION.ps1' `
              -PythonExe $Python -VerifierId $VerifierId -VerifiedAt $VerifiedAt
            if ($LASTEXITCODE -ne 0) { throw '最终结果签名验证失败。' }
            & '.\scripts\18_STAGE_SCORING_HANDOFF.ps1'
            if ($LASTEXITCODE -ne 0) { throw '评分交接白名单 staging 失败。' }

            $Handoff = '.\release\scoring_handoff_v2'
            Get-ChildItem -LiteralPath $Handoff -File | Sort-Object Name |
              Select-Object Name,Length,@{Name='SHA256';Expression={
                (Get-FileHash -LiteralPath $_.FullName -Algorithm SHA256).Hash.ToLowerInvariant()
              }}
            """
        ).strip(),
        "registration_message": dedent(
            """
            请确认以下事实并使用你控制的 Ed25519 密钥签署登记：
            1. 你的真实/稳定 ID 与项目关系；
            2. 未参与 M3-PI Guard 代码实现；
            3. 未参与安全阈值、策略或正式实验调参；
            4. 在 execution completion 前不向开发者释放标签；
            5. 私钥、authored pairs 和标签不进入云同步、Git 或开发者目录；
            6. 只按哈希锁定 toolkit 与 release 白名单操作；
            7. 同意在 3,600 条终态后独立评分并签署结果。
            请同时提供：custodian_id、关系说明、公钥 PEM、线下身份材料 SHA-256、确认时间和独立验证人 ID。不要提供私钥、口令、API key 或标签。
            """
        ).strip(),
    }


def main() -> int:
    path = build_document()
    print(
        json.dumps(
            {
                "status": "built",
                "path": str(path),
                "intended_final_path": str(FINAL_PATH),
                "bytes": path.stat().st_size,
                "sha256": _sha256_file(path),
                "network_calls_made": 0,
                "credentials_read": False,
                "blind_labels_read": 0,
            },
            ensure_ascii=False,
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
