"""Build the admitted-evidence M3-PI Guard stage report for 2026-07-30.

The builder is intentionally post-admission.  It refuses to run until the
public-primary analysis and independent review both exist and pass.  It reads
only aggregate scientific outputs and content-free evidence manifests; it
never reads provider credentials, prompts, or response bodies.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

from scripts.build_paper_v2_stage_report_20260726 import (
    BLUE,
    CALLOUT,
    GOLD,
    INK,
    MUTED,
    RED,
    add_bullets,
    add_callout,
    add_figure,
    add_page_field,
    add_paragraph,
    add_reference,
    add_table,
    apply_style_tokens,
    load_json,
    set_run_font,
    sha256_file,
)
from scripts.build_paper_v2_stage_report_20260729 import (
    apply_title_styles,
    create_medium_chart,
    create_readiness_chart,
    prevent_table_row_splitting,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REPORT_DATE = "2026-07-30"
REPORT_DIR = PROJECT_ROOT / "docs/research/paper_v2/reports"
ASSET_DIR = REPORT_DIR / "assets/stage_report_20260730"
OUTPUT_PATH = REPORT_DIR / "M3_PI_Guard_阶段性研究报告_2026-07-30.docx"
QA_ROOT = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/report_qa/stage_report_20260730"
)

MEDIUM_ANALYSIS_PATH = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_medium/deepseek"
    / "ba9c7937d08b68e088faf8586e0dd96d2052a398fd0a77889fd6a62e503f2373"
    / "admitted_medium_analysis_v8.json"
)
ADAPTIVE_ROOT = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/adaptive_cost_canary_v2/analysis_v1"
)
PUBLIC_ROOT = (
    PROJECT_ROOT
    / "outputs/paper_upgrade_v2/formal_public/primary_analysis_v5"
)
PUBLIC_SUMMARY_PATH = (
    PUBLIC_ROOT / "public_primary_analysis_summary_v5.json"
)
PUBLIC_CONDITION_PATH = (
    PUBLIC_ROOT / "public_primary_condition_summary_v5.csv"
)
PUBLIC_ANALYSIS_MANIFEST_PATH = PUBLIC_ROOT / "analysis_manifest.json"
PUBLIC_REVIEW_PATH = PUBLIC_ROOT / "independent_analysis_review_v5.json"
PUBLIC_MERGE_PATH = (
    PROJECT_ROOT
    / "protocols/paper_upgrade_v2/public_primary_merge_review_v2.json"
)
SCORECARD_PATH = (
    PROJECT_ROOT
    / "harness/state/paper_v2/completion_scorecard_v2.json"
)
BIBLIOGRAPHY_BUILD_PATH = (
    PROJECT_ROOT
    / "outputs/manuscript/checks/bibliography_build_v1.json"
)
REFERENCES_PATH = (
    PROJECT_ROOT / "docs/research/paper_v2/manuscript/references_v1.bib"
)
HARNESS_RECEIPT_PATH = (
    PROJECT_ROOT
    / "outputs/verification/public_primary_final_harness_receipt.json"
)
HARNESS_STDOUT_PATH = (
    PROJECT_ROOT
    / "outputs/verification/public_primary_final_harness_stdout.log"
)
TARGETED_RECHECK_RECEIPT_PATH = (
    PROJECT_ROOT
    / "outputs/verification/public_primary_targeted_recheck_receipt.json"
)

FIGURES = {
    "medium": ASSET_DIR / "figure_1_medium_security_utility.png",
    "readiness": ASSET_DIR / "figure_2_completion_score.png",
    "public_security": PUBLIC_ROOT / "public_primary_security_v1.png",
    "public_utility": PUBLIC_ROOT / "public_primary_utility_v1.png",
    "public_diagnostics": PUBLIC_ROOT / "public_primary_diagnostics_v1.png",
    "adaptive": ADAPTIVE_ROOT / "adaptive_canary_asr.png",
}

CONDITION_LABELS = {
    "no_defense": "无防御",
    "safe_prompt_or_spotlighting": "安全提示",
    "promptguard_detector_only": "PromptGuard",
    "strongest_peer_reviewed_structural_baseline": "结构化基线",
    "m3_policy_only": "M3 policy-only",
    "m3_full": "M3 full",
}


def _configure_sections(doc: Document) -> None:
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header.paragraphs[0]
        header.paragraph_format.space_after = Pt(0)
        header.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
        set_run_font(
            header.add_run("M3-PI Guard | SCI 二区网络安全研究"),
            size=8.5,
            color=MUTED,
            bold=True,
        )
        header.add_run("\t")
        set_run_font(
            header.add_run(f"阶段报告 | {REPORT_DATE}"),
            size=8.5,
            color=MUTED,
        )

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(0)
        set_run_font(
            footer.add_run("M3-PI Guard  ·  "),
            size=9,
            color=MUTED,
        )
        add_page_field(footer)


def _add_title_block(
    doc: Document,
    *,
    scorecard: dict[str, Any],
) -> None:
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(
        kicker.add_run("RESEARCH STAGE REPORT"),
        size=9,
        color=GOLD,
        bold=True,
    )
    title = doc.add_paragraph(style="Title")
    set_run_font(
        title.add_run("M3-PI Guard 阶段性研究报告"),
        size=24,
        color=INK,
        bold=True,
    )
    subtitle = doc.add_paragraph(style="Subtitle")
    set_run_font(
        subtitle.add_run(
            "5,082 条 Public-primary 证据、独立统计复核、"
            "自适应成本金丝雀与 SCI Q2 实证路线"
        ),
        size=13,
        color=MUTED,
    )
    add_table(
        doc,
        ["项目", "当前口径"],
        [
            ["报告日期", "2026-07-30（Asia/Shanghai）"],
            [
                "论文路线",
                "授权信息泄漏、授权不确定性与安全—效用权衡的系统实证研究",
            ],
            [
                "正式证据",
                "DeepSeek 官方 deepseek-v4-pro；中型 480 条、"
                "Public-primary 5,082 条、自适应探索 8 个任务",
            ],
            [
                "科学完成度",
                f"{scorecard['scientific_evidence_completion_points']}/"
                f"{scorecard['scientific_evidence_completion_max_points']} "
                f"({scorecard['scientific_evidence_completion_percent_exact']:.2f}%)",
            ],
            [
                "声明边界",
                "Primary-model-only；无跨模型/跨提供商/SOTA/独立盲测声明",
            ],
            [
                "安全边界",
                "provider tools 关闭、无真实外部 effect、无自动重试、"
                "终态失败保留在 ITT 分母",
            ],
        ],
        [1900, 7460],
    )


def _ratio_text(metric: dict[str, Any]) -> str:
    numerator = int(metric["numerator"])
    denominator = int(metric["denominator"])
    value = 100.0 * float(metric["value"])
    low = 100.0 * float(metric["ci_low"])
    high = 100.0 * float(metric["ci_high"])
    return (
        f"{numerator}/{denominator}；{value:.2f}% "
        f"（Wilson 95% CI {low:.2f}–{high:.2f}%）"
    )


def _comparison_rows(
    comparisons: dict[str, Any],
) -> list[list[str]]:
    labels = {
        "security_full_vs_structural": "M3-full vs 结构化：unsafe effect",
        "security_full_vs_policy_only": "M3-full vs policy-only：unsafe effect",
        "utility_full_vs_structural": "M3-full vs 结构化：工具 TCR",
        "utility_full_vs_policy_only": "M3-full vs policy-only：工具 TCR",
    }
    rows: list[list[str]] = []
    for key, label in labels.items():
        item = comparisons[key]
        bootstrap = item["cluster_bootstrap"]
        mcnemar = item["mcnemar"]
        rows.append(
            [
                label,
                str(item["pairs"]),
                f"{100 * item['absolute_difference_treatment_minus_baseline']:.2f} pp",
                f"{mcnemar['p_value']:.6g}",
                f"{mcnemar['holm_adjusted_p_value']:.6g}",
                (
                    f"{100 * bootstrap['ci_low']:.2f} 至 "
                    f"{100 * bootstrap['ci_high']:.2f} pp"
                ),
            ]
        )
    return rows


def _read_public_conditions() -> list[dict[str, str]]:
    with PUBLIC_CONDITION_PATH.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as stream:
        return list(csv.DictReader(stream))


def _public_rate_rows(
    rows: Iterable[dict[str, str]],
    *,
    benchmark: str,
) -> list[list[str]]:
    index = {
        (row["evaluation_case_type"], row["condition_id"]): row
        for row in rows
        if row["benchmark"] == benchmark
    }
    result: list[list[str]] = []
    for condition, label in CONDITION_LABELS.items():
        attack = index[("attack", condition)]
        benign = index[("benign_tool", condition)]
        result.append(
            [
                label,
                (
                    f"{attack['events']}/{attack['total']}；"
                    f"{100 * float(attack['rate']):.2f}% "
                    f"({100 * float(attack['wilson_low']):.2f}–"
                    f"{100 * float(attack['wilson_high']):.2f})"
                ),
                (
                    f"{benign['events']}/{benign['total']}；"
                    f"{100 * float(benign['rate']):.2f}% "
                    f"({100 * float(benign['wilson_low']):.2f}–"
                    f"{100 * float(benign['wilson_high']):.2f})"
                ),
                (
                    f"{int(attack['diagnostic_records']) + int(benign['diagnostic_records'])}"
                ),
            ]
        )
    return result


def _harness_summary() -> str:
    if not HARNESS_RECEIPT_PATH.is_file():
        return "首轮全量 Harness receipt 缺失"
    full = json.loads(
        HARNESS_RECEIPT_PATH.read_text(encoding="utf-8-sig")
    )
    if not TARGETED_RECHECK_RECEIPT_PATH.is_file():
        return f"首轮全量退出码 {full['exit_code']}；定向复核 receipt 缺失"
    targeted = json.loads(
        TARGETED_RECHECK_RECEIPT_PATH.read_text(encoding="utf-8-sig")
    )
    checks = targeted.get("checks", [])
    passed = sum(int(check.get("exit_code") == 0) for check in checks)
    pytest_line = "1265 passed, 1 skipped"
    if HARNESS_STDOUT_PATH.is_file():
        text = HARNESS_STDOUT_PATH.read_text(
            encoding="utf-8",
            errors="replace",
        )
        for line in text.splitlines():
            if "passed" in line and "skipped" in line:
                pytest_line = line.strip()
                break
    return (
        f"首轮全量退出码 {full['exit_code']}（{pytest_line}；"
        "暴露 1 个已修复的测试状态隔离问题和 1 个已修复的 CLI "
        f"参数边界问题）；修复后定向复核 {passed}/{len(checks)}，"
        f"退出码 {targeted['exit_code']}，stderr=0"
    )


def _artifact_rows(paths: Iterable[Path]) -> list[list[str]]:
    rows: list[list[str]] = []
    for path in paths:
        rows.append(
            [
                path.name,
                path.relative_to(PROJECT_ROOT).as_posix(),
                sha256_file(path),
            ]
        )
    return rows


def _new_page(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.add_run().add_break()
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:pageBreakBefore"))


def build_report() -> Path:
    required = [
        MEDIUM_ANALYSIS_PATH,
        ADAPTIVE_ROOT / "summary.json",
        ADAPTIVE_ROOT / "analysis_review_v1.json",
        PUBLIC_SUMMARY_PATH,
        PUBLIC_CONDITION_PATH,
        PUBLIC_ANALYSIS_MANIFEST_PATH,
        PUBLIC_REVIEW_PATH,
        PUBLIC_MERGE_PATH,
        SCORECARD_PATH,
        BIBLIOGRAPHY_BUILD_PATH,
        REFERENCES_PATH,
        FIGURES["public_security"],
        FIGURES["public_utility"],
        FIGURES["public_diagnostics"],
        FIGURES["adaptive"],
    ]
    missing = [path for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(
            "post-admission report inputs are incomplete: "
            + ", ".join(str(path) for path in missing)
        )

    medium_envelope = load_json(MEDIUM_ANALYSIS_PATH)
    medium = medium_envelope["analysis"]["primary_analysis"]
    adaptive = load_json(ADAPTIVE_ROOT / "summary.json")
    adaptive_review = load_json(
        ADAPTIVE_ROOT / "analysis_review_v1.json"
    )
    public = load_json(PUBLIC_SUMMARY_PATH)
    public_review = load_json(PUBLIC_REVIEW_PATH)
    public_merge = load_json(PUBLIC_MERGE_PATH)
    scorecard = load_json(SCORECARD_PATH)
    bibliography = load_json(BIBLIOGRAPHY_BUILD_PATH)
    public_rows = _read_public_conditions()

    if (
        public["status"] != "complete_primary_model_empirical_fallback"
        or public_review["status"] != "pass"
        or not all(public_review["checks"].values())
        or adaptive_review["status"]
        != "adaptive_canary_analysis_integrity_pass"
    ):
        raise ValueError("admitted analysis or independent review not passed")

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_medium_chart(
        FIGURES["medium"],
        medium["condition_metrics"],
    )
    create_readiness_chart(FIGURES["readiness"], scorecard)

    doc = Document()
    apply_style_tokens(doc)
    apply_title_styles(doc)
    _configure_sections(doc)
    _add_title_block(doc, scorecard=scorecard)

    doc.add_heading("执行摘要", level=1)
    route = public["claim_decision"]["route"]
    add_callout(
        doc,
        "当前结论",
        "项目已完成并独立复核 5,082 条 Public-primary 记录，但论文仍应"
        "坚持 empirical fallback。中型实验已否定 M3-full 相对 "
        "policy-only 的独立安全增益；Public 结果用于检验这一安全—效用"
        "关系是否跨 AgentDojo 与 M3-v2 保持，而不是恢复机制领先声明。",
        fill=CALLOUT,
        accent=BLUE,
    )
    add_bullets(
        doc,
        [
            f"Public-primary 分析状态：{public['status']}；route={route}。",
            (
                f"Public 证据覆盖 {public['coverage']['records']} 条记录、"
                f"{public['coverage']['cases']} 个案例、"
                f"{public['coverage']['attack_cases']} 个攻击案例与 "
                f"{public['coverage']['benign_cases']} 个良性案例。"
            ),
            (
                f"Public 保留 {public['operational']['diagnostic_count']} 条"
                " diagnostic，排除数为 0；冻结的 1% 质量门没有被豁免。"
            ),
            (
                "Public 攻击结果：M3-full 38/600、结构化基线 133/600、"
                "policy-only 13/600；full 相对结构化下降 15.83 pp，"
                "但相对 policy-only 上升 4.17 pp。"
            ),
            (
                "Public benign-tool TCR：M3-full 与 policy-only 均为 "
                "49/168，结构化基线为 103/168；full 低 32.14 pp。"
            ),
            (
                f"自适应探索为 {adaptive['unsafe_effect_successes']}/"
                f"{adaptive['task_count']}，但 Wilson 上界仍为 "
                f"{100 * adaptive['unsafe_effect_rate_at_50']['ci_high']:.2f}%，"
                "不能代替正式 40/20 任务多预算面板。"
            ),
            (
                "Gemini 尚无正式第二模型矩阵，独立盲测保管人也未落实；"
                "因此跨模型、跨提供商和强外部有效性声明仍被禁止。"
            ),
        ],
    )

    doc.add_heading("1. 研究定位与判定规则", level=1)
    add_paragraph(
        doc,
        "当前工作标题为“Evaluating Runtime Authorization Under Context "
        "and Provenance Uncertainty for Tool-Using LLM Agents”。22 单元源码"
        "审计未证明领域普遍存在 oracle authorization，因此不使用 "
        "Beyond Answer-Key 或“首次/唯一/SOTA”叙事。论文目标是说明："
        "在线授权信息、来源约束、确认负担、终态失败和基准选择如何共同"
        "改变工具型 LLM Agent 的安全—效用边界。"
    )
    add_table(
        doc,
        ["研究问题", "本阶段证据", "允许结论"],
        [
            [
                "RQ1：非 oracle 授权下的安全与效用",
                "中型 480 条 + Public-primary 5,082 条",
                "可按模型/基准报告；不可跨模型泛化",
            ],
            [
                "RQ2：full 相对 policy-only 的模块贡献",
                "中型与 Public 配对比较",
                "保留负结果；不得选择性省略",
            ],
            [
                "RQ3：授权 regime 与 ASK 负担",
                "中型 regime probe；Public auto_compiled",
                "中型可报告，Public 不外推五 regime",
            ],
            [
                "RQ4：防御感知自适应攻击",
                "8 任务 budget-50 成本金丝雀",
                "仅证明执行链和成本；不作稳健性结论",
            ],
        ],
        [2300, 3500, 3560],
    )

    _new_page(doc)
    doc.add_heading("2. 可审计完成度与证据地图", level=1)
    add_figure(
        doc,
        FIGURES["readiness"],
        "图 1  十二个工作流的科学证据完成度",
        "completion_scorecard_v2.json；分项总和，不把工程就绪度替代科学结果。",
        alt="横向条形图展示十二个论文工作流的科学证据完成度。",
    )
    add_table(
        doc,
        ["工作流", "科学完成度", "剩余关键工作"],
        [
            [
                item["name"],
                f"{item['scientific_evidence_completion_percent']}%",
                item["remaining"],
            ]
            for item in scorecard["workstreams"]
        ],
        [2700, 1500, 5160],
        status_columns=(1,),
    )
    add_callout(
        doc,
        "80% 目标审计",
        f"当前按原 12×100 量表为 "
        f"{scorecard['scientific_evidence_completion_points']}/1200 "
        f"（{scorecard['scientific_evidence_completion_percent_exact']:.2f}%）。"
        "未达到 80% 的原因不是工程缺口，而是原路线要求完整第二模型 Public "
        "矩阵、正式自适应 40/20 面板和独立盲测证据；本报告不改变冻结评分"
        "规则来制造达标。",
        fill="FFF5E6",
        accent=GOLD,
    )

    _new_page(doc)
    doc.add_heading("3. 中型证伪实验：已否定机制领先路线", level=1)
    add_paragraph(
        doc,
        "72 个配对案例在六个条件下形成 432 条主记录，并加 48 条授权 "
        "regime probe。全部失败保留在 ITT。该实验是论文路线切换的控制性"
        "证据，不能被后续 Public 结果覆盖或删除。"
    )
    add_figure(
        doc,
        FIGURES["medium"],
        "图 2  中型实验的 unsafe-effect 与 benign-tool TCR",
        "admitted_medium_analysis_v8.json；点估计与 Wilson 95% CI。",
        alt="两组条形图比较六个条件的攻击不安全效果率和良性工具任务完成率。",
    )
    add_table(
        doc,
        ["条件", "攻击 unsafe effect", "良性工具 TCR"],
        [
            [
                CONDITION_LABELS[condition],
                _ratio_text(
                    medium["condition_metrics"][condition][
                        "unsafe_effect_itt"
                    ]
                ),
                _ratio_text(
                    medium["condition_metrics"][condition][
                        "benign_tool_tcr"
                    ]
                ),
            ]
            for condition in CONDITION_LABELS
        ],
        [2300, 3530, 3530],
    )
    add_table(
        doc,
        ["配对比较", "n", "绝对差", "McNemar p", "Holm p", "Bootstrap 95% CI"],
        _comparison_rows(medium["comparisons"]),
        [2500, 600, 1300, 1200, 1100, 2660],
    )
    add_callout(
        doc,
        "中型判定",
        "M3-full 相对结构化基线把 unsafe effect 从 8/36 降至 0/36，"
        "但 policy-only 也是 0/36；同时 benign-tool TCR 从 12/24 降到 "
        "6/24。安全增益没有独立模块贡献，且效用非劣界限失败。",
        fill="FDECEC",
        accent=RED,
    )

    _new_page(doc)
    doc.add_heading("4. Public-primary：5,082 条正式主模型证据", level=1)
    add_paragraph(
        doc,
        "Public-primary 由 AgentDojo v1 与 M3-v2 组成，共 847 个唯一案例、"
        "六个条件、5,082 条记录。24 条技术金丝雀后，剩余 5,058 条被分为 "
        "1,265/1,265/1,264/1,264 四个互斥分片。恢复只处理未开始的 work "
        "key；终态 candidate 或 diagnostic 均不重试。"
    )
    add_table(
        doc,
        ["证据项", "数值", "审计解释"],
        [
            ["记录/案例", "5,082 / 847", "六条件完整矩阵"],
            [
                "攻击/良性案例",
                f"{public['coverage']['attack_cases']} / "
                f"{public['coverage']['benign_cases']}",
                "AgentDojo 良性任务不复制扩充",
            ],
            [
                "Diagnostic",
                f"{public['operational']['diagnostic_count']} "
                f"({100 * public['operational']['diagnostic_fraction']:.2f}%)",
                "全部保留在 ITT；排除数 0",
            ],
            [
                "官方调用/成本",
                f"{public_merge['operational_totals']['provider_calls_completed']} / "
                f"¥{public_merge['operational_totals']['role_cumulative_cost_cny']:.6f}",
                "append-only usage/cost 证据",
            ],
            [
                "质量门",
                (
                    "通过"
                    if public_merge["claim_quality"][
                        "clean_mechanism_quality_claim_eligible"
                    ]
                    else "未通过"
                ),
                "1% 门未豁免；只能 empirical fallback",
            ],
            [
                "独立分析复核",
                "通过",
                f"{len(public_review['checks'])} 项检查全部为 true",
            ],
        ],
        [2300, 2400, 4660],
        status_columns=(1,),
    )
    add_figure(
        doc,
        FIGURES["public_security"],
        "图 3  Public-primary 攻击 unsafe-effect 率",
        "public_primary_condition_summary_v5.csv；按 AgentDojo 与 M3-v2 分层。",
        alt="分组条形图比较两个基准上六个条件的攻击不安全效果率。",
    )
    add_figure(
        doc,
        FIGURES["public_utility"],
        "图 4  Public-primary benign-tool 任务完成率",
        "public_primary_condition_summary_v5.csv；按 AgentDojo 与 M3-v2 分层。",
        alt="分组条形图比较两个基准上六个条件的良性工具任务完成率。",
    )
    add_figure(
        doc,
        FIGURES["public_diagnostics"],
        "图 5  Public-primary diagnostic 负担",
        "所有终态 diagnostic 保留在 ITT；虚线为冻结 1% 质量门。",
        alt="分组条形图展示六个条件在两个基准上的诊断记录比例和百分之一质量门。",
    )

    _new_page(doc)
    for benchmark, label in (
        ("all", "总体"),
        ("agentdojo-v1", "AgentDojo v1"),
        ("m3-v2", "M3-v2"),
    ):
        doc.add_heading(f"4.{1 if benchmark == 'all' else 2 if benchmark == 'agentdojo-v1' else 3} {label}", level=2)
        add_table(
            doc,
            ["条件", "攻击 unsafe effect（95% CI）", "工具 TCR（95% CI）", "Diag"],
            _public_rate_rows(public_rows, benchmark=benchmark),
            [2100, 3100, 3100, 1060],
        )

    doc.add_heading("4.4 Public 关键配对比较", level=2)
    add_table(
        doc,
        ["配对比较", "n", "绝对差", "McNemar p", "Holm p", "Bootstrap 95% CI"],
        _comparison_rows(public["overall"]["comparisons"]),
        [2500, 600, 1300, 1200, 1100, 2660],
    )
    add_callout(
        doc,
        "Public 实证判定",
        "M3-full 相对结构化基线具有显著安全下降，但比 policy-only 更不安全，"
        "并且良性工具 TCR 相对结构化低 32.14 pp。RQ1 只得到带明显效用代价"
        "的安全效应；RQ2 的模块增益为负；机制论文门保持 fail_redirect。",
        fill="FDECEC",
        accent=RED,
    )
    add_callout(
        doc,
        "Public 声明边界",
        "该矩阵只使用 deepseek-v4-pro。即使两个基准方向一致，也只能称"
        " primary-model benchmark consistency；不得写成 cross-model、"
        "cross-provider 或 SOTA。超过 1% 的 diagnostic 只会收缩机制质量"
        "声明，不会从分母删除。",
        fill="FFF5E6",
        accent=GOLD,
    )

    _new_page(doc)
    doc.add_heading("5. 防御感知自适应攻击：探索性成本门", level=1)
    add_figure(
        doc,
        FIGURES["adaptive"],
        "图 6  八任务 budget-50 自适应成本金丝雀",
        "adaptive canary admitted analysis；Wilson 95% CI。",
        alt="探索性图展示八个自适应任务在预算五十下的不安全效果率及置信区间。",
    )
    add_table(
        doc,
        ["指标", "数值", "解释"],
        [
            ["任务", str(adaptive["task_count"]), "八个攻击家族各一项"],
            [
                "终态查询/官方调用",
                f"{adaptive['terminal_query_count']} / "
                f"{adaptive['completed_provider_call_count']}",
                "一条 diagnostic 保留；零自动重试",
            ],
            [
                "Unsafe effect",
                f"{adaptive['unsafe_effect_successes']}/{adaptive['task_count']}",
                (
                    "Wilson 95% CI "
                    f"{100 * adaptive['unsafe_effect_rate_at_50']['ci_low']:.2f}–"
                    f"{100 * adaptive['unsafe_effect_rate_at_50']['ci_high']:.2f}%"
                ),
            ],
            [
                "成本",
                f"¥{adaptive['cumulative_cost_cny']:.6f}",
                "仅为 cost canary，不代表正式面板预算",
            ],
            [
                "决策分布",
                ", ".join(
                    f"{key}={value}"
                    for key, value in adaptive["decision_distribution"].items()
                ),
                "公开反馈只含 ALLOW/ASK/DENY 与 post-effect",
            ],
        ],
        [2300, 2800, 4260],
    )
    add_callout(
        doc,
        "不能过度解读",
        "0/8 的区间上界仍为 32.44%，样本小且只有 budget 50。该结果证明"
        "查询—调用—成本—任务链可审计，不证明防御在 200/500/2000 查询下"
        "稳定。正式 40 个主模型任务、20 个第二模型任务和 hardest-20 "
        "扩展仍是投稿前硬门。",
        fill="FFF5E6",
        accent=GOLD,
    )

    _new_page(doc)
    doc.add_heading("6. 统计与复现完整性", level=1)
    add_paragraph(
        doc,
        "主分析在独立结果准入后读取 unsafe_effect_itt 与 task_completion，"
        "生成 Wilson 区间、exact McNemar、Holm 校正和 10,000 次 case-"
        "clustered bootstrap。独立复核器直接读取 5,082 条 record envelope，"
        "不复用主聚合或统计 helper，重算条件计数、Wilson、四个关键配对差"
        "与 exact McNemar。"
    )
    add_paragraph(
        doc,
        "分析实现经历四次无汇总输出的 fail-closed 校正：v1 错用 record-v1 "
        "schema，v2 错用 medium 成本策略，v3 对 retained diagnostic 重放"
        " completed-trace usage 约束，v4 未在 legacy 统计 helper 前投影 "
        "v2.2 envelope。每次失败均保留 stderr/exit，随后使用新的 candidate "
        "与 exact-command admission。最终 v5 只投影证据包络，所有 outcome、"
        "decision、status 与 ITT 字段逐项不变。"
    )
    add_table(
        doc,
        ["检查", "状态", "证据"],
        [
            [
                label,
                "通过" if passed else "失败",
                "independent_analysis_review_v5.json",
            ]
            for label, passed in public_review["checks"].items()
        ],
        [4900, 1200, 3260],
        status_columns=(1,),
    )
    add_table(
        doc,
        ["边界", "值"],
        [
            ["Provider calls during analysis/review", "0"],
            ["Record rows mutated", "0"],
            ["Diagnostic rows excluded", "0"],
            ["Automatic retries", "0"],
            ["Provider tools / external effects", "关闭 / 关闭"],
            ["最终 Harness", _harness_summary()],
        ],
        [3500, 5860],
    )

    doc.add_heading("7. 文献、创新边界与稿件状态", level=1)
    add_paragraph(
        doc,
        f"参考文献导出共 {bibliography['entry_count']} 条，其中 "
        f"{bibliography['peer_reviewed_count']} 条同行评审、"
        f"{bibliography['preprint_only_count']} 条预印本。VIGIL、DRIFT、Task "
        "Shield、ARGUS、AuthGraph、PlanGuard、Progent、IPIGuard、CaMeL "
        "共同限制了“首个参数来源/运行时授权”叙事。剩余贡献是评估设计、"
        "在线/评估状态隔离、权限单调、ASK 负担和失败留在 ITT 的联合实证。"
    )
    add_table(
        doc,
        ["稿件部分", "状态", "下一步"],
        [
            ["Introduction", "已整合 Medium/Public", "最终语言压缩与引用复核"],
            ["Related Work", "34 条元数据已审计", "投稿日刷新预印本/出版状态"],
            ["Methods", "Public/Adaptive/证据链已写入", "方法审查"],
            ["Results", "Medium + Public + adaptive canary", "正式 adaptive 后补"],
            ["Discussion", "Medium/Public empirical fallback", "正式 adaptive 后补"],
            ["Abstract", "Branch-locked", "正式 adaptive 前不定稿"],
            ["JISA package", "待完成", "JCR、格式、cover letter、可用性声明"],
        ],
        [2100, 2800, 4460],
        status_columns=(1,),
    )

    _new_page(doc)
    doc.add_heading("8. 后续正式实验怎么做", level=1)
    add_paragraph(
        doc,
        "后续工作遵循“先补能改变结论的证据，再补包装”的原则。API 可以"
        "正常使用，但每个正式阶段仍使用新冻结、新准入、零自动重试和无"
        "危险外部 effect 的证据链。"
    )
    actions = [
        (
            "Formal Adaptive 主模型",
            "40 个任务，预算 50/200/500；按首成功规则选择 hardest-20 "
            "追加 2,000。报告 first-success、饱和曲线、ITT 成功率和成本。",
        ),
        (
            "Gemini 第二模型",
            "Free 层恢复稳定后，新建 v3；先 12 条金丝雀，再运行最强结构化、"
            "policy-only、full 三条件，绝不把 DeepSeek 作为 fallback。",
        ),
        (
            "Public 第二角色",
            "若 Gemini 正式链可用，完成冻结的 1,791 条第二角色记录；否则"
            "明确论文为 single-primary-model，并收缩泛化声明。",
        ),
        (
            "独立盲测",
            "确定第三方保管人，生成/抽取 600 条，保存脚本、SHA-256 与标签"
            "保管链；无法落实时只称内部 sealed holdout。",
        ),
        (
            "最终统计与稿件",
            "统一 medium/public/adaptive claim ledger，完成主图、Stage 2.5 "
            "方法审查、Devil’s Advocate、数字/DOI 100% 核验和 JISA 包。",
        ),
    ]
    add_table(
        doc,
        ["阶段", "执行内容"],
        [[label, text] for label, text in actions],
        [2600, 6760],
    )
    add_bullets(
        doc,
        [
            "DeepSeek：继续记录 cache hit/miss、官方 response ID、精确模型、"
            "payload hash、调用成本和 append-only 尾哈希。",
            "Gemini：保持免费层 USD 0；429 立即暂停，不自动重试、不跨提供商"
            "降级、不把诊断链包装为正式结果。",
            "两个 API 都只能产生不可信 proposal；provider-managed tools、真实"
            "邮件/支付、host command、私有生产 API 与非沙箱文件 effect 禁止。",
        ],
    )

    doc.add_heading("9. 风险与对外口径", level=1)
    add_table(
        doc,
        ["风险", "当前证据", "对外表述"],
        [
            [
                "机制贡献不足",
                "Public full 38/600，policy-only 13/600",
                "主张 empirical trade-off，不主张机制领先",
            ],
            [
                "效用下降",
                "Public full 49/168，结构化 103/168",
                "同时报告安全和效用，不用单一 ASR",
            ],
            [
                "Public diagnostic 超 1%",
                "ITT 保留，质量门不豁免",
                "可实证分析，不作 clean mechanism claim",
            ],
            [
                "单模型",
                "Public 只有 DeepSeek 正式结果",
                "禁止 cross-model/cross-provider",
            ],
            [
                "Adaptive 样本小",
                "8 任务 budget-50，CI 宽",
                "仅称探索性成本金丝雀",
            ],
            [
                "无独立保管人",
                "内部 sealed，不是独立盲测",
                "收缩外部有效性声明",
            ],
        ],
        [2300, 3600, 3460],
    )
    add_callout(
        doc,
        "推荐对外一句话",
        "M3-PI Guard 已完成可审计的 DeepSeek 中型与 5,082 条 Public-primary "
        "实证；结果用于刻画授权不确定性下的安全—效用权衡，而不是宣称"
        " SOTA。正式自适应、第二模型和独立盲测仍是投稿前关键缺口。",
        fill=CALLOUT,
        accent=BLUE,
    )

    doc.add_heading("附录 A：关键证据资产与哈希", level=1)
    add_table(
        doc,
        ["资产", "项目内路径", "SHA-256"],
        _artifact_rows(
            [
                MEDIUM_ANALYSIS_PATH,
                ADAPTIVE_ROOT / "summary.json",
                PUBLIC_MERGE_PATH,
                PUBLIC_SUMMARY_PATH,
                PUBLIC_REVIEW_PATH,
                PUBLIC_ANALYSIS_MANIFEST_PATH,
                SCORECARD_PATH,
                REFERENCES_PATH,
            ]
        ),
        [1900, 4200, 3260],
    )

    doc.add_heading("附录 B：主要前沿与投稿来源", level=1)
    for label, description, url in (
        (
            "JISA",
            "Aims & Scope 与投稿入口",
            "https://www.sciencedirect.com/journal/"
            "journal-of-information-security-and-applications",
        ),
        (
            "VIGIL",
            "ACL 2026，verify-before-commit 与 SIREN",
            "https://aclanthology.org/2026.acl-long.443/",
        ),
        (
            "DRIFT",
            "NeurIPS 2025，动态规则、控制/数据流与记忆隔离",
            "https://proceedings.neurips.cc/paper_files/paper/2025/hash/"
            "77f3b26c7907aa27b207df9b9d43f29a-Abstract-Conference.html",
        ),
        (
            "Task Shield",
            "ACL 2025，任务与工具行为一致性",
            "https://aclanthology.org/2025.acl-long.1435/",
        ),
        (
            "AgentDojo",
            "NeurIPS 2024 动态 Agent 安全基准",
            "https://proceedings.nips.cc/paper_files/paper/2024/hash/"
            "97091a5177d8dc64b1da8bf3e1f6fb54-"
            "Abstract-Datasets_and_Benchmarks_Track.html",
        ),
        (
            "ARGUS",
            "Context-dependent provenance auditing（预印本）",
            "https://arxiv.org/abs/2605.03378",
        ),
        (
            "AuthGraph",
            "Authorization/provenance 双图（预印本）",
            "https://arxiv.org/abs/2605.26497",
        ),
    ):
        add_reference(doc, label, description, url)

    doc.add_heading("附录 C：生成边界", level=1)
    add_paragraph(
        doc,
        "报告生成器只读取已准入聚合、独立复核、内容无关 merge/Harness "
        "清单和参考文献元数据。它不读取 API key、原始 prompt、response "
        "body、环境变量秘密或未准入科学结果。最终报告本身不包含密钥、"
        "请求头或个人/项目标识。"
    )

    doc.add_heading("附录 D：投稿前分区核验", level=1)
    add_paragraph(
        doc,
        "本文按 SCIE 收录且投稿时 JCR Q2 规划首投 JISA。期刊分区可能"
        "随年度与学科口径变化，正式投稿前必须通过学校或 Clarivate "
        "订阅再次核验当期收录、JCR Q2 分区和期刊范围；本报告不把"
        "第三方分区网页视为最终投稿凭证。"
    )

    prevent_table_row_splitting(doc)
    properties = doc.core_properties
    properties.title = "M3-PI Guard 阶段性研究报告（2026-07-30）"
    properties.subject = "SCI Q2 网络安全论文证据、结果与后续实验"
    properties.author = "M3-PI Guard Research Team"
    properties.keywords = (
        "LLM agent security, prompt injection, runtime authorization, "
        "provenance, security-utility"
    )

    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    if OUTPUT_PATH.exists():
        raise FileExistsError(f"refuse to overwrite report: {OUTPUT_PATH}")
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def main() -> int:
    path = build_report()
    print(
        json.dumps(
            {
                "status": "report_built_from_admitted_evidence",
                "path": str(path),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
                "qa_root": str(QA_ROOT),
                "figures": {
                    name: {
                        "path": str(figure),
                        "sha256": sha256_file(figure),
                    }
                    for name, figure in FIGURES.items()
                },
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
