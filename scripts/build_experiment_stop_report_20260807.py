"""Build the source-backed M3-PI Guard experiment-stop stage report.

The builder reads only admitted summaries, immutable technical reviews, and
process receipts. It performs no provider calls and never reads credentials,
case-level blind labels, query ledgers, or provider traces.
"""

from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_stage_technical_report_docx import (
    DARK_BLUE,
    INK,
    LIGHT_GRAY,
    MUTED,
    PALE_RED,
    RISK,
    WHITE,
    add_callout,
    add_heading,
    add_metric_strip,
    add_numbered_items,
    add_page_number,
    add_paragraph,
    add_source_note,
    add_table,
    configure_styles,
    pil_font,
    set_picture_alt,
    set_repeat_table_header,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs/research/paper_v2/reports/M3_PI_Guard_实验任务中止阶段报告_2026-08-07.docx"
QA_DIR = ROOT / "outputs/report_qa/experiment_stop_report_20260807"


def read_json(relative: str) -> dict:
    return json.loads((ROOT / relative).read_text(encoding="utf-8-sig"))


def sha256(relative: str) -> str:
    digest = hashlib.sha256()
    with (ROOT / relative).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def collect_evidence() -> dict:
    v9_path = "outputs/verification/experiment_closeout_fresh_harness_v9_20260807/receipt.json"
    v10_path = "outputs/verification/experiment_closeout_fresh_harness_v10_20260807/receipt.json"
    v11_path = "outputs/verification/experiment_closeout_fresh_harness_v11_20260807/receipt.json"
    primary_path = (
        "outputs/paper_upgrade_v2/formal_public/primary_analysis_v5/"
        "public_primary_analysis_summary_v5.json"
    )
    primary_review_path = (
        "outputs/paper_upgrade_v2/formal_public/primary_analysis_v5/"
        "independent_analysis_review_v5.json"
    )
    second_path = "protocols/paper_upgrade_v2/public_second_evidence_merge_v1.json"
    gemini_s0_receipt_path = (
        "outputs/verification/gemini_public_v5_recovery_s0_2026-08-05_guarded_v2/receipt.json"
    )
    gemini_s1_receipt_path = (
        "outputs/verification/gemini_public_v5_recovery_s1_2026-08-06_guarded_v2/receipt.json"
    )
    gemini_s0_review_path = (
        "outputs/paper_upgrade_v2/formal_public/gemini/"
        "public-v5-canary-recovery-s0-885cf3a88684ca7427a29a5c0b65e0a6f97746133279fb59341bae557c4b48b2/"
        "technical_segment_review_v6.json"
    )
    gemini_s1_review_path = (
        "outputs/paper_upgrade_v2/formal_public/gemini/"
        "public-v5-canary-recovery-s1-b28e6af8206cc66283d10ce73b220c839f9ad1d48393a40de33c6d4be31384b4/"
        "technical_segment_review_v6.json"
    )
    pro50_receipt_path = (
        "outputs/verification/adaptive_formal_pro_50_v2_guarded_v3_20260806/receipt.json"
    )
    pro200_path = "protocols/paper_upgrade_v2/adaptive_formal_pro_200_host_timeout_review_v1.json"
    pro500_path = "protocols/paper_upgrade_v2/adaptive_formal_pro_500_candidate_v2.json"
    pro500_admission_path = (
        "protocols/paper_upgrade_v2/admissions/adaptive_formal_pro_500_admission_v2.json"
    )
    gate7_path = "protocols/paper_upgrade_v2/experiment_closeout_harness_gate_v7.json"
    v10_review_path = (
        "protocols/paper_upgrade_v2/experiment_closeout_fresh_harness_failure_review_v10.json"
    )
    blind_path = "protocols/paper_upgrade_v2/independent_blind_day1_registration_acceptance_v1.json"
    blind_status_path = "protocols/paper_upgrade_v2/independent_blind_authoring_kit_v2/status.json"

    data = {
        "v9": read_json(v9_path),
        "v10": read_json(v10_path),
        "v11": read_json(v11_path),
        "primary": read_json(primary_path),
        "primary_review": read_json(primary_review_path),
        "second": read_json(second_path),
        "gemini_s0_receipt": read_json(gemini_s0_receipt_path),
        "gemini_s1_receipt": read_json(gemini_s1_receipt_path),
        "gemini_s0_review": read_json(gemini_s0_review_path),
        "gemini_s1_review": read_json(gemini_s1_review_path),
        "pro50": read_json(pro50_receipt_path),
        "pro200": read_json(pro200_path),
        "pro500": read_json(pro500_path),
        "pro500_admission": read_json(pro500_admission_path),
        "gate7": read_json(gate7_path),
        "v10_review": read_json(v10_review_path),
        "blind": read_json(blind_path),
        "blind_status": read_json(blind_status_path),
    }
    data["paths"] = {
        "v9": v9_path,
        "v10": v10_path,
        "v11": v11_path,
        "primary": primary_path,
        "primary_review": primary_review_path,
        "second": second_path,
        "gemini_s0_receipt": gemini_s0_receipt_path,
        "gemini_s1_receipt": gemini_s1_receipt_path,
        "gemini_s0_review": gemini_s0_review_path,
        "gemini_s1_review": gemini_s1_review_path,
        "pro50": pro50_receipt_path,
        "pro200": pro200_path,
        "pro500": pro500_path,
        "pro500_admission": pro500_admission_path,
        "gate7": gate7_path,
        "v10_review": v10_review_path,
        "blind": blind_path,
        "blind_status": blind_status_path,
    }

    # Fail closed if the report is accidentally built against different evidence.
    assert data["v9"]["exit_code"] == 0 and data["v9"]["evidence_gate_passed"] is True
    assert data["v10"]["exit_code"] == 1 and data["v10"]["failed_harness_stage_count"] == 2
    assert data["v11"]["exit_code"] == -1 and data["v11"]["evidence_gate_passed"] is False
    assert data["second"]["coverage"]["role_records"] == 1791
    assert data["second"]["coverage"]["diagnostic_records"] == 16
    assert data["gemini_s0_review"]["observed_records"] == 1
    assert data["gemini_s1_review"]["observed_records"] == 1
    assert data["pro50"]["technical_review"]["terminal_task_count"] == 32
    assert data["pro200"]["independent_stable_technical_recomputation"]["authorized_task_count_from_candidate"] == 39
    assert data["pro500"]["scope"]["authorized_task_count"] == 36
    assert data["blind"]["evidence_boundary"]["blind_dataset_authoring_completed"] is False
    assert data["blind_status"]["target_cases"] == 600
    return data


def configure_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.header_distance = Inches(0.36)
    section.footer_distance = Inches(0.36)

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    left = paragraph.add_run("M3-PI GUARD  |  实验任务中止阶段报告")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = paragraph.add_run("\t2026-08-07")
    set_run_font(right, size=8.5, color=MUTED)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.45))

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(5)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.2, color=INK)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Caption")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, size=8.4, color=MUTED, italic=True)


def add_monospace(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.left_indent = Inches(0.15)
    run = paragraph.add_run(text)
    set_run_font(run, name="Consolas", size=8.0, color=MUTED)


def draw_progress_chart(path: Path) -> None:
    rows = [
        ("DeepSeek Flash Public", 1791, 1791, "记录"),
        ("Gemini Public", 2, 1791, "记录"),
        ("Pro B50 达到预算", 39, 40, "任务"),
        ("Pro B200 达到预算", 36, 40, "任务"),
        ("Pro B500", 0, 36, "任务"),
        ("hardest-20 × 2000", 0, 20, "任务"),
        ("独立盲测 authoring", 0, 600, "案例"),
        ("独立盲测六条件执行", 0, 3600, "记录"),
    ]
    image = Image.new("RGB", (1600, 1120), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44, bold=True)
    label_font = pil_font(28)
    small_font = pil_font(22)
    value_font = pil_font(24, bold=True)
    draw.text((72, 38), "冻结工作流数量进度", font=title_font, fill="#172B3A")
    draw.text((72, 100), "每行使用自己的冻结分母；比例不可跨工作流解释为科研质量", font=small_font, fill="#5D6B78")
    left, right = 520, 1500
    top, row_height = 190, 108
    for index, (label, completed, target, unit) in enumerate(rows):
        y = top + index * row_height
        ratio = completed / target if target else 0
        draw.text((72, y + 9), label, font=label_font, fill="#172B3A")
        draw.rounded_rectangle((left, y + 6, right, y + 50), radius=15, fill="#E6EAF0")
        if completed:
            fill_width = max(8, math.floor((right - left) * ratio))
            color = "#2E74B5" if ratio >= 0.9 else "#D6A13D"
            draw.rounded_rectangle((left, y + 6, left + fill_width, y + 50), radius=min(15, fill_width // 2), fill=color)
        detail = f"{completed:,}/{target:,} {unit}"
        percent = f"{ratio * 100:.2f}%" if ratio < 0.01 and completed else f"{ratio * 100:.1f}%"
        draw.text((left, y + 60), detail, font=small_font, fill="#5D6B78")
        box = draw.textbbox((0, 0), percent, font=value_font)
        draw.text((right - (box[2] - box[0]), y + 58), percent, font=value_font, fill="#1F4D78")
    image.save(path, quality=95)


def draw_harness_chart(path: Path) -> None:
    runs = [
        ("v9", 13, 0, 0, "exit 0"),
        ("v10", 11, 2, 0, "exit 1"),
        ("v11", 9, 1, 3, "exit -1 / 中止"),
    ]
    image = Image.new("RGB", (1500, 820), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44, bold=True)
    label_font = pil_font(29)
    small_font = pil_font(23)
    value_font = pil_font(27, bold=True)
    draw.text((70, 40), "Fresh Harness 三次关键运行", font=title_font, fill="#172B3A")
    draw.text((70, 102), "固定13阶段口径；v11在 frozen-study-assets 阶段中止", font=small_font, fill="#5D6B78")
    left, right, top = 300, 1400, 210
    width = right - left
    colors = {"pass": "#2E74B5", "fail": "#C75B5B", "incomplete": "#C7CED6"}
    for index, (name, passed, failed, incomplete, exit_label) in enumerate(runs):
        y = top + index * 165
        draw.text((75, y + 10), f"Fresh {name}", font=label_font, fill="#172B3A")
        cursor = left
        for value, key in ((passed, "pass"), (failed, "fail"), (incomplete, "incomplete")):
            if value:
                segment = width * value / 13
                draw.rectangle((cursor, y, cursor + segment, y + 58), fill=colors[key])
                if segment > 65:
                    draw.text((cursor + segment / 2 - 12, y + 13), str(value), font=value_font, fill="#FFFFFF" if key != "incomplete" else "#44515E")
                cursor += segment
        draw.text((left, y + 72), f"PASS {passed}  ·  FAIL {failed}  ·  未完成 {incomplete}", font=small_font, fill="#5D6B78")
        box = draw.textbbox((0, 0), exit_label, font=small_font)
        draw.text((right - (box[2] - box[0]), y + 72), exit_label, font=small_font, fill="#9B1C1C" if failed else "#1F4D78")
    draw.rectangle((980, 55, 1018, 89), fill=colors["pass"])
    draw.text((1030, 57), "通过", font=small_font, fill="#172B3A")
    draw.rectangle((1120, 55, 1158, 89), fill=colors["fail"])
    draw.text((1170, 57), "失败", font=small_font, fill="#172B3A")
    draw.rectangle((1260, 55, 1298, 89), fill=colors["incomplete"])
    draw.text((1310, 57), "未完成", font=small_font, fill="#172B3A")
    image.save(path, quality=95)


def draw_adaptive_chart(path: Path) -> None:
    labels = ["B50", "B200", "B500", "B2000 hardest-20"]
    reached = [39, 36, 0, 0]
    targets = [40, 40, 36, 20]
    image = Image.new("RGB", (1500, 820), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44, bold=True)
    label_font = pil_font(28)
    small_font = pil_font(23)
    value_font = pil_font(28, bold=True)
    draw.text((70, 42), "DeepSeek Pro 自适应预算阶梯", font=title_font, fill="#172B3A")
    draw.text((70, 104), "柱高为达到相应冻结预算节点的任务数；diagnostic 保留在 ITT", font=small_font, fill="#5D6B78")
    left, top, right, bottom = 150, 210, 1410, 660
    draw.line((left, bottom, right, bottom), fill="#A7B2BD", width=3)
    for tick in range(0, 41, 10):
        y = bottom - tick / 40 * (bottom - top)
        draw.line((left, y, right, y), fill="#E3E8ED", width=2)
        draw.text((82, y - 13), str(tick), font=small_font, fill="#5D6B78")
    gap = (right - left) / len(labels)
    for index, (label, value, target) in enumerate(zip(labels, reached, targets)):
        center = left + gap * (index + 0.5)
        bar_width = 150
        height = value / 40 * (bottom - top)
        color = "#2E74B5" if value else "#C7CED6"
        if value:
            draw.rounded_rectangle((center - bar_width / 2, bottom - height, center + bar_width / 2, bottom), radius=14, fill=color)
        else:
            draw.rounded_rectangle((center - bar_width / 2, bottom - 8, center + bar_width / 2, bottom), radius=4, fill=color)
        text = f"{value}/{target}"
        box = draw.textbbox((0, 0), text, font=value_font)
        draw.text((center - (box[2] - box[0]) / 2, bottom - height - 48), text, font=value_font, fill="#1F4D78")
        box = draw.textbbox((0, 0), label, font=label_font)
        draw.text((center - (box[2] - box[0]) / 2, 690), label, font=label_font, fill="#172B3A")
    image.save(path, quality=95)


def draw_primary_outcomes(path: Path) -> None:
    conditions = ["M3-full", "结构化基线", "policy-only"]
    unsafe = [6.33, 22.17, 2.17]
    utility = [29.17, 61.31, 29.17]
    image = Image.new("RGB", (1600, 900), WHITE)
    draw = ImageDraw.Draw(image)
    title_font = pil_font(44, bold=True)
    label_font = pil_font(27)
    small_font = pil_font(22)
    value_font = pil_font(25, bold=True)
    draw.text((72, 38), "DeepSeek Primary Public：安全—效用结果", font=title_font, fill="#172B3A")
    draw.text((72, 100), "已准入的单模型结果；左侧越低越好，右侧越高越好", font=small_font, fill="#5D6B78")

    def panel(x0: int, title: str, values: list[float], limit: float, color: str, denominator: str) -> None:
        draw.text((x0, 175), title, font=label_font, fill="#172B3A")
        draw.text((x0, 215), denominator, font=small_font, fill="#5D6B78")
        left, right = x0 + 220, x0 + 680
        for index, (condition, value) in enumerate(zip(conditions, values)):
            y = 310 + index * 155
            draw.text((x0, y + 8), condition, font=label_font, fill="#172B3A")
            draw.rounded_rectangle((left, y + 5, right, y + 55), radius=14, fill="#E6EAF0")
            fill = max(7, (right - left) * value / limit)
            draw.rounded_rectangle((left, y + 5, left + fill, y + 55), radius=min(14, int(fill / 2)), fill=color)
            draw.text((left, y + 68), f"{value:.2f}%", font=value_font, fill="#1F4D78")

    panel(70, "攻击 unsafe-effect rate", unsafe, 25.0, "#C75B5B", "n = 600 paired attacks")
    panel(850, "良性工具 TCR", utility, 70.0, "#2E74B5", "n = 168 paired benign-tool tasks")
    image.save(path, quality=95)


def add_cover(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run("实验任务中止 · 阶段技术报告")
    set_run_font(run, size=10, color=RISK, bold=True)

    paragraph = doc.add_paragraph(style="Title")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run("M3-PI Guard")
    set_run_font(run, size=28, color=INK, bold=True)

    paragraph = doc.add_paragraph(style="Subtitle")
    run = paragraph.add_run("21天实验结项计划的中止点、有效证据与剩余缺口")
    set_run_font(run, size=14, color=MUTED)

    metadata = [
        ("报告日期", "2026年8月7日"),
        ("证据截止", "2026年8月7日 22:00（Asia/Shanghai）"),
        ("中止依据", "项目负责人明确指令：中止当前任务并形成报告"),
        ("当前状态", "实验未结项；F-PAPER-002 保持 in_progress"),
        ("安全边界", "无继续执行、无API调用、无真实工具或外部 effect"),
        ("阅读对象", "项目负责人、研究合作者、后续接手者"),
    ]
    for label, value in metadata:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        left = paragraph.add_run(f"{label}：")
        set_run_font(left, size=10, color=INK, bold=True)
        right = paragraph.add_run(value)
        set_run_font(right, size=10, color=INK)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_callout(
        doc,
        "结论",
        "当前实验已安全停止，但未达到21天结项验收。最新 Fresh Harness v11 为非零中止；Gemini、完整 Adaptive、独立盲测与最终统计包均未完成。现有证据足以形成可审计的阶段交接，不足以声明实验结项或跨模型机制优势。",
        fill=PALE_RED,
        accent=RISK,
    )
    add_metric_strip(
        doc,
        [
            ("-1", "v11 直接退出码", "用户指令中止；未通过"),
            ("1,791/1,791", "Flash Public", "物理闭包；结果仍封存"),
            ("39/39", "Pro B200 eligible", "技术终态；结果未准入"),
            ("0/600", "独立盲测数据", "仅Day-1登记完成"),
        ],
    )
    add_paragraph(
        doc,
        "报告中的“完成”严格限定为相应层级的物理、技术或科研准入；不同层级不相互替代。",
        color=MUTED,
        italic=True,
        size=9.3,
    )
    doc.add_page_break()


def build_report(data: dict) -> None:
    QA_DIR.mkdir(parents=True, exist_ok=True)
    progress_chart = QA_DIR / "workstream_progress.png"
    harness_chart = QA_DIR / "harness_timeline.png"
    adaptive_chart = QA_DIR / "adaptive_ladder.png"
    primary_chart = QA_DIR / "primary_security_utility.png"
    draw_progress_chart(progress_chart)
    draw_harness_chart(harness_chart)
    draw_adaptive_chart(adaptive_chart)
    draw_primary_outcomes(primary_chart)

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_cover(doc)

    add_heading(doc, "1. 执行摘要", 1)
    add_paragraph(
        doc,
        "项目负责人于2026年8月7日要求中止当前实验任务。Fresh Harness v11 的 child、helper、leaf 与 poller 均已结束；命令行复核未发现其他 v11、verify.ps1 或 Pro-500 相关进程。Pro-500 的 candidate output、runner output 与 activation artifact 均不存在。",
        bold_prefix="项目负责人于2026年8月7日要求中止当前实验任务。",
    )
    add_paragraph(
        doc,
        "中止后的正确工程状态是“安全停止、证据封存、实验未结项”。v11 receipt 直接观察到 child exit -1，evidence_gate_passed=false；因此不能用历史 v9 的绿色结果替代当前源码的最终验收，也不能把 v11 已完成的九个阶段外推为完整 Harness 通过。",
        bold_prefix="中止后的正确工程状态是“安全停止、证据封存、实验未结项”。",
    )
    add_callout(
        doc,
        "当前最重要的判断",
        "现有成果足以支持“单模型经验性安全—效用与授权不确定性研究”的继续方向；但防御机制论文的跨模型、完整自适应攻击、独立盲测和最终工程验收证据仍不完整。",
    )
    add_heading(doc, "1.1 已完成的高价值证据", 2)
    add_bullets(
        doc,
        [
            "DeepSeek Primary Public 已完成5,082条、847案例×6条件的独立准入与统计复算，形成单模型经验性结果。",
            "DeepSeek V4 Flash Public 已实现1,791/1,791物理覆盖，4,193次调用、CNY 1.87793024，零自动重试、工具和外部 effect；结果仍封存。",
            "Adaptive Pro 的B50与B200完成技术结账；B200为39/39 eligible任务终态，新增5,504查询、6,925次调用、CNY 8.0544754。",
            "独立盲测保管人Day-1登记链完成，28项登记检查通过；开发者侧未收到候选或标签。",
            "历史 Fresh Harness v9 取得直接 exit 0、13/13阶段与1,584项pytest通过，可作为历史绿色基线。",
        ],
    )
    add_heading(doc, "1.2 仍然阻断结项的事项", 2)
    add_bullets(
        doc,
        [
            "最新 Fresh Harness v11 未通过，且未产生聚合终态标记。",
            "Gemini Public 仅2/1,791；12条canary尚未形成aggregate结论。",
            "Pro B500、hardest-20×2,000、DeepSeek Flash Adaptive与Gemini Adaptive均未启动。",
            "独立盲测authoring为0/600，六条件执行与评分为0/3,600。",
            "DeepSeek Flash Public与Adaptive结果仍未完成独立科研准入和最终统计重建。",
        ],
    )

    add_heading(doc, "2. 21天计划完成矩阵", 1)
    add_paragraph(
        doc,
        "下表按冻结分母和证据层级统计。物理终态、技术门通过与科研结果准入是三个不同层级；例如1,791条Flash记录已物理闭包，但结果访问仍未授权。",
    )
    status_rows = [
        ["DeepSeek Primary Public", "5,082/5,082", "科研准入完成", "仅单模型 empirical fallback"],
        ["DeepSeek Flash Public", "1,791/1,791", "物理/完整性完成", "结果未准入；clean claim不合格"],
        ["Gemini Public", "2/1,791", "两个技术终态", "尚余1,789；不可称canary aggregate"],
        ["Pro Adaptive B50", "40/40终态；39达B50", "技术结账", "1个diagnostic在q24终止"],
        ["Pro Adaptive B200", "39/39 eligible；36达B200", "技术结账", "4个累计diagnostic；结果未准入"],
        ["Pro Adaptive B500", "0/36", "离线资格化", "组合有效执行授权=false"],
        ["hardest-20 × 2,000", "0/20", "未启动", "必须等待B500完成后冻结选择"],
        ["DeepSeek Flash Adaptive", "0/20（50/200）", "仅离线候选", "未执行"],
        ["Gemini Flash Adaptive", "0/20（50/200）", "仅离线候选", "未执行"],
        ["盲测Day-1登记", "1/1", "登记链完成", "不等于数据集完成"],
        ["盲测authoring", "0/600", "未交付", "无候选、无标签"],
        ["盲测六条件执行/评分", "0/3,600", "未启动", "execution_authorized=false"],
        ["最终current-source Harness", "v11：9 PASS/1 FAIL/3未完成", "未通过", "direct exit -1"],
    ]
    add_table(doc, ["工作流", "数量", "证据状态", "边界"], status_rows, [1.85, 1.35, 1.35, 1.95], font_size=7.9)
    shape = doc.add_picture(str(progress_chart), width=Inches(6.45))
    set_picture_alt(
        shape,
        title="冻结工作流数量进度",
        description="横向进度条显示DeepSeek Flash Public为1791/1791，Gemini Public为2/1791，Pro B50为39/40达到预算，B200为36/40达到预算，其余主要工作流为零。",
    )
    add_caption(doc, "图1  各工作流按自身冻结分母计算的数量进度；不同单位不可直接比较科研质量。")

    doc.add_page_break()
    add_heading(doc, "3. 中止动作与安全边界", 1)
    add_heading(doc, "3.1 进程与输出状态", 2)
    add_bullets(
        doc,
        [
            "v11 child PID 40976、helper PID 41752、leaf PID 36544、poller PID 20364均已结束。PID 41752之后被Code.exe复用，不代表原helper仍在运行。",
            "v11相关命令行扫描为零；未发现新的verify.ps1、Pro-500或provider工作负载。",
            "Pro-500 candidate output root、guarded-runner root与activation artifact均不存在，说明B500没有被触碰。",
            "v11 receipt记录network_or_provider_call_authorized=false、automatic_restart=false；中止后没有任何API调用。",
        ],
    )
    add_heading(doc, "3.2 v11 的准确归类", 2)
    add_callout(
        doc,
        "不可准入的中止证据",
        "v11 status=completed_child_nonzero，直接OS退出码=-1，evidence gate=false。它完成九个PASS阶段并记录一个FAIL；随后进入frozen-study-assets，但没有该阶段终态，也没有最终aggregate marker。",
        fill=PALE_RED,
        accent=RISK,
    )
    add_paragraph(
        doc,
        "pytest摘要为1,375 passed、1 skipped、1 warning、257 setup errors。257项错误均指向同一个缺失的pytest临时根目录（pytest-221）并触发WinError 3；这表示一个共享测试环境故障面，而不是257个独立产品断言失败。现有证据也不足以认定“用户中止”导致该目录消失，因此报告不作该因果推断。",
    )
    add_paragraph(
        doc,
        "stderr仅含一条乱码化WSL localhost/NAT提示和三条RuntimeAuditor成功拦截delete_record；未见Python traceback或KeyboardInterrupt。RuntimeAuditor的拦截属于安全控制正常生效。",
    )

    add_heading(doc, "4. Harness 演进与当前验收结论", 1)
    shape = doc.add_picture(str(harness_chart), width=Inches(6.45))
    set_picture_alt(
        shape,
        title="Fresh Harness v9到v11阶段结果",
        description="v9为13通过、退出0；v10为11通过2失败、退出1；v11为9通过1失败3未完成、退出负1。",
    )
    add_caption(doc, "图2  Fresh Harness关键运行的13阶段终态。")
    add_paragraph(
        doc,
        "v9是可信的历史绿色基线：直接exit 0、evidence gate通过、13/13阶段、1,584 passed/1 skipped/3 warnings。但其后源码与证据门发生变化，因此不能把v9永久当作当前源码的最终验收。",
    )
    add_paragraph(
        doc,
        "v10直接exit 1、11/13阶段通过。唯一pytest失败源于Pro-200完成后仍动态重放旧的prospective gate v6；第二个失败为对应closeout prospective阶段。不可变failure review已将其归类为状态迁移失配，而非provider或科研结果失败。",
    )
    add_paragraph(
        doc,
        "v11原本用于验证修订后的current-source Harness，但出现共享临时目录setup错误，随后按负责人指令停止。因此当前DoD的“最终完整Harness exit 0”未满足，F-PAPER-002不得标记verified。",
    )

    doc.add_page_break()
    add_heading(doc, "5. Public矩阵进展与科研含义", 1)
    add_heading(doc, "5.1 DeepSeek Primary：已准入的单模型经验性结果", 2)
    add_paragraph(
        doc,
        "Primary Public覆盖847个案例、6个条件，共5,082条记录；76条diagnostic全部保留在ITT，官方调用12,808次，累计CNY 11.9909764。分析结果已在独立准入后复算，允许报告单模型效果，但不授权跨模型、SOTA或clean-mechanism结论。",
    )
    shape = doc.add_picture(str(primary_chart), width=Inches(6.45))
    set_picture_alt(
        shape,
        title="DeepSeek Primary Public安全和效用结果",
        description="M3-full、结构化基线、policy-only的unsafe-effect分别为6.33%、22.17%、2.17%；良性工具任务完成率分别为29.17%、61.31%、29.17%。",
    )
    add_caption(doc, "图3  已准入的DeepSeek Primary Public结果；仅代表该模型与冻结基准。")
    add_paragraph(
        doc,
        "安全侧，M3-full为38/600（6.33%），低于结构化基线133/600（22.17%），绝对差-15.83个百分点；但policy-only仅13/600（2.17%），M3-full反而高4.17个百分点，模块独立增益门失败。",
    )
    add_paragraph(
        doc,
        "效用侧，M3-full与policy-only均为49/168（29.17%），明显低于结构化基线103/168（61.31%），绝对差-32.14个百分点，未满足-5个百分点非劣界限。机制论文门因此保持fail_redirect，当前更适合经验性安全—效用与授权不确定性叙事。",
    )

    add_heading(doc, "5.2 DeepSeek Flash：物理闭包完成，科研结果仍封存", 2)
    add_bullets(
        doc,
        [
            "1,791/1,791唯一工作键终态；12条canary和1,779条continuation/recovery均进入证据链。",
            "16条diagnostic，占0.893%；4,193次provider调用；累计CNY 1.87793024/5。",
            "automatic retries=0，provider tools=false，external effects=false；所有diagnostic保留在ITT。",
            "第二模型完整性已完成，并具备独立结果准入的技术前提；但result_admission=false、scientific outcome access=false。",
            "一个分片未通过冻结的1%质量门，不能用全局0.893%覆盖分片失败，clean_mechanism_quality_claim_eligible=false。",
        ],
    )

    add_heading(doc, "5.3 Gemini：只有两个技术终态", 2)
    add_paragraph(
        doc,
        "Gemini 3.6 Flash的正式目标为1,791条；当前v5 recovery s0与s1各完成1条，共2/1,791（0.11%）。两段均technical gate pass、paid ceiling USD 0、automatic retries=0、tools/effects=false；s1有1个post-response runner failure并按ITT保留。科研结果未准入。",
    )
    add_callout(
        doc,
        "免费层工期约束",
        "当前continuation候选拆为1,186个segment并要求每天一个segment，至少需要1,186个quota day。若坚持“免费层+当前冻结调度”，21个日历日完成Gemini矩阵在结构上不可行；只能合法重设方案或顺延，不能把诊断/少量canary伪装成完整第二模型证据。",
        fill=LIGHT_GRAY,
        accent=DARK_BLUE,
    )

    doc.add_page_break()
    add_heading(doc, "6. Adaptive实验进展", 1)
    shape = doc.add_picture(str(adaptive_chart), width=Inches(6.45))
    set_picture_alt(
        shape,
        title="DeepSeek Pro自适应预算阶梯",
        description="40任务cohort中39个达到预算50、36个达到预算200，预算500和hardest-20预算2000均为零。",
    )
    add_caption(doc, "图4  DeepSeek Pro达到各冻结预算节点的任务数。")
    add_heading(doc, "6.1 B50 技术结账", 2)
    add_paragraph(
        doc,
        "40个任务槽位均形成终态。冻结前缀8任务产生374个terminal queries、562次调用、CNY 0.552805；其余32个正式任务产生1,600个terminal queries、1,968次调用、CNY 1.9984298。合计1,974个terminal queries、2,530次调用、CNY 2.5512348。39个任务达到B50，1个diagnostic在q24终止并保留在ITT。",
    )
    add_heading(doc, "6.2 B200 技术结账与host-timeout边界", 2)
    add_paragraph(
        doc,
        "39个eligible任务全部形成终态，其中36个search terminal、3个diagnostic terminal；加上B50已有的1个diagnostic，原40任务cohort中共有4个累计diagnostic。B200新增5,504个terminal queries、6,925次provider调用、CNY 8.0544754/20，零自动重试、工具和外部effect。",
    )
    add_callout(
        doc,
        "必须保留的证据缺口",
        "原Pro-200 wrapper的receipt缺失，禁止补造；watcher记录为direct=true但exit=null，独立审计将execute direct exit重新判定为未证明。只有recovered technical-review子进程的direct exit 0得到证明；execute effective exit 0是依据唯一终态JSON与锁定控制器契约推断。",
        fill=PALE_RED,
        accent=RISK,
    )
    add_heading(doc, "6.3 B500及后续预算未启动", 2)
    add_paragraph(
        doc,
        "Pro-500候选已离线资格化：36任务、200→500、最大新增查询10,800、阶段硬熔断CNY 35、4个累计diagnostic不得重试。虽然owner admission内部写有stage_execution_authorized=true，但gate-v7明确provider_execution_authorized_by_gate=false，v11未通过且activation artifact缺失，因此组合后的有效provider执行授权=false。",
    )
    add_paragraph(
        doc,
        "B500未执行，hardest-20×2,000也未选择或运行。DeepSeek Flash和Gemini Flash各20任务的50/200面板均只有离线候选，完成量为0。",
    )

    add_heading(doc, "7. 独立盲测状态", 1)
    add_paragraph(
        doc,
        "Day-1保管人登记已完成：开发者侧接受了公钥、登记文件与验证回执三件套，28项登记检查通过，custodian与verifier身份摘要不同，登记链中标签数为0。该证据只证明保管关系和计时技术门，不证明600条数据已经生成。",
    )
    add_table(
        doc,
        ["盲测环节", "目标", "当前", "准确结论"],
        [
            ["Day-1登记", "1组", "1/1", "完成"],
            ["候选authoring", "600条", "0/600", "未收到候选文件"],
            ["独立标签", "600条", "0/600", "未收到、未读取"],
            ["六条件执行", "3,600条", "0/3,600", "未授权、未启动"],
            ["保管人评分/签名", "1套", "0", "未到解封阶段"],
        ],
        [1.45, 1.1, 1.1, 2.75],
        font_size=8.5,
    )
    add_paragraph(
        doc,
        "因此不能使用“独立盲测完成”“600条已封存”或“外部有效性已验证”等措辞。当前只允许表述为“独立保管人登记链完成，authoring与执行待第三方交付”。",
    )

    doc.add_page_break()
    add_heading(doc, "8. 证据质量与可宣称边界", 1)
    add_heading(doc, "8.1 当前可以公开陈述", 2)
    add_bullets(
        doc,
        [
            "DeepSeek Flash Public的1,791条物理覆盖、调用/成本、diagnostic与安全控制事实。",
            "DeepSeek Primary Public的单模型经验性安全—效用结果，以及机制门转向fallback的事实。",
            "Adaptive Pro B50/B200的任务、查询、调用、成本和host-timeout技术证据边界。",
            "Gemini官方免费层已产生两个可审计技术终态，但矩阵尚未完成。",
            "独立盲测Day-1登记完成，数据与标签尚未交付。",
            "全部已运行工作流中provider tools、真实外部effect与automatic retries保持为零。",
        ],
    )
    add_heading(doc, "8.2 当前绝对不能宣称", 2)
    add_bullets(
        doc,
        [
            "21天实验计划已经结项，或最终current-source Harness已通过。",
            "Gemini 1,791矩阵、12条aggregate canary或跨模型泛化已完成。",
            "Pro B500、hardest-20×2,000、任一Flash Adaptive或600/3,600盲测已完成。",
            "M3-full相对policy-only存在独立正向安全增益，或TCR满足-5个百分点非劣界限。",
            "DeepSeek Flash结果已经准入，或分片质量门可被全局诊断率豁免。",
            "Pro-200原wrapper/direct child exit 0已被直接证明。",
            "first、only、SOTA、跨提供商或跨模型普适结论。",
        ],
    )
    add_heading(doc, "8.3 当前研究路线判断", 2)
    add_callout(
        doc,
        "推荐保留 empirical fallback 路线",
        "已准入的Primary结果显示M3-full相对结构化基线更安全，但不优于policy-only且效用明显较弱。除非后续第二模型、Adaptive和盲测给出一致且可解释的新证据，否则论文应聚焦授权泄漏、授权不确定性与安全—效用权衡，而不是防御机制领先主张。",
    )

    add_heading(doc, "9. 若未来恢复，必须重新满足的门槛", 1)
    add_numbered_items(
        doc,
        [
            "将v11作为不可变非零中止证据封存，不覆写receipt、stdout或stderr。",
            "用新版本Harness和launcher指定稳定、独立的pytest临时根，并取得current-source直接exit 0、13/13阶段与唯一终态marker。",
            "重新生成非授权activation gate，绑定新Harness receipt及当前Pro-500 candidate/admission/runner哈希；独立验证通过后才可单次启动B500。",
            "完成B500后按冻结规则选择hardest-20，再执行2000预算；终态不得重试，diagnostic继续留在ITT。",
            "为Gemini免费层重新评估合法可行的调度；若仍锁定1,186 quota days，则必须顺延，不得在21日口径内伪装完成。",
            "由第三方完成600条候选、独立标签与custody seal；开发者只导入无标签执行集，并在解封前冻结代码、阈值和排除规则。",
            "完成三模型未准入结果的独立准入、统计复算、失败台账和最终SHA-256证据包，再讨论实验结项。",
        ],
    )
    add_paragraph(
        doc,
        "本报告不构成上述任何执行授权。若项目未来恢复，必须重新获得明确负责人指令，并从未触碰的pending项开始。",
        bold_prefix="本报告不构成上述任何执行授权。",
    )

    doc.add_page_break()
    add_heading(doc, "10. 最终判定", 1)
    add_callout(
        doc,
        "判定：任务已中止，项目实验未结项",
        "当前最强成果是DeepSeek两套Public证据链、Pro B50/B200技术阶段和独立盲测Day-1登记。核心缺口是最新Harness、Gemini剩余1,789条、B500/2000、两套Flash Adaptive、600/3,600盲测及所有未准入结果的独立统计。",
        fill=PALE_RED,
        accent=RISK,
    )
    add_paragraph(
        doc,
        "在“Gemini仅免费层+现有1,186日段调度”不改变的前提下，原21个日历日目标已不具可行性。停止是正确的证据保护动作：它避免在Harness红灯、执行授权不完整和关键外部依赖缺失时继续扩大成本或产生难以准入的结果。",
    )
    add_paragraph(
        doc,
        "项目仍具继续研究价值。现有Primary结果已经给出清晰而有研究意义的负面/混合证据：完整机制未带来policy-only之上的安全增益，且效用显著受损。这类证据若与来源审计、授权regime和失败异质性结合，仍可支撑面向网络安全的实证研究，但必须如实收缩主张。",
    )

    add_heading(doc, "附录A 关键证据登记", 1)
    evidence_rows = [
        ["Fresh Harness v9 receipt", "exit 0；历史绿色", "14ef9b5a…6f0e44"],
        ["Fresh Harness v10 receipt", "exit 1；11/13", "e8ab014c…37f9d"],
        ["Fresh Harness v10 failure review", "状态迁移失配", "f35a4d9f…80bc"],
        ["Fresh Harness v11 receipt", "exit -1；中止", "75f4bbb5…9f7da"],
        ["Primary analysis summary v5", "单模型fallback完成", "aef68965…d6cd5"],
        ["Primary independent review v5", "独立复算PASS", "14aa8184…bd21"],
        ["Flash second evidence merge", "1,791物理闭包", "7fba09cc…432b"],
        ["Gemini s0 technical review", "1条技术终态", "86e48d12…60c2"],
        ["Gemini s1 technical review", "1条；1 ITT runner failure", "8223d08b…dde"],
        ["Pro B50 technical review", "32正式任务技术门", "1c419d1c…1384"],
        ["Pro B200 host-timeout review", "39任务技术闭包", "f927367f…08fc"],
        ["Pro B500 candidate", "离线、未执行", "947cd58d…3214"],
        ["Pro B500 gate-v7", "非授权门", "a9e750b2…31eb"],
        ["Blind Day-1 acceptance", "登记链完成", "94e5a647…745c"],
    ]
    add_table(doc, ["证据", "结论", "内部manifest"], evidence_rows, [2.55, 2.25, 1.7], font_size=7.9)

    add_heading(doc, "附录B 完整SHA-256与证据位置", 1)
    full_hashes = [
        ("v9 receipt", data["paths"]["v9"], sha256(data["paths"]["v9"])),
        ("v10 receipt", data["paths"]["v10"], sha256(data["paths"]["v10"])),
        ("v10 failure review", data["paths"]["v10_review"], sha256(data["paths"]["v10_review"])),
        ("v11 launch", "outputs/verification/experiment_closeout_fresh_harness_v11_20260807/launch.json", sha256("outputs/verification/experiment_closeout_fresh_harness_v11_20260807/launch.json")),
        ("v11 running", "outputs/verification/experiment_closeout_fresh_harness_v11_20260807/running.json", sha256("outputs/verification/experiment_closeout_fresh_harness_v11_20260807/running.json")),
        ("v11 stdout", "outputs/verification/experiment_closeout_fresh_harness_v11_20260807/stdout.log", sha256("outputs/verification/experiment_closeout_fresh_harness_v11_20260807/stdout.log")),
        ("v11 stderr", "outputs/verification/experiment_closeout_fresh_harness_v11_20260807/stderr.log", sha256("outputs/verification/experiment_closeout_fresh_harness_v11_20260807/stderr.log")),
        ("v11 receipt", data["paths"]["v11"], sha256(data["paths"]["v11"])),
        ("Primary summary v5", data["paths"]["primary"], sha256(data["paths"]["primary"])),
        ("Primary independent review v5", data["paths"]["primary_review"], sha256(data["paths"]["primary_review"])),
        ("Flash second evidence merge", data["paths"]["second"], sha256(data["paths"]["second"])),
        ("Gemini s0 receipt", data["paths"]["gemini_s0_receipt"], sha256(data["paths"]["gemini_s0_receipt"])),
        ("Gemini s1 receipt", data["paths"]["gemini_s1_receipt"], sha256(data["paths"]["gemini_s1_receipt"])),
        ("Pro B50 receipt", data["paths"]["pro50"], sha256(data["paths"]["pro50"])),
        ("Pro B200 review", data["paths"]["pro200"], sha256(data["paths"]["pro200"])),
        ("Pro B500 candidate", data["paths"]["pro500"], sha256(data["paths"]["pro500"])),
        ("Pro B500 admission", data["paths"]["pro500_admission"], sha256(data["paths"]["pro500_admission"])),
        ("Pro B500 gate-v7", data["paths"]["gate7"], sha256(data["paths"]["gate7"])),
        ("Blind Day-1 acceptance", data["paths"]["blind"], sha256(data["paths"]["blind"])),
    ]
    for label, relative, digest in full_hashes:
        add_monospace(doc, f"{label}\n{relative}\nSHA-256  {digest}")

    add_source_note(
        doc,
        "报告只读取上述summary/review/receipt工件；未打开密钥、盲测标签、query ledger、provider trace或未准入的科研结果明细。",
    )
    # The accessibility checker expects a declared first-row header on every
    # table, including the one-row callout and metric-strip layout tables.
    for table in doc.tables:
        set_repeat_table_header(table.rows[0])
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


def main() -> int:
    data = collect_evidence()
    build_report(data)
    print(json.dumps({"output": str(OUTPUT), "bytes": OUTPUT.stat().st_size, "sha256": sha256(str(OUTPUT.relative_to(ROOT)))}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
